from __future__ import annotations

import io
import hashlib
import unittest
import unicodedata
from unittest import mock

from openpyxl import load_workbook

import shortdrama_mcp.file_io as file_io
from shortdrama_mcp.file_io import (
    extract_xlsx,
    extract_xlsx_for_controlled_handoff_revalidation,
    storyboard_to_xlsx,
    workbook_to_storyboard,
)
from shortdrama_mcp.materializer import (
    COMPACT_SHOT_FIELDS,
    FORMAL_MERGED_FIELD,
    GPT_KEYFRAME_FIELD,
    LEGACY_FORMAL_MERGED_FIELD,
    MOTHER_TEMPLATE_COLUMNS,
    legacy_v050_mother_template_rows,
    legacy_v053_mother_template_rows,
)


def _compact_storyboard() -> dict:
    shot = {field: f"{field}-EP01-001" for field in COMPACT_SHOT_FIELDS}
    shot.update(
        {
            "镜头编号": "EP01-001",
            "时长": 4.0,
            "累计时长": 4.0,
            "功能类型": "对白推进",
            "场景连续性ID": "E01-S01-HALL-MORNING",
            "镜头类型": "中近景",
            "可见主体": "甲；乙",
            "空间站位描述": "甲在左侧，乙在右侧两步外。",
            "起始状态": "甲闭口看向乙，双手收在身前。",
            "机位": "左前方平视50mm。",
            "运镜方式": "固定镜头。",
            "画面内容": (
                "0.0–1.0秒：甲抬眼；"
                "1.0–2.5秒：【甲】第一句。\n"
                "2.5–4.0秒：【甲】第二句。"
            ),
            "音色": "平静克制。",
            "心理音轨": "不适用：本镜无心理音轨。",
            "环境音效": "衣料声与室内呼吸声。",
            "环境动态": "窗纸轻动。",
            "光影": "人物面部保持自然明暗层次。",
            "光源来源": "左侧木格窗晨光。",
            "光源方向": "左向右衰减。",
            "色温倾向": "中性微暖。",
            "生成限制": "9:16竖屏；禁止字幕、水印、额外配乐。",
            "转场类型": "动作切",
            "切镜/转场提示词": "甲收住视线时切。",
            "参考图锚定": "CHAR-A",
            "覆盖原文": "是",
            "原文依据类型": "直接对白",
            "原文支撑句": "【甲】第一句。【甲】第二句。",
        }
    )
    return {
        "project_version": "V16.3-148-R10-StableRuntime",
        "schema_version": "R10.1",
        "workflow_profile": "R10-PG1",
        "field_order": list(COMPACT_SHOT_FIELDS),
        "source_lock": {"aspect_ratio": "9:16"},
        "prevention_evidence": {},
        "shots": [shot],
        "total_duration": 4.0,
    }


def _saved_workbook(workbook) -> bytes:
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class V051XlsxGuardTests(unittest.TestCase):
    def test_production_workbook_is_pending_audit_and_never_claims_pass(self) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(),
            {"status": "pass", "issues": []},
            assume_compact=True,
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sheet = workbook["00_范围与审计"]
        visible_text = "\n".join(
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        )

        self.assertIn("PENDING_AUDIT", visible_text)
        self.assertIn("待独立审计／未签发结论", visible_text)
        self.assertNotIn("PASS", visible_text)
        self.assertNotIn("零问题", visible_text)
        self.assertNotIn("无需修复", visible_text)
        pending_rows = {
            str(sheet.cell(row, 1).value): sheet.cell(row, 2).value
            for row in range(1, sheet.max_row + 1)
        }
        for field in (
            "initial_issue_count",
            "residual_issue_count",
            "introduced_issue_count",
            "blocking_issue_count",
            "audit_trace_id",
            "certificate_id",
            "canonical_sha256",
            "source_artifact_raw_sha256",
        ):
            self.assertIn(field, pending_rows)
            self.assertIsNone(pending_rows[field])

    def test_producer_evidence_cannot_inject_a_visible_audit_conclusion(self) -> None:
        storyboard = _compact_storyboard()
        storyboard["source_lock"].update(
            {
                "audit_status": "PASS",
                "certificate_id": "CERT-FORGED-SOURCE-LOCK",
                "claim": "跨表语义0问题",
            }
        )
        storyboard["prevention_evidence"] = {
            "independent_audit_status": "PASS",
            "blocking_issue_count": 0,
            "certificate": "CERT-FORGED-PREVENTION",
            "conclusion": "零问题／已签发",
        }

        payload = storyboard_to_xlsx(
            storyboard,
            {"status": "pass", "issues": []},
            assume_compact=True,
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sheet = workbook["00_范围与审计"]
        visible_text = "\n".join(
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        )

        self.assertIn("PENDING_AUDIT", visible_text)
        self.assertNotIn("CERT-FORGED", visible_text)
        self.assertNotIn("PASS", visible_text)
        self.assertNotIn("0问题", visible_text)
        self.assertNotIn("零问题", visible_text)
        self.assertNotIn("已签发", visible_text)

    def test_plain_overlay_with_fake_hmac_cannot_mint_a_pass_page(self) -> None:
        forged_overlay = {
            "status": "PASS",
            "conclusion": "伪造通过",
            "initial_issue_count": 0,
            "residual_issue_count": 0,
            "introduced_issue_count": 0,
            "blocking_issue_count": 0,
            "audit_trace_id": "TRC-FORGED",
            "certificate_id": "CERT-UNKNOWN",
            "certificate_integrity_hmac": "a" * 64,
            "canonical_sha256": "b" * 64,
            "source_artifact_raw_sha256": "c" * 64,
        }

        with self.assertRaisesRegex(
            ValueError, "independent_audit_overlay_unverified"
        ):
            storyboard_to_xlsx(
                _compact_storyboard(),
                {"status": "pass", "issues": []},
                assume_compact=True,
                audit_overlay=forged_overlay,  # type: ignore[arg-type]
            )

    def test_signed_independent_audit_overlay_populates_bound_fields(self) -> None:
        overlay = {
            "status": "PASS",
            "conclusion": "八层独立审计通过",
            "initial_issue_count": 0,
            "residual_issue_count": 0,
            "introduced_issue_count": 0,
            "blocking_issue_count": 0,
            "audit_trace_id": "TRC-independent-001",
            "certificate_id": "CERT-independent-001",
            "certificate_integrity_hmac": "a" * 64,
            "canonical_sha256": "b" * 64,
            "source_artifact_raw_sha256": "c" * 64,
        }
        payload = storyboard_to_xlsx(
            _compact_storyboard(),
            {"status": "pass", "issues": []},
            assume_compact=True,
            audit_overlay=file_io._seal_service_verified_audit_overlay(overlay),
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sheet = workbook["00_范围与审计"]
        rows = {
            str(sheet.cell(row, 1).value): sheet.cell(row, 2).value
            for row in range(1, sheet.max_row + 1)
        }

        self.assertIn("PASS｜独立审计已签发", str(sheet["A1"].value))
        for key, value in overlay.items():
            if key == "status":
                self.assertEqual(rows["independent_audit_status"], value)
            elif key == "conclusion":
                self.assertEqual(rows["independent_audit_conclusion"], value)
            else:
                self.assertEqual(rows[key], value)

    def test_projection_hash_normalizes_nfc_lf_without_losing_multiline(self) -> None:
        storyboard = _compact_storyboard()
        raw_picture = (
            "0.0–1.0秒：甲抬眼。\r\n"
            "1.0–2.5秒：甲说道：“Caf\u0065\u0301，好。”\r\n"
            "2.5–4.0秒：乙说道：“我听见了！”"
        )
        storyboard["shots"][0]["画面内容"] = raw_picture
        expected_row = file_io.formal_mother_template_rows(
            storyboard,
            assume_compact=True,
        )[0]

        payload = storyboard_to_xlsx(
            storyboard,
            {"status": "pass", "issues": []},
            assume_compact=True,
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sheet = workbook["正式生产母版"]
        formal_column = list(MOTHER_TEMPLATE_COLUMNS).index(FORMAL_MERGED_FIELD) + 1
        actual_formal = sheet.cell(2, formal_column).value

        expected_identity = file_io._projection_cell_identity(
            expected_row[FORMAL_MERGED_FIELD]
        )
        actual_identity = file_io._projection_cell_identity(actual_formal)
        self.assertEqual(
            hashlib.sha256(expected_identity.encode("utf-8")).hexdigest(),
            hashlib.sha256(actual_identity.encode("utf-8")).hexdigest(),
        )
        normalized = unicodedata.normalize(
            "NFC",
            str(actual_formal).replace("\r\n", "\n").replace("\r", "\n"),
        )
        self.assertIn("Café，好。”\n乙说道：“我听见了！”", normalized)
        self.assertNotIn("2.5–4.0秒", normalized)
        recovered = extract_xlsx(payload)["canonical_storyboard"]
        self.assertEqual(recovered["shots"][0]["画面内容"], raw_picture)

    def test_repeated_dialogue_and_voice_labels_are_preserved_in_order(self) -> None:
        merged = "\n".join(
            (
                "镜头序号：EP01-001（4.0s；累计4.0s）",
                "镜头类型：中近景",
                "可见主体：甲；乙",
                "空间站位描述：甲在左侧，乙在右侧。",
                "起始状态：甲闭口看向乙。",
                "机位：左前方平视50mm。",
                "运镜方式：固定镜头。",
                "画面内容：0.0–1.0秒：甲抬眼。",
                "台词：【甲】第一句。",
                "台词:【甲】第二句。",
                "发声方式：第一句轻声。",
                "发声方式:第二句压低尾音。",
                "音色：平静克制。",
                "心理音轨：不适用。",
                "环境音效：衣料声。",
                "环境动态：窗纸轻动。",
                "光影：面部自然明暗层次。",
                "光源：左侧窗光。",
                "光线方向：左向右。",
                "色温：中性微暖。",
                "生成限制：9:16竖屏；禁止字幕。",
                "转场类型：动作切",
            )
        )
        alternate_header = (
            "正式分镜母版"
            if FORMAL_MERGED_FIELD != "正式分镜母版"
            else "正式分镜母版·字段合并"
        )
        headers = [
            alternate_header if value == FORMAL_MERGED_FIELD else value
            for value in MOTHER_TEMPLATE_COLUMNS
        ]
        row = {
            "分镜序号": "EP01-001",
            "时长": 4.0,
            "累计时长": 4.0,
            "功能类型/场景连续性ID": "功能类型：对白推进\n场景连续性ID：E01-S01",
            alternate_header: merged,
            GPT_KEYFRAME_FIELD: "旧投影不进入canonical",
            "参考图锚定": "CHAR-A",
            "覆盖原文": "是",
            "原文依据类型": "直接对白",
            "原文支撑句": "第一句。第二句。",
        }

        canonical = workbook_to_storyboard(
            [{"name": "旧版正式母版", "headers": headers, "rows": [row]}]
        )
        picture = canonical["shots"][0]["画面内容"]

        self.assertLess(picture.index("第一句"), picture.index("第二句"))
        voice = canonical["shots"][0]["音色"]
        self.assertIn("平静克制", voice)
        self.assertIn("第一句轻声", voice)
        self.assertIn("第二句压低尾音", voice)
        self.assertLess(voice.index("第一句轻声"), voice.index("第二句压低尾音"))

    def test_visible_projection_drift_is_rejected_with_intact_sidecar(self) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        formal_sheet = next(
            sheet
            for sheet in workbook.worksheets
            if [cell.value for cell in sheet[1]] == list(MOTHER_TEMPLATE_COLUMNS)
        )
        formal_column = list(MOTHER_TEMPLATE_COLUMNS).index(FORMAL_MERGED_FIELD) + 1
        formal_sheet.cell(2, formal_column).value = (
            str(formal_sheet.cell(2, formal_column).value) + "\n未经canonical授权的漂移"
        )

        with self.assertRaisesRegex(
            ValueError, "visible formal-master projection mismatch"
        ):
            extract_xlsx(_saved_workbook(workbook))

    def test_v050_legacy_header_sidecar_remains_readable(self) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        formal_sheet = workbook["正式生产母版"]
        formal_column = list(MOTHER_TEMPLATE_COLUMNS).index(FORMAL_MERGED_FIELD) + 1
        formal_sheet.cell(1, formal_column).value = "正式分镜母版·字段合并"
        legacy_row = legacy_v050_mother_template_rows(_compact_storyboard())[0]
        for column, header in enumerate(MOTHER_TEMPLATE_COLUMNS, 1):
            formal_sheet.cell(2, column).value = legacy_row.get(header)
        sidecar = workbook[file_io.CANONICAL_SIDECAR_SHEET]
        sidecar.delete_cols(9, 3)

        recovered = extract_xlsx(_saved_workbook(workbook))["canonical_storyboard"]

        self.assertEqual(recovered["shots"][0]["镜头编号"], "EP01-001")

    def test_controlled_revalidation_accepts_only_v053_rows_under_short_header(
        self,
    ) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        formal_sheet = workbook["正式生产母版"]
        transitional_row = legacy_v053_mother_template_rows(
            _compact_storyboard()
        )[0]
        for column, header in enumerate(MOTHER_TEMPLATE_COLUMNS, 1):
            formal_sheet.cell(2, column).value = transitional_row.get(
                header
            )
        sidecar = workbook[file_io.CANONICAL_SIDECAR_SHEET]
        sidecar.delete_cols(9, 3)
        transitional_payload = _saved_workbook(workbook)

        with self.assertRaisesRegex(
            ValueError,
            "visible header does not match signed projection contract",
        ):
            extract_xlsx(transitional_payload)

        recovered = extract_xlsx_for_controlled_handoff_revalidation(
            transitional_payload
        )
        proof = recovered["controlled_projection_revalidation"]
        self.assertEqual(
            proof["mode"],
            "legacy_envelope_short_header_v053_rows",
        )
        self.assertTrue(proof["compatibility_applied"])
        self.assertEqual(
            proof["selected_projector"],
            "shortdrama-projector/0.5.3-frozen",
        )
        self.assertTrue(
            proof["visible_rows_match_selected_projector"]
        )
        self.assertFalse(
            proof["visible_projection_matches_sidecar"]
        )
        self.assertEqual(proof["sidecar_sheet_state"], "veryHidden")
        self.assertEqual(
            recovered["canonical_storyboard"], _compact_storyboard()
        )

    def test_controlled_revalidation_rejects_current_rows_under_legacy_envelope(
        self,
    ) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sidecar = workbook[file_io.CANONICAL_SIDECAR_SHEET]
        sidecar.delete_cols(9, 3)

        with self.assertRaisesRegex(
            ValueError, "visible formal-master projection mismatch"
        ):
            extract_xlsx_for_controlled_handoff_revalidation(
                _saved_workbook(workbook)
            )

    def test_controlled_revalidation_rejects_v050_long_header_projection(
        self,
    ) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        formal_sheet = workbook["正式生产母版"]
        formal_sheet.cell(1, 5).value = LEGACY_FORMAL_MERGED_FIELD
        formal_sheet.cell(2, 5).value = legacy_v050_mother_template_rows(
            _compact_storyboard()
        )[0][file_io.FORMAL_MERGED_FIELD]
        sidecar = workbook[file_io.CANONICAL_SIDECAR_SHEET]
        sidecar.delete_cols(9, 3)

        with self.assertRaisesRegex(
            ValueError, "supports only the exact v0.5.3 unsigned projection"
        ):
            extract_xlsx_for_controlled_handoff_revalidation(
                _saved_workbook(workbook)
            )

    def test_controlled_revalidation_accepts_signed_v054_projection_only(
        self,
    ) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sidecar = workbook[file_io.CANONICAL_SIDECAR_SHEET]
        for row in range(2, sidecar.max_row + 1):
            sidecar.cell(row, 9).value = "shortdrama-projector/0.5.4"
        historical_payload = _saved_workbook(workbook)

        with self.assertRaisesRegex(
            ValueError, "projection contract marker mismatch"
        ):
            extract_xlsx(historical_payload)

        recovered = extract_xlsx_for_controlled_handoff_revalidation(
            historical_payload
        )
        proof = recovered["controlled_projection_revalidation"]
        self.assertEqual(proof["mode"], "strict_projection_contract")
        self.assertEqual(
            proof["selected_projector"],
            "shortdrama-projector/0.5.4",
        )
        self.assertTrue(proof["visible_projection_matches_sidecar"])
        self.assertTrue(proof["visible_audit_page_matches_sidecar"])

    def test_controlled_revalidation_rejects_non_very_hidden_sidecar(
        self,
    ) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        sidecar = workbook[file_io.CANONICAL_SIDECAR_SHEET]
        transitional_row = legacy_v053_mother_template_rows(
            _compact_storyboard()
        )[0]
        formal_sheet = workbook["正式生产母版"]
        for column, header in enumerate(MOTHER_TEMPLATE_COLUMNS, 1):
            formal_sheet.cell(2, column).value = transitional_row.get(
                header
            )
        sidecar.delete_cols(9, 3)
        sidecar.sheet_state = "hidden"

        with self.assertRaisesRegex(ValueError, "must remain veryHidden"):
            extract_xlsx_for_controlled_handoff_revalidation(
                _saved_workbook(workbook)
            )

    def test_current_sidecar_cannot_be_downgraded_to_legacy_projection(self) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        formal_sheet = workbook["正式生产母版"]
        formal_column = list(MOTHER_TEMPLATE_COLUMNS).index(FORMAL_MERGED_FIELD) + 1
        formal_sheet.cell(1, formal_column).value = "正式分镜母版·字段合并"
        legacy_row = legacy_v050_mother_template_rows(_compact_storyboard())[0]
        for column, header in enumerate(MOTHER_TEMPLATE_COLUMNS, 1):
            formal_sheet.cell(2, column).value = legacy_row.get(header)

        with self.assertRaisesRegex(
            ValueError, "visible header does not match signed projection contract"
        ):
            extract_xlsx(_saved_workbook(workbook))

    def test_audit_page_and_visible_sheet_set_are_sidecar_bound(self) -> None:
        payload = storyboard_to_xlsx(
            _compact_storyboard(), {"issues": []}, assume_compact=True
        )
        for mutation in ("status", "certificate", "extra_sheet"):
            with self.subTest(mutation=mutation):
                workbook = load_workbook(io.BytesIO(payload), data_only=False)
                scope = workbook["00_范围与审计"]
                if mutation == "status":
                    scope["B4"] = "PASS"
                elif mutation == "certificate":
                    scope["B12"] = "CERT-forged"
                else:
                    workbook.create_sheet("最终审计PASS")["A1"] = "PASS"
                with self.assertRaisesRegex(
                    ValueError,
                    "visible audit page SHA-256 mismatch|unexpected visible sheets",
                ):
                    extract_xlsx(_saved_workbook(workbook))

    def test_signed_overlay_report_must_be_self_consistent(self) -> None:
        base_overlay = {
            "status": "PASS",
            "conclusion": "通过",
            "initial_issue_count": 0,
            "residual_issue_count": 0,
            "introduced_issue_count": 0,
            "blocking_issue_count": 0,
            "audit_trace_id": "TRC-1",
            "certificate_id": "CERT-1",
            "certificate_integrity_hmac": "a" * 64,
            "canonical_sha256": "b" * 64,
            "source_artifact_raw_sha256": "c" * 64,
        }
        bad_cases = (
            (
                {**base_overlay, "status": "PASS"},
                {"issues": [{"error_code": "E-TEST", "severity": "fail"}]},
            ),
            (
                {
                    **base_overlay,
                    "status": "FAIL_CURRENT",
                    "initial_issue_count": 1,
                    "blocking_issue_count": 1,
                },
                {"issues": []},
            ),
            (
                {
                    **base_overlay,
                    "status": "BLOCKED",
                    "initial_issue_count": 1,
                    "blocking_issue_count": 1,
                },
                {"issues": []},
            ),
        )
        for overlay, report in bad_cases:
            with self.subTest(status=overlay["status"]):
                with self.assertRaisesRegex(
                    ValueError,
                    "independent_audit_overlay_",
                ):
                    storyboard_to_xlsx(
                        _compact_storyboard(),
                        report,
                        assume_compact=True,
                        audit_overlay=file_io._seal_service_verified_audit_overlay(
                            overlay
                        ),
                    )

    def test_save_time_roundtrip_blocks_nondeterministic_projection(self) -> None:
        storyboard = _compact_storyboard()
        original_write_table = file_io._write_table

        def write_drifted_table(*args, **kwargs):
            sheet = original_write_table(*args, **kwargs)
            formal_column = list(MOTHER_TEMPLATE_COLUMNS).index(FORMAL_MERGED_FIELD) + 1
            sheet.cell(2, formal_column).value = (
                str(sheet.cell(2, formal_column).value) + "\n写入阶段投影漂移"
            )
            return sheet

        with mock.patch.object(
            file_io,
            "_write_table",
            side_effect=write_drifted_table,
        ):
            with self.assertRaisesRegex(
                ValueError, "visible formal-master projection mismatch"
            ):
                storyboard_to_xlsx(storyboard, {"issues": []}, assume_compact=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
