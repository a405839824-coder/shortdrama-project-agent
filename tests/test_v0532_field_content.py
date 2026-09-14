"""Field ownership, valid counterexamples, fixed projections and actual XLSX."""
from __future__ import annotations
import copy
import hashlib
import io
import json
from pathlib import Path
import sys
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'server'), str(ROOT/'tests'), str(ROOT/'skills/shortdrama-storyboard/scripts')]
from shortdrama_mcp.field_content import field_content_findings, field_policy
from shortdrama_mcp.materializer import compile_formal_merged_field, compile_gpt_keyframe_instruction
from shortdrama_mcp.file_io import storyboard_to_xlsx, extract_xlsx
from shortdrama_mcp.written_payload import context_instruction_text
from shortdrama_mcp.r5_workbook import WIDTHS
from openpyxl import load_workbook
from test_strict_prompt_contract import strict_fixture
from validate_storyboard import validate_project, project_formal_master, project_gpt_keyframe
from test_v050_compact_master_io import compact_storyboard
from test_v0527_source_decision_independent_audit import complete_source_input
from production_input import inspect_production_input
import test_v0527_mcp_production_authoring as mcp_fixtures


def faults(field, text):
    return field_content_findings({field:text}, mask=context_instruction_text)


class FieldContentTests(unittest.TestCase):
    def test_policy_covers_all_canonical_fields_and_voice_labels(self):
        p = field_policy()
        self.assertEqual(list(p['fields']), strict_fixture()['field_order'])
        self.assertEqual(len(p['voice_subfields']), 11)
        self.assertEqual(sum(not r['optional'] for r in p['keyframe_sections']), 10)
        self.assertIsNone(p['remaining_problem_plan']['direct_design_coverage']['measured_resolution'])

    def test_three_previously_accepted_cross_field_probes_are_rejected(self):
        for field, appendix in [('机位','江氏皱眉，右手握紧袖口。'),
                                ('机位','江氏的声音清亮，语速放慢。'),
                                ('音色','摄影机位于门口，平视50mm。')]:
            obj = strict_fixture(); shot = obj['shots'][0]
            if field == '音色':
                shot[field] = shot[field].replace('基础声线：','基础声线：'+appendix)
            else:
                shot[field] += '；'+appendix
            report = validate_project(obj)
            self.assertTrue(any(i['error_code']=='E-FIELD-CONTENT-001' for i in report['issues']))
            for renderer in (compile_formal_merged_field, compile_gpt_keyframe_instruction,
                             project_formal_master, project_gpt_keyframe):
                with self.subTest(field=field,renderer=renderer.__name__), self.assertRaisesRegex(ValueError,'field_content_violation'):
                    renderer(shot)

    def test_voice_subfields_reject_wrong_delivery_dimensions(self):
        cases = [('基础声线','语速放慢。'), ('音量','中低音区。'), ('音量','语速放慢。'),
                 ('语速与节奏','音量降低。'), ('呼吸','尾音上扬。'), ('停顿','重音落在“没来晚”。'),
                 ('重音','逗号处短停。'), ('尾音','开口前轻吸气。'), ('禁止语气','禁止字幕、水印。'),
                 ('发声方式','摄影机位于门口。')]
        for label,value in cases:
            voice = strict_fixture()['shots'][0]['音色']
            voice = voice.replace(label+'：',label+'：'+value)
            with self.subTest(label=label,value=value):
                self.assertTrue(faults('音色',voice))

    def test_voice_preamble_cannot_hide_a_new_instruction(self):
        self.assertTrue(faults('音色','临时说明：镜头推近。\n'+strict_fixture()['shots'][0]['音色']))

    def test_wrong_static_and_global_policy_roles_are_rejected(self):
        cases = [('机位','焦点落在江氏转身后的眼睛。'), ('起始状态','承接上一个镜头。'),
                 ('画面内容','江氏站定；未列人物不得出现。'),
                 ('画面内容','本轮仅核对三人站位。'), ('光源方向','光源离地2米。'),
                 ('空间站位描述','摄影机在门内，平视50mm。')]
        for field,text in cases:
            with self.subTest(field=field,text=text):self.assertTrue(faults(field,text))

    def test_legitimate_role_relations_and_plot_actions_are_not_rejected(self):
        cases = [('机位','摄影机在门内，平视50mm；焦点落在江氏已经握住的袖口。'),
                 ('空间站位描述','江氏在木案北侧，她的目光略偏向摄影机右侧的邹氏。'),
                 ('画面内容','江氏没有开口。她推开窗，吹灭油灯。'),
                 ('画面内容','江氏在“没来晚”重音期间不眨眼。'),
                 ('光源来源','北侧木格窗关闭，光源离地2米。'),
                 ('心理音轨','内心独白：江氏心想：“我该怎么做？”人物不张嘴，不驱动口型，不生成字幕。'),
                 ('生成限制','9:16竖屏；禁止字幕、水印；禁止额外配乐。')]
        for field,text in cases:
            with self.subTest(field=field,text=text):self.assertEqual(faults(field,text),[])

    def test_source_and_real_spoken_literals_are_not_instructions(self):
        literal='摄影机位于门口，平视50mm；语速放慢；承接上一个镜头。'
        for field,text in [('原文支撑句',literal),('画面内容',f'江氏说道：“{literal}”'),
                           ('起始状态',f'纸面现有提示：“{literal}”')]:
            with self.subTest(field=field):self.assertEqual(faults(field,text),[])
        self.assertTrue(faults('起始状态',f'拍摄要求：“{literal}”'))

    def test_new_picture_does_not_need_voice_prose_to_pass_punctuation(self):
        obj=strict_fixture(); result=validate_project(obj)
        self.assertEqual(result['issues'],[])
        self.assertNotIn('尾音克制',obj['shots'][0]['画面内容'])
        obj['shots'][0]['音色']=obj['shots'][0]['音色'].replace('带试探意味的问候','普通问候').replace('试探','确认').replace('克制上抬','未说明')
        self.assertTrue(any(i['error_code']=='E-PUNCT-004' for i in validate_project(obj)['issues']))

    def test_static_projection_has_fixed_r5_routes_and_no_binding_status(self):
        shot=strict_fixture()['shots'][0]
        shot['参考图锚定']='文字造型锚：江氏穿绛红常服，乌发圆髻。\n图片绑定状态：未附图\n来源依据：人物设定表'
        a=compile_gpt_keyframe_instruction(shot); b=project_gpt_keyframe(shot)
        self.assertEqual(a,b)
        labels=[line.split('：',1)[0] for line in a.splitlines()[1:]]
        self.assertEqual(labels,[r['label'] for r in field_policy()['keyframe_sections']])
        for value in ('环境动态：','音色：','运镜方式：','切镜/转场提示词：','未附图','来源依据：'):
            self.assertNotIn(value,a)
        self.assertIn('江氏穿绛红常服',a)
        self.assertEqual(compile_formal_merged_field(shot),project_formal_master(shot))

    def test_legitimate_environment_motion_stays_out_of_static_frame(self):
        shot=strict_fixture()['shots'][0];shot['环境动态']='窗纸随实际风力轻动。'
        self.assertIn('窗纸随实际风力轻动',compile_formal_merged_field(shot))
        self.assertNotIn('窗纸随实际风力轻动',compile_gpt_keyframe_instruction(shot))

    def test_camera_payload_cannot_override_typed_output_fields(self):
        decisions,graph,_=complete_source_input()
        decisions['shots'][0]['camera']['机位']='她皱眉，右手握紧袖口。'
        result=inspect_production_input(decisions,graph_bytes=graph,expected_graph_sha256=hashlib.sha256(graph).hexdigest())
        self.assertFalse(result.get('status')=='pass')
        self.assertIn('camera',str(result))

    def test_r5_saved_workbook_and_readback_preserve_every_field(self):
        obj=compact_storyboard();before=copy.deepcopy(obj)
        data=storyboard_to_xlsx(obj,{'status':'fail','issues':[]})
        w=load_workbook(io.BytesIO(data))
        self.assertEqual(w.sheetnames[:3],['6集_逐镜阅读','正式分镜母版','使用与来源'])
        master=w['正式分镜母版']
        self.assertEqual([master.column_dimensions[c].width for c in 'ABCDEFGHIJ'],list(WIDTHS))
        self.assertEqual(master.max_column,10)
        self.assertEqual(master.cell(2,6).value,compile_gpt_keyframe_instruction(obj['shots'][0]))
        restored=extract_xlsx(data)['canonical_storyboard']
        self.assertEqual(obj,before)
        self.assertEqual(restored['shots'][0]['原文支撑句'],obj['shots'][0]['原文支撑句'])
        w.close()

    def test_reading_page_edit_cannot_hide_behind_unchanged_canonical(self):
        data=storyboard_to_xlsx(compact_storyboard(),{'status':'fail','issues':[]})
        w=load_workbook(io.BytesIO(data));s=w['6集_逐镜阅读'];s['C9']='擅自改变站位。'
        out=io.BytesIO();w.save(out);w.close()
        with self.assertRaisesRegex(ValueError,'reading/source projection mismatch'):
            extract_xlsx(out.getvalue())

    def test_painted_appearance_cannot_carry_attachment_status(self):
        shot=strict_fixture()['shots'][0];shot['参考图锚定']='文字造型锚：未附图，江氏穿红衣。'
        with self.assertRaisesRegex(ValueError,'field_content_violation'):
            compile_gpt_keyframe_instruction(shot)

    def test_previous_workbook_uses_its_frozen_projection_without_rewriting(self):
        from historical_workbook_fixtures import previous_workbook
        from shortdrama_mcp import file_io
        original = previous_workbook(compact_storyboard())
        digest = hashlib.sha256(original).hexdigest()
        restored = extract_xlsx(original)['canonical_storyboard']
        self.assertEqual(restored['shots'][0]['镜头编号'], compact_storyboard()['shots'][0]['镜头编号'])
        self.assertEqual(hashlib.sha256(original).hexdigest(), digest)
        book = load_workbook(io.BytesIO(original))
        sidecar = book[file_io.CANONICAL_SIDECAR_SHEET]
        for row in range(2, sidecar.max_row + 1):
            sidecar.cell(row, 9).value = file_io.CURRENT_PROJECTION_CONTRACT
        output = io.BytesIO(); book.save(output); book.close()
        with self.assertRaises(ValueError):
            extract_xlsx(output.getvalue())

    def test_conflicting_aspect_in_reference_cannot_escape_by_omission_from_gpt(self):
        shot = strict_fixture()['shots'][0]
        shot['参考图锚定'] = '参考16:9横屏构图'
        for projector in (compile_gpt_keyframe_instruction, project_gpt_keyframe):
            with self.assertRaisesRegex(ValueError, 'aspect_ratio_projection_violation'):
                projector(shot)

    def test_foreign_heading_cannot_bypass_roles_as_a_plain_paragraph(self):
        for field, value in [('机位', '平视50mm。\n当前表达：迟疑。'),
                             ('光源方向', '从右侧入射；音色：柔和。'),
                             ('环境动态', '窗纸轻动。\n起始状态：甲站门前。')]:
            shot = strict_fixture()['shots'][0]; shot[field] = value
            for projector in (compile_formal_merged_field, compile_gpt_keyframe_instruction,
                              project_formal_master, project_gpt_keyframe):
                with self.subTest(field=field, projector=projector):
                    with self.assertRaisesRegex(ValueError, 'field_content_violation'):
                        projector(shot)
        shot = strict_fixture()['shots'][0]
        shot['起始状态'] = '纸面当前已有题头：“当前表达：迟疑”；纸页平放。'
        self.assertIn('当前表达：迟疑', compile_gpt_keyframe_instruction(shot))


class FieldContentMcpTests(unittest.TestCase):
    def setUp(self):
        self.fixture=mcp_fixtures.McpProductionAuthoringTests()
        self.fixture.setUp()

    def tearDown(self):
        self.fixture.tearDown()

    def test_raw_preflight_cannot_skip_field_checks(self):
        for bad in ('江氏皱眉，右手握紧袖口。','江氏的声音清亮，语速放慢。'):
            candidate=mcp_fixtures.valid_storyboard()
            candidate['shots'][0]['机位']+='；'+bad
            result=self.fixture.direct(candidate)
            self.assertFalse(result.data['commit_ready'])
            self.assertIn('E-FIELD-CONTENT-001',str(result.data['report']))
            self.assertEqual(result.data['field_content_validation']['input_route'], 'canonical_compatibility')
        self.fixture.assert_no_publication()

    def test_current_status_exposes_the_actual_field_contract(self):
        result = self.fixture.service.agent_status(self.fixture.namespace, self.fixture.project)
        contract = result.data['field_content_contract']
        self.assertEqual(contract['contract'], field_policy()['contract'])
        self.assertEqual(len(contract['fields']), 27)
        self.assertEqual(len(contract['voice_subfields']), 11)
        self.assertEqual(contract['workbook']['formal_sheet'], '正式分镜母版')

    def test_bad_editorial_is_rejected_without_silent_fix_or_publication(self):
        request=copy.deepcopy(self.fixture.request)
        request['decisions']['shots'][0]['editorial']['参考图锚定']='摄影机位于门口，平视50mm。'
        result=self.fixture.direct(request)
        self.assertEqual(result.status,'authoring_rejected')
        self.assertIsNone(result.data['compiled_storyboard'])
        self.assertIsNone(result.data['quality'])
        folder,=self.fixture.records()
        self.assertEqual(json.loads((folder/'request.original.json').read_text(encoding='utf-8')),request)
        self.fixture.assert_no_publication()


if __name__=='__main__':unittest.main()
