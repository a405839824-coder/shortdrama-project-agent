"""External bytes and incomplete reviews cannot acquire production provenance."""
import copy
import hashlib
import io
import inspect
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from jsonschema import Draft202012Validator

from shortdrama_mcp.audit_coverage import build_obligations, coverage_summary, validate_coverage
from shortdrama_mcp.file_io import extract_xlsx
from shortdrama_mcp.materializer import COMPACT_SHOT_FIELDS, MOTHER_TEMPLATE_COLUMNS, LEGACY_MOTHER_TEMPLATE_COLUMNS
from shortdrama_mcp.models import SEMANTIC_AUDIT_LAYER_NAMES
from shortdrama_mcp.service import StoryboardService, _draft_quality_metrics, canonical_storyboard_sha256
from shortdrama_mcp.storage import ProjectStore
from semantic_audit_helpers import synthetic_coverage, semantic_audit_for_artifact
from test_efficiency import valid_storyboard
import test_efficiency as efficiency_support
import validate_storyboard


CONTEXT = {key: 'e' * 64 for key in ('rules', 'schema', 'validator', 'runtime')}


def workbook_bytes(*, legacy=False, reading='复制文字不作为内容权威', frame='单一静态；甲站在门内'):
    book = Workbook()
    sheet = book.active
    sheet.title = '正式分镜母版'
    headers = LEGACY_MOTHER_TEMPLATE_COLUMNS if legacy else MOTHER_TEMPLATE_COLUMNS
    sheet.append(headers)
    shot = {field: '测试内容' for field in COMPACT_SHOT_FIELDS}
    shot.update({'镜头编号': 'EP10-001', '时长': 4, '累计时长': 4,
                 '镜头类型': '中景', '可见主体': '甲', '起始状态': '甲站在门内，双手空着。',
                 '生成限制': '不加字幕、水印、黑边或额外配乐。', '音色': '甲：自然声线。'})
    body = '镜头类型：中景｜甲\n' + '\n'.join(
        f'{field}：{shot[field]}' for field in COMPACT_SHOT_FIELDS[7:23])
    row = {'分镜序号': 'EP10-001', '时长': 4, '累计时长': 4,
           '功能类型/场景连续性ID': '对话 / SCENE-A', '正式分镜母版': body,
           'GPT关键帧生成指令': frame, '光影': '\n'.join(f'{f}：{shot[f]}' for f in COMPACT_SHOT_FIELDS[16:20]),
           '参考图锚定': '甲门内', '覆盖原文': 'P001', '原文依据类型': '原文', '原文支撑句': '甲站在门内。'}
    sheet.append([row.get(header, body if '字段合并' in header else '') for header in headers])
    copy_sheet = book.create_sheet('10集_逐镜阅读')
    copy_sheet.append(headers)
    copy_sheet.append(['FORGED', None, None, None, reading])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class ExternalAuditTests(unittest.TestCase):
    def test_audit_mapping_is_lossless_and_production_parse_stays_strict(self):
        for legacy in (False, True):
            raw = workbook_bytes(legacy=legacy)
            result = extract_xlsx(raw, audit_only=True)
            candidate = result['canonical_storyboard']
            evidence = candidate['audit_workbook_evidence']
            self.assertFalse(result['canonical_sidecar_present'])
            self.assertEqual(evidence['raw_sha256'], hashlib.sha256(raw).hexdigest())
            self.assertEqual(evidence['column_contract'], 10 if legacy else 11)
            self.assertFalse(evidence['production_eligible'])
            self.assertFalse(evidence['historical_canonical_available'])
            self.assertEqual(len(candidate['shots']), 1)
            self.assertEqual(candidate['shots'][0]['生成限制'], '不加字幕、水印、黑边或额外配乐。')
            self.assertEqual(evidence['rows'][0]['cells']['GPT关键帧生成指令']['address'], 'F2')
            if not legacy:
                with self.assertRaisesRegex(ValueError, 'requires its canonical sidecar'):
                    extract_xlsx(raw)

    def test_present_invalid_sidecar_cannot_fall_back_to_unsigned_audit(self):
        book = load_workbook(io.BytesIO(workbook_bytes()))
        book.create_sheet('__shortdrama_canonical_v1').append(['forged sidecar'])
        stream = io.BytesIO(); book.save(stream)
        with self.assertRaises(ValueError):
            extract_xlsx(stream.getvalue(), audit_only=True)

    def test_original_row_coordinates_and_extra_content_are_not_silently_lost(self):
        book = load_workbook(io.BytesIO(workbook_bytes()))
        master = book['正式分镜母版']
        master.insert_rows(1, 2)
        stream = io.BytesIO(); book.save(stream)
        candidate = extract_xlsx(stream.getvalue(), audit_only=True)['canonical_storyboard']
        self.assertEqual(candidate['audit_workbook_evidence']['rows'][0]['cells']['GPT关键帧生成指令']['address'], 'F4')
        master['L4'] = '未经登记的额外执行指令'
        stream = io.BytesIO(); book.save(stream)
        with self.assertRaisesRegex(ValueError, 'unmapped_master_cells: L4'):
            extract_xlsx(stream.getvalue(), audit_only=True)

    def test_signed_legacy_actual_cells_reach_current_audit_without_changing_canonical(self):
        path = Path(__file__).resolve().parent / 'fixtures/v0555_four_visible_original.xlsx'
        raw = path.read_bytes()
        normal = extract_xlsx(raw)
        audited = extract_xlsx(raw, audit_only=True)
        canonical = audited['canonical_storyboard']
        self.assertEqual(canonical, normal['canonical_storyboard'])
        self.assertEqual(canonical_storyboard_sha256(canonical), canonical_storyboard_sha256(normal['canonical_storyboard']))
        self.assertNotIn('audit_workbook_evidence', normal)
        evidence = audited['audit_workbook_evidence']
        self.assertEqual(evidence['projection_contract'], 'shortdrama-projector/0.5.46')
        self.assertTrue(evidence['projection_integrity_verified'])
        self.assertEqual(evidence['rows'][0]['cells']['GPT关键帧生成指令']['address'], 'F2')
        with tempfile.TemporaryDirectory() as directory:
            service = StoryboardService.__new__(StoryboardService)
            service.store = ProjectStore(Path(directory), 'https://test.invalid')
            service.store.create_project('test', {'project_id': 'P', 'work_orders': []}, {'project_id': 'P', 'events': []})
            record = service.store.save_artifact_bytes('test', 'P', raw, file_name='old-signed.xlsx',
                artifact_type='source:storyboard', version='old')
            self.assertEqual(service._saved_workbook_audit_evidence('test', 'P', record['artifact_id']), evidence)
        checked = validate_storyboard.validate_project(canonical, mode='audit', saved_workbook_evidence=evidence)
        checks = checked['quality_metrics']['semantic_checks']
        self.assertTrue(any(row['check_id'] == 'saved_visible_capacity' and row['status'] == 'fail' for row in checks))
        self.assertTrue(any(row['check_id'] == 'saved_projection_binding' and row['status'] == 'pass' for row in checks))
        rejected = validate_storyboard.validate_project(canonical, mode='production', saved_workbook_evidence=evidence)
        self.assertTrue(any(issue['error_code'] == 'E-GATE-001' and issue['location'] == 'saved_workbook_evidence' for issue in rejected['issues']))
        self.assertEqual(path.read_bytes(), raw)

    def test_actual_f_changes_obligations_and_reading_copy_is_not_a_source(self):
        before = extract_xlsx(workbook_bytes(), audit_only=True)['canonical_storyboard']
        after = extract_xlsx(workbook_bytes(frame='甲从站到坐的全过程清楚可见'), audit_only=True)['canonical_storyboard']
        self.assertEqual(before['shots'], after['shots'])
        a, b = build_obligations(before, CONTEXT), build_obligations(after, CONTEXT)
        self.assertEqual(a['obligation_count'], 14)
        self.assertNotEqual(a['manifest_sha256'], b['manifest_sha256'])
        self.assertTrue(any('F2' in ref for row in a['obligations'] for ref in row['required_observation_refs']))
        changed_copy = extract_xlsx(workbook_bytes(reading='恶意副本'), audit_only=True)['canonical_storyboard']
        self.assertEqual(before['shots'], changed_copy['shots'])
        self.assertEqual(before['audit_workbook_evidence']['rows'], changed_copy['audit_workbook_evidence']['rows'])
        self.assertNotEqual(before['audit_workbook_evidence']['raw_sha256'], changed_copy['audit_workbook_evidence']['raw_sha256'])

    def test_external_service_reads_registered_bytes_without_recovering_canonical(self):
        with tempfile.TemporaryDirectory() as directory:
            service = StoryboardService.__new__(StoryboardService)
            service.store = ProjectStore(Path(directory), 'https://test.invalid')
            service.store.create_project('test', {'project_id': 'P', 'work_orders': []}, {'project_id': 'P', 'events': []})
            original = service.store.save_artifact_bytes('test', 'P', workbook_bytes(), file_name='external.xlsx',
                artifact_type='source:storyboard', version='v1', work_order_id='IMPORT-ONLY')
            count = len(service.store.list_artifacts('test', 'P'))
            candidate, record = service._storyboard_from_artifact('test', 'P', original['artifact_id'], require_bound_projection=True)
            self.assertEqual(candidate['audit_workbook_evidence']['raw_sha256'], original['fingerprint'])
            self.assertEqual(count, len(service.store.list_artifacts('test', 'P')))
            self.assertNotIn('canonical_json_artifact_id', record.get('metadata', {}))
            class IntakeReached(Exception): pass
            with patch.object(service, '_require_runtime'), patch.object(service, '_ensure_audit_work_order', return_value={'work_order_id': 'AUDIT'}), \
                 patch.object(service, '_require_runtime_test_certificate'), patch.object(service, '_matching_handoff', return_value=None), \
                 patch.object(service, '_persist_audit_intake', side_effect=IntakeReached):
                with self.assertRaises(IntakeReached):
                    inspect.unwrap(StoryboardService.audit_storyboard)(service, 'test', project_id='P', artifact_id=original['artifact_id'], audit_cycle_id=None)
            _, path = service.store.get_artifact('test', 'P', original['artifact_id'])
            path.write_bytes(workbook_bytes(frame='changed actual F'))
            with self.assertRaisesRegex(ValueError, 'fingerprint_mismatch'):
                service._storyboard_from_artifact('test', 'P', original['artifact_id'], require_bound_projection=True)

    @unittest.skipUnless(os.environ.get('V0556_AUDIT_WORKBOOK_MANIFEST'), 'optional real originals provided by audit workspace')
    def test_real_registered_delivery_originals(self):
        rows = json.loads(Path(os.environ['V0556_AUDIT_WORKBOOK_MANIFEST']).read_text(encoding='utf-8-sig'))
        for entry in rows:
            with self.subTest(episode=entry['episode']):
                path = Path(entry['xlsx']); raw = path.read_bytes()
                self.assertEqual(hashlib.sha256(raw).hexdigest(), entry['xlsx_sha256'])
                candidate = extract_xlsx(raw, audit_only=True)['canonical_storyboard']
                evidence = candidate['audit_workbook_evidence']
                book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
                master = book[evidence['master_sheet']]
                actual_cells = {cell.coordinate: cell.value for row in master.iter_rows() for cell in row if cell.value is not None}
                for row in evidence['rows']:
                    for cell in row['cells'].values():
                        self.assertEqual(cell['value'], actual_cells.get(cell['address']))
                book.close()
                manifest = build_obligations(candidate, CONTEXT)
                self.assertEqual(manifest['obligation_count'], 14 * len(candidate['shots']))
                self.assertEqual(path.read_bytes(), raw)


class SignedWorkbookAuditFlowTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.flow = efficiency_support.EfficiencyContractTests()
        await self.flow.asyncSetUp()

    async def asyncTearDown(self):
        await self.flow.asyncTearDown()

    async def audit_original(self, *, paged):
        h = self.flow
        project = await h.create_project('旧投影实际原件审计')
        raw = (Path(__file__).resolve().parent / 'fixtures/v0555_four_visible_original.xlsx').read_bytes()
        original = h.service.store.save_artifact_bytes(h.namespace, project, raw,
            file_name='original-0555.xlsx', artifact_type='source:storyboard', version='0.5.55')
        identity = canonical_storyboard_sha256(extract_xlsx(raw)['canonical_storyboard'])
        order = await h.begin_storyboard_task(project, operation='storyboard_audit', mode='audit_only')
        submission = semantic_audit_for_artifact(h.service, h.settings, project, original['artifact_id'])
        args = {'project_id': project, 'artifact_id': original['artifact_id'], 'work_order_id': order}
        if paged:
            manifest = submission.pop('coverage')
            pages, collection = [], None
            for offset in range(0, len(manifest['checks']), 64):
                page = {'contract': 'semantic-obligation-pages-v1', 'action': 'append',
                        'collection_key': 'old-sidecar-review', 'manifest_sha256': manifest['manifest_sha256'],
                        'offset': offset, 'checks': manifest['checks'][offset:offset + 64]}
                if collection:
                    page['collection_id'] = collection
                staged = await h.call('audit_storyboard', {**args, 'semantic_audit': {**submission, 'coverage_page': page}})
                self.assertEqual(staged['status'], 'coverage_pending')
                collection = staged['data']['collection_id']
                pages.append(staged['data']['page_id'])
            submission['coverage_page'] = {'contract': 'semantic-obligation-pages-v1', 'action': 'finalize',
                'manifest_sha256': manifest['manifest_sha256'], 'collection_id': collection, 'page_ids': pages}
        audited = await h.call('audit_storyboard', {**args, 'semantic_audit': submission})
        self.assertEqual(audited['status'], 'fail_current')
        cert = h.service.store.get_project_record(h.namespace, project, kind='certificates',
            record_id=audited['data']['certificate']['certificate_id'])
        self.assertEqual(cert['canonical_sha256'], identity)
        self.assertEqual(cert['artifact_raw_sha256'], hashlib.sha256(raw).hexdigest())
        metrics = cert['report']['quality_metrics']
        self.assertEqual(metrics['saved_workbook_evidence']['raw_sha256'], hashlib.sha256(raw).hexdigest())
        failed = [check for check in metrics['semantic_checks'] if check['check_id'] == 'saved_visible_capacity' and check['status'] == 'fail']
        self.assertTrue(failed)
        self.assertTrue(all(check['field_sha256'].get('GPT关键帧生成指令') for check in failed))
        self.assertEqual(cert['semantic_obligation_coverage']['expected_count'], 14 * len(extract_xlsx(raw)['canonical_storyboard']['shots']))
        _, original_path = h.service.store.get_artifact(h.namespace, project, original['artifact_id'])
        self.assertEqual(original_path.read_bytes(), raw)

    async def test_fresh_audit_uses_old_saved_F_and_retains_original_identity(self):
        await self.audit_original(paged=False)

    async def test_paged_finalization_uses_old_saved_F_and_retains_original_identity(self):
        await self.audit_original(paged=True)


class CoverageAccountingTests(unittest.TestCase):
    def test_unknown_cannot_leave_layer_and_gate_pass_or_break_receipt_schema(self):
        candidate = valid_storyboard()
        service = StoryboardService.__new__(StoryboardService)
        for name in ('ruleset', 'validator', 'runtime', 'schema'):
            setattr(service, '_' + name + '_fingerprint', 'e' * 64)
        with patch.object(service, '_storyboard_from_artifact', return_value=(candidate, {})), \
             patch('semantic_audit_helpers.request_namespace', return_value='test'):
            submission = semantic_audit_for_artifact(service, None, 'P', 'A')
        for check in submission['coverage']['checks']:
            if check['obligation_id'].endswith('/source_fidelity'):
                check['status'] = 'unknown'
        # Synthetic independent declarations test state propagation, not story acceptance.
        report = validate_storyboard.validate_project(candidate, mode='audit')
        report['issues'] = []
        report['quality_metrics']['semantic_checks'] = []
        enriched = service._attach_independent_evidence_matrix(report, candidate, submission)
        matrix = enriched['independent_evidence_matrix']
        self.assertEqual(enriched['status'], 'fail_current')
        self.assertEqual(matrix['audit_layer_results']['源文对白']['status'], 'unknown')
        self.assertFalse(any(row['status'] == 'pass' for row in matrix['preventive_gate_results'].values()))
        self.assertIsNone(matrix['quality_thresholds']['source_and_key_dialogue_mapping_percent'])
        schema_path = Path(__file__).resolve().parents[1] / 'skills/shortdrama-storyboard/references/agent_ledger_schema.json'
        schema = json.loads(schema_path.read_text(encoding='utf-8'))
        receipt_schema = schema['$defs']['independent_audit_certificate']['properties']['semantic_obligation_coverage']
        validator = Draft202012Validator({**receipt_schema, '$defs': schema['$defs']})
        self.assertEqual(list(validator.iter_errors(matrix['semantic_obligation_coverage'])), [])
        for name in ('check_result', 'semantic_check_result'):
            rows = matrix['preventive_gate_results'] if name == 'check_result' else matrix['audit_layer_results']
            for row in rows.values():
                self.assertEqual(list(Draft202012Validator(schema['$defs'][name]).iter_errors(row)), [])

    def test_execution_failure_and_raw_diagnostic_have_distinct_problem_accounting(self):
        candidate = {'shots': [{'镜头编号': 'EP10-001'}]}
        coverage = {'completed_shot_ids': ['EP10-001'], 'passed_shot_ids': ['EP10-001']}
        issues = [{'error_code': 'E-FIELD-CONTENT-001', 'shot_ids': ['EP10-001'], 'severity': 'fail', 'blocking': True}]
        diagnostic = _draft_quality_metrics(candidate, issues, coverage=coverage)
        self.assertEqual(diagnostic['confirmed_problem_shot_count'], 0)
        self.assertEqual(diagnostic['failed_completed_shot_count'], 1)
        checks = [{'status': 'fail', 'shot_ids': ['EP10-001']}]
        confirmed = _draft_quality_metrics(candidate, issues, coverage=coverage, execution_checks=checks)
        self.assertEqual(confirmed['confirmed_problem_shot_count'], 1)
        incomplete = _draft_quality_metrics(candidate, issues, execution_checks=checks)
        self.assertEqual(incomplete['known_problem_uncompleted_shot_ids'], ['EP10-001'])
        self.assertEqual(incomplete['completed_shot_count'], 0)

    def test_complete_semantic_failure_changes_episode_grade_without_raw_diagnostics(self):
        candidate = {'shots': [{'镜头编号': sid, 'episode_id': sid[:4]} for sid in ('EP10-001', 'EP11-001')]}
        coverage = {'completed_shot_ids': ['EP10-001', 'EP11-001'], 'passed_shot_ids': ['EP11-001'],
                    'confirmed_problem_shot_ids': ['EP10-001']}
        quality = _draft_quality_metrics(candidate, [], coverage=coverage)
        self.assertEqual(quality['first_draft_correctness_percent'], 50)
        self.assertEqual(quality['episode_correctness_percent'], {'EP10': 0, 'EP11': 100})
        self.assertEqual(quality['minimum_episode_correctness_percent'], 0)
        self.assertEqual(quality['grade'], 'B')
        self.assertEqual(quality['diagnostic_summary']['diagnostic_grade'], 'A')

    def test_removing_complete_marker_does_not_select_legacy_for_new_audit(self):
        service = StoryboardService.__new__(StoryboardService)
        arguments = dict(project_id='P', artifact_record={'artifact_id': 'A', 'fingerprint': 'raw'},
            storyboard={'shots': [{'镜头编号': 'EP10-001', 'episode_id': 'EP10'}]},
            report={'issues': [], 'independent_evidence_matrix': {'semantic_obligation_coverage': {}}},
            ledger={'events': []}, cycle_id='C', audit_class='independent_audit', audit_cache_key='K',
            prior_events=[], audit_round=1, known_pattern_keys=set())
        new = service._audit_event_content(**arguments)
        self.assertIsNone(new['first_draft_quality']['first_draft_correctness_percent'])
        old = service._audit_event_content(**arguments, historical_quality_compatibility=True)
        self.assertEqual(old['first_draft_quality']['first_draft_correctness_percent'], 100)
        self.assertNotIn('completed_shot_count', old['first_draft_quality'])

    def test_missing_author_tags_never_reduce_obligations_or_authorize_pass(self):
        candidate = extract_xlsx(workbook_bytes(), audit_only=True)['canonical_storyboard']
        manifest = build_obligations(candidate, CONTEXT)
        tagged = copy.deepcopy(candidate)
        tagged['workflow_profile'] = 'R10-PG1'
        self.assertEqual({x['obligation_id'] for x in manifest['obligations']},
                         {x['obligation_id'] for x in build_obligations(tagged, CONTEXT)['obligations']})
        value = {'coverage': synthetic_coverage(manifest), 'issues': [],
                 'layers': [{'layer': layer, 'status': 'pass'} for layer in SEMANTIC_AUDIT_LAYER_NAMES]}
        with self.assertRaisesRegex(ValueError, 'unresolved'):
            validate_coverage(candidate, value, CONTEXT, allow_unresolved=True)
        for check in value['coverage']['checks']:
            if check['obligation_id'].endswith('/source_fidelity'):
                check['status'] = 'unknown'
        receipt = validate_coverage(candidate, value, CONTEXT, allow_unresolved=True)
        self.assertFalse(receipt['complete'])
        self.assertEqual(receipt['completed_shot_count'], 0)
        self.assertEqual(receipt['unknown_shot_count'], 1)

    def test_partial_failure_is_also_unknown_not_completed(self):
        obligations = [{'obligation_id': f'{sid}/{rule}', 'shot_id': sid,
                        'allowed_statuses': ['pass', 'fail_current', 'unknown']}
                       for sid in ('EP10-001', 'EP10-002', 'EP10-003') for rule in ('a', 'b')]
        checks = [{'obligation_id': 'EP10-001/a', 'status': 'fail_current'},
                  {'obligation_id': 'EP10-002/a', 'status': 'pass'},
                  {'obligation_id': 'EP10-002/b', 'status': 'pass'},
                  {'obligation_id': 'EP10-003/a', 'status': 'fail_current'},
                  {'obligation_id': 'EP10-003/b', 'status': 'pass'}]
        receipt = coverage_summary(obligations, checks)
        self.assertEqual(receipt['completed_shot_count'], 2)
        self.assertEqual(receipt['passed_shot_count'], 1)
        self.assertEqual(receipt['unknown_shot_ids'], ['EP10-001'])
        self.assertEqual(receipt['confirmed_problem_shot_count'], 2)
        candidate = {'shots': [{'镜头编号': sid} for sid in ('EP10-001', 'EP10-002', 'EP10-003')]}
        score = _draft_quality_metrics(candidate, [], coverage=receipt)
        self.assertEqual(score['review_coverage_percent'], 66.67)
        self.assertEqual(score['completed_review_accuracy_percent'], 50)
        self.assertIsNone(score['first_draft_correctness_percent'])
        self.assertEqual(score['known_problem_uncompleted_shot_ids'], ['EP10-001'])
        unknown = _draft_quality_metrics(candidate, [])
        self.assertEqual(unknown['completed_shot_count'], 0)
        self.assertIsNone(unknown['first_draft_correctness_percent'])
        self.assertFalse(unknown['commit_eligible'])


if __name__ == '__main__':
    unittest.main()
