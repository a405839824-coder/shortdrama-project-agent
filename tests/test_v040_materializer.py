from __future__ import annotations

import copy
import io
import os
import tempfile
import unittest

from openpyxl import load_workbook

from shortdrama_mcp.app import create_mcp
from shortdrama_mcp.auth import request_namespace
from shortdrama_mcp.canonical import (
    COMPACT_SHOT_FIELDS,
    LEGACY_DERIVED_SHOT_FIELDS,
    LEGACY_REMOVED_SHOT_FIELDS,
    NO_DUPLICATE_GENERATION_EVIDENCE,
    compact_storyboard_canonical,
)
from shortdrama_mcp.config import Settings
from shortdrama_mcp.delta import ensure_shot_uids, merge_storyboard_delta
from shortdrama_mcp.file_io import extract_xlsx, storyboard_to_xlsx
from shortdrama_mcp.materializer import (
    FORMAL_MERGED_FIELD,
    GPT_KEYFRAME_FIELD,
    MOTHER_TEMPLATE_COLUMNS,
    compile_formal_merged_field,
    compile_gpt_keyframe_instruction,
    formal_mother_template_rows,
    materialize_storyboard,
)
from test_efficiency import legacy_storyboard, valid_storyboard


EXPECTED_COMPACT_COLUMNS = [
    "分镜序号",
    "时长",
    "累计时长",
    "功能类型/场景连续性ID",
    "正式分镜母版",
    "GPT关键帧生成指令",
    "参考图锚定",
    "覆盖原文",
    "原文依据类型",
    "原文支撑句",
]


class CompactMasterMaterializerTests(unittest.TestCase):
    def test_commit_materialization_is_lean_idempotent_and_discards_duplicate_inputs(self) -> None:
        source = legacy_storyboard()
        source["shots"][0][FORMAL_MERGED_FIELD] = "调用方伪造合并字段"
        source["shots"][0][GPT_KEYFRAME_FIELD] = "调用方伪造关键帧"
        source["shots"][0]["Seedance 2.5独立执行提示词"] = "调用方伪造Seedance"
        source["batches"][0]["独立执行提示词"] = "调用方伪造批次提示词"
        untouched = copy.deepcopy(source)

        first = materialize_storyboard(source)
        second = materialize_storyboard(first)

        self.assertEqual(source, untouched, "normalization must not mutate its input")
        self.assertEqual(first, second, "compact normalization must be idempotent")
        self.assertEqual(first["schema_version"], "R10.1")
        self.assertEqual(first["field_order"], list(COMPACT_SHOT_FIELDS))
        self.assertEqual(len(first["field_order"]), 27)
        self.assertNotIn("batches", first)
        self.assertNotIn("mother_template_columns", first)
        shot = first["shots"][0]
        self.assertFalse(LEGACY_REMOVED_SHOT_FIELDS.intersection(shot))
        self.assertFalse(LEGACY_DERIVED_SHOT_FIELDS.intersection(shot))
        self.assertNotIn("调用方伪造", str(first))

    def test_projection_contains_dialogue_performance_tone_and_four_light_fields(self) -> None:
        shot = valid_storyboard()["shots"][0]
        shot["画面内容"] = (
            "[0.0–0.8s] 苏清柳先看向右侧主位，垂眸整理袖口；"
            "[0.8–3.1s] 她重新抬眼，眉眼稳定，以平静语气说："
            "【苏清柳】母亲，儿媳失手，请见谅。口型与精确台词一致；"
            "[3.1–4.0s] 邹氏延后约0.6秒收紧呼吸，苏清柳保持站位。"
        )
        shot["音色"] = (
            "基础声线：二十余岁青年女声，中音区，声线偏薄而清亮，长期保持温和克制的质感。\n\n"
            "言语行为：承认失手并向母亲请求见谅。\n\n"
            "当前表达：平静克制、温和带锋，致歉时保持自身分寸。\n\n"
            "发声方式：\n"
            "音量：声量中等，仅够近处母亲听清。\n"
            "语速与节奏：语速平稳，逻辑关键词清楚，不拖长字音。\n"
            "呼吸：开口前轻吸气，句中保持稳定，不额外抢气。\n"
            "停顿：在“母亲”后短停，在“失手”后稍停。\n"
            "重音：明确落在原文中的“失手”和“见谅”。\n"
            "尾音：克制落稳。\n\n"
            "禁止语气：撒娇、讨好、夸张惊慌。"
        )

        merged = compile_formal_merged_field(shot)
        keyframe = compile_gpt_keyframe_instruction(shot)

        self.assertIn("空间站位描述：", merged)
        self.assertIn("起始状态：", merged)
        self.assertIn("画面内容：[0.0–0.8s]", merged)
        self.assertIn("【苏清柳】母亲，儿媳失手，请见谅。", merged)
        self.assertIn("口型与精确台词一致", merged)
        self.assertIn(f"音色：\n\n{shot['音色']}", merged)
        for field in ("光影", "光源来源", "光源方向", "色温倾向"):
            self.assertIn(f"{field}：{shot[field]}", merged)
        for removed in ("人物调度：", "动作与表演：", "台词：", "标点语气："):
            self.assertNotIn(removed, merged)
        self.assertNotIn("发声方式", shot)
        self.assertNotIn("Seedance 2.5独立执行", merged)
        self.assertNotIn("【苏清柳】", keyframe)
        self.assertNotIn("0.0–0.8s", keyframe)
        self.assertIn("光源来源：", keyframe)

    def test_delta_changes_only_affected_compact_shot_without_batches(self) -> None:
        base, identity = ensure_shot_uids(valid_storyboard())
        shot = copy.deepcopy(base["shots"][0])
        shot["画面内容"] = (
            "【0.0—0.8秒】女主推门；"
            "【0.8—3.3秒】她停步后确认桌上的新信件，本镜无对白；"
            "【3.3—4.0秒】她保持视线与稳定呼吸。"
        )
        performance = copy.deepcopy(base["prevention_evidence"]["performance_chains"][0])
        performance["body_or_breath"] = "停步后呼吸轻收，再随视线确认恢复稳定"

        merged, manifest = merge_storyboard_delta(
            base,
            affected_shot_uids=[identity["shot_uids"][0]["shot_uid"]],
            replacement_shots=[shot],
            evidence_updates={"performance_chains": [performance]},
        )

        self.assertIn("桌上的新信件", merged["shots"][0]["画面内容"])
        self.assertNotIn("batches", merged)
        self.assertFalse(LEGACY_DERIVED_SHOT_FIELDS.intersection(merged["shots"][0]))
        self.assertEqual(manifest["discarded_legacy_replacement_batch_count"], 0)
        self.assertEqual(
            manifest.get("derived_fields_recomputed", []),
            ["累计时长", "total_duration"],
        )

    def test_xlsx_has_exact_compact_ten_columns_and_no_batch_sheet(self) -> None:
        storyboard = compact_storyboard_canonical(valid_storyboard())
        workbook_bytes = storyboard_to_xlsx(storyboard, {"issues": []})
        workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=False)

        self.assertEqual(MOTHER_TEMPLATE_COLUMNS, EXPECTED_COMPACT_COLUMNS)
        self.assertIn("正式生产母版", workbook.sheetnames)
        headers = [cell.value for cell in workbook["正式生产母版"][1]]
        self.assertEqual(headers, EXPECTED_COMPACT_COLUMNS)
        self.assertNotIn("分镜", workbook.sheetnames)
        self.assertNotIn("生成批次", workbook.sheetnames)
        row = formal_mother_template_rows(storyboard)[0]
        self.assertEqual(workbook["正式生产母版"].cell(2, 4).value, row["功能类型/场景连续性ID"])
        self.assertEqual(workbook["正式生产母版"].cell(2, 5).value, row[FORMAL_MERGED_FIELD])
        canonical = extract_xlsx(workbook_bytes)["canonical_storyboard"]
        self.assertEqual(canonical, storyboard)
        self.assertNotIn(FORMAL_MERGED_FIELD, canonical["shots"][0])
        self.assertNotIn(GPT_KEYFRAME_FIELD, canonical["shots"][0])
        self.assertNotIn("batches", canonical)

    def test_no_duplicate_generation_evidence_is_zeroed(self) -> None:
        self.assertEqual(
            NO_DUPLICATE_GENERATION_EVIDENCE,
            {
                "model_authored_derived_chars": 0,
                "legacy_generation_calls": 0,
                "seedance_batches_generated": 0,
                "duplicate_generated_bytes": 0,
            },
        )


class FullCommitCompactMasterIntegrationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="shortdrama-compact-master-test-")
        os.environ["SHORTDRAMA_DATA_DIR"] = self.temp.name
        os.environ["PUBLIC_BASE_URL"] = "http://127.0.0.1:8765"
        os.environ["SHORTDRAMA_DEPLOYMENT_MODE"] = "local"
        os.environ["SHORTDRAMA_AUTH_MODE"] = "none"
        self.settings = Settings.from_env()
        self.server, self.service = create_mcp(self.settings)
        self.namespace = request_namespace(self.settings)

    async def asyncTearDown(self) -> None:
        self.temp.cleanup()

    async def call(self, name: str, arguments: dict) -> dict:
        result = await self.server.call_tool(name, arguments)
        self.assertFalse(result.is_error, result.content)
        self.assertIsNotNone(result.structured_content)
        return result.structured_content

    async def test_full_commit_persists_only_compact_canonical_and_projects_on_export(self) -> None:
        created = await self.call("create_project", {"project_name": "紧凑母版专项测试"})
        project_id = str(created["project_id"])
        await self.call("get_agent_status", {"project_id": project_id, "full_self_test": True})
        task = await self.call(
            "begin_project_task",
            {
                "project_id": project_id,
                "operation": "storyboard_create",
                "mode": "execute",
                "authorization": "explicit_execute",
                "request_text": "提交紧凑分镜并仅在导出时投影正式母版",
                "targets": ["EP01"],
                "authoritative_source": "锁定测试源文",
                "latest_version": "new",
                "target_boundary": "EP01",
                "must_keep": ["剧情、表演、对白、声光语义"],
                "must_remove": ["重复长字段与Seedance批次"],
                "protected_scope": [],
                "expected_outputs": ["正式分镜", "正式生产母版"],
                "window_key": "production-compact-master",
            },
        )
        storyboard = valid_storyboard()
        committed = await self.call(
            "commit_storyboard",
            {
                "project_id": project_id,
                "work_order_id": task["work_order_id"],
                "storyboard": storyboard,
                "file_name": "紧凑母版.xlsx",
            },
        )

        self.assertEqual(committed["status"], "production_preflight_pass")
        canonical, record = self.service._storyboard_from_artifact(
            self.namespace, project_id, str(committed["artifact_id"])
        )
        self.assertEqual(canonical["schema_version"], "R10.1")
        self.assertNotIn("batches", canonical)
        self.assertFalse(LEGACY_DERIVED_SHOT_FIELDS.intersection(canonical["shots"][0]))
        _artifact, path = self.service.store.get_artifact(
            self.namespace, project_id, str(record["artifact_id"])
        )
        workbook = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=False)
        self.assertIn("正式生产母版", workbook.sheetnames)
        self.assertEqual(
            [cell.value for cell in workbook["正式生产母版"][1]],
            EXPECTED_COMPACT_COLUMNS,
        )
        self.assertNotIn("生成批次", workbook.sheetnames)


if __name__ == "__main__":
    unittest.main(verbosity=2)
