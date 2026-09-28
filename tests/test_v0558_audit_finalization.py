"""Synthetic protocol regressions; never reads production storyboards."""
import copy
import hashlib
import io
import unittest
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from shortdrama_mcp.file_io import (
    audit_saved_workbook_overlay, extract_xlsx, _seal_service_verified_audit_overlay,
)
from shortdrama_mcp.service import canonical_storyboard_sha256
from semantic_audit_helpers import semantic_audit_for_artifact
import test_efficiency as efficiency_support
from test_v0556_audit_intake_coverage import workbook_bytes


class OverlayPreservationTests(unittest.TestCase):
    def test_explicit_empty_string_in_reading_page_survives_audit(self):
        source = io.BytesIO(workbook_bytes())
        output = io.BytesIO()
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(output, 'w') as after:
            for entry in before.infolist():
                payload = before.read(entry.filename)
                if entry.filename == 'xl/worksheets/sheet2.xml':
                    sheet = ET.fromstring(payload)
                    row = sheet.find(f'{ns}sheetData/{ns}row[@r="2"]')
                    cell = ET.Element(ns + 'c', {'r': 'B2', 't': 'inlineStr'})
                    ET.SubElement(ET.SubElement(cell, ns + 'is'), ns + 't').text = ''
                    row.insert(1, cell)
                    payload = ET.tostring(sheet, encoding='utf-8')
                after.writestr(entry, payload)
        raw = output.getvalue()
        book = load_workbook(io.BytesIO(raw))
        self.assertEqual(book['10集_逐镜阅读']['B2'].value, '')
        book.close()
        canonical = extract_xlsx(raw, audit_only=True)['canonical_storyboard']
        overlay = _seal_service_verified_audit_overlay({
            'status': 'FAIL_CURRENT', 'conclusion': 'synthetic incomplete review',
            'initial_issue_count': 1, 'residual_issue_count': 1,
            'introduced_issue_count': 0, 'blocking_issue_count': 1,
            'audit_trace_id': 'TRACE-test', 'certificate_id': 'CERT-test',
            'certificate_integrity_hmac': 'sealed-test',
            'canonical_sha256': canonical_storyboard_sha256(canonical),
            'source_artifact_raw_sha256': hashlib.sha256(raw).hexdigest(),
        })
        result = audit_saved_workbook_overlay(raw, canonical,
            {'issues': [{'error_code': 'synthetic-missing-evidence'}]}, audit_overlay=overlay)
        restored = extract_xlsx(result, audit_only=True)
        self.assertEqual(restored['canonical_storyboard']['shots'], canonical['shots'])
        with zipfile.ZipFile(io.BytesIO(raw)) as original, zipfile.ZipFile(io.BytesIO(result)) as final:
            self.assertEqual(original.read('xl/worksheets/sheet2.xml'), final.read('xl/worksheets/sheet2.xml'))


class AuditFinalizationFlowTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.flow = efficiency_support.EfficiencyContractTests()
        await self.flow.asyncSetUp()

    async def asyncTearDown(self):
        await self.flow.asyncTearDown()

    async def audit_external(self, project, order, name, *, shot_count=1):
        h = self.flow
        raw = workbook_bytes(reading=name)
        if shot_count > 1:
            book = load_workbook(io.BytesIO(raw))
            master = book['正式分镜母版']
            for index in range(2, shot_count + 1):
                values = [cell.value for cell in master[2]]
                values[0], values[2] = f'EP10-{index:03}', 4 * index
                master.append(values)
            buffer = io.BytesIO()
            book.save(buffer)
            book.close()
            raw = buffer.getvalue()
        original = h.service.store.save_artifact_bytes(h.namespace, project, raw,
            file_name=name + '.xlsx', artifact_type='source:storyboard', version='external')
        submission = semantic_audit_for_artifact(h.service, h.settings, project, original['artifact_id'])
        submission['conclusion']['status'] = 'fail_current'
        for check in submission['coverage']['checks']:
            check['status'] = 'unknown'
        for layer in submission['layers']:
            layer['status'] = 'fail_current'
        args = dict(project_id=project, artifact_id=original['artifact_id'], work_order_id=order,
                    semantic_audit=submission)
        result = await h.call('audit_storyboard', args)
        return original, args, result

    async def test_online_readback_uses_audit_route_without_forging_sidecar(self):
        h = self.flow
        project = await h.create_project('synthetic readback')
        order = await h.begin_storyboard_task(project, operation='storyboard_audit', mode='audit_only')
        original, _args, result = await self.audit_external(project, order, 'external-A')
        audited = result['data']['audited_workbook']['artifact_id']
        _debug_record, debug_path = h.service.store.get_artifact(h.namespace, project, audited)
        debug_candidate = extract_xlsx(debug_path.read_bytes(), audit_only=True)['canonical_storyboard']
        original_candidate, _ = h.service._storyboard_from_artifact(h.namespace, project, original['artifact_id'])
        self.assertNotEqual(debug_candidate['audit_workbook_evidence']['raw_sha256'],
                            original_candidate['audit_workbook_evidence']['raw_sha256'])
        self.assertEqual(debug_candidate['shots'], original_candidate['shots'])
        reread = await h.call('read_project_artifact', dict(project_id=project, artifact_id=audited,
                            section='正式分镜母版', offset=0, limit=1, automatic_bulk=False))
        self.assertEqual(reread['status'], 'ok')
        self.assertEqual(len(reread['data']['content']['rows']), 1)
        _record, path = h.service.store.get_artifact(h.namespace, project, audited)
        self.assertFalse(extract_xlsx(path.read_bytes(), audit_only=True)['canonical_sidecar_present'])
        source, source_path = h.service.store.get_artifact(h.namespace, project, original['artifact_id'])
        self.assertEqual(hashlib.sha256(source_path.read_bytes()).hexdigest(), source['fingerprint'])
        overlay_record, overlay_path = h.service.store.get_artifact(h.namespace, project, audited)
        for field, value, expected in (
            ('00_范围与审计!B4', 'PASS', 'visible_claim_mismatch'),
            ('正式分镜母版!E2', 'tampered execution', 'content_mapping_changed'),
            ('10集_逐镜阅读!E2', 'tampered preserved reading copy', 'preservation_mismatch'),
        ):
            with self.subTest(field=field):
                workbook = load_workbook(io.BytesIO(overlay_path.read_bytes()))
                sheet, address = field.split('!')
                workbook[sheet][address] = value
                buffer = io.BytesIO()
                workbook.save(buffer)
                workbook.close()
                forged = h.service.store.save_artifact_bytes(h.namespace, project, buffer.getvalue(),
                    file_name='forged.xlsx', artifact_type='storyboard_audit_overlay', version='external',
                    source_artifact_id=original['artifact_id'], metadata=copy.deepcopy(overlay_record['metadata']))
                with self.assertRaisesRegex(Exception, expected):
                    await h.call('read_project_artifact', dict(project_id=project, artifact_id=forged['artifact_id'],
                                section='正式分镜母版', automatic_bulk=False))

    async def test_same_work_order_keeps_artifact_histories_and_reaudit_separate(self):
        h = self.flow
        project = await h.create_project('synthetic multiple artifact audit')
        order = await h.begin_storyboard_task(project, operation='storyboard_audit', mode='audit_only')
        first, args, result_a = await self.audit_external(project, order, 'external-A')
        _second, _args_b, result_b = await self.audit_external(project, order, 'external-B', shot_count=2)
        ledger = h.service.store.load_ledger(h.namespace, project)
        events = [x for x in ledger['events'] if x.get('audit_class') == 'independent_audit']
        self.assertEqual([x['audit_round'] for x in events[-2:]], [1, 1])
        self.assertNotEqual(events[-1]['scope_key'], events[-2]['scope_key'])
        self.assertEqual(events[-1]['first_draft_quality']['total_shot_count'], 2)
        self.assertEqual(events[-2]['first_draft_quality']['total_shot_count'], 1)
        self.assertIsNone(events[-1]['first_draft_quality']['first_draft_correctness_percent'])
        changed = copy.deepcopy(args)
        changed['force_machine_recheck'] = True
        changed['semantic_audit']['layers'][0]['summary'] += ' Independent supplemental current review.'
        repeated = await h.call('audit_storyboard', changed)
        self.assertNotEqual(repeated['data']['certificate']['certificate_id'],
                            result_a['data']['certificate']['certificate_id'])
        ledger = h.service.store.load_ledger(h.namespace, project)
        events = [x for x in ledger['events'] if x.get('audit_class') == 'independent_audit']
        self.assertEqual(events[-1]['audit_round'], 2)
        self.assertEqual(events[-1]['scope_key'], events[-3]['scope_key'])


if __name__ == '__main__':
    unittest.main()
