"""Full generic sheet reads using synthetic workbooks, never user files."""
import hashlib
import io
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from shortdrama_mcp.file_io import (
    _rows_from_sheet, _sheet_value_matrix, extract_xlsx,
    preview_extracted, preview_extracted_bulk, storyboard_to_xlsx,
)
from shortdrama_mcp.materializer import MOTHER_TEMPLATE_COLUMNS
from test_efficiency import valid_storyboard
from test_v0556_audit_intake_coverage import workbook_bytes


class SheetReadTests(unittest.TestCase):
    def test_merged_title_preserves_later_columns_and_column_only_rows(self):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Synthetic report'
        sheet.merge_cells('A1:C1')
        sheet['A1'] = 'Audit result'
        sheet.append(['Status', 'UNKNOWN', 'Missing source evidence'])
        sheet['B3'] = 'Only B is populated'
        sheet['C4'] = False
        sheet['E5'] = 0
        sheet['D6'] = '=1+1'
        sheet['Z99'].font = Font(bold=True)  # Formatting is not report content.
        matrix = _sheet_value_matrix(sheet)
        headers, rows = _rows_from_sheet(sheet)
        self.assertEqual(headers, ['Audit result', '列2', '列3', '列4', '列5'])
        self.assertEqual([list(row.values()) for row in rows], matrix[1:])
        value = {'kind': 'workbook', 'sheets': [
            {'name': sheet.title, 'headers': headers, 'rows': rows, 'matrix': matrix}]}
        preview = preview_extracted(value, sheet.title, 0, 30)
        bulk = preview_extracted_bulk(value, [
            {'section': sheet.title, 'unit': 'rows', 'offset': 0, 'limit': 30}])
        self.assertEqual(preview['rows'], rows)
        self.assertEqual(bulk['ranges'][0]['rows'], rows)
        self.assertEqual(bulk['ranges'][0]['headers'], headers)
        self.assertTrue(bulk['ranges'][0]['complete'])
        book.close()

    def test_generated_headers_keep_existing_names_without_collisions(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['列2', None, '列2_2'])
        sheet.append(['A', 'B', 'C', 'D'])
        headers, rows = _rows_from_sheet(sheet)
        self.assertEqual(headers, ['列2', '列2_2', '列2_2_2', '列4'])
        self.assertEqual(list(rows[0].values()), ['A', 'B', 'C', 'D'])
        book.close()

    def test_empty_styled_sheet_has_no_invented_headers_or_rows(self):
        book = Workbook()
        book.active['Q20'].font = Font(bold=True)
        self.assertEqual(_rows_from_sheet(book.active), ([], []))
        book.close()

    def test_signed_master_and_canonical_stay_exact_while_scope_is_readable(self):
        canonical = valid_storyboard()
        raw = storyboard_to_xlsx(canonical, {'issues': []}, assume_compact=True)
        before = hashlib.sha256(raw).hexdigest()
        extracted = extract_xlsx(raw)
        self.assertEqual(extracted['canonical_storyboard'], canonical)
        master = preview_extracted(extracted, None, 0, 30)
        self.assertEqual(master['headers'], MOTHER_TEMPLATE_COLUMNS)
        scope = next(s for s in extracted['sheets'] if s['name'] == '00_范围与审计')
        self.assertGreater(len(scope['headers']), 1)
        # The generic row view must contain every report cell after its title.
        self.assertEqual([list(row.values()) for row in scope['rows']],
                         [row for row in scope['matrix'][1:] if any(v not in (None, '') for v in row)])
        bulk = preview_extracted_bulk(extracted, [
            {'section': scope['name'], 'unit': 'rows', 'offset': 0, 'limit': 100}])
        self.assertEqual(bulk['ranges'][0]['rows'], scope['rows'])
        self.assertEqual(hashlib.sha256(raw).hexdigest(), before)
        self.assertEqual(extract_xlsx(raw)['canonical_storyboard'], canonical)

    def test_formal_extra_column_still_rejected_instead_of_losing_master_identity(self):
        for legacy in (False, True):
            with self.subTest(legacy=legacy):
                book = load_workbook(io.BytesIO(workbook_bytes(legacy=legacy)))
                sheet = book['正式分镜母版']
                address = 'K2' if legacy else 'L2'
                sheet[address] = 'Unmapped execution instruction'
                output = io.BytesIO()
                book.save(output)
                book.close()
                with self.assertRaisesRegex(ValueError, 'unmapped_master_cells: ' + address):
                    extract_xlsx(output.getvalue(), audit_only=True)

    def test_signed_master_tampering_remains_rejected(self):
        raw = storyboard_to_xlsx(valid_storyboard(), {'issues': []}, assume_compact=True)
        book = load_workbook(io.BytesIO(raw))
        master = next(s for s in book if s['A1'].value == MOTHER_TEMPLATE_COLUMNS[0])
        master['F2'] = 'Tampered synthetic keyframe'
        output = io.BytesIO()
        book.save(output)
        book.close()
        with self.assertRaisesRegex(ValueError, 'projection.*mismatch'):
            extract_xlsx(output.getvalue())


if __name__ == '__main__':
    unittest.main()
