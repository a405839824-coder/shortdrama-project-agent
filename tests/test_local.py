from __future__ import annotations

import asyncio
import os
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from shortdrama_mcp.app import create_mcp
from shortdrama_mcp.config import Settings
from shortdrama_mcp.file_io import _validate_zip_archive, extract_xlsx, storyboard_to_xlsx


PLUGIN_ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = PLUGIN_ROOT / "skills" / "shortdrama-storyboard"

import sys

sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from test_public_scene_ecology import public_fixture as historical_public_fixture  # noqa: E402
from semantic_audit_helpers import inject_semantic_audit, semantic_audit_for_artifact
from test_efficiency import valid_storyboard as compact_valid_storyboard


def valid_storyboard() -> dict:
    return compact_valid_storyboard()


def public_fixture() -> dict:
    """Submit the ecology example with an explicit current production look."""
    from prompt_libraries import style_fragment
    value = historical_public_fixture()
    value.setdefault("source_lock", {}).update(compact_valid_storyboard()["source_lock"])
    for shot in value["shots"]:
        shot["生成限制"] += style_fragment(value["source_lock"])
    return value


class LocalMcpTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="shortdrama-plugin-test-")
        os.environ["SHORTDRAMA_DATA_DIR"] = self.temp.name
        os.environ["PUBLIC_BASE_URL"] = "http://127.0.0.1:8765"
        os.environ["SHORTDRAMA_DEPLOYMENT_MODE"] = "local"
        os.environ["SHORTDRAMA_AUTH_MODE"] = "none"
        os.environ.pop("ALLOW_LOCAL_FILE_URLS", None)
        self.settings = Settings.from_env()
        self.server, self.service = create_mcp(self.settings)
        from historical_entry_fixture import freeze_legacy_entry
        freeze_legacy_entry(self.service)

    async def asyncTearDown(self) -> None:
        self.temp.cleanup()

    async def call(self, name: str, arguments: dict) -> dict:
        arguments = inject_semantic_audit(self.service, self.settings, name, arguments)
        result = await self.server.call_tool(name, arguments)
        self.assertFalse(result.is_error, result.content)
        self.assertIsNotNone(result.structured_content)
        return result.structured_content

    async def create_project(self, name: str = "插件自动化测试") -> str:
        result = await self.call("create_project", {"project_name": name})
        self.assertEqual(result["status"], "created")
        project_id = result["project_id"]
        checked = await self.call(
            "get_agent_status", {"project_id": project_id, "full_self_test": True}
        )
        self.assertTrue(checked["data"]["runtime_test_certificate_valid"])
        return project_id

    async def test_tool_contracts_and_file_params(self) -> None:
        tools = await self.server.list_tools()
        self.assertEqual(len(tools), 24)
        by_name = {tool.name: tool.model_dump(by_alias=True, exclude_none=True) for tool in tools}
        self.assertTrue({"preflight_narrative_plan", "review_narrative_artifact",
                         "record_narrative_attempt", "read_narrative_record"}.issubset(by_name))
        self.assertEqual(by_name["get_agent_status"]["annotations"]["readOnlyHint"], False)
        self.assertEqual(by_name["get_agent_status"]["annotations"]["idempotentHint"], False)
        self.assertEqual(by_name["mark_performance_stage"]["annotations"]["readOnlyHint"], False)
        self.assertEqual(by_name["audit_storyboard"]["annotations"]["readOnlyHint"], False)
        file_tool = by_name["import_source_file"]
        self.assertEqual(file_tool["_meta"]["openai/fileParams"], ["file"])
        file_schema = file_tool["inputSchema"]["$defs"]["OpenAIFile"]
        self.assertEqual(set(file_schema["properties"]), {"download_url", "file_id", "mime_type", "file_name"})
        self.assertEqual(file_schema["required"], ["download_url", "file_id"])
        gap_schema = by_name["record_capability_gap"]["inputSchema"]
        self.assertIn("work_order_id", gap_schema["properties"])
        self.assertIn("work_order_id", gap_schema["required"])
        self.assertNotIn("ctx", gap_schema["properties"])
        for descriptor in by_name.values():
            self.assertIn("outputSchema", descriptor)
            self.assertIn("securitySchemes", descriptor["_meta"])

    async def test_plan_artifact_validation_and_signoff(self) -> None:
        project_id = await self.create_project("方案流程测试")
        task = await self.call(
            "begin_project_task",
            {
                "project_id": project_id,
                "operation": "rewrite_plan",
                "mode": "plan_only",
                "authorization": "plan_only",
                "request_text": "给出A/B改写方案，不修改原稿",
                "targets": ["第一集"],
                "authoritative_source": "原剧本v1",
                "latest_version": "v1",
                "target_boundary": "第一集",
                "must_keep": ["主角动机"],
                "must_remove": [],
                "protected_scope": ["第二集及以后"],
                "expected_outputs": ["改写方案"],
                "window_key": "test-plan",
            },
        )
        work_order_id = task["work_order_id"]
        saved = await self.call(
            "commit_project_artifact",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "artifact_type": "rewrite_plan",
                "output_format": "markdown",
                "file_name": "第一集改写方案.md",
                "content_text": "# 改写方案\n\n推荐A方案，保留主角动机。",
                "change_summary": "仅输出方案，未修改原稿",
            },
        )
        self.assertEqual(saved["status"], "saved_pending_validation")
        gates = {gate: "pass" for gate in saved["data"]["required_validation_gates"]}
        validated = await self.call(
            "complete_task_validation",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "gate_results": gates,
                "residual_issue_count": 0,
                "introduced_issue_count": 0,
                "blockers": [],
            },
        )
        self.assertEqual(validated["status"], "pass")
        signed = await self.call(
            "signoff_task",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "decision": "approve",
                "signer": "自动化测试员",
            },
        )
        self.assertEqual(signed["status"], "archived")
        status = await self.call("get_agent_status", {"project_id": project_id})
        self.assertTrue(status["data"]["state_valid"])

    async def test_storyboard_fail_repair_pass_and_duplicate_audit(self) -> None:
        project_id = await self.create_project("分镜闭环测试")
        task = await self.call(
            "begin_project_task",
            {
                "project_id": project_id,
                "operation": "storyboard_create",
                "mode": "execute",
                "authorization": "explicit_execute",
                "request_text": "生成正式分镜并全链路审计",
                "targets": ["第一集"],
                "authoritative_source": "测试剧本",
                "latest_version": "v1",
                "target_boundary": "第一集",
                "must_keep": [],
                "must_remove": [],
                "protected_scope": [],
                "expected_outputs": ["正式分镜", "正式分镜母版"],
                "window_key": "test-storyboard",
            },
        )
        work_order_id = task["work_order_id"]
        broken = valid_storyboard()
        broken["shots"][0]["光源来源"] = ""
        first = await self.call(
            "commit_storyboard",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "storyboard": broken,
                "file_name": "第一集分镜.xlsx",
            },
        )
        self.assertEqual(first["status"], "fail_current")
        self.assertGreater(len(first["data"]["report"]["issues"]), 0)
        repaired = await self.call(
            "commit_storyboard",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "storyboard": valid_storyboard(),
                "file_name": "第一集分镜_修复版.xlsx",
                "source_artifact_id": first["artifact_id"],
            },
        )
        self.assertEqual(repaired["status"], "production_preflight_pass")
        self.assertEqual(repaired["data"]["report"]["issues"], [])
        with self.assertRaisesRegex(Exception, "semantic_audit_required"):
            await self.server.call_tool(
                "audit_storyboard",
                {"project_id": project_id, "artifact_id": repaired["artifact_id"]},
            )
        invalid_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, repaired["artifact_id"]
        )
        invalid_semantic["layers"][0]["evidence_ids"] = ["foo:bar"]
        with self.assertRaisesRegex(Exception, "evidence kind must be one of"):
            await self.server.call_tool(
                "audit_storyboard",
                {
                    "project_id": project_id,
                    "artifact_id": repaired["artifact_id"],
                    "semantic_audit": invalid_semantic,
                },
            )
        independent = await self.call(
            "audit_storyboard",
            {"project_id": project_id, "artifact_id": repaired["artifact_id"]},
        )
        self.assertEqual(independent["status"], "pass")
        self.assertEqual(independent["data"]["audit"]["status"], "recorded")
        self.assertFalse(independent["data"]["cache_hit"])
        self.assertEqual(
            independent["data"]["production_trace_id"], repaired["trace_id"]
        )
        self.assertEqual(independent["data"]["audit_trace_id"], independent["trace_id"])
        self.assertFalse(independent["data"]["fast_path_evidence"]["cache_hit"])
        repeated = await self.call(
            "audit_storyboard",
            {"project_id": project_id, "artifact_id": repaired["artifact_id"]},
        )
        self.assertEqual(repeated["status"], "pass")
        self.assertEqual(repeated["data"]["audit"]["status"], "duplicate")
        self.assertTrue(repeated["data"]["cache_hit"])
        status = await self.call("get_agent_status", {"project_id": project_id})
        self.assertTrue(status["data"]["state_valid"])
        # The rejected draft creates no artifact, but is retained as learning;
        # its clean closure, production preflight and independent audit remain
        # separate events.
        self.assertEqual(status["data"]["issues"]["event_count"], 4)
        self.assertEqual(status["data"]["issues"]["open_issue_count"], 0)

    async def test_public_scene_people_audio_gate_and_repair(self) -> None:
        project_id = await self.create_project("公共场景生态测试")
        task = await self.call(
            "begin_project_task",
            {
                "project_id": project_id,
                "operation": "storyboard_create",
                "mode": "execute",
                "authorization": "explicit_execute",
                "request_text": "检查古代市集普通人物和环境音并修正",
                "targets": ["市集场景"],
                "authoritative_source": "测试源文",
                "latest_version": "v1",
                "target_boundary": "市集场景",
                "must_keep": ["主角动作"],
                "must_remove": ["泛称环境声"],
                "protected_scope": [],
                "expected_outputs": ["正式分镜", "正式分镜母版"],
                "window_key": "test-public-ecology",
            },
        )
        work_order_id = task["work_order_id"]
        broken = public_fixture()
        broken["shots"][0]["环境音效"] = "街道嘈杂声"
        first = await self.call(
            "commit_storyboard",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "storyboard": broken,
                "file_name": "市集分镜_错误版.xlsx",
            },
        )
        self.assertEqual(first["status"], "fail_current")
        codes = {item["error_code"] for item in first["data"]["report"]["issues"]}
        self.assertIn("E-PUBLIC-AUDIO-001", codes)
        self.assertNotIn("validation", first["data"])
        self.assertFalse(first["data"]["artifact_created"])
        self.assertFalse(first["data"]["handoff_created"])

        repaired = await self.call(
            "commit_storyboard",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "storyboard": public_fixture(),
                "file_name": "市集分镜_修复版.xlsx",
                "source_artifact_id": first["artifact_id"],
            },
        )
        self.assertEqual(repaired["status"], "production_preflight_pass")
        self.assertEqual(repaired["data"]["report"]["issues"], [])

    async def test_file_import_and_workbook_roundtrip(self) -> None:
        project_id = await self.create_project("文件导入测试")
        intake = await self.call(
            "begin_project_task",
            {
                "project_id": project_id,
                "operation": "project_intake",
                "mode": "execute",
                "authorization": "explicit_execute",
                "request_text": "导入用户提供的正式分镜来源文件",
                "targets": ["测试分镜.xlsx"],
                "authoritative_source": "用户上传文件",
                "latest_version": "current",
                "target_boundary": "全文件",
                "must_keep": [],
                "must_remove": [],
                "protected_scope": [],
                "expected_outputs": ["来源文件与提取结果"],
                "window_key": "test-source-intake",
            },
        )
        workbook = storyboard_to_xlsx(valid_storyboard(), {"issues": []})

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self) -> None:  # noqa: N802
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Length", str(len(workbook)))
                self.end_headers()
                self.wfile.write(workbook)

            def log_message(self, _format: str, *_args: object) -> None:
                return

        httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        try:
            os.environ["ALLOW_LOCAL_FILE_URLS"] = "1"
            settings = Settings.from_env()
            server, _service = create_mcp(settings)
            from historical_entry_fixture import freeze_legacy_entry
            freeze_legacy_entry(_service)
            result = await server.call_tool(
                "import_source_file",
                {
                    "project_id": project_id,
                    "work_order_id": intake["work_order_id"],
                    "file": {
                        "download_url": f"http://127.0.0.1:{httpd.server_port}/fixture.xlsx",
                        "file_id": "file_test_fixture",
                        "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        "file_name": "测试分镜.xlsx",
                    },
                    "source_role": "storyboard",
                },
            )
            self.assertFalse(result.is_error, result.content)
            imported = result.structured_content
            self.assertEqual(imported["status"], "imported")
            read = await server.call_tool(
                "read_project_artifact",
                {
                    "project_id": project_id,
                    "artifact_id": imported["artifact_id"],
                    "section": "正式分镜母版",
                    "offset": 0,
                    "limit": 10,
                    "automatic_bulk": False,
                },
            )
            self.assertFalse(read.is_error, read.content)
            self.assertEqual(read.structured_content["data"]["content"]["total"], 1)
        finally:
            httpd.shutdown()
            httpd.server_close()
            os.environ.pop("ALLOW_LOCAL_FILE_URLS", None)

    async def test_xlsx_formula_escape_and_archive_limit(self) -> None:
        from io import BytesIO
        from zipfile import ZIP_DEFLATED, ZipFile

        from openpyxl import load_workbook

        storyboard = valid_storyboard()
        storyboard["shots"][0]["画面内容"] = '=HYPERLINK("https://example.test","打开")'
        workbook = storyboard_to_xlsx(storyboard, {"issues": []})
        loaded = load_workbook(BytesIO(workbook), data_only=False, read_only=True)
        sheet = loaded["正式分镜母版"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        content_index = headers.index("正式分镜母版") + 1
        projected = str(sheet.cell(2, content_index).value)
        self.assertFalse(projected.startswith(("=", "+", "-", "@")))
        self.assertIn("画面内容：=HYPERLINK", projected)
        canonical = extract_xlsx(workbook)["canonical_storyboard"]
        self.assertEqual(
            canonical["shots"][0]["画面内容"],
            '=HYPERLINK("https://example.test","打开")',
        )

        archive_stream = BytesIO()
        with ZipFile(archive_stream, "w", ZIP_DEFLATED) as archive:
            archive.writestr("large.txt", b"0" * 2048)
        with self.assertRaisesRegex(ValueError, "safe limit"):
            _validate_zip_archive(archive_stream.getvalue(), max_uncompressed_bytes=1024)


if __name__ == "__main__":
    unittest.main(verbosity=2)
