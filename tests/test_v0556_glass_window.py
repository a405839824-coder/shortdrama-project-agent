"""Real glass daylight through the common lower/render/export path."""
import copy
import io
import unittest

from openpyxl import load_workbook
from test_v0527_authoring_contract_blocks import light_args
from test_v0527_production_decisions import production_fixture
from authoring_contract_blocks import scene_lighting_blocks, ContractBlockError
from production_decisions import lower_production_decisions, ProductionDecisionError
from production_actions import effects
from author_storyboard import _render_decisions
from shortdrama_mcp.file_io import storyboard_to_xlsx
from validate_storyboard import _lighting_context_faults
from lighting_checks import lighting_context_faults


def glass_args(state='closed'):
    args = light_args()
    args['source'].update(kind='glass_window', name='东玻璃窗', state=state)
    return args


def glass_production(state='closed'):
    data = production_fixture()
    scene = data['initial_world']['scenes']['ROOM']
    scene['label'] = '室内维修室'
    scene['light'].update(kind='glass_window', name='东玻璃窗', state=state)
    scene['door_window_state'] = {'东玻璃窗': state}
    scene['opening_labels'] = {'东玻璃窗': '玻璃窗'}
    return data


class GlassWindowTests(unittest.TestCase):
    def test_closed_and_open_are_distinct_real_material(self):
        for state in ('closed', 'open'):
            result = scene_lighting_blocks(**glass_args(state))
            self.assertNotIn('窗纸', str(result))
            self.assertNotIn('漫射', str(result))
            self.assertIn('日间自然光', result['lighting_fields']['光源来源'])
            self.assertIn('关闭' if state == 'closed' else '敞开', result['spatial_state_append'])
        paper = scene_lighting_blocks(**light_args())
        self.assertEqual(paper['spatial_state_append'], '室内日间；纸窗甲关闭，窗纸透入漫射天光。')

    def test_night_unknown_and_invalid_height_remain_rejected(self):
        for change in ({'period': 'night'}, {'state': 'unknown'}, {'height_m': 0}, {'height_m': True}):
            args = glass_args(); args['source'].update(change)
            with self.subTest(change=change), self.assertRaises(ContractBlockError):
                scene_lighting_blocks(**args)

    def test_lower_canonical_actual_xlsx_preserve_closed_glass(self):
        data = glass_production(); before = copy.deepcopy(data)
        lowered = lower_production_decisions(data)
        canonical = _render_decisions(lowered['typed_decisions'])
        self.assertEqual(data, before)
        row = canonical['shots'][0]
        self.assertEqual(lighting_context_faults(row, legacy_checker=_lighting_context_faults), [])
        self.assertIn('闭合玻璃', row['光源来源'])
        raw = storyboard_to_xlsx(canonical, {'status': 'fail', 'issues': []}, assume_compact=True)
        book = load_workbook(io.BytesIO(raw), read_only=True)
        formal = next(s for s in book if s.max_column == 11 and s.max_row == 2)
        for col in (5, 6):
            self.assertIn('玻璃', formal.cell(2, col).value)
            self.assertNotIn('窗纸', formal.cell(2, col).value)
        book.close()

    def test_conflicting_registry_and_night_fail_actual_lower(self):
        for mode in ('mismatch', 'night'):
            data = glass_production()
            scene = data['initial_world']['scenes']['ROOM']
            if mode == 'mismatch': scene['door_window_state']['东玻璃窗'] = 'open'
            else: scene['light']['period'] = 'night'
            with self.subTest(mode=mode), self.assertRaises(ProductionDecisionError):
                lower_production_decisions(data)

    def test_opening_effect_updates_exact_window_light_at_end(self):
        data = glass_production()
        action = {'kind': 'set_opening', 'actor': '甲', 'opening': '东玻璃窗',
                  'to': 'open', 'start': 0, 'end': 10}
        before = effects(data['initial_world'], action, 5)
        after = effects(data['initial_world'], action, 10)
        self.assertEqual(before[('scenes', 'ROOM', 'light')]['state'], 'closed')
        self.assertEqual(after[('scenes', 'ROOM', 'light')]['state'], 'open')
        self.assertIn('东玻璃窗开启', after[('scenes', 'ROOM', 'door_window_state')])

    def test_glass_does_not_authorize_visible_beams_or_opening(self):
        row = _render_decisions(lower_production_decisions(glass_production())['typed_decisions'])['shots'][0]
        for change in ({'光影': '窗边出现明亮光柱。'}, {'光源来源': '日间自然光通过敞开的东玻璃窗进入。'}):
            bad = {**row, **change}
            with self.subTest(change=change):
                self.assertTrue(lighting_context_faults(bad, legacy_checker=_lighting_context_faults))


if __name__ == '__main__':
    unittest.main()
