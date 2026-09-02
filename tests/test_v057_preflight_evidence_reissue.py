from __future__ import annotations

import copy
import io
import unittest
from pathlib import Path
from typing import Any
from unittest.mock import patch

from openpyxl import load_workbook

from shortdrama_mcp.file_io import (
    CANONICAL_SIDECAR_SHEET,
    extract_xlsx,
)
from shortdrama_mcp.materializer import (
    MOTHER_TEMPLATE_COLUMNS,
    legacy_v053_mother_template_rows,
)
import shortdrama_mcp.service as service_module

import test_v056_handoff_recovery as recovery_tests


class V057PreflightEvidenceReissueTests(
    recovery_tests.V056HandoffRecoveryTests
):
    test_batch_recovers_each_handoff_with_one_shared_trace = None
    test_execute_requires_exact_confirmation_before_mutation = None
    test_execute_supersedes_preserves_history_and_replays_idempotently = None
    test_interrupted_commit_fails_closed_and_retry_recovers = None
    test_tool_contract_and_plan_are_controlled_and_read_only = None

    async def asyncSetUp(self) -> None:
        await super().asyncSetUp()
        original = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        self.clean_peer_handoff_id = "HND-v057-clean-peer"
        clean_peer = self.service._handoff_record_projection(original)
        clean_peer["handoff_id"] = self.clean_peer_handoff_id
        clean_peer_record = self.service.store.save_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.clean_peer_handoff_id,
            value=clean_peer,
        )
        with self.service.store.mutate_project(
            self.namespace, self.project_id
        ) as (hub, _ledger):
            hub["handoffs"].append(
                self.service._handoff_record_projection(clean_peer_record)
            )

        self.original_trace_id = str(original["stage_trace_id"])
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        ) as incomplete:
            for field in (
                "production_preflight_status",
                "production_preflight_issue_count",
                "independent_audit_status",
            ):
                incomplete.pop(field, None)
        with self.service.store.mutate_project(
            self.namespace, self.project_id
        ) as (hub, _ledger):
            row = next(
                item
                for item in hub["handoffs"]
                if item["handoff_id"] == self.old_handoff_id
            )
            for field in (
                "production_preflight_status",
                "production_preflight_issue_count",
                "independent_audit_status",
            ):
                row.pop(field, None)
        response = self.service.begin_task(
            self.namespace,
            project_id=self.project_id,
            operation="rule_improvement",
            mode="execute",
            authorization="explicit_execute",
            request_text="authorized controlled preflight evidence reissue",
            targets=[
                self.gap_id,
                self.old_handoff_id,
                "controlled_phase:preflight_evidence_backfill",
            ],
            authoritative_source="signed project records",
            latest_version="0.5.9",
            target_boundary="one exact legacy handoff",
            must_keep=["history", "content", "independent audit separation"],
            must_remove=[],
            protected_scope=["storyboard", "canonical", "XLSX", "audit history"],
            expected_outputs=["one append-only production preflight evidence successor"],
            window_key=self.session_id,
            request_session_id=self.session_id,
        )
        self.backfill_work_order_id = str(response.work_order_id)

    def _backfill_plan(self) -> Any:
        return self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=self.backfill_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=[self.old_handoff_id],
            action="plan",
            expected_plan_sha256=None,
            confirmation=None,
            request_session_id=self.session_id,
            phase="preflight_evidence_backfill",
        )

    def _backfill_execute(self, plan_sha256: str, confirmation: str | None = None) -> Any:
        return self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=self.backfill_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=[self.old_handoff_id],
            action="execute",
            expected_plan_sha256=plan_sha256,
            confirmation=(
                service_module.PREFLIGHT_EVIDENCE_REISSUE_CONFIRMATION
                if confirmation is None
                else confirmation
            ),
            request_session_id=self.session_id,
            phase="preflight_evidence_backfill",
        )

    def _artifact_bytes(self) -> bytes:
        handoff = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        _record, path = self.service.store.get_artifact(
            self.namespace,
            self.project_id,
            str(handoff["storyboard_artifact_id"]),
        )
        return Path(path).read_bytes()

    def _install_v053_transitional_workbook_fixture(
        self,
    ) -> dict[str, Any]:
        """Make the fixture look exactly like one signed v0.5.3 export.

        The base fixture is first produced with today's writer.  Rebinding all
        historical authorities here models the same workbook having originally
        been emitted by v0.5.3; the recovery operation itself must still leave
        every old record and authored byte untouched.
        """

        handoff = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        artifact_id = str(handoff["storyboard_artifact_id"])
        storyboard, artifact_record = self.service._storyboard_from_artifact(
            self.namespace,
            self.project_id,
            artifact_id,
            persist_recovered_canonical=False,
            require_bound_projection=True,
        )
        _verified_record, artifact_path = self.service.store.get_artifact(
            self.namespace,
            self.project_id,
            artifact_id,
        )
        workbook = load_workbook(artifact_path, data_only=False)
        formal_sheet = workbook["正式生产母版"]
        self.assertEqual(
            [cell.value for cell in formal_sheet[1]],
            list(MOTHER_TEMPLATE_COLUMNS),
        )
        expected_rows = legacy_v053_mother_template_rows(storyboard)
        self.assertEqual(formal_sheet.max_row, len(expected_rows) + 1)
        for row_index, row in enumerate(expected_rows, start=2):
            for column_index, header in enumerate(
                MOTHER_TEMPLATE_COLUMNS, start=1
            ):
                formal_sheet.cell(row_index, column_index).value = row.get(
                    header
                )
        sidecar = workbook[CANONICAL_SIDECAR_SHEET]
        self.assertEqual(sidecar.sheet_state, "veryHidden")
        sidecar.delete_cols(9, 3)
        payload_buffer = io.BytesIO()
        workbook.save(payload_buffer)
        workbook.close()
        payload = payload_buffer.getvalue()

        with self.assertRaisesRegex(
            ValueError,
            "visible header does not match signed projection contract",
        ):
            extract_xlsx(payload)

        artifact_raw_sha256 = service_module.sha256_bytes(payload)
        self.service.store.atomic_write_bytes(artifact_path, payload)
        artifact_record["fingerprint"] = artifact_raw_sha256
        artifact_record["size_bytes"] = len(payload)
        artifact_meta_path = (
            self.service.store.project_dir(
                self.namespace, self.project_id
            )
            / "artifacts"
            / f"{artifact_id}.meta.json"
        )
        self.service.store.atomic_write_json(
            artifact_meta_path, artifact_record
        )

        cache_components = copy.deepcopy(handoff["cache_components"])
        cache_components["artifact"] = artifact_raw_sha256
        cache_key = self.service._cache_key_from_components(
            cache_components
        )
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        ) as rebound_handoff:
            rebound_handoff["fingerprints"][
                "artifact_raw"
            ] = artifact_raw_sha256
            rebound_handoff["cache_components"] = copy.deepcopy(
                cache_components
            )
            rebound_handoff["cache_key"] = cache_key
            for cache_entry in rebound_handoff["cache_manifest"]:
                cache_entry["cache_key"] = cache_key
                cache_entry["dependency_fingerprints"][
                    "artifact_raw"
                ] = artifact_raw_sha256
            workbook_manifest = rebound_handoff["objective_evidence"][
                "workbook_manifest"
            ]
            workbook_manifest["raw_sha256"] = artifact_raw_sha256
            workbook_manifest["size_bytes"] = len(payload)

        rebound_handoff = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        original_certificate_id = str(
            rebound_handoff["certificate_id"]
        )
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=original_certificate_id,
        ) as rebound_certificate:
            rebound_certificate["cache_components"] = copy.deepcopy(
                cache_components
            )
            rebound_certificate["cache_key"] = cache_key

        original_trace_id = str(rebound_handoff["stage_trace_id"])
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="telemetry",
            record_id=original_trace_id,
        ) as rebound_trace:
            rebound_trace["input_fingerprints"] = copy.deepcopy(
                rebound_handoff["fingerprints"]
            )
            rebound_trace["input_payload_bytes"] = len(payload)
            rebound_trace["output_payload_bytes"] = len(payload)
            for stage in rebound_trace.get("stages", []):
                stage["payload_bytes"] = len(payload)

        self._replace_hub_handoff_rows(
            self.old_handoff_id,
            [self.service._handoff_record_projection(rebound_handoff)],
        )
        return {
            "artifact_bytes": payload,
            "artifact_raw_sha256": artifact_raw_sha256,
            "old_handoff": rebound_handoff,
            "old_certificate": self.service.store.get_project_record(
                self.namespace,
                self.project_id,
                kind="certificates",
                record_id=original_certificate_id,
            ),
            "original_trace": self.service.store.get_project_record(
                self.namespace,
                self.project_id,
                kind="telemetry",
                record_id=original_trace_id,
            ),
        }

    def _remove_handoff_preflight_summary(self, handoff_id: str) -> None:
        fields = (
            "production_preflight_status",
            "production_preflight_issue_count",
            "independent_audit_status",
        )
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=handoff_id,
        ) as handoff:
            for field in fields:
                handoff.pop(field, None)
        with self.service.store.mutate_project(
            self.namespace, self.project_id
        ) as (hub, _ledger):
            row = next(
                item
                for item in hub["handoffs"]
                if item["handoff_id"] == handoff_id
            )
            for field in fields:
                row.pop(field, None)

    def _replace_hub_handoff_rows(
        self,
        handoff_id: str,
        rows: list[dict[str, Any]],
    ) -> None:
        """Replace one handoff's Hub rows without changing its signed record."""

        with self.service.store.mutate_project(
            self.namespace, self.project_id
        ) as (hub, _ledger):
            hub["handoffs"] = [
                row
                for row in hub.get("handoffs", [])
                if not (
                    isinstance(row, dict)
                    and row.get("handoff_id") == handoff_id
                )
            ]
            hub["handoffs"].extend(copy.deepcopy(rows))

    def _recovery_authority_snapshot(self) -> dict[str, Any]:
        """Capture all authority stores that a rejected plan must not mutate."""

        records: dict[str, list[tuple[str, str]]] = {}
        for kind in (
            "handoffs",
            "preflight_evidence_reissues",
            "certificates",
            "telemetry",
            "handoff_recoveries",
        ):
            rows = self.service.store.list_project_records(
                self.namespace,
                self.project_id,
                kind=kind,
            )
            records[kind] = [
                (
                    str(row.get("record_id") or ""),
                    self.service.store.project_state_sha256(row),
                )
                for row in rows
            ]
        hub = self.service.store.load_hub(
            self.namespace, self.project_id
        )
        return {
            "hub_sha256": self.service.store.project_state_sha256(hub),
            "records": records,
        }

    def _controlled_work_order(
        self, *, phase: str, handoff_ids: list[str], operation: str = "rule_improvement"
    ) -> str:
        response = self.service.begin_task(
            self.namespace,
            project_id=self.project_id,
            operation=operation,
            mode="execute",
            authorization="explicit_execute",
            request_text=f"authorized controlled {phase}",
            targets=[
                self.gap_id,
                *handoff_ids,
                f"controlled_phase:{phase}",
            ],
            authoritative_source="signed project records",
            latest_version="0.5.9",
            target_boundary="exact authorized handoff batch",
            must_keep=["history", "content", "independent audit separation"],
            must_remove=[],
            protected_scope=["storyboard", "canonical", "XLSX", "audit history"],
            expected_outputs=[phase],
            window_key=self.session_id,
            request_session_id=self.session_id,
        )
        return str(response.work_order_id)

    async def test_tool_count_stays_nineteen_and_phase_is_advertised(self) -> None:
        tools = await self.server.list_tools()
        self.assertEqual(len(tools), 19)
        recovery = next(
            tool.model_dump(mode="json", by_alias=True)
            for tool in tools
            if tool.name == "recover_storyboard_handoff"
        )
        phase = recovery["inputSchema"]["properties"]["phase"]
        phase_variants = phase.get("anyOf", [phase])
        advertised = {
            value
            for variant in phase_variants
            for value in variant.get("enum", [])
        }
        self.assertEqual(
            advertised,
            {"handoff_recovery", "preflight_evidence_backfill"},
        )
        self.assertNotIn("phase", recovery["inputSchema"].get("required", []))
        self.assertNotEqual(phase.get("default"), "handoff_recovery")

    async def test_omitted_phase_infers_backfill_and_replays_exact_plan(self) -> None:
        snapshot_before = self._recovery_authority_snapshot()
        explicit = self._backfill_plan().data["preflight_evidence_plan"]
        inferred_response = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=self.backfill_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=[self.old_handoff_id],
            action="plan",
            expected_plan_sha256=None,
            confirmation=None,
            request_session_id=self.session_id,
        )
        inferred = inferred_response.data["preflight_evidence_plan"]
        self.assertEqual(inferred["phase"], "preflight_evidence_backfill")
        self.assertEqual(inferred["plan_sha256"], explicit["plan_sha256"])
        self.assertEqual(self._recovery_authority_snapshot(), snapshot_before)

        result = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=self.backfill_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=[self.old_handoff_id],
            action="execute",
            expected_plan_sha256=inferred["plan_sha256"],
            confirmation=service_module.PREFLIGHT_EVIDENCE_REISSUE_CONFIRMATION,
            request_session_id=self.session_id,
        )
        replay = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=self.backfill_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=[self.old_handoff_id],
            action="execute",
            expected_plan_sha256=inferred["plan_sha256"],
            confirmation=service_module.PREFLIGHT_EVIDENCE_REISSUE_CONFIRMATION,
            request_session_id=self.session_id,
        )
        self.assertEqual(
            replay.data["mappings"][0]["preflight_evidence_id"],
            result.data["mappings"][0]["preflight_evidence_id"],
        )
        self.assertTrue(replay.data["mappings"][0]["idempotent_replay"])

    async def test_public_tool_accepts_stale_schema_call_without_phase(self) -> None:
        result = await self.server.call_tool(
            "recover_storyboard_handoff",
            {
                "project_id": self.project_id,
                "work_order_id": self.backfill_work_order_id,
                "gap_id": self.gap_id,
                "handoff_ids": [self.old_handoff_id],
                "action": "plan",
            },
        )
        self.assertFalse(result.is_error, result.content)
        self.assertIsNotNone(result.structured_content)
        response = result.structured_content or {}
        self.assertEqual(
            response["status"], "preflight_evidence_reissue_plan_ready"
        )
        self.assertEqual(
            response["data"]["preflight_evidence_plan"]["phase"],
            "preflight_evidence_backfill",
        )

    async def test_omitted_phase_rejects_ambiguous_tokens_without_mutation(self) -> None:
        malformed_targets = (
            [self.gap_id, self.old_handoff_id],
            [
                self.gap_id,
                self.old_handoff_id,
                "controlled_phase:preflight_evidence_backfill",
                "controlled_phase:preflight_evidence_backfill",
            ],
            [
                self.gap_id,
                self.old_handoff_id,
                "controlled_phase:preflight_evidence_backfill",
                "controlled_phase:handoff_recovery",
            ],
            [
                self.gap_id,
                self.old_handoff_id,
                "controlled_phase:unknown",
            ],
            [
                self.gap_id,
                self.old_handoff_id,
                "controlled_phase:preflight_evidence_backfill",
                "controlled_phase:unknown",
            ],
        )
        for index, targets in enumerate(malformed_targets):
            with self.subTest(index=index):
                response = self.service.begin_task(
                    self.namespace,
                    project_id=self.project_id,
                    operation="rule_improvement",
                    mode="execute",
                    authorization="explicit_execute",
                    request_text="malformed omitted-phase authority",
                    targets=targets,
                    authoritative_source="signed project records",
                    latest_version="0.5.9",
                    target_boundary="exact authorized handoff batch",
                    must_keep=["history"],
                    must_remove=[],
                    protected_scope=["storyboard", "audit history"],
                    expected_outputs=["controlled recovery"],
                    window_key=self.session_id,
                    request_session_id=self.session_id,
                )
                snapshot_before = self._recovery_authority_snapshot()
                for action in ("plan", "execute"):
                    with self.subTest(index=index, action=action):
                        with self.assertRaisesRegex(
                            ValueError, "handoff_recovery_phase_required"
                        ):
                            self.service.recover_storyboard_handoff(
                                self.namespace,
                                project_id=self.project_id,
                                work_order_id=str(response.work_order_id),
                                gap_id=self.gap_id,
                                handoff_ids=[self.old_handoff_id],
                                action=action,
                                expected_plan_sha256=(
                                    "0" * 64 if action == "execute" else None
                                ),
                                confirmation=(
                                    service_module.PREFLIGHT_EVIDENCE_REISSUE_CONFIRMATION
                                    if action == "execute"
                                    else None
                                ),
                                request_session_id=self.session_id,
                            )
                self.assertEqual(
                    self._recovery_authority_snapshot(), snapshot_before
                )

    async def test_explicit_phase_rejects_duplicate_targets_without_mutation(self) -> None:
        base_targets = [
            self.gap_id,
            self.old_handoff_id,
            "controlled_phase:preflight_evidence_backfill",
        ]
        for duplicate_index in range(len(base_targets)):
            with self.subTest(duplicate_index=duplicate_index):
                targets = [*base_targets, base_targets[duplicate_index]]
                response = self.service.begin_task(
                    self.namespace,
                    project_id=self.project_id,
                    operation="rule_improvement",
                    mode="execute",
                    authorization="explicit_execute",
                    request_text="duplicate exact-scope target",
                    targets=targets,
                    authoritative_source="signed project records",
                    latest_version="0.5.9",
                    target_boundary="exact authorized handoff batch",
                    must_keep=["history"],
                    must_remove=[],
                    protected_scope=["storyboard", "audit history"],
                    expected_outputs=["controlled recovery"],
                    window_key=self.session_id,
                    request_session_id=self.session_id,
                )
                snapshot_before = self._recovery_authority_snapshot()
                with self.assertRaisesRegex(
                    ValueError, "handoff_recovery_scope_mismatch"
                ):
                    self.service.recover_storyboard_handoff(
                        self.namespace,
                        project_id=self.project_id,
                        work_order_id=str(response.work_order_id),
                        gap_id=self.gap_id,
                        handoff_ids=[self.old_handoff_id],
                        action="plan",
                        expected_plan_sha256=None,
                        confirmation=None,
                        request_session_id=self.session_id,
                        phase="preflight_evidence_backfill",
                    )
                self.assertEqual(
                    self._recovery_authority_snapshot(), snapshot_before
                )

    async def test_omitted_phase_checks_session_before_target_shape(self) -> None:
        response = self.service.begin_task(
            self.namespace,
            project_id=self.project_id,
            operation="rule_improvement",
            mode="execute",
            authorization="explicit_execute",
            request_text="session ordering for omitted phase",
            targets=[self.gap_id, self.old_handoff_id],
            authoritative_source="signed project records",
            latest_version="0.5.9",
            target_boundary="exact authorized handoff batch",
            must_keep=["history"],
            must_remove=[],
            protected_scope=["storyboard", "audit history"],
            expected_outputs=["controlled recovery"],
            window_key=self.session_id,
            request_session_id=self.session_id,
        )
        snapshot_before = self._recovery_authority_snapshot()
        with self.assertRaisesRegex(ValueError, "window_session_mismatch"):
            self.service.recover_storyboard_handoff(
                self.namespace,
                project_id=self.project_id,
                work_order_id=str(response.work_order_id),
                gap_id=self.gap_id,
                handoff_ids=[self.old_handoff_id],
                action="plan",
                expected_plan_sha256=None,
                confirmation=None,
                request_session_id="another-production-session",
            )
        self.assertEqual(
            self._recovery_authority_snapshot(), snapshot_before
        )

    async def test_backfill_plan_is_zero_write_and_binds_exact_evidence(self) -> None:
        old_before = self.service.store.project_state_sha256(
            self.service.store.get_project_record(
                self.namespace,
                self.project_id,
                kind="handoffs",
                record_id=self.old_handoff_id,
            )
        )
        hub_before = self.service.store.project_state_sha256(
            self.service.store.load_hub(self.namespace, self.project_id)
        )
        certificate_count = len(
            self.service.store.list_project_records(
                self.namespace, self.project_id, kind="certificates"
            )
        )
        artifact_before = self._artifact_bytes()
        response = self._backfill_plan()
        self.assertEqual(response.status, "preflight_evidence_reissue_plan_ready")
        plan = response.data["preflight_evidence_plan"]
        item = plan["items"][0]
        self.assertTrue(item["preflight_evidence_id"].startswith("PFE-"))
        self.assertEqual(item["production_preflight_issue_count"], 0)
        self.assertFalse(item["first_draft_grade_reconstructed"])
        self.assertEqual(
            self.service.store.list_project_records(
                self.namespace,
                self.project_id,
                kind="preflight_evidence_reissues",
            ),
            [],
        )
        self.assertEqual(
            len(
                self.service.store.list_project_records(
                    self.namespace, self.project_id, kind="certificates"
                )
            ),
            certificate_count,
        )
        self.assertEqual(
            self.service.store.project_state_sha256(
                self.service.store.get_project_record(
                    self.namespace,
                    self.project_id,
                    kind="handoffs",
                    record_id=self.old_handoff_id,
                )
            ),
            old_before,
        )
        self.assertEqual(
            self.service.store.project_state_sha256(
                self.service.store.load_hub(self.namespace, self.project_id)
            ),
            hub_before,
        )
        self.assertEqual(self._artifact_bytes(), artifact_before)

    async def test_backfill_is_append_only_and_idempotent(self) -> None:
        old_before = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        old_sha = self.service.store.project_state_sha256(old_before)
        old_certificate = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=str(old_before["certificate_id"]),
        )
        old_certificate_sha = self.service.store.project_state_sha256(
            old_certificate
        )
        artifact_before = self._artifact_bytes()
        plan = self._backfill_plan().data["preflight_evidence_plan"]
        result = self._backfill_execute(plan["plan_sha256"])
        self.assertEqual(result.status, "preflight_evidence_reissued")
        mapping = result.data["mappings"][0]
        evidence = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="preflight_evidence_reissues",
            record_id=mapping["preflight_evidence_id"],
        )
        certificate = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=mapping["certificate_id"],
        )
        self.assertEqual(evidence["status"], "committed")
        self.assertTrue(self.service.store.project_record_is_signed(evidence))
        self.assertEqual(
            certificate["evidence_origin"],
            "controlled_current_revalidation",
        )
        self.assertEqual(certificate["preflight_issue_count"], 0)
        self.assertNotIn("result", certificate)
        self.assertEqual(
            self.service.store.project_state_sha256(
                self.service.store.get_project_record(
                    self.namespace,
                    self.project_id,
                    kind="handoffs",
                    record_id=self.old_handoff_id,
                )
            ),
            old_sha,
        )
        self.assertEqual(
            self.service.store.project_state_sha256(
                self.service.store.get_project_record(
                    self.namespace,
                    self.project_id,
                    kind="certificates",
                    record_id=str(old_before["certificate_id"]),
                )
            ),
            old_certificate_sha,
        )
        self.assertEqual(self._artifact_bytes(), artifact_before)
        replay = self._backfill_execute(plan["plan_sha256"])
        self.assertEqual(
            replay.data["mappings"][0]["preflight_evidence_id"],
            mapping["preflight_evidence_id"],
        )
        self.assertTrue(replay.data["mappings"][0]["idempotent_replay"])
        self.assertEqual(
            len(
                self.service.store.list_project_records(
                    self.namespace,
                    self.project_id,
                    kind="preflight_evidence_reissues",
                )
            ),
            1,
        )

    async def test_v053_transitional_workbook_backfill_full_chain(
        self,
    ) -> None:
        fixture = self._install_v053_transitional_workbook_fixture()
        old_handoff = fixture["old_handoff"]
        old_certificate = fixture["old_certificate"]
        original_trace = fixture["original_trace"]
        old_handoff_sha256 = self.service.store.project_state_sha256(
            old_handoff
        )
        old_certificate_sha256 = self.service.store.project_state_sha256(
            old_certificate
        )
        original_trace_sha256 = self.service.store.project_state_sha256(
            original_trace
        )
        _canonical_record, canonical_path = self.service.store.get_artifact(
            self.namespace,
            self.project_id,
            str(old_handoff["canonical_artifact_id"]),
        )
        canonical_bytes_before = canonical_path.read_bytes()
        authority_before_plan = self._recovery_authority_snapshot()

        response = self._backfill_plan()
        self.assertEqual(
            response.status, "preflight_evidence_reissue_plan_ready"
        )
        self.assertEqual(
            self._recovery_authority_snapshot(), authority_before_plan
        )
        self.assertEqual(
            self._artifact_bytes(), fixture["artifact_bytes"]
        )
        self.assertEqual(canonical_path.read_bytes(), canonical_bytes_before)

        plan = response.data["preflight_evidence_plan"]
        item = plan["items"][0]
        proof = item["workbook_projection_revalidation"]
        self.assertEqual(
            proof["mode"],
            "legacy_envelope_short_header_v053_rows",
        )
        self.assertTrue(proof["compatibility_applied"])
        self.assertEqual(
            proof["sidecar_projection_contract"],
            "shortdrama-projector/0.5.0",
        )
        self.assertEqual(
            proof["selected_projector"],
            "shortdrama-projector/0.5.3-frozen",
        )
        self.assertTrue(proof["visible_rows_match_selected_projector"])
        self.assertFalse(proof["visible_projection_matches_sidecar"])
        self.assertFalse(proof["visible_audit_page_matches_sidecar"])
        self.assertIsNone(proof["sidecar_projection_sha256"])
        self.assertIsNone(proof["sidecar_audit_page_sha256"])
        self.assertEqual(
            proof["audit_page_integrity_source"],
            "signed_original_artifact_raw_sha256",
        )
        self.assertTrue(
            proof["artifact_raw_bound_by_original_certificate"]
        )
        self.assertEqual(
            proof["artifact_raw_sha256"],
            fixture["artifact_raw_sha256"],
        )
        self.assertEqual(
            proof["canonical_sha256"], old_handoff["canonical_sha256"]
        )
        self.assertEqual(
            proof["old_handoff_sha256"], old_handoff_sha256
        )
        self.assertEqual(
            proof["old_certificate_id"], old_certificate["certificate_id"]
        )
        self.assertEqual(
            proof["old_certificate_sha256"], old_certificate_sha256
        )
        self.assertEqual(
            proof["original_production_trace_id"],
            original_trace["trace_id"],
        )
        self.assertEqual(
            proof["original_production_trace_sha256"],
            original_trace_sha256,
        )
        proof_without_identity = copy.deepcopy(proof)
        proof_sha256 = proof_without_identity.pop("proof_sha256")
        self.assertEqual(
            proof_sha256,
            service_module.sha256_bytes(
                service_module._stable_json_bytes(proof_without_identity)
            ),
        )

        result = self._backfill_execute(plan["plan_sha256"])
        self.assertEqual(result.status, "preflight_evidence_reissued")
        mapping = result.data["mappings"][0]
        self.assertEqual(
            mapping["workbook_projection_mode"], proof["mode"]
        )
        self.assertEqual(
            mapping["workbook_projection_proof_sha256"], proof_sha256
        )
        self.assertFalse(mapping["idempotent_replay"])
        evidence = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="preflight_evidence_reissues",
            record_id=str(mapping["preflight_evidence_id"]),
        )
        certificate = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=str(mapping["certificate_id"]),
        )
        evidence_trace = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="telemetry",
            record_id=str(mapping["trace_id"]),
        )
        for record in (evidence, certificate, evidence_trace):
            self.assertTrue(
                self.service.store.project_record_is_signed(record)
            )
        self.assertEqual(evidence["status"], "committed")
        self.assertEqual(
            evidence["workbook_projection_revalidation"], proof
        )
        self.assertEqual(
            certificate["workbook_projection_revalidation"], proof
        )
        self.assertEqual(
            evidence["original_production_trace_id"],
            original_trace["trace_id"],
        )
        self.assertEqual(
            certificate["original_certificate_id"],
            old_certificate["certificate_id"],
        )

        self.assertEqual(
            self.service.store.project_state_sha256(
                self.service.store.get_project_record(
                    self.namespace,
                    self.project_id,
                    kind="handoffs",
                    record_id=self.old_handoff_id,
                )
            ),
            old_handoff_sha256,
        )
        self.assertEqual(
            self.service.store.project_state_sha256(
                self.service.store.get_project_record(
                    self.namespace,
                    self.project_id,
                    kind="certificates",
                    record_id=str(old_certificate["certificate_id"]),
                )
            ),
            old_certificate_sha256,
        )
        self.assertEqual(
            self.service.store.project_state_sha256(
                self.service.store.get_project_record(
                    self.namespace,
                    self.project_id,
                    kind="telemetry",
                    record_id=str(original_trace["trace_id"]),
                )
            ),
            original_trace_sha256,
        )
        self.assertEqual(
            self._artifact_bytes(), fixture["artifact_bytes"]
        )
        self.assertEqual(canonical_path.read_bytes(), canonical_bytes_before)

        authority_after_execute = self._recovery_authority_snapshot()
        replay = self._backfill_execute(plan["plan_sha256"])
        replay_mapping = replay.data["mappings"][0]
        for field in (
            "preflight_evidence_id",
            "certificate_id",
            "trace_id",
            "workbook_projection_mode",
            "workbook_projection_proof_sha256",
        ):
            self.assertEqual(replay_mapping[field], mapping[field])
        self.assertTrue(replay_mapping["idempotent_replay"])
        self.assertEqual(
            self._recovery_authority_snapshot(), authority_after_execute
        )
        self.assertEqual(
            self._artifact_bytes(), fixture["artifact_bytes"]
        )
        self.assertEqual(canonical_path.read_bytes(), canonical_bytes_before)

    async def test_exact_legacy_signed_hub_envelope_supports_full_chain(
        self,
    ) -> None:
        signed_old = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        self.assertEqual(
            {
                "record_id",
                "project_id",
                "integrity_hmac",
            },
            {
                field
                for field in signed_old
                if field
                in {"record_id", "project_id", "integrity_hmac"}
            },
        )
        self._replace_hub_handoff_rows(
            self.old_handoff_id,
            [signed_old],
        )
        legacy_hub_before = self.service.store.load_hub(
            self.namespace, self.project_id
        )
        legacy_row_before = next(
            row
            for row in legacy_hub_before["handoffs"]
            if row.get("handoff_id") == self.old_handoff_id
        )
        self.assertEqual(legacy_row_before, signed_old)

        old_handoff_sha256 = self.service.store.project_state_sha256(
            signed_old
        )
        original_certificate = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=str(signed_old["certificate_id"]),
        )
        original_certificate_sha256 = (
            self.service.store.project_state_sha256(original_certificate)
        )
        artifact_before = self._artifact_bytes()

        backfill_plan = self._backfill_plan().data[
            "preflight_evidence_plan"
        ]
        backfill_result = self._backfill_execute(
            backfill_plan["plan_sha256"]
        )
        self.assertEqual(
            backfill_result.status, "preflight_evidence_reissued"
        )
        backfill_mapping = backfill_result.data["mappings"][0]
        backfill_replay = self._backfill_execute(
            backfill_plan["plan_sha256"]
        )
        self.assertEqual(
            backfill_replay.data["mappings"][0][
                "preflight_evidence_id"
            ],
            backfill_mapping["preflight_evidence_id"],
        )
        self.assertTrue(
            backfill_replay.data["mappings"][0]["idempotent_replay"]
        )

        signed_after_backfill = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        self.assertEqual(
            self.service.store.project_state_sha256(
                signed_after_backfill
            ),
            old_handoff_sha256,
        )
        hub_after_backfill = self.service.store.load_hub(
            self.namespace, self.project_id
        )
        legacy_rows_after_backfill = [
            row
            for row in hub_after_backfill["handoffs"]
            if row.get("handoff_id") == self.old_handoff_id
        ]
        self.assertEqual(legacy_rows_after_backfill, [signed_old])
        self.assertEqual(
            self.service.store.project_state_sha256(
                self.service.store.get_project_record(
                    self.namespace,
                    self.project_id,
                    kind="certificates",
                    record_id=str(signed_old["certificate_id"]),
                )
            ),
            original_certificate_sha256,
        )
        self.assertEqual(self._artifact_bytes(), artifact_before)

        recovery_plan = self._plan().data["recovery_plan"]
        self.assertEqual(
            recovery_plan["items"][0]["preflight_evidence_source"],
            "controlled_current_revalidation",
        )
        recovery_result = self._execute(
            recovery_plan["plan_sha256"]
        )
        self.assertEqual(recovery_result.status, "handoffs_recovered")
        recovery_mapping = recovery_result.data["mappings"][0]
        recovery_replay = self._execute(recovery_plan["plan_sha256"])
        self.assertEqual(
            recovery_replay.data["mappings"][0]["new_handoff_id"],
            recovery_mapping["new_handoff_id"],
        )
        self.assertTrue(
            recovery_replay.data["mappings"][0]["idempotent_replay"]
        )

        old_after = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        new_after = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=str(recovery_mapping["new_handoff_id"]),
        )
        self.assertEqual(old_after["status"], "superseded")
        self.assertEqual(
            old_after["superseded_by_handoff_id"],
            new_after["handoff_id"],
        )
        self.assertEqual(new_after["status"], "prepared")
        self.assertEqual(
            new_after["independent_audit_status"], "PENDING_AUDIT"
        )
        self.assertEqual(
            new_after["supersedes_handoff_id"], old_after["handoff_id"]
        )

        final_hub = self.service.store.load_hub(
            self.namespace, self.project_id
        )
        final_old_rows = [
            row
            for row in final_hub["handoffs"]
            if row.get("handoff_id") == old_after["handoff_id"]
        ]
        final_new_rows = [
            row
            for row in final_hub["handoffs"]
            if row.get("handoff_id") == new_after["handoff_id"]
        ]
        self.assertEqual(
            final_old_rows,
            [self.service._handoff_record_projection(old_after)],
        )
        self.assertEqual(
            final_new_rows,
            [self.service._handoff_record_projection(new_after)],
        )
        for row in (*final_old_rows, *final_new_rows):
            self.assertTrue(
                {"record_id", "project_id", "integrity_hmac"}.isdisjoint(
                    row
                )
            )
        self.assertEqual(self._artifact_bytes(), artifact_before)

    async def test_legacy_signed_hub_envelope_poison_is_zero_write(
        self,
    ) -> None:
        signed_old = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )

        poison_cases = {
            "integrity_hmac_drift": lambda row: row.__setitem__(
                "integrity_hmac", "0" * 64
            ),
            "record_id_drift": lambda row: row.__setitem__(
                "record_id", "HND-v058-envelope-drift"
            ),
            "project_id_drift": lambda row: row.__setitem__(
                "project_id", "SBP-v058-envelope-drift"
            ),
            "extra_key": lambda row: row.__setitem__(
                "legacy_shadow", True
            ),
            "partial_envelope": lambda row: row.pop(
                "integrity_hmac"
            ),
            "missing_business_field": lambda row: row.pop(
                "canonical_sha256"
            ),
            "semantic_drift": lambda row: row.__setitem__(
                "canonical_sha256", "1" * 64
            ),
        }
        for label, poison in poison_cases.items():
            with self.subTest(case=label):
                poisoned = copy.deepcopy(signed_old)
                poison(poisoned)
                self._replace_hub_handoff_rows(
                    self.old_handoff_id,
                    [poisoned],
                )
                before = self._recovery_authority_snapshot()
                with self.assertRaisesRegex(
                    ValueError, "handoff_recovery_projection_invalid"
                ):
                    self._backfill_plan()
                self.assertEqual(
                    self._recovery_authority_snapshot(), before
                )

        with self.subTest(case="duplicate_row"):
            self._replace_hub_handoff_rows(
                self.old_handoff_id,
                [signed_old, signed_old],
            )
            before = self._recovery_authority_snapshot()
            with self.assertRaisesRegex(
                ValueError, "handoff_recovery_projection_invalid"
            ):
                self._backfill_plan()
            self.assertEqual(self._recovery_authority_snapshot(), before)

    async def test_wrong_confirmation_is_zero_write(self) -> None:
        plan = self._backfill_plan().data["preflight_evidence_plan"]
        with self.assertRaisesRegex(
            ValueError, "preflight_evidence_reissue_confirmation_required"
        ):
            self._backfill_execute(plan["plan_sha256"], confirmation="wrong")
        self.assertEqual(
            self.service.store.list_project_records(
                self.namespace,
                self.project_id,
                kind="preflight_evidence_reissues",
            ),
            [],
        )

    async def test_interrupted_certificate_write_recovers_same_identity(self) -> None:
        plan = self._backfill_plan().data["preflight_evidence_plan"]
        certificate_id = plan["items"][0]["new_certificate_id"]
        original = self.service.store.atomic_write_json
        tripped = False

        def fail_once(path: Path, value: dict[str, Any]) -> None:
            nonlocal tripped
            if (
                not tripped
                and path.parent.name == "certificates"
                and path.name == f"{certificate_id}.json"
            ):
                tripped = True
                raise OSError("simulated preflight certificate interruption")
            original(path, value)

        with patch.object(
            self.service.store, "atomic_write_json", side_effect=fail_once
        ):
            with self.assertRaisesRegex(
                OSError, "simulated preflight certificate interruption"
            ):
                self._backfill_execute(plan["plan_sha256"])
        evidence_id = plan["items"][0]["preflight_evidence_id"]
        prepared = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="preflight_evidence_reissues",
            record_id=evidence_id,
        )
        self.assertEqual(prepared["status"], "prepared")
        retry = self._backfill_execute(plan["plan_sha256"])
        self.assertEqual(
            retry.data["mappings"][0]["preflight_evidence_id"], evidence_id
        )
        committed = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="preflight_evidence_reissues",
            record_id=evidence_id,
        )
        self.assertEqual(committed["status"], "committed")

    async def test_explicit_failure_and_content_drift_are_rejected(self) -> None:
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        ) as failed:
            failed["production_preflight_status"] = (
                service_module.PRODUCTION_PREFLIGHT_FAIL
            )
            failed["production_preflight_issue_count"] = 1
        with self.service.store.mutate_project(
            self.namespace, self.project_id
        ) as (hub, _ledger):
            row = next(
                item
                for item in hub["handoffs"]
                if item["handoff_id"] == self.old_handoff_id
            )
            row["production_preflight_status"] = (
                service_module.PRODUCTION_PREFLIGHT_FAIL
            )
            row["production_preflight_issue_count"] = 1
        with self.assertRaisesRegex(
            ValueError, "preflight_evidence_reissue_forbidden"
        ):
            self._backfill_plan()
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        ) as drifted:
            drifted.pop("production_preflight_status", None)
            drifted.pop("production_preflight_issue_count", None)
            drifted["fingerprints"]["artifact_raw"] = "1" * 64
        with self.service.store.mutate_project(
            self.namespace, self.project_id
        ) as (hub, _ledger):
            row = next(
                item
                for item in hub["handoffs"]
                if item["handoff_id"] == self.old_handoff_id
            )
            row.pop("production_preflight_status", None)
            row.pop("production_preflight_issue_count", None)
            row["fingerprints"]["artifact_raw"] = "1" * 64
        with self.assertRaisesRegex(
            ValueError,
            "handoff_recovery_(?:certificate_invalid|content_changed)",
        ):
            self._backfill_plan()

    async def test_signed_legacy_certificate_report_failure_is_rejected(self) -> None:
        handoff = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=str(handoff["certificate_id"]),
        ) as certificate:
            for field in (
                "preflight_status",
                "preflight_issue_count",
                "independent_audit_status",
            ):
                certificate.pop(field, None)
            report = copy.deepcopy(certificate["report"])
            report["issues"] = [
                {
                    "error_code": "E-EXPLICIT-LEGACY-FAILURE",
                    "message": "signed legacy report contains a failure",
                }
            ]
            certificate["report"] = report
            certificate["report_sha256"] = "0" * 64
        with self.assertRaisesRegex(
            ValueError, "handoff_recovery_certificate_invalid"
        ):
            self._backfill_plan()

    async def test_visible_workbook_projection_drift_is_rejected(self) -> None:
        handoff = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        artifact_id = str(handoff["storyboard_artifact_id"])
        record, path = self.service.store.get_artifact(
            self.namespace, self.project_id, artifact_id
        )
        workbook = load_workbook(path)
        workbook["正式生产母版"].cell(row=2, column=5).value = (
            "tampered visible projection"
        )
        workbook.save(path)
        workbook.close()
        payload = path.read_bytes()
        meta_path = (
            self.service.store.project_dir(
                self.namespace, self.project_id
            )
            / "artifacts"
            / f"{artifact_id}.meta.json"
        )
        record["fingerprint"] = service_module.sha256_bytes(payload)
        record["size_bytes"] = len(payload)
        self.service.store.atomic_write_json(meta_path, record)
        with self.assertRaisesRegex(
            ValueError, "canonical sidecar invalid"
        ):
            self._backfill_plan()

    async def test_phase_bound_work_orders_cannot_be_reused(self) -> None:
        with self.assertRaisesRegex(
            ValueError, "handoff_recovery_scope_mismatch"
        ):
            self.service.recover_storyboard_handoff(
                self.namespace,
                project_id=self.project_id,
                work_order_id=self.recovery_work_order_id,
                gap_id=self.gap_id,
                handoff_ids=[self.old_handoff_id],
                action="plan",
                expected_plan_sha256=None,
                confirmation=None,
                request_session_id=self.session_id,
                phase="preflight_evidence_backfill",
            )
        backfill_plan = self._backfill_plan().data[
            "preflight_evidence_plan"
        ]
        self._backfill_execute(backfill_plan["plan_sha256"])
        with self.assertRaisesRegex(
            ValueError, "handoff_recovery_scope_mismatch"
        ):
            self.service.recover_storyboard_handoff(
                self.namespace,
                project_id=self.project_id,
                work_order_id=self.backfill_work_order_id,
                gap_id=self.gap_id,
                handoff_ids=[self.old_handoff_id],
                action="plan",
                expected_plan_sha256=None,
                confirmation=None,
                request_session_id=self.session_id,
                phase="handoff_recovery",
            )
        old = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=self.old_handoff_id,
        )
        self.assertEqual(old["status"], "prepared")

    async def test_signed_but_semantically_misbound_reissue_is_rejected(self) -> None:
        plan = self._backfill_plan().data["preflight_evidence_plan"]
        result = self._backfill_execute(plan["plan_sha256"])
        mapping = result.data["mappings"][0]
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=str(mapping["certificate_id"]),
        ) as certificate:
            certificate["source_fingerprint"] = "0" * 64
            certificate["report"]["issues"] = [
                {"error_code": "E-MISBOUND", "message": "tampered"}
            ]
        tampered_certificate = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=str(mapping["certificate_id"]),
        )
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="preflight_evidence_reissues",
            record_id=str(mapping["preflight_evidence_id"]),
        ) as evidence:
            evidence["certificate_after"] = copy.deepcopy(
                tampered_certificate
            )
            evidence["certificate_after_sha256"] = (
                self.service.store.project_state_sha256(
                    tampered_certificate
                )
            )
        with self.assertRaisesRegex(
            ValueError,
            "(handoff_preflight_evidence_reissue_invalid|lifecycle_schema_mismatch)",
        ):
            self.service.recover_storyboard_handoff(
                self.namespace,
                project_id=self.project_id,
                work_order_id=self.recovery_work_order_id,
                gap_id=self.gap_id,
                handoff_ids=[self.old_handoff_id],
                action="plan",
                expected_plan_sha256=None,
                confirmation=None,
                request_session_id=self.session_id,
            )

    async def test_batch_backfill_authority_is_consumable_by_batch_recovery(self) -> None:
        second_id = self.clean_peer_handoff_id
        self._remove_handoff_preflight_summary(second_id)
        handoff_ids = [self.old_handoff_id, second_id]
        backfill_work_order_id = self._controlled_work_order(
            phase="preflight_evidence_backfill",
            handoff_ids=handoff_ids,
        )
        backfill_plan_response = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=backfill_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=handoff_ids,
            action="plan",
            expected_plan_sha256=None,
            confirmation=None,
            request_session_id=self.session_id,
            phase="preflight_evidence_backfill",
        )
        backfill_plan = backfill_plan_response.data[
            "preflight_evidence_plan"
        ]
        self.assertEqual(
            backfill_plan["authorized_handoff_ids"], handoff_ids
        )
        backfill_result = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=backfill_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=handoff_ids,
            action="execute",
            expected_plan_sha256=backfill_plan["plan_sha256"],
            confirmation=service_module.PREFLIGHT_EVIDENCE_REISSUE_CONFIRMATION,
            request_session_id=self.session_id,
            phase="preflight_evidence_backfill",
        )
        self.assertEqual(len(backfill_result.data["mappings"]), 2)
        for mapping in backfill_result.data["mappings"]:
            evidence = self.service.store.get_project_record(
                self.namespace,
                self.project_id,
                kind="preflight_evidence_reissues",
                record_id=mapping["preflight_evidence_id"],
            )
            self.assertEqual(evidence["authorized_handoff_ids"], handoff_ids)

        recovery_work_order_id = self._controlled_work_order(
            phase="handoff_recovery", handoff_ids=handoff_ids
        )
        recovery_plan_response = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=recovery_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=handoff_ids,
            action="plan",
            expected_plan_sha256=None,
            confirmation=None,
            request_session_id=self.session_id,
        )
        recovery_plan = recovery_plan_response.data["recovery_plan"]
        recovery_result = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=recovery_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=handoff_ids,
            action="execute",
            expected_plan_sha256=recovery_plan["plan_sha256"],
            confirmation=service_module.HANDOFF_RECOVERY_CONFIRMATION,
            request_session_id=self.session_id,
        )
        self.assertEqual(
            [row["old_handoff_id"] for row in recovery_result.data["mappings"]],
            handoff_ids,
        )

    async def test_version_management_cannot_authorize_either_phase(self) -> None:
        evidence_count_before = len(
            self.service.store.list_project_records(
                self.namespace,
                self.project_id,
                kind="preflight_evidence_reissues",
            )
        )
        for phase in (
            "preflight_evidence_backfill",
            "handoff_recovery",
        ):
            work_order_id = self._controlled_work_order(
                phase=phase,
                handoff_ids=[self.old_handoff_id],
                operation="version_management",
            )
            with self.assertRaisesRegex(
                ValueError, "handoff_recovery_authority_required"
            ):
                self.service.recover_storyboard_handoff(
                    self.namespace,
                    project_id=self.project_id,
                    work_order_id=work_order_id,
                    gap_id=self.gap_id,
                    handoff_ids=[self.old_handoff_id],
                    action="plan",
                    expected_plan_sha256=None,
                    confirmation=None,
                    request_session_id=self.session_id,
                    phase=phase,
                )
        self.assertEqual(
            len(
                self.service.store.list_project_records(
                    self.namespace,
                    self.project_id,
                    kind="preflight_evidence_reissues",
                )
            ),
            evidence_count_before,
        )

    async def test_controlled_certificate_provenance_poison_is_rejected(self) -> None:
        plan = self._backfill_plan().data["preflight_evidence_plan"]
        result = self._backfill_execute(plan["plan_sha256"])
        mapping = result.data["mappings"][0]
        certificate_id = str(mapping["certificate_id"])
        evidence_id = str(mapping["preflight_evidence_id"])
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=certificate_id,
        ) as certificate:
            certificate["evidence_reissue_gap_id"] = "GAP-semantic-poison"
            certificate["original_certificate_id"] = "CERT-semantic-poison"
            certificate["authored_artifacts_modified"] = True
        poisoned = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=certificate_id,
        )
        with self.service.store.mutate_project_record(
            self.namespace,
            self.project_id,
            kind="preflight_evidence_reissues",
            record_id=evidence_id,
        ) as evidence:
            evidence["certificate_after"] = copy.deepcopy(poisoned)
            evidence["certificate_after_sha256"] = (
                self.service.store.project_state_sha256(poisoned)
            )
        with self.assertRaisesRegex(
            ValueError,
            "(handoff_preflight_evidence_reissue_invalid|lifecycle_schema_mismatch)",
        ):
            self.service.recover_storyboard_handoff(
                self.namespace,
                project_id=self.project_id,
                work_order_id=self.recovery_work_order_id,
                gap_id=self.gap_id,
                handoff_ids=[self.old_handoff_id],
                action="plan",
                expected_plan_sha256=None,
                confirmation=None,
                request_session_id=self.session_id,
            )

    async def test_revalidation_traces_are_measured_and_claim_lifecycle_stays_published(self) -> None:
        backfill_plan = self._backfill_plan().data[
            "preflight_evidence_plan"
        ]
        backfill_result = self._backfill_execute(
            backfill_plan["plan_sha256"]
        )
        evidence_trace = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="telemetry",
            record_id=backfill_result.data["mappings"][0]["trace_id"],
        )
        self.assertGreater(evidence_trace["wall_duration_ms"], 0)
        self.assertGreater(evidence_trace["stages"][0]["duration_ms"], 0)

        recovery_plan = self._plan().data["recovery_plan"]
        recovery_result = self._execute(recovery_plan["plan_sha256"])
        mapping = recovery_result.data["mappings"][0]
        recovery_trace = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="telemetry",
            record_id=mapping["trace_id"],
        )
        self.assertGreater(recovery_trace["wall_duration_ms"], 0)
        self.assertGreater(recovery_trace["stages"][0]["duration_ms"], 0)
        recovered_handoff = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=mapping["new_handoff_id"],
        )
        audit_session_id = "v057-independent-audit"
        audit_task = self.service.begin_task(
            self.namespace,
            project_id=self.project_id,
            operation="storyboard_audit",
            mode="audit_only",
            authorization="read_only",
            request_text="v0.5.7 recovered handoff lifecycle audit fixture",
            targets=[str(recovered_handoff["storyboard_artifact_id"])],
            authoritative_source="signed recovered production handoff",
            latest_version="current",
            target_boundary="EP01",
            must_keep=["read-only independent audit"],
            must_remove=[],
            protected_scope=["storyboard", "canonical", "XLSX"],
            expected_outputs=["independent_audit_certificate"],
            window_key=audit_session_id,
            request_session_id=audit_session_id,
        )
        audit_work_order_id = str(audit_task.work_order_id)
        claimed = self.service.store.claim_handoff(
            self.namespace,
            self.project_id,
            mapping["new_handoff_id"],
            consumer_work_order_id=audit_work_order_id,
            expected_to_role="independent_audit",
            expected_types={"production_to_audit"},
        )
        self.assertTrue(
            self.service.store.handoff_is_published(
                self.namespace, self.project_id, claimed
            )
        )
        with self.assertRaisesRegex(
            ValueError, "handoff_state_conflict: handoff is being consumed"
        ):
            self.service.store.claim_handoff(
                self.namespace,
                self.project_id,
                mapping["new_handoff_id"],
                consumer_work_order_id=audit_work_order_id,
                expected_to_role="independent_audit",
                expected_types={"production_to_audit"},
            )
        released = self.service.store.release_handoff_claim(
            self.namespace,
            self.project_id,
            mapping["new_handoff_id"],
            consumer_work_order_id=audit_work_order_id,
            claim_token=claimed["claim_token"],
            reason="retry fixture",
        )
        self.assertTrue(
            self.service.store.handoff_is_published(
                self.namespace, self.project_id, released
            )
        )
        reclaimed = self.service.store.claim_handoff(
            self.namespace,
            self.project_id,
            mapping["new_handoff_id"],
            consumer_work_order_id=audit_work_order_id,
            expected_to_role="independent_audit",
            expected_types={"production_to_audit"},
        )
        self.assertEqual(reclaimed["claim_generation"], 2)
        audit_certificate_id = "AUD-v057-finished"
        audit_certificate = self.service.store.save_project_record(
            self.namespace,
            self.project_id,
            kind="certificates",
            record_id=audit_certificate_id,
            value={
                "certificate_id": audit_certificate_id,
                "certificate_type": "independent_audit",
                "audit_class": "independent_audit",
                "handoff_id": mapping["new_handoff_id"],
                "work_order_id": audit_work_order_id,
                "artifact_id": str(
                    recovered_handoff["storyboard_artifact_id"]
                ),
            },
        )
        self.assertTrue(
            self.service.store.project_record_is_signed(audit_certificate)
        )
        with self.service.store.mutate_project(
            self.namespace, self.project_id
        ) as (hub, _ledger):
            hub_handoff = next(
                row
                for row in hub["handoffs"]
                if row["handoff_id"] == mapping["new_handoff_id"]
            )
            hub_handoff["status"] = "consumed"
            hub_handoff["audit_certificate_id"] = audit_certificate_id
        finished = self.service.store.finish_handoff(
            self.namespace,
            self.project_id,
            mapping["new_handoff_id"],
            consumer_work_order_id=audit_work_order_id,
            audit_certificate_id=audit_certificate_id,
            claim_token=reclaimed["claim_token"],
        )
        self.assertTrue(
            self.service.store.handoff_is_published(
                self.namespace, self.project_id, finished
            )
        )
        consumed_replay = self.service.store.claim_handoff(
            self.namespace,
            self.project_id,
            mapping["new_handoff_id"],
            consumer_work_order_id=audit_work_order_id,
            expected_to_role="independent_audit",
            expected_types={"production_to_audit"},
        )
        self.assertTrue(consumed_replay["idempotent_replay"])

    async def test_mixed_original_and_reissued_evidence_recovers_atomically(self) -> None:
        backfill_plan = self._backfill_plan().data[
            "preflight_evidence_plan"
        ]
        self._backfill_execute(backfill_plan["plan_sha256"])
        mixed_session = "v057-mixed-recovery"
        response = self.service.begin_task(
            self.namespace,
            project_id=self.project_id,
            operation="rule_improvement",
            mode="execute",
            authorization="explicit_execute",
            request_text="authorized mixed evidence recovery",
            targets=[
                self.gap_id,
                self.clean_peer_handoff_id,
                self.old_handoff_id,
                "controlled_phase:handoff_recovery",
            ],
            authoritative_source="signed project records",
            latest_version="0.5.9",
            target_boundary="two exact handoffs",
            must_keep=["history", "content", "independent audit separation"],
            must_remove=[],
            protected_scope=["storyboard", "canonical", "XLSX", "audit history"],
            expected_outputs=["two prepared PENDING_AUDIT successors"],
            window_key=mixed_session,
            request_session_id=mixed_session,
        )
        work_order_id = str(response.work_order_id)
        ids = [self.clean_peer_handoff_id, self.old_handoff_id]
        plan_response = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=work_order_id,
            gap_id=self.gap_id,
            handoff_ids=ids,
            action="plan",
            expected_plan_sha256=None,
            confirmation=None,
            request_session_id=mixed_session,
        )
        plan = plan_response.data["recovery_plan"]
        self.assertEqual(
            [row["preflight_evidence_source"] for row in plan["items"]],
            ["original_handoff", "controlled_current_revalidation"],
        )
        result = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=work_order_id,
            gap_id=self.gap_id,
            handoff_ids=ids,
            action="execute",
            expected_plan_sha256=plan["plan_sha256"],
            confirmation=service_module.HANDOFF_RECOVERY_CONFIRMATION,
            request_session_id=mixed_session,
        )
        self.assertEqual(result.status, "handoffs_recovered")
        self.assertEqual(len(result.data["mappings"]), 2)
        handoff_count = len(
            self.service.store.list_project_records(
                self.namespace, self.project_id, kind="handoffs"
            )
        )
        for mapping in result.data["mappings"]:
            old = self.service.store.get_project_record(
                self.namespace,
                self.project_id,
                kind="handoffs",
                record_id=mapping["old_handoff_id"],
            )
            new = self.service.store.get_project_record(
                self.namespace,
                self.project_id,
                kind="handoffs",
                record_id=mapping["new_handoff_id"],
            )
            self.assertEqual(old["status"], "superseded")
            self.assertEqual(new["status"], "prepared")
            self.assertEqual(new["independent_audit_status"], "PENDING_AUDIT")
        independent = [
            row
            for row in self.service.store.list_project_records(
                self.namespace, self.project_id, kind="certificates"
            )
            if row.get("certificate_type") == "independent_audit"
        ]
        self.assertEqual(independent, [])
        replay = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=work_order_id,
            gap_id=self.gap_id,
            handoff_ids=ids,
            action="execute",
            expected_plan_sha256=plan["plan_sha256"],
            confirmation=service_module.HANDOFF_RECOVERY_CONFIRMATION,
            request_session_id=mixed_session,
        )
        self.assertTrue(
            all(row["idempotent_replay"] for row in replay.data["mappings"])
        )
        self.assertEqual(
            len(
                self.service.store.list_project_records(
                    self.namespace, self.project_id, kind="handoffs"
                )
            ),
            handoff_count,
        )

    async def test_recovery_successor_is_unclaimable_until_journal_commits(self) -> None:
        backfill_plan = self._backfill_plan().data[
            "preflight_evidence_plan"
        ]
        self._backfill_execute(backfill_plan["plan_sha256"])
        recovery_plan = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=self.recovery_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=[self.old_handoff_id],
            action="plan",
            expected_plan_sha256=None,
            confirmation=None,
            request_session_id=self.session_id,
        ).data["recovery_plan"]
        new_handoff_id = recovery_plan["items"][0]["new_handoff_id"]
        original = self.service.store.atomic_write_json
        tripped = False

        def fail_after_hub_write(path: Path, value: dict[str, Any]) -> None:
            nonlocal tripped
            original(path, value)
            if not tripped and path.name == "hub.json":
                tripped = True
                raise OSError("simulated recovery journal finalization interruption")

        with patch.object(
            self.service.store,
            "atomic_write_json",
            side_effect=fail_after_hub_write,
        ):
            with self.assertRaisesRegex(
                OSError, "simulated recovery journal finalization"
            ):
                self.service.recover_storyboard_handoff(
                    self.namespace,
                    project_id=self.project_id,
                    work_order_id=self.recovery_work_order_id,
                    gap_id=self.gap_id,
                    handoff_ids=[self.old_handoff_id],
                    action="execute",
                    expected_plan_sha256=recovery_plan["plan_sha256"],
                    confirmation=service_module.HANDOFF_RECOVERY_CONFIRMATION,
                    request_session_id=self.session_id,
                )
        partial = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=new_handoff_id,
        )
        self.assertFalse(
            self.service.store.handoff_is_published(
                self.namespace, self.project_id, partial
            )
        )
        with self.assertRaisesRegex(ValueError, "uncommitted successor"):
            self.service.store.claim_handoff(
                self.namespace,
                self.project_id,
                new_handoff_id,
                consumer_work_order_id="WO-v057-auditor",
                expected_to_role="independent_audit",
                expected_types={"production_to_audit"},
            )
        retry = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=self.recovery_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=[self.old_handoff_id],
            action="execute",
            expected_plan_sha256=recovery_plan["plan_sha256"],
            confirmation=service_module.HANDOFF_RECOVERY_CONFIRMATION,
            request_session_id=self.session_id,
        )
        self.assertEqual(
            retry.data["mappings"][0]["new_handoff_id"], new_handoff_id
        )
        committed = self.service.store.get_project_record(
            self.namespace,
            self.project_id,
            kind="handoffs",
            record_id=new_handoff_id,
        )
        self.assertTrue(
            self.service.store.handoff_is_published(
                self.namespace, self.project_id, committed
            )
        )
        replay = self.service.recover_storyboard_handoff(
            self.namespace,
            project_id=self.project_id,
            work_order_id=self.recovery_work_order_id,
            gap_id=self.gap_id,
            handoff_ids=[self.old_handoff_id],
            action="execute",
            expected_plan_sha256=recovery_plan["plan_sha256"],
            confirmation=service_module.HANDOFF_RECOVERY_CONFIRMATION,
            request_session_id=self.session_id,
        )
        replay_mapping = replay.data["mappings"][0]
        retry_mapping = retry.data["mappings"][0]
        for field in (
            "new_handoff_id",
            "recovery_id",
            "trace_id",
            "workbook_projection_mode",
            "workbook_projection_proof_sha256",
        ):
            self.assertEqual(replay_mapping[field], retry_mapping[field])
        self.assertTrue(replay_mapping["idempotent_replay"])


if __name__ == "__main__":
    unittest.main()
