"""Actual source/plan/render checks for thought, report and ordered dialogue."""
import copy
import hashlib
import json
import io
from pathlib import Path
import sys
import subprocess
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'server'),str(ROOT/'tests'),str(ROOT/'skills/shortdrama-storyboard/scripts')]
from test_v0527_production_decisions import production_fixture,action,voice_spec
from test_v0527_source_decisions import fact,build
from production_plan_test_support import frozen_plan
from production_decisions import lower_production_decisions,ProductionDecisionError
from source_decisions import SourceDecisionError
from author_storyboard import _render_decisions
from production_derivation import attach_derivation,inspect_derivation
from field_content import field_content_findings
from validate_storyboard import _voice_template_errors
from voice_sequence import build_voice_sequence


def fixture(kind='inner_voice'):
    data=production_fixture()
    text='甲心想：明日再议。' if kind=='inner_voice' else '乙听完仆从禀报，得知茶水洒了。'
    actor='甲' if kind=='inner_voice' else '乙'
    f=fact('F1',text,subject=actor,channel='thought' if kind=='inner_voice' else 'report',
           content='明日再议。' if kind=='inner_voice' else '',availability='given' if kind=='inner_voice' else 'unstated',occurrence='occurred')
    f['expression_status']='asserted' if kind=='inner_voice' else 'not_applicable'
    graph={'contract':'source-facts-v1','sources':{'DOC':hashlib.sha256(text.encode()).hexdigest()},'facts':[f]}
    # Keep the established contract identifier authoritative.
    from source_decisions import GRAPH_CONTRACT,DECISIONS_CONTRACT
    graph['contract']=GRAPH_CONTRACT
    use={'id':'U1','fact_ids':['F1'],'mode':'inner_voice' if kind=='inner_voice' else 'objective_narration',
         'wording':'verbatim' if kind=='inner_voice' else 'authored','actor':actor if kind=='inner_voice' else None,
         'text':'明日再议。' if kind=='inner_voice' else text}
    decisions={'contract':DECISIONS_CONTRACT,'uses':[use],'written_objects':[]}
    evidence=data['project']['prevention_evidence'];evidence['source_decisions']=decisions
    evidence['source_ledger']=[{'source_id':'DOC','source_text':text}]
    shot=data['shots'][0];shot['source_unit_ids']=['DOC'];shot['editorial']['原文支撑句']=text
    a=action(kind if kind=='inner_voice' else 'narrate',id='audio',source_use_id='U1');a.pop('actor',None)
    shot['actions']=[a]
    return data,graph,{'DOC':text}


def compile_render(data,graph,texts):
    c=build(graph,data['project']['prevention_evidence']['source_decisions'],texts)
    lower=lower_production_decisions(data,source_compilation=c)
    return lower,_render_decisions(lower['typed_decisions'])


class AudioBridgeTests(unittest.TestCase):
    def test_isolated_server_export_uses_same_ordered_speaker_parser(self):
        rows=[('甲','明日再议。'),('乙','明日再议。')]
        block=build_voice_sequence([(a,t,voice_spec()) for a,t in rows])
        payload={'画面内容':'甲说道：“明日再议。”；乙说道：“明日再议。”','音色':block}
        code="""
import json,sys
sys.path.insert(0,sys.argv[1])
from shortdrama_mcp.field_content import field_content_findings
shot=json.load(sys.stdin)
assert 'validate_storyboard' not in sys.modules
assert field_content_findings(shot)==[]
shot['音色']=shot['音色'].replace('〔2:乙〕','〔2:甲〕',1)
assert field_content_findings(shot)
assert 'validate_storyboard' not in sys.modules
"""
        result=subprocess.run([sys.executable,'-I','-X','utf8','-c',code,str(ROOT/'server')],
                              input=json.dumps(payload,ensure_ascii=False),encoding='utf-8',capture_output=True)
        self.assertEqual(result.returncode,0,result.stderr)

    def test_full_episode_thought_reaches_real_validator_and_xlsx(self):
        from test_v0535_production_contract import complete_episode,generate
        from validate_storyboard import validate_project
        from shortdrama_mcp.file_io import storyboard_to_xlsx,extract_file
        from shortdrama_mcp.materializer import MOTHER_TEMPLATE_COLUMNS
        from openpyxl import load_workbook
        data,graph=complete_episode();e=data['project']['prevention_evidence']
        text='甲心想：编号十九可以开始。'
        graph['sources']['DOC19']=hashlib.sha256(text.encode()).hexdigest()
        f=next(f for f in graph['facts'] if f['id']=='F20');f['channel']='thought'
        f['supports']=[{'source_id':'DOC19','start':0,'end':len(text),'text':text}]
        use=next(u for u in e['source_decisions']['uses'] if u['id']=='U20');use['mode']='inner_voice'
        source=next(s for s in e['source_ledger'] if s['source_id']=='DOC19')
        source.update(source_text=text,source_type='心理')
        shot=data['shots'][18];shot['actions'][0]['kind']='inner_voice'
        shot['actions'][0].pop('voice')
        shot['editorial']['原文支撑句']=text
        beat=next(b for b in e['episode_contract']['episodes'][0]['beat_events'] if b['id']=='B19')
        beat['delivery']['field']='心理音轨'
        with tempfile.TemporaryDirectory() as folder:
            canonical=generate(folder,data,graph)
        report=validate_project(canonical)
        self.assertEqual(report['issues'],[],report['issues'])
        row=canonical['shots'][18]
        self.assertIn('甲内心独白：“编号十九可以开始。”',row['心理音轨'])
        self.assertNotIn('编号十九可以开始',row['画面内容'])
        raw=storyboard_to_xlsx(canonical,report,assume_compact=True)
        wb=load_workbook(io.BytesIO(raw));master=wb['正式分镜母版']
        self.assertEqual([c.value for c in master[1]],list(MOTHER_TEMPLATE_COLUMNS))
        self.assertIn('甲内心独白：“编号十九可以开始。”',str(master.cell(20,5).value))
        self.assertTrue(extract_file(raw,'synthetic-thought.xlsx'))
        beat['delivery']['field']='画面内容'
        with tempfile.TemporaryDirectory() as folder:
            wrong_field=generate(folder,data,graph)
        self.assertTrue(any(i['error_code']=='E-EPISODE-BEAT-001' for i in validate_project(wrong_field)['issues']))

    def test_thought_is_exact_psychological_audio_without_lips_or_narrator_truth(self):
        data,g,t=fixture();g['production_plan']=frozen_plan(data,g,t)
        u=g['production_plan']['knowledge']['uses'][0]
        u.update(stance='suspected',evidence_fact_ids=['F1']);g['production_plan']['knowledge']['initial']=[]
        lower,canonical=compile_render(data,g,t);row=canonical['shots'][0]
        self.assertIn('甲内心独白：“明日再议。”',row['心理音轨'])
        self.assertNotIn('明日再议',row['画面内容']);self.assertNotIn('口型与对白一致',row['画面内容'])
        timing=canonical['prevention_evidence']['audio_workflow']['timing']['utterances'][0]
        self.assertEqual((timing['kind'],timing['speaker']),('psychological','甲'))
        self.assertEqual(lower['derivation']['production_prevention']['semantic_status'],'pending_independent_review')
        raw=json.dumps(g,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()
        attached=attach_derivation(canonical,data,raw,production=lower)
        proof=inspect_derivation(attached)
        self.assertTrue(proof['valid'],proof)

    def test_thought_cannot_change_source_or_speaker_or_be_audible_dialogue(self):
        for change in ['text','speaker','oral','invented','narrator']:
            data,g,t=fixture();use=data['project']['prevention_evidence']['source_decisions']['uses'][0]
            if change=='text':use['text']='今日就议。'
            if change=='speaker':use['actor']='乙'
            if change=='oral':use['mode']='spoken'
            if change=='invented':g['facts'][0]['content_text']=use['text']='根本没有这句。'
            if change=='narrator':
                g['production_plan']=frozen_plan(data,g,t);g['production_plan']['knowledge']['uses'][0]['stance']='narrator'
            with self.subTest(change=change),self.assertRaises((SourceDecisionError,ProductionDecisionError)):
                compile_render(data,g,t)

    def report(self):
        data,g,t=fixture('report');data['initial_world']['actors']['乙']=copy.deepcopy(data['initial_world']['actors']['甲'])
        row=data['shots'][0];row['ticks']=[0,10,30,40];row['actions'][0].update(start=0,end=30)
        row['actions'].append(action('hold',id='reaction',actor='乙',start=30,end=38,after=['audio']))
        g['production_plan']=frozen_plan(data,g,t);k=g['production_plan']['knowledge'];k['initial']=[]
        k['uses'][-1].update(fact_ids=['F1'],stance='reported',evidence_fact_ids=['F1'])
        k['acquisitions']=[{'id':'receipt','actor':'乙','fact_ids':['F1'],'evidence_fact_ids':['F1'],'action_id':'audio','mode':'source_report'}]
        return data,g,t

    def test_exact_report_can_establish_receipt_but_never_verified_truth(self):
        data,g,t=self.report();compile_render(data,g,t)
        g['production_plan']['knowledge']['uses'][-1]['stance']='known'
        with self.assertRaisesRegex(ProductionDecisionError,'before its acquisition'):compile_render(data,g,t)

    def test_report_rejects_wrong_recipient_future_nonreport_and_authored_words(self):
        for change in ['recipient','future','channel','narration','before','heard_thought']:
            data,g,t=self.report()
            if change=='recipient':g['facts'][0]['subject']='甲'
            if change=='future':g['facts'][0]['occurrence']='future'
            if change=='channel':g['facts'][0]['channel']='narration'
            if change=='narration':data['project']['prevention_evidence']['source_decisions']['uses'][0]['text']='乙听到全部真相。'
            if change=='before':
                data['shots'][0]['actions'][-1].update(start=25,after=[])
                g['production_plan']['decisions']=frozen_plan(data,g,t)['decisions']
            if change=='heard_thought':g['production_plan']['knowledge']['acquisitions'][0]['mode']='heard'
            with self.subTest(change=change),self.assertRaises((SourceDecisionError,ProductionDecisionError)):compile_render(data,g,t)

    def test_inner_voice_does_not_allow_a_listener_to_hear_it(self):
        data,g,t=fixture();data['initial_world']['actors']['乙']=copy.deepcopy(data['initial_world']['actors']['甲'])
        g['production_plan']=frozen_plan(data,g,t)
        g['production_plan']['knowledge']['acquisitions']=[{'id':'hear','actor':'乙','fact_ids':['F1'],'evidence_fact_ids':['F1'],'action_id':'audio','mode':'heard'}]
        with self.assertRaisesRegex(ProductionDecisionError,'actual spoken event'):compile_render(data,g,t)

    def test_ordered_voices_are_checked_per_speaker_by_field_ownership(self):
        rows=[('甲','明日再议。'),('乙','明日再议。')];voice=voice_spec()
        block=build_voice_sequence([(a,text,voice) for a,text in rows])
        self.assertEqual(field_content_findings({'音色':block},dialogue_rows=rows),[])
        self.assertTrue(field_content_findings({'音色':block.replace('平静','全景',1)},dialogue_rows=rows))
        for tag in ['〔0:甲〕','〔1:乙〕','〔1:甲','〔1:甲〕〔bad〕']:
            bad=block.replace('〔1:甲〕',tag,1)
            self.assertTrue(_voice_template_errors(bad,''.join(x[1] for x in rows),rows))

    def test_multiple_actual_speakers_reach_the_same_typed_renderer(self):
        data,g,t=fixture();text='甲说：明日再议。乙说：明日再议。';t={'DOC':text}
        g['sources']={'DOC':hashlib.sha256(text.encode()).hexdigest()};g['facts']=[]
        uses=data['project']['prevention_evidence']['source_decisions']['uses'];uses.clear()
        row=data['shots'][0];row['ticks']=[0,20,40,60];row['actions']=[]
        data['project']['prevention_evidence']['source_ledger']=[{'source_id':'DOC','source_text':text}]
        for i,name in enumerate(['甲','乙']):
            left=i*8;right=left+8;f=fact('F'+str(i+1),text,subject=name,event='E'+str(i+1))
            f['supports']=[{'source_id':'DOC','start':left,'end':right,'text':text[left:right]}];g['facts'].append(f)
            uses.append({'id':'U'+str(i+1),'fact_ids':[f['id']],'mode':'spoken','wording':'verbatim','actor':name,'text':'明日再议。'})
            data['initial_world']['actors'][name]=copy.deepcopy(data['initial_world']['actors']['甲'])
            row['actions'].append(action('speak',id='say'+str(i+1),source_use_id='U'+str(i+1),voice=voice_spec(),start=i*30,end=i*30+20))
        lower,canonical=compile_render(data,g,t)
        self.assertIn('〔2:乙〕',canonical['shots'][0]['音色'])
        self.assertEqual([u['speaker'] for u in canonical['prevention_evidence']['audio_workflow']['timing']['utterances']],['甲','乙'])


if __name__=='__main__':unittest.main()
