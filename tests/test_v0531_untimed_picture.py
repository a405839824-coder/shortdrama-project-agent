"""User-visible clocks are removed without deleting dialogue or timing checks."""
import copy
import io
from pathlib import Path
import sys
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "skills/shortdrama-storyboard/scripts"))
from picture_presentation import (
    FIELD, PictureTimingError, clock_annotations, display_picture,
    separate_picture_clock, separate_shot_clock, timed_picture,
)
from timing_model import validate_timing
from timing_events import validate_events
from performance_checks import eye_contract_faults, listener_reaction_faults
from validate_storyboard import project_formal_master, validate_project
from shortdrama_mcp.materializer import compile_formal_merged_field
from shortdrama_mcp.file_io import storyboard_to_xlsx, extract_xlsx
from openpyxl import load_workbook
from test_v0527_authoring_serializer import decisions, render_decisions, speaking_decisions
from test_v051_template_projection import compact_storyboard


class UntimedPictureTests(unittest.TestCase):
    def test_clock_spellings_are_removed_and_roundtrip_exactly(self):
        for clock in ("【0.0—0.4秒】", "[0.0-0.4s]", "【0—1秒】  ",
                      "于0.4—3.4秒", "在第0.2秒", "0.2秒", "约半秒", "t=0.2s",
                      "0.0–0.4秒：", "【0.0—0.4秒】： ", "0.2秒: "):
            with self.subTest(clock=clock):
                text = clock + "甲抬眼。"
                shot = separate_shot_clock({"画面内容": text})
                self.assertEqual(shot["画面内容"], "甲抬眼。")
                self.assertEqual(timed_picture(shot), text)

    def test_bounded_visibility_remains_limited_and_reversible(self):
        text = "甲眼睛仅在0.0—1.2秒清楚可见。"
        shot = separate_shot_clock({"画面内容": text})
        self.assertEqual(shot["画面内容"], "甲眼睛一度清楚可见。")
        self.assertEqual(timed_picture(shot), text)

    def test_relative_order_is_retained_without_seconds(self):
        text = "三秒后，甲看向门口。"
        shot = separate_shot_clock({"画面内容": text})
        self.assertEqual(shot["画面内容"], "随后，甲看向门口。")
        self.assertEqual(timed_picture(shot), text)

    def test_protected_literal_times_and_story_dates_are_unchanged(self):
        text = '【0.0—0.4秒】甲说道：“等我三秒，明日再来。”纸上写着「0—3秒」。三日后赴约。道具编号props3s。'
        shot = separate_shot_clock({"画面内容": text})
        self.assertEqual(shot["画面内容"], text.removeprefix("【0.0—0.4秒】"))
        self.assertEqual(timed_picture(shot), text)
        self.assertEqual(clock_annotations(shot["画面内容"]), [])

    def test_new_author_output_is_plain_and_input_is_not_mutated(self):
        data = decisions(); before = copy.deepcopy(data)
        candidate = render_decisions(data)
        shot = candidate["shots"][0]
        self.assertNotIn("秒】", shot["画面内容"])
        self.assertIn(FIELD, shot)
        self.assertEqual(shot["时长"], 4.0)
        self.assertEqual(shot["累计时长"], 4.0)
        self.assertEqual(len(shot[FIELD]["annotations"]), 3)
        self.assertEqual(data, before)
        self.assertEqual(validate_timing(candidate, candidate["shots"])["faults"], [])

    def test_separation_is_idempotent(self):
        shot = separate_shot_clock({"画面内容": "【0.0—0.4秒】甲于0.2秒抬眼。"})
        self.assertEqual(separate_shot_clock(shot), shot)

    def test_stale_clock_blocks_display_and_timing_validation(self):
        candidate = render_decisions(decisions())
        candidate["shots"][0]["画面内容"] += "甲转身。"
        with self.assertRaisesRegex(PictureTimingError, "binding"):
            display_picture(candidate["shots"][0])
        self.assertTrue(any("binding" in issue["message"]
                            for issue in validate_timing(candidate, candidate["shots"])["faults"]))

    def test_direct_performance_consumers_read_clocks_and_reject_stale_bindings(self):
        from test_v0527_dynamic_eye_decisions import compile_shot
        result, raw = compile_shot()
        plain = separate_shot_clock(raw)
        project = {"prevention_evidence": {"audio_workflow": {"timing": {"events": result["events"]}}}}
        before = copy.deepcopy(plain)
        self.assertEqual(eye_contract_faults(plain, duration=4),
                         eye_contract_faults(raw, duration=4))
        self.assertEqual(validate_events(project, [plain], []),
                         validate_events(project, [raw], []))
        self.assertEqual(plain, before)
        plain["画面内容"] += "甲抬手。"
        self.assertTrue(any("binding" in item for item in eye_contract_faults(plain, duration=4)))
        self.assertTrue(any("binding" in item for item in listener_reaction_faults(
            plain, duration=4, dialogue_rows=[])))
        self.assertTrue(any("binding" in item["message"]
                            for item in validate_events(project, [plain], [])["faults"]))

    def test_annotation_cannot_insert_a_new_action(self):
        shot = separate_shot_clock({"画面内容": "【0.0—0.4秒】甲抬眼。"})
        shot[FIELD]["annotations"][0]["token"] = "【0.0—0.4秒】乙夺走信封。"
        with self.assertRaisesRegex(PictureTimingError, "token"):
            timed_picture(shot)

    def test_annotation_cannot_enter_a_quoted_payload(self):
        shot = separate_shot_clock({"画面内容": '甲说道：“随后再来。”'})
        shot[FIELD]["annotations"] = [{"offset": shot["画面内容"].index("随后"), "token": "3秒后", "kind": "clock"}]
        with self.assertRaisesRegex(PictureTimingError, "roundtrip"):
            timed_picture(shot)

    def test_wrong_offsets_and_duplicate_keys_are_not_accepted(self):
        for change in ({"offset": True}, {"offset": -1}, {"offset": 900}, {"extra": "通过"}):
            shot = separate_shot_clock({"画面内容": "【0.0—0.4秒】甲抬眼。"})
            shot[FIELD]["annotations"][0].update(change)
            with self.subTest(change=change), self.assertRaises(PictureTimingError):
                timed_picture(shot)

    def test_clock_only_is_not_a_valid_picture(self):
        with self.assertRaises(PictureTimingError):
            separate_picture_clock("【0.0—0.4秒】")

    def test_internal_phase_capacity_still_rejects_overlong_speech(self):
        candidate = render_decisions(speaking_decisions())
        self.assertEqual(validate_timing(candidate, candidate["shots"])["faults"], [])
        # A longer line still has only the original 0.7-second speech phase.
        timed = timed_picture(candidate["shots"][0]).replace("放这里。", "这封信暂时放在这里，谁都不能拿走。")
        candidate["shots"][0].pop(FIELD)
        candidate["shots"][0]["画面内容"] = timed
        candidate["shots"][0] = separate_shot_clock(candidate["shots"][0])
        candidate["prevention_evidence"]["audio_workflow"]["timing"]["utterances"] = []
        result = validate_timing(candidate, candidate["shots"])
        self.assertTrue(any(issue["code"] == "E-DIALOGUE-003" for issue in result["faults"]))

    def test_both_formal_renderers_use_same_plain_picture(self):
        shot = compact_storyboard()["shots"][0]
        shot["画面内容"] = '【0.0—0.4秒】苏清柳抬眼。\n【0.4—2.6秒】苏清柳于0.4—2.6秒说道：“等我三秒。”\n【2.6—3.0秒】苏清柳站稳。'
        packed = separate_shot_clock(shot)
        expected = display_picture(packed)
        self.assertEqual(compile_formal_merged_field(packed), project_formal_master(packed))
        self.assertIn("画面内容：" + expected, compile_formal_merged_field(shot))
        self.assertIn("等我三秒。", expected)
        self.assertNotIn("秒】", expected)

    def test_known_cross_reference_variants_are_blocked_at_formal_renderers(self):
        for text in ("承接上一个镜头，苏清柳站定。", "沿用上个镜头的人物站位。", "保持上一个关键帧。"):
            for renderer in (compile_formal_merged_field, project_formal_master):
                shot = compact_storyboard()["shots"][0]; shot["起始状态"] = text
                with self.subTest(text=text, renderer=renderer.__name__), self.assertRaisesRegex(ValueError, "context_violation"):
                    renderer(shot)

    def test_literal_dialogue_is_preserved_but_quoted_instructions_still_fail(self):
        text = '苏清柳说道：“上一个镜头要保留。”'
        for renderer in (compile_formal_merged_field, project_formal_master):
            shot = compact_storyboard()["shots"][0]
            shot["画面内容"] = text
            with self.subTest(renderer=renderer.__name__):
                self.assertIn(text, renderer(shot))
                shot["起始状态"] = '站位要求：“沿用上个镜头的人物站位。”'
                with self.assertRaisesRegex(ValueError, "context_violation"):
                    renderer(shot)
        project = compact_storyboard()
        project["shots"][0]["画面内容"] = text
        self.assertFalse(any(issue["error_code"] == "E-I2V-INDEP-001"
                             for issue in validate_project(project)["issues"]))
        project["shots"][0]["画面内容"] += '拍摄要求：“沿用上个镜头的位置。”'
        self.assertTrue(any(issue["error_code"] == "E-I2V-INDEP-001"
                            for issue in validate_project(project)["issues"]))

    def test_plain_and_legacy_timelines_have_same_validation_codes(self):
        project = compact_storyboard()
        before = validate_project(project)
        plain = copy.deepcopy(project)
        plain["shots"] = [separate_shot_clock(shot) for shot in plain["shots"]]
        after = validate_project(plain)
        self.assertEqual([(x["error_code"], x["location"]) for x in before["issues"]],
                         [(x["error_code"], x["location"]) for x in after["issues"]])

    def test_serialized_xlsx_has_plain_picture_and_preserves_timing_columns(self):
        project = compact_storyboard()
        project["shots"] = [separate_shot_clock(shot) for shot in project["shots"]]
        payload = storyboard_to_xlsx(project, {"status": "fail", "issues": []})
        book = load_workbook(io.BytesIO(payload), data_only=False)
        sheet = book["正式分镜母版"]
        formal = sheet.cell(2, 5).value
        picture = formal.split("画面内容：", 1)[1].split("\n\n音色：", 1)[0]
        self.assertEqual(clock_annotations(picture), [])
        self.assertEqual(float(sheet.cell(2, 2).value), float(project["shots"][0]["时长"]))
        self.assertEqual(float(sheet.cell(2, 3).value), float(project["shots"][0]["累计时长"]))
        book.close()

    def test_current_xlsx_sidecar_preserves_internal_clocks_on_readback(self):
        project = compact_storyboard()
        project["shots"] = [separate_shot_clock(shot) for shot in project["shots"]]
        payload = storyboard_to_xlsx(project, {"status": "fail", "issues": []})
        extracted = extract_xlsx(payload)
        self.assertTrue(extracted["canonical_sidecar_present"])
        restored = extracted["canonical_storyboard"]
        for expected, actual in zip(project["shots"], restored["shots"], strict=True):
            self.assertEqual(actual["画面内容"], expected["画面内容"])
            self.assertEqual(actual[FIELD], expected[FIELD])
            self.assertEqual(timed_picture(actual), timed_picture(expected))


if __name__ == "__main__":
    unittest.main()
