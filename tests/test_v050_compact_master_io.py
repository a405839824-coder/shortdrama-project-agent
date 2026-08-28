from __future__ import annotations

import io
import re
import unittest
from unittest import mock

from openpyxl import load_workbook

import shortdrama_mcp.file_io as file_io_module
import shortdrama_mcp.materializer as materializer_module
from shortdrama_mcp.file_io import (
    extract_xlsx,
    storyboard_to_xlsx,
    workbook_to_storyboard,
)
from shortdrama_mcp.materializer import (
    COMPACT_SHOT_FIELDS,
    FORMAL_MERGED_FIELD,
    GPT_KEYFRAME_FIELD,
    MOTHER_TEMPLATE_COLUMNS,
    formal_mother_template_rows,
    materialize_storyboard,
)


def compact_storyboard() -> dict:
    return {
        "project_version": "V16.3-148-R10-StableRuntime",
        "schema_version": "R10.1",
        "workflow_profile": "R10-PG1",
        "prevention_evidence": {},
        "source_lock": {},
        "field_order": list(COMPACT_SHOT_FIELDS),
        "total_duration": 3.6,
        "shots": [
            {
                "镜头编号": "EP06-004",
                "时长": 3.6,
                "累计时长": 12.2,
                "功能类型": "礼数致歉",
                "场景连续性ID": "E06-S01-XIANGYA-MAIN-HALL-MORNING",
                "镜头类型": "中近景",
                "可见主体": "苏清柳；邹氏肩缘；荣婶儿",
                "空间站位描述": "苏清柳在中央偏左，邹氏在右前景，荣婶儿在左后两步。",
                "起始状态": "苏清柳微欠身，双手已经离开茶盏。",
                "机位": "主位左前方平视65mm，焦点在苏清柳。",
                "运镜方式": "极缓慢推近苏清柳，聚焦礼数致歉。",
                "画面内容": (
                    "0.0–0.8秒：苏清柳先看邹氏，再垂眸整理袖口；"
                    "0.8–2.9秒：【苏清柳】母亲，儿媳失手，请见谅。"
                    "她神色平静，口型与对白一致；"
                    "2.9–3.6秒：她重新抬眼，邹氏呼吸收紧，镜尾不抢下一动作。"
                ),
                "音色": "平静克制、温和带锋，声量中等。",
                "心理音轨": "不适用：本镜无心理音轨。",
                "环境音效": "衣料、呼吸与瓷盏细响，远处晨鸟压低。",
                "环境动态": "同一清晨私宅礼仪空间，仅出现有源家人与仆役。",
                "光影": "晨光向画面右侧衰减，面部不过曝。",
                "光源来源": "左侧木格窗自然晨光与右前浅墙反射。",
                "光源方向": "左侧向右侧衰减，右前弱补。",
                "色温倾向": "中性微暖晨光，低饱和木色。",
                "生成限制": "9:16真人古装短剧；禁止字幕、水印、额外配乐。",
                "转场类型": "视线与动作切",
                "切镜/转场提示词": "人物动作与呼吸稳定后切。",
                "参考图锚定": "CHAR-SU-QINGLIU",
                "覆盖原文": "是",
                "原文依据类型": "直接依据",
                "原文支撑句": "苏清柳敛袖赔礼。",
                "shot_uid": "shot-ep06-004",
                "episode_id": "EP06",
                "scene_uid": "scene-xiangya-main",
            }
        ],
    }


class CompactMaterializerTests(unittest.TestCase):
    def test_canonical_stores_atoms_once(self) -> None:
        source = compact_storyboard()

        compact = materialize_storyboard(source)

        self.assertEqual(compact["schema_version"], "R10.1")
        self.assertNotIn("batches", compact)
        self.assertEqual(compact["field_order"], COMPACT_SHOT_FIELDS)
        for field in (
            FORMAL_MERGED_FIELD,
            GPT_KEYFRAME_FIELD,
            "Seedance 2.5独立正式执行分镜",
            "生成批次",
        ):
            self.assertNotIn(field, compact["shots"][0])

    def test_r101_rejects_legacy_derived_and_batch_fields(self) -> None:
        source = compact_storyboard()
        source["batches"] = []
        source["shots"][0][FORMAL_MERGED_FIELD] = "不得保留"
        source["shots"][0][GPT_KEYFRAME_FIELD] = "不得保留"
        source["shots"][0]["Seedance 2.5独立正式执行分镜"] = "不得保留"
        source["shots"][0]["人物调度"] = "不得保留"

        with self.assertRaisesRegex(ValueError, "compact_contract_violation"):
            materialize_storyboard(source)

        label_source = compact_storyboard()
        label_source["shots"][0]["画面内容"] += "；动作与表演：旧字段标签。"
        with self.assertRaisesRegex(ValueError, "deprecated-label"):
            materialize_storyboard(label_source)

        movement_label_source = compact_storyboard()
        movement_label_source["shots"][0]["运镜方式"] += (
            "；运镜主体：苏清柳；叙事功能：聚焦致歉"
        )
        with self.assertRaisesRegex(ValueError, "deprecated-label"):
            materialize_storyboard(movement_label_source)

        unknown_source = compact_storyboard()
        unknown_source["shots"][0]["后台缓存字段"] = "不得进入canonical"
        with self.assertRaisesRegex(ValueError, "unknown"):
            materialize_storyboard(unknown_source)

        nested_light_source = compact_storyboard()
        nested_light_source["shots"][0]["光影"] = "光源来源：窗光；面部不过曝。"
        with self.assertRaisesRegex(ValueError, "nested-light-label"):
            materialize_storyboard(nested_light_source)

        underivable_legacy = compact_storyboard()
        underivable_legacy["schema_version"] = "R10.0"
        underivable_legacy["shots"][0]["光影"] = ""
        underivable_legacy["shots"][0]["光源方向"] = "左侧向右侧"
        with self.assertRaisesRegex(ValueError, "cannot derive 光影"):
            materialize_storyboard(underivable_legacy)

        unknown_legacy = compact_storyboard()
        unknown_legacy["schema_version"] = "R10.0"
        unknown_legacy["shots"][0]["未映射剧情事实"] = "不得静默丢失"
        with self.assertRaisesRegex(ValueError, "require explicit mapping"):
            materialize_storyboard(unknown_legacy)

        direct_unknown = compact_storyboard()["shots"][0]
        direct_unknown["未映射剧情事实"] = "不得静默丢失"
        with self.assertRaisesRegex(ValueError, "require explicit mapping"):
            materializer_module.migrate_legacy_shot(direct_unknown, legacy=True)

        direct_residual_label = compact_storyboard()["shots"][0]
        direct_residual_label["环境动态"] += "；运镜主体：旧结构。"
        with self.assertRaisesRegex(ValueError, "retired structural labels"):
            materializer_module.migrate_legacy_shot(
                direct_residual_label,
                legacy=True,
            )

        misplaced_label = compact_storyboard()
        misplaced_label["shots"][0]["环境动态"] += "；叙事功能：旧标签。"
        with self.assertRaisesRegex(ValueError, "deprecated-label"):
            materialize_storyboard(misplaced_label)

    def test_projection_has_exact_columns_and_visual_only_gpt(self) -> None:
        row = formal_mother_template_rows(compact_storyboard())[0]

        self.assertEqual(list(row), MOTHER_TEMPLATE_COLUMNS)
        self.assertTrue(
            row[FORMAL_MERGED_FIELD].startswith(
                "镜头序号：EP06-004（3.6s；累计12.2s）\n"
                "功能类型：礼数致歉\n"
                "场景连续性ID：E06-S01-XIANGYA-MAIN-HALL-MORNING"
            )
        )
        self.assertIn("光源来源：", row[FORMAL_MERGED_FIELD])
        self.assertIn("光源方向：", row[FORMAL_MERGED_FIELD])
        self.assertIn("色温倾向：", row[FORMAL_MERGED_FIELD])
        self.assertIn("母亲，儿媳失手，请见谅", row[FORMAL_MERGED_FIELD])
        gpt = row[GPT_KEYFRAME_FIELD]
        self.assertNotIn("母亲，儿媳失手，请见谅", gpt)
        self.assertIsNone(re.search(r"\d+(?:\.\d+)?\s*[–—-]\s*\d", gpt))
        self.assertNotIn("音色：", gpt)
        self.assertNotIn("环境音效：", gpt)

    def test_xlsx_has_no_legacy_or_batch_sheet_and_round_trips_sidecar(self) -> None:
        payload = storyboard_to_xlsx(compact_storyboard(), {"issues": []})
        workbook = load_workbook(io.BytesIO(payload), data_only=False)

        visible = [
            sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"
        ]
        self.assertEqual(visible, ["00_范围与审计", "正式生产母版"])
        self.assertNotIn("分镜", workbook.sheetnames)
        self.assertNotIn("生成批次", workbook.sheetnames)
        headers = [cell.value for cell in workbook["正式生产母版"][1]]
        self.assertEqual(headers, MOTHER_TEMPLATE_COLUMNS)
        self.assertEqual(
            [workbook["正式生产母版"].column_dimensions[letter].width for letter in "ABCDEFGHIJ"],
            [16, 9, 11, 23.5, 62, 60, 42, 20, 13, 46],
        )

        canonical = extract_xlsx(payload)["canonical_storyboard"]
        self.assertEqual(canonical["schema_version"], "R10.1")
        self.assertNotIn("batches", canonical)
        self.assertEqual(canonical["field_order"], COMPACT_SHOT_FIELDS)
        self.assertNotIn(FORMAL_MERGED_FIELD, canonical["shots"][0])
        self.assertNotIn(GPT_KEYFRAME_FIELD, canonical["shots"][0])

    def test_source_evidence_with_retired_words_round_trips_verbatim(self) -> None:
        source = compact_storyboard()
        support = "原文写道：“台词：母亲……”人物调度：仍按礼位。"
        anchor = "参考原文标签‘人物调度：’，不得改写"
        source["shots"][0]["原文支撑句"] = support
        source["shots"][0]["参考图锚定"] = anchor

        compact = materialize_storyboard(source)
        self.assertEqual(compact["shots"][0]["原文支撑句"], support)
        self.assertEqual(compact["shots"][0]["参考图锚定"], anchor)
        row = formal_mother_template_rows(compact)[0]
        self.assertEqual(row["原文支撑句"], support)
        self.assertEqual(row["参考图锚定"], anchor)

        payload = storyboard_to_xlsx(compact, {"issues": []}, assume_compact=True)
        restored = extract_xlsx(payload)["canonical_storyboard"]
        self.assertEqual(restored["shots"][0]["原文支撑句"], support)
        self.assertEqual(restored["shots"][0]["参考图锚定"], anchor)

        legacy = compact_storyboard()
        legacy["schema_version"] = "R10.0"
        legacy["shots"][0]["原文支撑句"] = support
        self.assertEqual(
            materialize_storyboard(legacy)["shots"][0]["原文支撑句"],
            support,
        )

    def test_legacy_atoms_collapse_without_reappearing(self) -> None:
        legacy = compact_storyboard()
        legacy["schema_version"] = "R10.0"
        shot = legacy["shots"][0]
        shot.pop("空间站位描述")
        shot.pop("起始状态")
        shot.pop("切镜/转场提示词")
        shot["人物调度"] = "人物调度：苏清柳中央偏左，邹氏右前景。"
        shot["运镜主体"] = "苏清柳"
        shot["运镜功能"] = "聚焦礼数致歉"
        shot["动作与表演"] = "动作与表演：苏清柳整理袖口后抬眼。"
        shot["台词"] = "台词：【苏清柳】母亲，儿媳失手，请见谅。"
        shot["发声方式"] = "发声方式：开口前呼吸稳定。"
        shot["标点语气"] = "标点语气：逗号短停，句号收住。"
        shot["心理音轨属性"] = "心理音轨属性：画外音：否。"
        shot["切镜转场提示词"] = "人物动作与呼吸稳定后切。"
        legacy["batches"] = [{"批次编号": "B001"}]

        compact = materialize_storyboard(legacy)
        migrated = compact["shots"][0]

        self.assertEqual(migrated["空间站位描述"], "苏清柳中央偏左，邹氏右前景。")
        self.assertIn("苏清柳", migrated["运镜方式"])
        self.assertIn("聚焦礼数致歉", migrated["运镜方式"])
        self.assertNotIn("运镜主体：", migrated["运镜方式"])
        self.assertNotIn("运镜功能：", migrated["运镜方式"])
        self.assertNotIn("叙事功能：", migrated["运镜方式"])
        self.assertTrue(migrated["起始状态"])
        self.assertIn("苏清柳中央偏左", migrated["起始状态"])
        self.assertIn("苏清柳整理袖口后抬眼", migrated["起始状态"])
        self.assertIn("苏清柳整理袖口后抬眼", migrated["画面内容"])
        self.assertIn("母亲，儿媳失手，请见谅", migrated["画面内容"])
        self.assertIn("画外音：否", migrated["心理音轨"])
        self.assertEqual(migrated["切镜/转场提示词"], shot["切镜转场提示词"])
        for legacy_field in (
            "人物调度",
            "动作与表演",
            "台词",
            "发声方式",
            "标点语气",
            "心理音轨属性",
            "切镜转场提示词",
        ):
            self.assertNotIn(legacy_field, migrated)
        rendered = formal_mother_template_rows(compact)[0][FORMAL_MERGED_FIELD]
        for label in (
            "运镜主体：",
            "运镜功能：",
            "叙事功能：",
            "人物调度：",
            "动作与表演：",
            "台词：",
            "发声方式：",
            "标点语气：",
            "心理音轨属性：",
        ):
            self.assertNotIn(label, rendered)

    def test_legacy_nested_light_atoms_split_without_formal_duplication(self) -> None:
        legacy = compact_storyboard()
        legacy["schema_version"] = "R10.0"
        shot = legacy["shots"][0]
        shot["光源来源"] = ""
        shot["光源方向"] = ""
        shot["色温倾向"] = ""
        shot["光影"] = (
            "光源来源：左侧木格窗晨光；"
            "光源方向：左向右衰减，人物面部不过曝；"
            "色温倾向：中性微暖"
        )

        compact = materialize_storyboard(legacy)
        migrated = compact["shots"][0]
        self.assertEqual(migrated["光源来源"], "左侧木格窗晨光")
        self.assertEqual(
            migrated["光源方向"],
            "左向右衰减，人物面部不过曝",
        )
        self.assertEqual(migrated["色温倾向"], "中性微暖")
        self.assertEqual(migrated["光影"], "人物面部不过曝")
        formal = formal_mother_template_rows(compact)[0][FORMAL_MERGED_FIELD]
        self.assertEqual(formal.count("光源来源："), 1)
        self.assertEqual(formal.count("光源方向："), 1)
        self.assertEqual(formal.count("色温倾向："), 1)

    def test_legacy_pipe_light_atoms_split_without_duplication(self) -> None:
        legacy = compact_storyboard()
        legacy["schema_version"] = "R10.0"
        shot = legacy["shots"][0]
        shot["光源来源"] = ""
        shot["光源方向"] = ""
        shot["色温倾向"] = ""
        shot["光影"] = (
            "人物面部保持自然明暗层次｜光源来源：左侧木格窗晨光｜"
            "光源方向：左向右衰减｜色温倾向：中性微暖"
        )

        compact = materialize_storyboard(legacy)
        migrated = compact["shots"][0]
        self.assertEqual(migrated["光影"], "人物面部保持自然明暗层次")
        self.assertEqual(migrated["光源来源"], "左侧木格窗晨光")
        self.assertEqual(migrated["光源方向"], "左向右衰减")
        self.assertEqual(migrated["色温倾向"], "中性微暖")
        formal = formal_mother_template_rows(compact)[0][FORMAL_MERGED_FIELD]
        self.assertEqual(formal.count("光源来源："), 1)
        self.assertEqual(formal.count("光源方向："), 1)
        self.assertEqual(formal.count("色温倾向："), 1)

    def test_xlsx_normalizes_once_and_projects_each_output_once(self) -> None:
        with (
            mock.patch.object(
                file_io_module,
                "migrate_storyboard_to_compact",
                wraps=materializer_module.migrate_storyboard_to_compact,
            ) as migrate,
            mock.patch.object(
                materializer_module,
                "migrate_legacy_shot",
                wraps=materializer_module.migrate_legacy_shot,
            ) as migrate_shot,
            mock.patch.object(
                materializer_module,
                "compile_formal_merged_field",
                wraps=materializer_module.compile_formal_merged_field,
            ) as formal,
            mock.patch.object(
                materializer_module,
                "compile_gpt_keyframe_instruction",
                wraps=materializer_module.compile_gpt_keyframe_instruction,
            ) as gpt,
        ):
            storyboard_to_xlsx(compact_storyboard(), {"issues": []})

        self.assertEqual(migrate.call_count, 1)
        self.assertEqual(migrate_shot.call_count, 1)
        self.assertEqual(formal.call_count, 1)
        self.assertEqual(gpt.call_count, 1)

        with mock.patch.object(
            file_io_module,
            "migrate_storyboard_to_compact",
            wraps=materializer_module.migrate_storyboard_to_compact,
        ) as already_compact_migrate:
            storyboard_to_xlsx(
                compact_storyboard(),
                {"issues": []},
                assume_compact=True,
            )
        self.assertEqual(already_compact_migrate.call_count, 0)

    def test_uploaded_sample_style_formal_cell_migrates_without_sidecar(self) -> None:
        merged = "\n".join(
            (
                "镜头序号：EP05-001（3.0s；累计3.0s）",
                "镜头类型：三人中近景",
                "可见主体：邹氏、江氏、许氏",
                "机位：门内侧平视50mm",
                "运镜方式：固定镜头｜运镜主体：邹氏、江氏、许氏｜叙事功能：固定空间关系",
                "人物调度：邹氏在中右主位，江氏在左侧客座。",
                "画面内容：",
                "0.0–0.8秒：院外脚步接近，江氏侧向邹氏。",
                "0.8–3.0秒：江氏保持视线落点。",
                "动作与表演：江氏偏头，嘴角带试探浅笑。",
                "台词：【江氏】大嫂，没来晚吧？",
                "音色：表面亲近、暗含试探。",
                "发声方式：开口前轻吸气，声量中等偏轻。",
                "标点语气：逗号短停，问号克制上扬。",
                "心理音轨：不适用。",
                "心理音轨属性：画外音：否。",
                "环境音效：脚步与衣料轻响。",
                "环境动态：窗纸轻动。",
                "光源：木格窗外清晨自然天光。",
                "光线方向：画面右上斜入，左侧弱补，人物面部保持自然明暗层次。",
                "色温：自然中性微暖。",
                "生成限制：禁止字幕、水印。",
                "转场类型：正切近景。",
            )
        )
        sheets = [
            {
                "name": "E05_正式分镜",
                "headers": list(MOTHER_TEMPLATE_COLUMNS),
                "rows": [
                    {
                        "分镜序号": "EP05-001",
                        "时长": 3.0,
                        "累计时长": 3.0,
                        "功能类型/场景连续性ID": (
                            "功能类型：跨集钩子兑现\n场景连续性ID：EP05-S01"
                        ),
                        FORMAL_MERGED_FIELD: merged,
                        GPT_KEYFRAME_FIELD: "旧派生字段不得导入canonical",
                        "参考图锚定": "ANCHOR",
                        "覆盖原文": "CH006-E01",
                        "原文依据类型": "直接对白",
                        "原文支撑句": "【江氏】大嫂，我们没来晚吧？",
                    }
                ],
            }
        ]

        canonical = workbook_to_storyboard(sheets)
        self.assertIsNotNone(canonical)
        shot = canonical["shots"][0]
        self.assertEqual(canonical["schema_version"], "R10.1")
        self.assertEqual(shot["功能类型"], "跨集钩子兑现")
        self.assertEqual(shot["场景连续性ID"], "EP05-S01")
        self.assertIn("邹氏在中右主位", shot["空间站位描述"])
        self.assertTrue(shot["起始状态"])
        self.assertIn("【江氏】大嫂，没来晚吧？", shot["画面内容"])
        self.assertIn("0.0–0.8秒", shot["画面内容"])
        self.assertEqual(shot["光源来源"], "木格窗外清晨自然天光。")
        self.assertEqual(
            shot["光源方向"],
            "画面右上斜入，左侧弱补，人物面部保持自然明暗层次。",
        )
        self.assertEqual(shot["色温倾向"], "自然中性微暖。")
        self.assertNotIn("光源来源：", shot["光影"])
        self.assertNotIn("光源方向：", shot["光影"])
        self.assertNotIn("色温倾向：", shot["光影"])
        self.assertEqual(shot["切镜/转场提示词"], "按正切近景执行。")
        self.assertNotIn(GPT_KEYFRAME_FIELD, shot)
        projected = formal_mother_template_rows(canonical)[0]
        self.assertTrue(
            projected[FORMAL_MERGED_FIELD].startswith(
                "镜头序号：EP05-001（3.0s；累计3.0s）"
            )
        )
        gpt = projected[GPT_KEYFRAME_FIELD]
        self.assertNotIn("大嫂，没来晚吧", gpt)
        self.assertNotIn("说话时", gpt)
        self.assertNotIn("声音强弱", gpt)
        self.assertNotIn("尾音变化", gpt)
        self.assertNotIn("问号尾音", gpt)
        self.assertNotIn("”", gpt)
        formal = projected[FORMAL_MERGED_FIELD]
        self.assertEqual(formal.count("光源来源："), 1)
        self.assertEqual(formal.count("光源方向："), 1)
        self.assertEqual(formal.count("色温倾向："), 1)
        self.assertNotIn("运镜主体：", formal)
        self.assertNotIn("运镜功能：", formal)
        self.assertNotIn("叙事功能：", formal)

    def test_gpt_scrubs_audio_dialogue_and_timing_from_every_visual_field(self) -> None:
        source = compact_storyboard()
        shot = source["shots"][0]
        shot["空间站位描述"] = "左前主门开启，远处脚步声压低。"
        shot["画面内容"] = (
            "[0.0–0.8s]苏清柳抬眼；【苏清柳】“母亲，请见谅。”；"
            "她手指收住，口型与对白一致；"
            "肩背随呼吸收紧，开口前呼吸状态稳定；"
            "她侧向邹氏试探开口。"
        )
        shot["环境动态"] = (
            "窗纸轻动，远处晨鸟声压低；帘角摆动并伴随脚步声；"
            "画内匾额“香雅院”清晰可见。"
        )
        shot["参考图锚定"] = "【参考图锚点】绛红衣饰固定；声线固定。"
        shot["生成限制"] = (
            "禁止字幕、水印、额外配乐、人物看镜头；"
            "保留木门与湿衣可见，环境音自然；"
            "仅保留角色对白、动作声与环境声。"
        )
        shot["运镜方式"] += "；问句前半平稳。"

        gpt = formal_mother_template_rows(source)[0][GPT_KEYFRAME_FIELD]

        for visible in (
            "左前主门开启",
            "苏清柳抬眼",
            "她手指收住",
            "肩背随呼吸收紧",
            "窗纸轻动",
            "帘角摆动",
            "【参考图锚点】绛红衣饰固定",
            "画内匾额“香雅院”清晰可见",
            "禁止字幕",
            "人物看镜头",
            "保留木门与湿衣可见",
        ):
            self.assertIn(visible, gpt)
        for removed in (
            "0.0–0.8",
            "母亲，请见谅",
            "口型",
            "对白",
            "脚步声",
            "晨鸟声",
            "声线",
            "配乐",
            "环境音",
            "呼吸状态",
            "问句前半平稳",
            "角色对白",
            "动作声",
            "环境声",
            "脚步",
            "开口",
        ):
            self.assertNotIn(removed, gpt)


if __name__ == "__main__":
    unittest.main()
