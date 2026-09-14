"""Construct authentic previous projections, not relabelled current R5 output."""
import io
from openpyxl import Workbook
from shortdrama_mcp import file_io, frozen_projector_0531
from shortdrama_mcp.materializer import MOTHER_TEMPLATE_COLUMNS


def use_previous_layout(book):
    for name in list(book.sheetnames):
        if name.endswith('_逐镜阅读') or name == '使用与来源':
            del book[name]
    if '正式分镜母版' in book.sheetnames:
        book['正式分镜母版'].title='正式生产母版'
    return book


def previous_workbook(storyboard, audit_report=None, **_options):
    book=Workbook();book.remove(book.active)
    audit=file_io._write_scope_audit_sheet(book,storyboard,audit_report or {'issues':[]})
    rows=frozen_projector_0531.formal_mother_template_rows(storyboard)
    file_io._write_table(book,'正式生产母版',list(MOTHER_TEMPLATE_COLUMNS),rows)
    file_io._write_canonical_sidecar(book,storyboard,rows,
                                   file_io._visible_sheet_sha256(audit.title,file_io._sheet_value_matrix(audit)))
    sidecar=book[file_io.CANONICAL_SIDECAR_SHEET]
    for row in range(2,sidecar.max_row+1):
        sidecar.cell(row,9).value=file_io.PREVIOUS_PROJECTION_CONTRACT
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()
