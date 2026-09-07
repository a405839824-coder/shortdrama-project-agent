from __future__ import annotations

import copy
import json
import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from typing import Any
from unittest.mock import patch

from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from shortdrama_mcp.app import create_mcp
from shortdrama_mcp.auth import request_namespace
from shortdrama_mcp.canonical import (
    COMPACT_SHOT_FIELDS,
    LEGACY_DERIVED_SHOT_FIELDS,
    LEGACY_REMOVED_SHOT_FIELDS,
)
from shortdrama_mcp.config import Settings
from shortdrama_mcp.delta import ensure_shot_uids
from shortdrama_mcp.materializer import materialize_storyboard
import shortdrama_mcp.service as service_module


PLUGIN_ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = PLUGIN_ROOT / "skills" / "shortdrama-storyboard"

import sys

sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from run_golden_tests import inject_pg1_fixture_evidence  # noqa: E402
from semantic_audit_helpers import (
    failing_semantic_audit_for_artifact,
    inject_semantic_audit,
)


def legacy_storyboard() -> dict[str, Any]:
    """Return the frozen R10.0 fixture only for explicit import migration tests."""

    suite = json.loads((SKILL_ROOT / "references" / "golden_tests.json").read_text(encoding="utf-8"))
    value = copy.deepcopy(suite["fixtures"]["valid_base"])
    inject_pg1_fixture_evidence(value)
    value["prevention_evidence"]["audit_layers"][5] = "正式母版单镜执行完整性"
    return value


def valid_storyboard() -> dict[str, Any]:
    """Return one R10.1 compact master with no persisted duplicate outputs."""

    value = materialize_storyboard(legacy_storyboard())
    value["schema_version"] = "R10.1"
    value["schema_profile"] = "compact_master_v1"
    value["field_order"] = list(COMPACT_SHOT_FIELDS)
    value.pop("batches", None)
    value.pop("mother_template_columns", None)
    value.setdefault("source_lock", {}).update(
        {
            "output_contract": "R10-SPC-1",
            "target_model": "Seedance 2.5",
            "aspect_ratio": "9:16",
        }
    )
    running = 0.0
    for index, shot in enumerate(value["shots"], 1):
        shot["镜头编号"] = f"EP01-{index:03d}"
        shot["episode_id"] = "EP01"
        shot["场景连续性ID"] = "EP01-SC01"
        running += float(shot.get("时长") or 0)
        shot["累计时长"] = running
        legacy_cut = shot.pop("切镜转场提示词", None)
        if not shot.get("切镜/转场提示词"):
            shot["切镜/转场提示词"] = legacy_cut or "手停在门框时切到桌面信件"
        shot["空间站位描述"] = (
            "同一室内清晨；女主位于左前门边，桌与信件在画面右中部，"
            "闭合木窗在左后，仅出现有源人物与道具。"
        )
        shot["起始状态"] = (
            "女主站在门外朝向桌面，右手已扶门框，视线尚未落到信件；"
            "木门半开，桌面信件未移动；女主双眼自然睁开，不处于半眨眼瞬间。"
        )
        shot["画面内容"] = (
            "【0.0—0.8秒】女主推门进入画面，右手扶住门框，视线从门框移向桌面信件，"
            "眼球先轻微换焦，头部随后轻跟，上眼睑略微放松，在0.4秒换焦后自然眨眼一次；"
            "【0.8—3.3秒】女主停在桌前，肩膀缓慢落下，嘴角自然放松，本镜无对白；"
            "【3.3—4.0秒】女主视线最终落在桌面信件上并保持稳定呼吸，手停在门框，动作收束。"
        )
        shot["音色"] = "无对白，不适用。"
        shot["心理音轨"] = (
            "不适用：本镜无心理音轨；不张嘴、不驱动口型、不生成字幕、不计入对白时长。"
        )
        shot["光影"] = (
            "女主面部与桌面信件保持自然明暗层次；高光不过曝，暗部保留皮肤、衣料和木构质感。"
            "空气清澈，背景结构清晰可辨，不泛白、不雾化。"
        )
        shot["光源来源"] = "左侧闭合木格窗透入自然晨光，右前浅色墙面提供弱反射补光。"
        shot["光源方向"] = "主光由画面左上向右下斜落，右前反射光弱补暗侧，向室内纵深逐级衰减。"
        shot["色温倾向"] = "中性微暖晨光；肤色自然，衣料、木构与纸张保持低饱和本色。"
        shot["生成限制"] = (
            "9:16竖屏；单镜单一主机位、单一核心动作阶段；"
            "眼神自然聚焦，保留轻微眼球焦点调整和眼睑松紧变化；"
            "禁止无依据全程不眨眼；禁止僵直瞪视；禁止高频机械眨眼；"
            "禁止字幕、水印、额外配乐、现代物件、人物看镜头、同步反应、复制脸、"
            "上下黑边、无来源光、重HDR、塑料皮肤与海报摆拍。"
        )
        for field in LEGACY_REMOVED_SHOT_FIELDS | LEGACY_DERIVED_SHOT_FIELDS:
            shot.pop(field, None)
        internal = {
            field: shot[field]
            for field in ("shot_uid", "episode_id", "scene_uid")
            if shot.get(field) not in (None, "")
        }
        compact = {field: shot.get(field) for field in COMPACT_SHOT_FIELDS}
        compact.update(internal)
        shot.clear()
        shot.update(compact)
    inject_pg1_fixture_evidence(value)
    value["prevention_evidence"]["audit_layers"][5] = "正式母版单镜执行完整性"
    return value


def _result_text(result: Any) -> str:
    """Return both structured and error tool output for rejection assertions."""

    values: list[str] = []
    structured = getattr(result, "structured_content", None)
    if structured is not None:
        values.append(json.dumps(structured, ensure_ascii=False, sort_keys=True))
    for item in getattr(result, "content", []) or []:
        text = getattr(item, "text", None)
        if text:
            values.append(str(text))
        else:
            values.append(str(item))
    return "\n".join(values)


class EfficiencyContractTests(unittest.IsolatedAsyncioTestCase):
    """No-loss latency contracts for the phase-1 and phase-2 runtime.

    These tests intentionally exercise public MCP behavior where possible.  The
    patch API tests are executable specifications for phase 2 and are expected
    to remain red until that tool is implemented.
    """

    async def asyncSetUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="shortdrama-efficiency-test-")
        os.environ["SHORTDRAMA_DATA_DIR"] = self.temp.name
        os.environ["PUBLIC_BASE_URL"] = "http://127.0.0.1:8765"
        os.environ["SHORTDRAMA_DEPLOYMENT_MODE"] = "local"
        os.environ["SHORTDRAMA_AUTH_MODE"] = "none"
        os.environ.pop("ALLOW_LOCAL_FILE_URLS", None)
        self.settings = Settings.from_env()
        self.server, self.service = create_mcp(self.settings)
        self.namespace = request_namespace(self.settings)

    async def asyncTearDown(self) -> None:
        self.temp.cleanup()

    async def call(self, name: str, arguments: dict[str, Any]) -> dict[str, Any]:
        arguments = inject_semantic_audit(self.service, self.settings, name, arguments)
        result = await self.server.call_tool(name, arguments)
        self.assertFalse(result.is_error, _result_text(result))
        self.assertIsNotNone(result.structured_content)
        return result.structured_content

    async def rejection_text(self, name: str, arguments: dict[str, Any]) -> str:
        """Normalize MCP ToolError and structured rejections for stable assertions."""

        arguments = inject_semantic_audit(self.service, self.settings, name, arguments)
        try:
            result = await self.server.call_tool(name, arguments)
        except Exception as exc:  # MCPServer raises ToolError for tool-side ValueError.
            return str(exc)
        return _result_text(result)

    async def create_project(self, name: str) -> str:
        result = await self.call("create_project", {"project_name": name})
        self.assertEqual(result["status"], "created")
        project_id = str(result["project_id"])
        checked = await self.call(
            "get_agent_status", {"project_id": project_id, "full_self_test": True}
        )
        self.assertTrue(checked["data"]["runtime_test_certificate_valid"])
        return project_id

    async def begin_storyboard_task(
        self,
        project_id: str,
        *,
        operation: str = "storyboard_create",
        mode: str = "execute",
        protected_scope: list[str] | None = None,
        window_key: str = "efficiency-test",
    ) -> str:
        result = await self.call(
            "begin_project_task",
            {
                "project_id": project_id,
                "operation": operation,
                "mode": mode,
                "authorization": "explicit_execute" if mode != "audit_only" else "read_only",
                "request_text": "效率优化等质回归测试",
                "targets": ["第一集"],
                "authoritative_source": "测试剧本",
                "latest_version": "v1",
                "target_boundary": "第一集",
                "must_keep": ["剧情、对白、人物认知与空间连续"],
                "must_remove": [],
                "protected_scope": protected_scope or [],
                "expected_outputs": ["正式分镜", "正式生产母版"],
                "window_key": window_key,
            },
        )
        return str(result["work_order_id"])

    def save_extracted_text(self, project_id: str, text: str, *, name: str = "长篇原文") -> dict[str, Any]:
        return self.service.store.save_artifact_json(
            self.namespace,
            project_id,
            {"kind": "text", "format": "txt", "text": text},
            file_name=f"{name}_提取内容.json",
            artifact_type="source_extraction",
            version="v1",
            metadata={"source_role": "novel"},
        )

    async def commit_fixture(
        self,
        project_id: str,
        work_order_id: str,
        storyboard: dict[str, Any],
        *,
        file_name: str = "效率测试分镜.xlsx",
    ) -> dict[str, Any]:
        return await self.call(
            "commit_storyboard",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "storyboard": storyboard,
                "file_name": file_name,
            },
        )

    async def test_bulk_text_ranges_are_lossless_and_non_overlapping(self) -> None:
        project_id = await self.create_project("批量读取无损测试")
        source = "".join(f"〔段落{i:04d}〕甲乙丙丁戊己庚辛壬癸。\n" for i in range(560))
        record = self.save_extracted_text(project_id, source)
        width = 4_000
        requested_ranges = [
            {
                "offset": offset,
                "limit": min(width, len(source) - offset),
                "unit": "characters",
            }
            for offset in range(0, len(source), width)
        ]

        result = await self.call(
            "read_project_artifact",
            {
                "project_id": project_id,
                "artifact_id": record["artifact_id"],
                "bulk_ranges": requested_ranges,
                "unit": "characters",
                "max_chars": 16_000,
            },
        )
        content = result["data"]["content"]
        returned = content["ranges"]
        self.assertEqual(len(returned), len(requested_ranges))
        self.assertEqual(content["returned_chars"], len(source))

        cursor = 0
        pieces: list[str] = []
        for requested, actual in zip(requested_ranges, returned, strict=True):
            self.assertEqual(actual["unit"], "characters")
            self.assertEqual(actual["offset"], requested["offset"])
            self.assertEqual(actual["offset"], cursor, "bulk ranges must not overlap or leave a gap")
            piece = actual["content"]
            pieces.append(piece)
            self.assertEqual(piece, source[actual["offset"] : actual["next_offset"]])
            cursor = actual["next_offset"]
        self.assertEqual(cursor, len(source))
        self.assertTrue(returned[-1]["complete"])
        self.assertEqual("".join(pieces), source)

    async def test_read_if_none_match_returns_no_payload(self) -> None:
        project_id = await self.create_project("条件读取测试")
        record = self.save_extracted_text(project_id, "原文内容" * 1_200)
        first = await self.call(
            "read_project_artifact",
            {
                "project_id": project_id,
                "artifact_id": record["artifact_id"],
                "offset": 0,
                "limit": 4_000,
                "unit": "characters",
            },
        )
        etag = first["data"]["etag"]
        self.assertEqual(etag, record["fingerprint"])

        with patch.object(
            self.service,
            "_load_extracted_or_payload",
            side_effect=AssertionError("matching if_none_match must not reload or parse the payload"),
        ):
            second = await self.call(
                "read_project_artifact",
                {
                    "project_id": project_id,
                    "artifact_id": record["artifact_id"],
                    "offset": 0,
                    "limit": 4_000,
                    "unit": "characters",
                    "if_none_match": etag,
                },
            )
        self.assertEqual(second["status"], "not_modified")
        self.assertTrue(second["data"]["not_modified"])
        self.assertEqual(second["data"]["etag"], etag)
        self.assertNotIn("content", second["data"])

    async def test_linked_canonical_json_is_preferred_over_xlsx_roundtrip(self) -> None:
        project_id = await self.create_project("Canonical优先测试")
        work_order_id = await self.begin_storyboard_task(project_id)
        committed = await self.commit_fixture(project_id, work_order_id, valid_storyboard())

        with patch(
            "shortdrama_mcp.service.extract_file",
            side_effect=AssertionError("linked canonical JSON must be loaded before parsing XLSX"),
        ):
            storyboard, record = self.service._storyboard_from_artifact(
                self.namespace,
                project_id,
                committed["artifact_id"],
            )

        self.assertEqual(record["artifact_id"], committed["artifact_id"])
        expected_storyboard, identity_manifest = ensure_shot_uids(valid_storyboard())
        expected_storyboard = materialize_storyboard(expected_storyboard)
        self.assertEqual(storyboard["shots"], expected_storyboard["shots"])
        self.assertNotIn("batches", storyboard)
        self.assertEqual(len(identity_manifest["shot_uids"]), len(storyboard["shots"]))
        self.assertTrue(all(shot.get("shot_uid") for shot in storyboard["shots"]))
        for shot in storyboard["shots"]:
            self.assertFalse(LEGACY_REMOVED_SHOT_FIELDS.intersection(shot))
            self.assertFalse(LEGACY_DERIVED_SHOT_FIELDS.intersection(shot))
        self.assertEqual(
            committed["data"]["canonical_sha256"],
            service_module.canonical_storyboard_sha256(storyboard),
        )

    async def test_production_preflight_and_independent_audit_have_separate_one_time_runs(self) -> None:
        project_id = await self.create_project("预检审计去重测试")
        work_order_id = await self.begin_storyboard_task(project_id)
        original_validator = service_module.validate_storyboard.validate_project

        with patch.object(
            service_module.validate_storyboard,
            "validate_project",
            wraps=original_validator,
        ) as validator:
            committed = await self.commit_fixture(project_id, work_order_id, valid_storyboard())
            self.assertEqual(validator.call_count, 1, "production preflight must validate once")

            first_audit = await self.call(
                "audit_storyboard",
                {"project_id": project_id, "artifact_id": committed["artifact_id"]},
            )
            self.assertEqual(first_audit["data"]["audit_class"], "independent_audit")
            self.assertFalse(first_audit["data"]["cache_hit"])
            self.assertEqual(validator.call_count, 2, "the first independent audit must make its own judgment")

            repeated_audit = await self.call(
                "audit_storyboard",
                {"project_id": project_id, "artifact_id": committed["artifact_id"]},
            )
            self.assertEqual(repeated_audit["data"]["audit_class"], "independent_audit")
            self.assertTrue(repeated_audit["data"]["cache_hit"])
            self.assertEqual(validator.call_count, 2, "same independent audit cache key must not validate twice")
            for audited in (first_audit, repeated_audit):
                self.assertEqual(
                    audited["data"]["repair_mode_recommendation"]["status"],
                    "no_repair",
                )

    async def test_audit_cache_key_invalidates_on_every_quality_input(self) -> None:
        base = {
            "canonical_sha256": "a" * 64,
            "source_fingerprint": "b" * 64,
            "boundary_hash": "c" * 64,
            "ruleset_version": "R10-rules-1",
            "validator_version": "validator-1",
            "audit_class": "independent_audit",
            "artifact_fingerprint": "1" * 64,
        }
        stable = service_module.build_audit_cache_key(**base)
        self.assertEqual(stable, service_module.build_audit_cache_key(**base))

        replacements = {
            "canonical_sha256": "d" * 64,
            "source_fingerprint": "e" * 64,
            "boundary_hash": "f" * 64,
            "ruleset_version": "R10-rules-2",
            "validator_version": "validator-2",
            "audit_class": "production_preflight",
            "artifact_fingerprint": "2" * 64,
        }
        invalidated = set()
        for component, changed_value in replacements.items():
            changed = dict(base)
            changed[component] = changed_value
            changed_key = service_module.build_audit_cache_key(**changed)
            self.assertNotEqual(stable, changed_key, f"{component} must invalidate the audit cache")
            invalidated.add(changed_key)
        self.assertEqual(len(invalidated), len(replacements))

    async def test_telemetry_never_contains_story_text_dialogue_url_or_subject(self) -> None:
        project_id = await self.create_project("遥测脱敏测试")
        sentinels = {
            "正文绝密标记-NEVER-LOG",
            "台词绝密标记-DIALOGUE-NEVER-LOG",
            "https://private.example.test/secret?token=NEVER-LOG",
            "auth0|private-user-subject-NEVER-LOG",
        }
        source = "｜".join(sorted(sentinels)) + "｜普通原文" * 400
        record = self.save_extracted_text(project_id, source)
        result = await self.call(
            "read_project_artifact",
            {
                "project_id": project_id,
                "artifact_id": record["artifact_id"],
                "offset": 0,
                "limit": 4_000,
                "unit": "characters",
            },
        )
        telemetry = result["data"]["telemetry"]
        serialized = json.dumps(telemetry, ensure_ascii=False, sort_keys=True)
        for sentinel in sentinels:
            self.assertNotIn(sentinel, serialized)

        forbidden_key_fragments = ("text", "content", "dialogue", "台词", "正文", "url", "subject")

        def check_keys(value: Any) -> None:
            if isinstance(value, dict):
                for key, item in value.items():
                    normalized = str(key).casefold()
                    self.assertFalse(
                        any(fragment in normalized for fragment in forbidden_key_fragments),
                        f"sensitive telemetry key is forbidden: {key}",
                    )
                    check_keys(item)
            elif isinstance(value, list):
                for item in value:
                    check_keys(item)

        check_keys(telemetry)
        self.assertIn("trace_id", telemetry)
        self.assertIn("duration_ms", telemetry)
        self.assertIn("stages", telemetry)

        # Commit/audit traces are persisted separately from the read response.
        # They must obey the same content-free contract even when sensitive
        # values participate in source locking and validation.
        work_order_id = await self.begin_storyboard_task(project_id, window_key="telemetry-persistence")
        storyboard = valid_storyboard()
        storyboard["source_lock"]["上集终点"] = source
        committed = await self.commit_fixture(project_id, work_order_id, storyboard)
        await self.call(
            "audit_storyboard",
            {"project_id": project_id, "artifact_id": committed["artifact_id"]},
        )
        persisted = self.service.store.list_project_records(
            self.namespace,
            project_id,
            kind="telemetry",
        )
        self.assertGreaterEqual(len(persisted), 2)
        for trace in persisted:
            trace_text = json.dumps(trace, ensure_ascii=False, sort_keys=True)
            for sentinel in sentinels:
                self.assertNotIn(sentinel, trace_text)
            check_keys(trace)

    async def test_artifacts_and_cache_records_are_isolated_by_project(self) -> None:
        project_a = await self.create_project("隔离项目A")
        project_b = await self.create_project("隔离项目B")
        record_a = self.save_extracted_text(project_a, "两个项目即使内容相同也不能越权读取")
        self.save_extracted_text(project_b, "两个项目即使内容相同也不能越权读取")

        cross_read = await self.rejection_text(
            "read_project_artifact",
            {
                "project_id": project_b,
                "artifact_id": record_a["artifact_id"],
                "if_none_match": record_a["fingerprint"],
            },
        )
        self.assertIn("unknown artifact_id", cross_read)

        work_order_a = await self.begin_storyboard_task(project_a, window_key="project-a")
        committed_a = await self.commit_fixture(project_a, work_order_a, valid_storyboard())
        cross_audit = await self.rejection_text(
            "audit_storyboard",
            {"project_id": project_b, "artifact_id": committed_a["artifact_id"]},
        )
        self.assertIn("unknown artifact_id", cross_audit)

    async def test_full_runtime_regression_certificate_is_reused_only_on_exact_fingerprints(self) -> None:
        project_id = await self.create_project("运行时回归证书测试")
        checked = await self.call(
            "get_agent_status",
            {"project_id": project_id, "full_self_test": True},
        )
        self.assertTrue(checked["data"]["runtime_test_certificate_valid"])
        certificate = checked["data"]["runtime_test_certificate"]
        self.assertEqual(certificate["certificate_type"], "runtime_full_regression")
        self.assertEqual(certificate["tests"]["golden_tests"]["test_pass_count"], 82)
        self.assertEqual(
            certificate["tests"]["strict_prompt_tests"]["test_pass_count"], 75
        )
        self.assertEqual(certificate["tests"]["public_scene_tests"]["test_pass_count"], 11)
        self.assertEqual(certificate["tests"]["agent_tests"]["test_pass_count"], 38)
        self.assertEqual(certificate["tests"]["project_agent_tests"]["test_pass_count"], 22)
        upgraded = certificate["tests"]["rule_upgrade_tests"]
        self.assertEqual(upgraded["status"], "pass")
        self.assertEqual(upgraded["test_fail_count"], 0)
        self.assertEqual(upgraded["test_count"], upgraded["test_pass_count"])
        self.assertGreater(upgraded["test_count"], 0)
        persisted_certificate = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="certificates",
            record_id=str(certificate["certificate_id"]),
        )
        self.assertEqual(persisted_certificate, certificate)
        self.assertTrue(
            self.service.store.project_record_is_signed(persisted_certificate)
        )
        self.assertNotIn(
            "runtime_test_certificates",
            self.service.store.load_hub(self.namespace, project_id),
        )
        self.assertNotIn(
            "runtime_test_certificates",
            self.service._project_schema["properties"],
        )
        self.assertNotIn(
            "runtime_test_certificate",
            self.service._project_schema["$defs"],
        )

        with patch.object(
            service_module.r10_bootstrap,
            "verify_bundled",
            side_effect=AssertionError("normal status must reuse the exact certificate, not rerun suites"),
        ):
            resumed = await self.call(
                "get_agent_status",
                {"project_id": project_id, "full_self_test": False},
            )
        self.assertTrue(resumed["data"]["runtime_test_certificate_valid"])
        self.assertEqual(
            resumed["data"]["runtime_test_certificate"]["certificate_id"],
            certificate["certificate_id"],
        )

    async def test_storyboard_signoff_requires_exact_clean_independent_audit(self) -> None:
        project_id = await self.create_project("独立终审签核门禁测试")
        work_order_id = await self.begin_storyboard_task(project_id)
        committed = await self.commit_fixture(project_id, work_order_id, valid_storyboard())

        blocked = await self.rejection_text(
            "signoff_task",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "decision": "approve",
                "signer": "测试签核人",
            },
        )
        self.assertIn("independent_audit_required", blocked)

        audited = await self.call(
            "audit_storyboard",
            {"project_id": project_id, "artifact_id": committed["artifact_id"]},
        )
        self.assertEqual(audited["status"], "pass")
        certificate = audited["data"]["certificate"]
        audit_summary = certificate["audit_summary"]
        self.assertEqual(audit_summary["residual_issue_count"], 0)
        self.assertEqual(audit_summary["introduced_issue_count"], 0)
        self.assertEqual(audit_summary["blocking_issue_count"], 0)
        self.assertEqual(audit_summary["final_status"], "pass")
        self.assertEqual(audit_summary["audit_trace_id"], audited["data"]["audit_trace_id"])
        self.assertEqual(
            audit_summary["canonical_sha256"], audited["data"]["canonical_sha256"]
        )
        self.assertRegex(audit_summary["artifact_raw_sha256"], r"^[0-9a-f]{64}$")
        record, _ = self.service.store.get_artifact(
            self.namespace, project_id, committed["artifact_id"]
        )
        self.assertEqual(record["metadata"]["independent_audit_status"], "PASS")
        self.assertEqual(record["metadata"]["audit_status"], "PASS")
        self.assertEqual(record["metadata"]["issue_count"], 0)
        audited_workbook = audited["data"]["audited_workbook"]
        self.assertEqual(
            record["metadata"]["independent_audit_overlay_artifact_id"],
            audited_workbook["artifact_id"],
        )
        overlay_record, overlay_path = self.service.store.get_artifact(
            self.namespace,
            project_id,
            audited_workbook["artifact_id"],
        )
        self.assertEqual(overlay_record["artifact_type"], "storyboard_audit_overlay")
        overlay_book = load_workbook(overlay_path, data_only=False, read_only=True)
        audit_sheet = overlay_book["00_范围与审计"]
        audit_text = "\n".join(
            str(cell.value)
            for row in audit_sheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        )
        overlay_book.close()
        self.assertIn("独立审计已签发", audit_text)
        self.assertIn("PASS", audit_text)
        self.assertIn(certificate["certificate_id"], audit_text)
        self.assertIn(certificate["audit_trace_id"], audit_text)
        self.assertIn(audit_summary["canonical_sha256"], audit_text)
        self.assertIn(audit_summary["artifact_raw_sha256"], audit_text)
        hub = self.service.store.load_hub(self.namespace, project_id)
        producer_order = next(
            row
            for row in hub["work_orders"]
            if row["work_order_id"] == work_order_id
        )
        self.assertEqual(producer_order["status"], "passed")
        self.assertEqual(
            producer_order["verified_independent_audit_certificate_id"],
            certificate["certificate_id"],
        )
        deliverable = next(
            row
            for row in hub["deliverables"]
            if row["deliverable_id"] == committed["artifact_id"]
        )
        self.assertEqual(deliverable["status"], "validated")
        self.assertEqual(
            hub["project_state"]["latest_artifacts"]["storyboard"]["deliverable_id"],
            audited_workbook["artifact_id"],
        )
        overlay_deliverable = next(
            row
            for row in hub["deliverables"]
            if row["deliverable_id"] == audited_workbook["artifact_id"]
        )
        self.assertEqual(overlay_deliverable["status"], "validated")
        signed = await self.call(
            "signoff_task",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "decision": "approve",
                "signer": "测试签核人",
            },
        )
        self.assertEqual(signed["status"], "archived")
        self.assertEqual(
            signed["data"]["independent_audit_certificate_id"],
            audited["data"]["certificate"]["certificate_id"],
        )

    async def _prepare_patch_base(
        self,
        *,
        project_name: str,
        protected_scope: list[str] | None = None,
    ) -> tuple[str, str, dict[str, Any], dict[str, Any]]:
        project_id = await self.create_project(project_name)
        create_work_order = await self.begin_storyboard_task(project_id, window_key=f"{project_name}-create")
        machine_clean = valid_storyboard()
        # This value remains structurally valid, so production preflight must
        # persist and hand off the artifact.  The left/right source conflict is
        # intentionally a semantic finding made only by the independent audit.
        machine_clean["shots"][0]["光源来源"] = "右侧木格窗透入的日光"
        committed = await self.commit_fixture(
            project_id,
            create_work_order,
            machine_clean,
            file_name="待修分镜.xlsx",
        )
        self.assertEqual(committed["status"], "production_preflight_pass")
        self.assertEqual(committed["data"]["report"]["issues"], [])
        self.assertIsNotNone(committed["artifact_id"])
        semantic_failure = failing_semantic_audit_for_artifact(
            self.service,
            self.settings,
            project_id,
            str(committed["artifact_id"]),
            shot_id="S001",
            field="光源来源",
            error_code="E-TEMPLATE-008",
            issue_id="shots[1].light_source_semantic_mismatch",
            evidence=(
                "独立审计核对场景锁后确认唯一自然光来自画面左侧木格窗；"
                "当前光源来源写为右侧，与光源方向和场景事实不一致。"
            ),
            repair_action="把光源来源修正为左侧木格窗透入的日光并完整回归。",
        )
        audited = await self.call(
            "audit_storyboard",
            {
                "project_id": project_id,
                "artifact_id": committed["artifact_id"],
                "semantic_audit": semantic_failure,
            },
        )
        self.assertEqual(audited["status"], "fail_current")
        self.assertIsNotNone(audited["data"]["repair_package"])
        repair_work_order = await self.begin_storyboard_task(
            project_id,
            operation="storyboard_repair",
            mode="audit_and_repair",
            protected_scope=protected_scope,
            window_key=f"{project_name}-repair",
        )
        return project_id, repair_work_order, committed, audited

    @staticmethod
    def _patch_arguments(
        project_id: str,
        work_order_id: str,
        committed: dict[str, Any],
        audited: dict[str, Any],
        *,
        expected_hash: str | None = None,
        field: str = "光源来源",
        value: str = "左侧木格窗透入的日光",
    ) -> dict[str, Any]:
        certificate = audited["data"]["certificate"]
        return {
            "project_id": project_id,
            "work_order_id": work_order_id,
            "base_artifact_id": committed["artifact_id"],
            "expected_canonical_hash": expected_hash or committed["data"]["canonical_sha256"],
            "audit_certificate_id": certificate["certificate_id"],
            "operations": [
                {
                    "op": "replace_field",
                    "shot_id": "EP01-001",
                    "field": field,
                    "value": value,
                    # Only the certificate row whose exact location is
                    # EP01-001.光源来源 may authorize this field replacement.
                    "issue_ids": ["E-TEMPLATE-008"],
                    "reason": "修复独立审计确认的光源语义错位，不改变剧情、认知或空间边界",
                }
            ],
            "impact_scope": ["EP01-001"],
            "file_name": "待修分镜_补丁修复版.xlsx",
        }

    async def test_storyboard_patch_rejects_stale_canonical_hash(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁哈希冲突测试"
        )
        arguments = self._patch_arguments(
            project_id,
            work_order_id,
            committed,
            audited,
            expected_hash="0" * 64,
        )
        result = await self.rejection_text("commit_storyboard_patch", arguments)
        self.assertIn("canonical_hash_conflict", result)

    async def test_storyboard_patch_cannot_modify_work_order_protected_scope(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁保护范围测试",
            protected_scope=["EP01-001"],
        )
        arguments = self._patch_arguments(project_id, work_order_id, committed, audited)
        result = await self.rejection_text("commit_storyboard_patch", arguments)
        self.assertIn("protected_scope_violation", result)

    async def test_storyboard_patch_rejects_invalid_audit_certificate(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁审计证书测试"
        )
        arguments = self._patch_arguments(project_id, work_order_id, committed, audited)
        arguments["audit_certificate_id"] = "CERT-does-not-exist"
        result = await self.rejection_text("commit_storyboard_patch", arguments)
        self.assertIn("invalid_audit_certificate", result)

    async def test_storyboard_patch_rejects_structural_or_boundary_field(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁结构字段测试"
        )
        arguments = self._patch_arguments(
            project_id,
            work_order_id,
            committed,
            audited,
            field="场景连续性ID",
            value="SC99",
        )
        result = await self.rejection_text("commit_storyboard_patch", arguments)
        self.assertIn("full_rebuild_required", result)

    async def test_storyboard_patch_runs_full_production_validation(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁完整生产校验测试"
        )
        arguments = self._patch_arguments(project_id, work_order_id, committed, audited)
        original_validator = service_module.validate_storyboard.validate_project
        with patch.object(
            service_module.validate_storyboard,
            "validate_project",
            wraps=original_validator,
        ) as validator:
            result = await self.call("commit_storyboard_patch", arguments)

        self.assertEqual(result["status"], "production_preflight_pass")
        self.assertTrue(result["data"]["full_production_validation"])
        # Compatibility alias for 0.3.x callers; not an independent-audit verdict.
        self.assertTrue(result["data"]["final_full_validation"])
        self.assertEqual(result["data"]["report"]["issues"], [])
        self.assertEqual(
            result["data"]["report"]["shot_count"],
            len(valid_storyboard()["shots"]),
        )
        self.assertGreaterEqual(validator.call_count, 1)
        validated_storyboard = validator.call_args.args[0]
        self.assertEqual(len(validated_storyboard["shots"]), len(valid_storyboard()["shots"]))
        self.assertEqual(
            validated_storyboard["shots"][0]["光源来源"],
            "左侧木格窗透入的日光",
        )

    async def test_repair_package_is_signed_consumed_and_bidirectionally_linked(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="不可变维修包闭环测试"
        )
        issued = audited["data"]["repair_package"]
        audit_to_repair = audited["data"]["audit_to_repair_handoff"]
        self.assertEqual(issued["status"], "issued")
        self.assertTrue(self.service.store.project_record_is_signed(issued))
        self.service._validate_schema_record(
            self.service._project_schema, "immutable_repair_package", issued
        )
        self.service._validate_schema_record(
            self.service._project_schema,
            "handoff_manifest",
            self.service.store.get_project_record(
                self.namespace,
                project_id,
                kind="handoffs",
                record_id=audit_to_repair["handoff_id"],
            ),
        )
        self.assertEqual(
            audit_to_repair["immutable_repair_package_id"],
            issued["repair_package_id"],
        )
        self.assertEqual(
            issued["audit_to_repair_handoff_id"], audit_to_repair["handoff_id"]
        )

        result = await self.call(
            "commit_storyboard_patch",
            self._patch_arguments(project_id, work_order_id, committed, audited),
        )
        consumed = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="repair_packages",
            record_id=issued["repair_package_id"],
        )
        repair_to_audit = result["data"]["handoff"]
        self.assertEqual(consumed["status"], "consumed")
        self.assertEqual(consumed["successor_artifact_id"], result["artifact_id"])
        self.assertEqual(
            consumed["repair_to_audit_handoff_id"], repair_to_audit["handoff_id"]
        )
        self.assertEqual(
            repair_to_audit["immutable_repair_package_id"],
            consumed["repair_package_id"],
        )
        self.service._validate_schema_record(
            self.service._project_schema, "immutable_repair_package", consumed
        )
        for field in ("consumed_by_work_order_id", "successor_artifact_id"):
            with self.subTest(empty_consumed_binding=field):
                invalid_consumed = copy.deepcopy(consumed)
                invalid_consumed[field] = ""
                with self.assertRaisesRegex(
                    ValueError,
                    rf"lifecycle_schema_mismatch: immutable_repair_package\.{field}",
                ):
                    self.service._validate_schema_record(
                        self.service._project_schema,
                        "immutable_repair_package",
                        invalid_consumed,
                    )
        self.service._validate_schema_record(
            self.service._project_schema,
            "handoff_manifest",
            self.service.store.get_project_record(
                self.namespace,
                project_id,
                kind="handoffs",
                record_id=repair_to_audit["handoff_id"],
            ),
        )
        hub = self.service.store.load_hub(self.namespace, project_id)
        Draft202012Validator(self.service._project_schema).validate(hub)
        hub_package = next(
            row
            for row in hub["repair_packages"]
            if row["repair_package_id"] == consumed["repair_package_id"]
        )
        self.assertEqual(hub_package["status"], "consumed")
        with self.assertRaisesRegex(ValueError, "repair_package_unavailable"):
            self.service.store.claim_repair_package(
                self.namespace,
                project_id,
                consumed["repair_package_id"],
                work_order_id=work_order_id,
            )

    async def test_consumed_repair_package_cannot_be_resurrected_after_deletion(
        self,
    ) -> None:
        project_id, work_order_id, committed, audited = (
            await self._prepare_patch_base(
                project_name="已消费维修包删除后禁止复活测试"
            )
        )
        result = await self.call(
            "commit_storyboard_patch",
            self._patch_arguments(
                project_id, work_order_id, committed, audited
            ),
        )
        package_id = str(
            audited["data"]["repair_package"]["repair_package_id"]
        )
        consumed = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="repair_packages",
            record_id=package_id,
        )
        self.assertEqual(consumed["status"], "consumed")
        self.assertEqual(
            consumed["successor_artifact_id"], result["artifact_id"]
        )
        project_dir = self.service.store.project_dir(
            self.namespace, project_id
        )
        package_path = (
            project_dir / "repair_packages" / f"{package_id}.json"
        )
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (hub, _ledger):
            hub["repair_packages"] = [
                row
                for row in hub.get("repair_packages", [])
                if row.get("repair_package_id") != package_id
            ]
        package_path.unlink()
        before = {
            str(path.relative_to(project_dir)): path.read_bytes()
            for path in sorted(project_dir.rglob("*"))
            if path.is_file()
        }
        certificate = audited["data"]["certificate"]

        # Bypass the service precheck to prove the storage-layer project-lock
        # CAS independently closes the same resurrection/TOCTOU window.
        with patch.object(
            self.service,
            "_repair_package_downstream_evidence",
            return_value=[],
        ):
            with self.assertRaisesRegex(
                ValueError,
                "consumed repair authorization cannot be resurrected",
            ):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=project_id,
                    artifact_id=str(committed["artifact_id"]),
                    audit_cycle_id=None,
                    work_order_id=str(certificate["work_order_id"]),
                    handoff_artifact_id=str(
                        committed["data"]["handoff"]["handoff_id"]
                    ),
                    force_machine_recheck=False,
                    semantic_audit=None,
                    request_session_id=None,
                )

        self.assertFalse(package_path.exists())
        self.assertEqual(
            sum(
                row.get("repair_package_id") == package_id
                for row in self.service.store.load_hub(
                    self.namespace, project_id
                ).get("repair_packages", [])
            ),
            0,
        )
        self.assertEqual(
            {
                str(path.relative_to(project_dir)): path.read_bytes()
                for path in sorted(project_dir.rglob("*"))
                if path.is_file()
            },
            before,
        )

    async def test_storyboard_patch_rejects_tampered_repair_package(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="维修包篡改阻断测试"
        )
        package_id = str(audited["data"]["repair_package"]["repair_package_id"])
        project_dir = self.service.store.project_dir(self.namespace, project_id)
        package_path = (
            project_dir
            / "repair_packages"
            / f"{package_id}.json"
        )
        tampered = self.service.store.read_json(package_path)
        tampered["authorized_scope"]["fields"] = ["任意字段"]
        self.service.store.atomic_write_json(package_path, tampered)
        tampered_bytes = package_path.read_bytes()
        before = {
            str(path.relative_to(project_dir)): path.read_bytes()
            for path in sorted(project_dir.rglob("*"))
            if path.is_file()
        }
        rejection = await self.rejection_text(
            "commit_storyboard_patch",
            self._patch_arguments(project_id, work_order_id, committed, audited),
        )
        self.assertIn("project_record_integrity_mismatch", rejection)
        self.assertIn(package_id, rejection)
        self.assertEqual(package_path.read_bytes(), tampered_bytes)
        self.assertEqual(
            {
                str(path.relative_to(project_dir)): path.read_bytes()
                for path in sorted(project_dir.rglob("*"))
                if path.is_file()
            },
            before,
        )

    async def test_failed_production_preflight_has_zero_publication_side_effects(self) -> None:
        project_id = await self.create_project("预检失败零发布副作用测试")
        work_order_id = await self.begin_storyboard_task(
            project_id,
            window_key="preflight-zero-side-effects-create",
        )
        invalid = valid_storyboard()
        invalid["shots"][0]["光源来源"] = ""
        before_artifacts = self.service.store.list_artifacts(
            self.namespace, project_id
        )
        before_hub = self.service.store.load_hub(self.namespace, project_id)
        original_export = service_module._export_compact_storyboard

        with patch.object(
            service_module,
            "_export_compact_storyboard",
            wraps=original_export,
        ) as exporter:
            rejected = await self.commit_fixture(
                project_id,
                work_order_id,
                invalid,
                file_name="不得生成的失败预检.xlsx",
            )

        self.assertEqual(rejected["status"], "fail_current")
        self.assertIsNone(rejected["artifact_id"])
        self.assertIsNone(rejected["download_url"])
        self.assertFalse(rejected["data"]["artifact_created"])
        self.assertFalse(rejected["data"]["handoff_created"])
        self.assertEqual(
            rejected["data"]["telemetry"]["output_payload_bytes"], 0
        )
        self.assertEqual(exporter.call_count, 0, "failed preflight must not export XLSX")

        after_artifacts = self.service.store.list_artifacts(
            self.namespace, project_id
        )
        after_hub = self.service.store.load_hub(self.namespace, project_id)
        self.assertEqual(
            [row["artifact_id"] for row in after_artifacts],
            [row["artifact_id"] for row in before_artifacts],
        )
        self.assertEqual(after_hub.get("deliverables", []), before_hub.get("deliverables", []))
        self.assertEqual(after_hub.get("handoffs", []), before_hub.get("handoffs", []))

        # A later valid commit must still allocate the first storyboard
        # version, proving the rejected attempt consumed no version number.
        # v0.5.5 binds a failed draft scope to its originating work order so
        # the correction cannot evade the learned prevention profile by
        # opening a fresh storyboard batch.
        committed = await self.commit_fixture(
            project_id,
            work_order_id,
            valid_storyboard(),
            file_name="首个有效分镜.xlsx",
        )
        self.assertEqual(committed["status"], "production_preflight_pass")
        record, _path = self.service.store.get_artifact(
            self.namespace,
            project_id,
            str(committed["artifact_id"]),
        )
        self.assertEqual(record["version"], "v1")

    async def test_production_preflight_never_marks_storyboard_work_order_passed(self) -> None:
        project_id = await self.create_project("生产预检不得终审测试")
        work_order_id = await self.begin_storyboard_task(project_id)
        committed = await self.commit_fixture(project_id, work_order_id, valid_storyboard())
        self.assertEqual(committed["status"], "production_preflight_pass")
        hub = self.service.store.load_hub(self.namespace, project_id)
        work_order = next(row for row in hub["work_orders"] if row["work_order_id"] == work_order_id)
        self.assertEqual(work_order["status"], "validating")
        self.assertTrue(work_order["awaiting_independent_audit"])
        self.assertFalse(committed["data"]["validation"]["final_verdict"])
        record, _ = self.service.store.get_artifact(
            self.namespace, project_id, committed["artifact_id"]
        )
        self.assertEqual(
            record["metadata"]["production_preflight_status"],
            "production_preflight_pass",
        )
        self.assertEqual(record["metadata"]["production_preflight_issue_count"], 0)
        self.assertEqual(record["metadata"]["independent_audit_status"], "PENDING_AUDIT")
        self.assertEqual(record["metadata"]["audit_status"], "PENDING_AUDIT")
        self.assertIsNone(record["metadata"]["issue_count"])
        certificate = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="certificates",
            record_id=record["metadata"]["production_certificate_id"],
        )
        self.assertEqual(
            certificate["preflight_status"], "production_preflight_pass"
        )
        self.assertEqual(certificate["preflight_issue_count"], 0)
        self.assertEqual(
            certificate["independent_audit_status"], "PENDING_AUDIT"
        )
        self.assertNotIn("result", certificate)
        self.assertNotIn("issue_count", certificate)
        handoff = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=record["metadata"]["handoff_id"],
        )
        self.assertEqual(
            handoff["production_preflight_status"],
            "production_preflight_pass",
        )
        self.assertEqual(handoff["independent_audit_status"], "PENDING_AUDIT")
        self.assertNotIn("production_result", handoff)
        self.assertNotIn("issue_count", handoff)
        deliverable = next(
            row
            for row in hub["deliverables"]
            if row["deliverable_id"] == committed["artifact_id"]
        )
        self.assertEqual(deliverable["status"], "draft")
        self.assertNotIn("storyboard", hub["project_state"]["latest_artifacts"])

    async def test_same_transport_session_cannot_self_sign_independent_audit(self) -> None:
        project_id = await self.create_project("同会话自签阻断测试")
        production = self.service.begin_task(
            self.namespace,
            project_id=project_id,
            operation="storyboard_create",
            mode="execute",
            authorization="explicit_execute",
            request_text="生产",
            targets=["第一集"],
            authoritative_source="测试原文",
            latest_version="v1",
            target_boundary="第一集",
            must_keep=["全部剧情"],
            must_remove=[],
            protected_scope=[],
            expected_outputs=["正式分镜"],
            window_key="production-window",
            request_session_id="transport-session-shared",
        )
        committed = self.service.commit_storyboard(
            self.namespace,
            project_id=project_id,
            work_order_id=str(production.work_order_id),
            storyboard=valid_storyboard(),
            file_name="同会话测试.xlsx",
            source_artifact_id=None,
            request_session_id="transport-session-shared",
        )
        hub_before = self.service.store.load_hub(self.namespace, project_id)
        with self.assertRaisesRegex(ValueError, "session_role_separation_required"):
            self.service.begin_task(
                self.namespace,
                project_id=project_id,
                operation="storyboard_audit",
                mode="audit_only",
                authorization="read_only",
                request_text="独立审计",
                targets=[str(committed.artifact_id)],
                authoritative_source="handoff",
                latest_version="current",
                target_boundary="第一集",
                must_keep=["只读"],
                must_remove=[],
                protected_scope=["全部源内容"],
                expected_outputs=["audit certificate"],
                window_key="audit-window",
                request_session_id="transport-session-shared",
            )
        self.assertEqual(
            self.service.store.load_hub(self.namespace, project_id), hub_before
        )

    async def test_storyboard_patch_blocks_shot_range_protected_scope(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁范围保护测试",
            protected_scope=["EP01-001—EP01-010"],
        )
        rejection = await self.rejection_text(
            "commit_storyboard_patch",
            self._patch_arguments(project_id, work_order_id, committed, audited),
        )
        self.assertIn("protected_scope_violation", rejection)

    async def test_storyboard_patch_rejects_removed_batch_topology_reference(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁拓扑字段测试"
        )
        arguments = self._patch_arguments(
            project_id,
            work_order_id,
            committed,
            audited,
            field="Seedance批次引用",
            value="BATCH-02",
        )
        rejection = await self.rejection_text("commit_storyboard_patch", arguments)
        self.assertIn("full_rebuild_required", rejection)

    async def test_storyboard_patch_issue_must_match_exact_shot_and_field(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁精确证据测试"
        )
        arguments = self._patch_arguments(project_id, work_order_id, committed, audited)
        arguments["operations"][0]["issue_ids"] = ["E-LIGHT-001"]
        rejection = await self.rejection_text("commit_storyboard_patch", arguments)
        self.assertIn("patch_evidence_mismatch", rejection)

    async def test_xlsx_byte_tampering_blocks_signoff(self) -> None:
        project_id = await self.create_project("XLSX字节篡改签核阻断测试")
        work_order_id = await self.begin_storyboard_task(project_id)
        committed = await self.commit_fixture(project_id, work_order_id, valid_storyboard())
        await self.call(
            "audit_storyboard", {"project_id": project_id, "artifact_id": committed["artifact_id"]}
        )
        _record, path = self.service.store.get_artifact(
            self.namespace, project_id, committed["artifact_id"]
        )
        path.write_bytes(path.read_bytes() + b"tampered")
        rejection = await self.rejection_text(
            "signoff_task",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "decision": "approve",
                "signer": "测试签核人",
            },
        )
        self.assertIn("artifact_raw_fingerprint_mismatch", rejection)

    async def test_certificate_report_tampering_blocks_signoff(self) -> None:
        project_id = await self.create_project("证书报告篡改签核阻断测试")
        work_order_id = await self.begin_storyboard_task(project_id)
        committed = await self.commit_fixture(project_id, work_order_id, valid_storyboard())
        audited = await self.call(
            "audit_storyboard", {"project_id": project_id, "artifact_id": committed["artifact_id"]}
        )
        certificate_id = audited["data"]["certificate"]["certificate_id"]
        certificate_path = (
            self.service.store.project_dir(self.namespace, project_id)
            / "certificates"
            / f"{certificate_id}.json"
        )
        certificate = json.loads(certificate_path.read_text(encoding="utf-8"))
        certificate["report"]["status"] = "tampered"
        self.service.store.atomic_write_json(certificate_path, certificate)
        rejection = await self.rejection_text(
            "signoff_task",
            {
                "project_id": project_id,
                "work_order_id": work_order_id,
                "decision": "approve",
                "signer": "测试签核人",
            },
        )
        self.assertIn("independent_audit_required", rejection)

    async def test_internal_storyboard_missing_handoff_blocks_audit(self) -> None:
        project_id = await self.create_project("内部产物缺交接阻断测试")
        work_order_id = await self.begin_storyboard_task(project_id)
        committed = await self.commit_fixture(project_id, work_order_id, valid_storyboard())
        handoff_id = committed["data"]["handoff"]["handoff_id"]
        handoff_path = (
            self.service.store.project_dir(self.namespace, project_id)
            / "handoffs"
            / f"{handoff_id}.json"
        )
        handoff_path.unlink()
        rejection = await self.rejection_text(
            "audit_storyboard", {"project_id": project_id, "artifact_id": committed["artifact_id"]}
        )
        self.assertIn("handoff_required", rejection)

    async def test_patch_successor_inherits_original_source_fingerprint(self) -> None:
        project_id, work_order_id, committed, audited = await self._prepare_patch_base(
            project_name="补丁来源链继承测试"
        )
        result = await self.call(
            "commit_storyboard_patch",
            self._patch_arguments(project_id, work_order_id, committed, audited),
        )
        base, _ = self.service.store.get_artifact(
            self.namespace, project_id, committed["artifact_id"]
        )
        successor, _ = self.service.store.get_artifact(
            self.namespace, project_id, result["artifact_id"]
        )
        self.assertEqual(
            successor["metadata"]["source_fingerprint"],
            base["metadata"]["source_fingerprint"],
        )
        self.assertNotEqual(
            successor["metadata"]["source_fingerprint"],
            committed["data"]["canonical_sha256"],
        )

    async def test_lifecycle_schemas_and_persisted_records_resolve_all_refs(self) -> None:
        Draft202012Validator.check_schema(self.service._project_schema)
        Draft202012Validator.check_schema(self.service._ledger_schema)
        self.assertIn("sha256", self.service._ledger_schema["$defs"])
        project_id = await self.create_project("生命周期Schema闭环测试")
        work_order_id = await self.begin_storyboard_task(project_id)
        committed = await self.commit_fixture(project_id, work_order_id, valid_storyboard())
        audited = await self.call(
            "audit_storyboard", {"project_id": project_id, "artifact_id": committed["artifact_id"]}
        )
        handoff = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=audited["data"]["certificate"]["handoff_id"],
        )
        certificate = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="certificates",
            record_id=audited["data"]["certificate"]["certificate_id"],
        )
        ledger = self.service.store.load_ledger(self.namespace, project_id)
        event = next(row for row in ledger["events"] if row["event_id"] == certificate["event_id"])
        trace = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="telemetry",
            record_id=certificate["stage_trace_id"],
        )
        self.service._validate_schema_record(self.service._project_schema, "handoff_manifest", handoff)
        self.service._validate_schema_record(
            self.service._ledger_schema, "independent_audit_certificate", certificate
        )
        self.service._validate_schema_record(self.service._ledger_schema, "audit_event", event)
        self.service._validate_schema_record(self.service._project_schema, "stage_timing_trace", trace)


if __name__ == "__main__":
    unittest.main(verbosity=2)
