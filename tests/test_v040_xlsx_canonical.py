from __future__ import annotations

import copy
import io
import json
import unittest

from openpyxl import Workbook, load_workbook

from shortdrama_mcp.canonical import compact_storyboard_canonical
from shortdrama_mcp.file_io import (
    CANONICAL_SIDECAR_PAYLOAD_PREFIX,
    CANONICAL_SIDECAR_SHEET,
    extract_xlsx,
    storyboard_to_xlsx,
)
from shortdrama_mcp.materializer import MOTHER_TEMPLATE_COLUMNS
from test_efficiency import legacy_storyboard, valid_storyboard


def _full_canonical_storyboard() -> dict:
    storyboard = compact_storyboard_canonical(valid_storyboard())
    storyboard["shots"][0].update(
        {
            "shot_uid": "shot-uid-ep06-001",
            "episode_id": "EP06",
            "scene_uid": "scene-uid-market-001",
            "画面内容": '=HYPERLINK("https://example.invalid","不执行")',
        }
    )
    storyboard.update(
        {
            "chapter_map": {
                "chapter_id": "CH06",
                "events": [
                    {"event_id": "EV06-01", "cause": "舒乔被跟踪", "result": "决定离开闹市"}
                ],
            },
            "source_mapping": [
                {
                    "source_unit_id": "SRC06-001",
                    "source_text": "她听见身后的脚步，借人群掩护转入巷口。",
                    "shot_uids": ["shot-uid-ep06-001"],
                }
            ],
            "protected_dialogue": [
                {
                    "dialogue_id": "D06-001",
                    "speaker": "舒乔",
                    "text": "有人跟着我。",
                    "mapped_shot_uids": ["shot-uid-ep06-001"],
                }
            ],
            "cognition_matrix": {
                "舒乔": {
                    "knows": ["有人跟踪"],
                    "does_not_know": ["跟踪者身份"],
                    "effective_from_shot_uid": "shot-uid-ep06-001",
                }
            },
            "scene_state": {
                "scene-uid-market-001": {
                    "start": {
                        "position": "布摊前",
                        "body_direction": "朝东",
                        "gaze_target": "铜镜倒影",
                        "held_prop": "药包",
                        "door_window_state": "不适用",
                        "action_progress": "尚未转身",
                        "light_direction": "右前方日光",
                    },
                    "end": {
                        "position": "巷口",
                        "body_direction": "朝北",
                        "gaze_target": "身后人流",
                        "held_prop": "药包",
                        "door_window_state": "不适用",
                        "action_progress": "进入巷口",
                        "light_direction": "右后方日光",
                    },
                }
            },
            "coverage": {
                "core_events": {"covered": 1, "total": 1, "rate": 1.0},
                "protected_dialogue": {"covered": 1, "total": 1, "rate": 1.0},
            },
            "source_ingestion": {
                "mode": "chapter_to_storyboard",
                "complete_read": True,
                "chunk_ids": ["CH06-P01", "CH06-P02"],
                "source_sha256": "a" * 64,
            },
            "anchors": {
                "character": {"舒乔": "青色窄袖衣、旧银簪"},
                "scene": {"scene-uid-market-001": "宋代石板街与连续木构铺面"},
                "prop": {"药包": "棕纸麻绳"},
            },
            "character_anchors": [{"anchor_id": "CHAR-SQ", "name": "舒乔"}],
            "scene_anchors": [{"anchor_id": "SC-MARKET", "scene_uid": "scene-uid-market-001"}],
            "prop_anchors": [{"anchor_id": "PROP-MED", "name": "药包"}],
            "costume_anchors": [{"anchor_id": "COSTUME-SQ-01", "description": "青色窄袖衣"}],
            "project_anchors": {"aspect_ratio": "9:16", "era": "古装"},
            "canonical_extension": {
                "nested": {"list": ["alpha", {"beta": 2}]},
                # Force more than one Excel-safe payload chunk.
                "large_text": "公共场景人物与环境声生态" * 8_000,
            },
        }
    )
    return storyboard


def _save_workbook(workbook) -> bytes:
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class CanonicalXlsxRoundtripTests(unittest.TestCase):
    def test_full_canonical_roundtrip_is_lossless_and_sidecar_is_invisible(self) -> None:
        storyboard = _full_canonical_storyboard()
        workbook_bytes = storyboard_to_xlsx(storyboard, {"issues": []})

        workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False, read_only=False)
        self.assertIn(CANONICAL_SIDECAR_SHEET, workbook.sheetnames)
        sidecar = workbook[CANONICAL_SIDECAR_SHEET]
        self.assertEqual(sidecar.sheet_state, "veryHidden")
        self.assertGreater(int(sidecar.cell(2, 7).value), 1)
        for row in range(3, sidecar.max_row + 1):
            payload = str(sidecar.cell(row, 8).value)
            self.assertTrue(payload.startswith(CANONICAL_SIDECAR_PAYLOAD_PREFIX))
            self.assertFalse(payload.startswith(("=", "+", "-", "@")))

        mother_headers = [cell.value for cell in workbook["正式分镜母版"][1]]
        self.assertEqual(mother_headers, MOTHER_TEMPLATE_COLUMNS)
        merged_value = str(workbook["正式分镜母版"].cell(2, 5).value)
        self.assertFalse(merged_value.startswith(("=", "+", "-", "@")))
        self.assertIn("画面内容：=HYPERLINK", merged_value)
        self.assertNotIn("分镜", workbook.sheetnames)
        self.assertNotIn("生成批次", workbook.sheetnames)

        extracted = extract_xlsx(workbook_bytes)
        self.assertNotIn(CANONICAL_SIDECAR_SHEET, [sheet["name"] for sheet in extracted["sheets"]])
        self.assertEqual(extracted["canonical_storyboard"], storyboard)

        roundtrip = extracted["canonical_storyboard"]
        for key in (
            "chapter_map",
            "source_mapping",
            "protected_dialogue",
            "cognition_matrix",
            "scene_state",
            "coverage",
            "source_ingestion",
            "anchors",
            "character_anchors",
            "scene_anchors",
            "prop_anchors",
            "costume_anchors",
            "project_anchors",
        ):
            self.assertEqual(roundtrip[key], storyboard[key], key)
        self.assertEqual(roundtrip["shots"][0]["shot_uid"], "shot-uid-ep06-001")
        self.assertEqual(roundtrip["shots"][0]["episode_id"], "EP06")
        self.assertEqual(roundtrip["shots"][0]["scene_uid"], "scene-uid-market-001")
        self.assertNotIn("batches", roundtrip)
        self.assertNotIn("正式分镜母版·字段合并", roundtrip["shots"][0])
        self.assertNotIn("GPT关键帧生成指令", roundtrip["shots"][0])

    def test_tampered_canonical_payload_is_rejected(self) -> None:
        workbook = load_workbook(
            io.BytesIO(storyboard_to_xlsx(_full_canonical_storyboard(), {"issues": []})),
            data_only=False,
            read_only=False,
        )
        sidecar = workbook[CANONICAL_SIDECAR_SHEET]
        payload = str(sidecar.cell(3, 8).value)
        position = len(CANONICAL_SIDECAR_PAYLOAD_PREFIX) + 17
        replacement = "A" if payload[position] != "A" else "B"
        sidecar.cell(3, 8).value = payload[:position] + replacement + payload[position + 1 :]

        with self.assertRaisesRegex(ValueError, "canonical payload SHA-256 mismatch"):
            extract_xlsx(_save_workbook(workbook))

    def test_missing_canonical_chunk_is_rejected_when_sidecar_exists(self) -> None:
        workbook = load_workbook(
            io.BytesIO(storyboard_to_xlsx(_full_canonical_storyboard(), {"issues": []})),
            data_only=False,
            read_only=False,
        )
        sidecar = workbook[CANONICAL_SIDECAR_SHEET]
        self.assertGreater(sidecar.max_row, 3)
        sidecar.delete_rows(sidecar.max_row, 1)

        with self.assertRaisesRegex(ValueError, "canonical payload chunk count mismatch"):
            extract_xlsx(_save_workbook(workbook))

    def test_workbook_without_sidecar_uses_legacy_parser(self) -> None:
        storyboard = copy.deepcopy(legacy_storyboard())
        storyboard["shots"][0].update(
            {"shot_uid": "legacy-shot-uid", "episode_id": "EP01", "scene_uid": "legacy-scene-uid"}
        )
        storyboard["batches"][0].update({"batch_uid": "legacy-batch-uid", "episode_id": "EP01"})
        workbook = Workbook()
        shot_sheet = workbook.active
        shot_sheet.title = "分镜"
        shot_headers = list(storyboard["shots"][0])
        shot_sheet.append(shot_headers)
        shot_sheet.append(
            [
                json.dumps(storyboard["shots"][0][key], ensure_ascii=False)
                if isinstance(storyboard["shots"][0][key], (dict, list))
                else storyboard["shots"][0][key]
                for key in shot_headers
            ]
        )
        batch_sheet = workbook.create_sheet("生成批次")
        batch_headers = list(storyboard["batches"][0])
        batch_sheet.append(batch_headers)
        batch_sheet.append(
            [
                json.dumps(storyboard["batches"][0][key], ensure_ascii=False)
                if isinstance(storyboard["batches"][0][key], (dict, list))
                else storyboard["batches"][0][key]
                for key in batch_headers
            ]
        )

        extracted = extract_xlsx(_save_workbook(workbook))
        canonical = extracted["canonical_storyboard"]
        self.assertIsNotNone(canonical)
        self.assertEqual(canonical["shots"][0]["shot_uid"], "legacy-shot-uid")
        self.assertEqual(canonical["shots"][0]["scene_uid"], "legacy-scene-uid")
        self.assertEqual(canonical["schema_version"], "R10.1")
        self.assertNotIn("batches", canonical)
        self.assertNotIn("生成批次", canonical["shots"][0])
        self.assertEqual(canonical["shots"][0]["镜头编号"], storyboard["shots"][0]["镜头编号"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
