"""Package-preserving overlays, using generated synthetic workbooks only."""
import hashlib
import io
import unittest
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from shortdrama_mcp.file_io import (
    CANONICAL_SIDECAR_SHEET, _seal_service_verified_audit_overlay,
    audit_saved_workbook_overlay, extract_xlsx, storyboard_to_xlsx,
)
from shortdrama_mcp.service import canonical_storyboard_sha256
from test_efficiency import valid_storyboard
from test_v0556_audit_intake_coverage import workbook_bytes


NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
DOC_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'


def package(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        return {p.filename: z.read(p) for p in z.infolist()}


def repack(parts):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        z.comment = b'original synthetic package comment'
        for name, value in parts.items():
            z.writestr(name, value)
    return stream.getvalue()


def with_shared_strings(parts, strings):
    sst = ET.Element(f'{{{NS}}}sst', {'count': str(len(strings)), 'uniqueCount': str(len(strings))})
    for value in strings:
        ET.SubElement(ET.SubElement(sst, f'{{{NS}}}si'), f'{{{NS}}}t').text = value
    parts['xl/sharedStrings.xml'] = ET.tostring(sst)
    rels = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
    ET.SubElement(rels, f'{{{REL}}}Relationship', {'Id': 'rIdSharedSynthetic',
        'Type': DOC_REL + '/sharedStrings', 'Target': 'sharedStrings.xml'})
    parts['xl/_rels/workbook.xml.rels'] = ET.tostring(rels)
    types = ET.fromstring(parts['[Content_Types].xml'])
    ET.SubElement(types, f'{{{CT}}}Override', {'PartName': '/xl/sharedStrings.xml',
        'ContentType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml'})
    parts['[Content_Types].xml'] = ET.tostring(types)


class OverlayPreservationTests(unittest.TestCase):
    def overlay(self, raw, *, canonical=None, payload_updates=None):
        canonical = canonical or extract_xlsx(raw, audit_only=True)['canonical_storyboard']
        payload = {
            'status': 'FAIL_CURRENT', 'conclusion': 'Synthetic missing evidence remains unknown.',
            'initial_issue_count': 1, 'residual_issue_count': 1,
            'introduced_issue_count': 0, 'blocking_issue_count': 1,
            'audit_trace_id': 'TRACE-synthetic', 'certificate_id': 'CERT-synthetic',
            'certificate_integrity_hmac': 'service-test-seal',
            'canonical_sha256': canonical_storyboard_sha256(canonical),
            'source_artifact_raw_sha256': hashlib.sha256(raw).hexdigest(),
        }
        payload.update(payload_updates or {})
        return audit_saved_workbook_overlay(raw, canonical,
            {'issues': [{'error_code': 'E-SYNTHETIC-UNKNOWN', 'evidence': 'Synthetic only.'}]},
            audit_overlay=_seal_service_verified_audit_overlay(payload))

    def assert_unchanged_parts(self, before, after, mutable):
        original, saved = package(before), package(after)
        for name, value in original.items():
            if name not in mutable:
                self.assertEqual(saved[name], value, name)
        with zipfile.ZipFile(io.BytesIO(before)) as a, zipfile.ZipFile(io.BytesIO(after)) as b:
            self.assertEqual(a.comment, b.comment)
        # Every original style definition remains at its original index.
        a = ET.fromstring(original['xl/styles.xml'])
        b = ET.fromstring(saved['xl/styles.xml'])
        for table in a:
            target = b.find(table.tag)
            self.assertIsNotNone(target, table.tag)
            for i, entry in enumerate(table):
                self.assertEqual(ET.tostring(entry), ET.tostring(target[i]), (table.tag, i))

    def test_empty_rich_shared_formula_cache_comments_and_format_bytes_survive(self):
        book = load_workbook(io.BytesIO(workbook_bytes()))
        reading = book['10集_逐镜阅读']
        reading['F2'] = 'styled content'
        reading['F2'].font = Font(name='Arial', bold=True, color='AB1234')
        reading['F2'].fill = PatternFill('solid', fgColor='FFFF00')
        reading['F2'].alignment = Alignment(wrap_text=True, text_rotation=17)
        reading['F2'].number_format = '0.0000'
        reading['F2'].comment = Comment('Keep original comment XML.', 'Synthetic')
        reading.row_dimensions[2].height = 61
        reading.column_dimensions['F'].width = 35
        reading.freeze_panes = 'B2'
        raw = io.BytesIO(); book.save(raw); book.close()
        parts = package(raw.getvalue())
        xml = ET.fromstring(parts['xl/worksheets/sheet2.xml'])
        row = xml.find(f'{{{NS}}}sheetData/{{{NS}}}row[@r="2"]')
        for address in ('B2', 'C2', 'D2', 'E2'):
            old = row.find(f'{{{NS}}}c[@r="{address}"]')
            if old is not None:
                row.remove(old)
        empty = ET.fromstring(f'<c xmlns="{NS}" r="B2" t="inlineStr"><is><t></t></is></c>')
        rich = ET.fromstring(f'<c xmlns="{NS}" r="C2" t="inlineStr"><is><r><rPr><b/></rPr><t>Bold</t></r><r><t xml:space="preserve"> plain</t></r></is></c>')
        formula = ET.fromstring(f'<c xmlns="{NS}" r="D2"><f>5+6</f><v>11</v></c>')
        shared = ET.fromstring(f'<c xmlns="{NS}" r="E2" t="s"><v>0</v></c>')
        row.extend([empty, rich, formula, shared])
        parts['xl/worksheets/sheet2.xml'] = ET.tostring(xml)
        with_shared_strings(parts, ['Shared literal'])
        raw = repack(parts)
        final = self.overlay(raw)
        self.assert_unchanged_parts(raw, final, {'xl/styles.xml', 'xl/workbook.xml',
            'xl/_rels/workbook.xml.rels', '[Content_Types].xml'})
        book = load_workbook(io.BytesIO(final), data_only=True)
        self.assertEqual(book['10集_逐镜阅读']['B2'].value, '')
        self.assertEqual(book['10集_逐镜阅读']['C2'].value, 'Bold plain')
        self.assertEqual(book['10集_逐镜阅读']['D2'].value, 11)
        self.assertEqual(book['10集_逐镜阅读']['E2'].value, 'Shared literal')
        self.assertEqual(book['10集_逐镜阅读']['F2'].font.name, 'Arial')
        self.assertEqual(book['10集_逐镜阅读']['F2'].number_format, '0.0000')
        self.assertEqual(book.sheetnames[-1], '00_范围与审计')
        self.assertEqual(book['00_范围与审计']['A1'].fill.fgColor.rgb, '001F6F5F')
        book.close()

    def test_existing_audit_custom_targets_vba_and_relationships_remain_bound(self):
        book = load_workbook(io.BytesIO(workbook_bytes()))
        book.create_sheet('00_范围与审计').append(['Old synthetic scope'])
        stream = io.BytesIO(); book.save(stream); book.close()
        parts = package(stream.getvalue())
        renames = {'xl/worksheets/sheet1.xml': 'xl/worksheets/protected/main.xml',
                   'xl/worksheets/sheet2.xml': 'xl/reading.xml',
                   'xl/worksheets/sheet3.xml': 'xl/audit/scope.xml'}
        rels = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        for item in rels:
            old = item.attrib['Target'].lstrip('/')
            if old in renames:
                item.set('Target', '/' + renames[old])
        ET.SubElement(rels, f'{{{REL}}}Relationship', {'Id': 'rIdVbaSynthetic',
            'Type': 'http://schemas.microsoft.com/office/2006/relationships/vbaProject',
            'Target': 'macros/vbaProject.bin'})
        parts['xl/_rels/workbook.xml.rels'] = ET.tostring(rels)
        types = ET.fromstring(parts['[Content_Types].xml'])
        for item in types:
            old = item.get('PartName', '').lstrip('/')
            if old in renames:
                item.set('PartName', '/' + renames[old])
        parts['[Content_Types].xml'] = ET.tostring(types)
        for old, new in renames.items():
            parts[new] = parts.pop(old)
        parts['xl/macros/vbaProject.bin'] = b'\x00synthetic opaque macro bytes\xff'
        raw = repack(parts)
        final = self.overlay(raw)
        self.assert_unchanged_parts(raw, final, {'xl/styles.xml', 'xl/audit/scope.xml'})
        self.assertEqual(set(package(raw)), set(package(final)))
        reread = extract_xlsx(final, audit_only=True)
        self.assertEqual(reread['canonical_storyboard']['shots'],
                         extract_xlsx(raw, audit_only=True)['canonical_storyboard']['shots'])
        second = self.overlay(final)
        self.assert_unchanged_parts(final, second, {'xl/audit/scope.xml'})

    def test_existing_sidecar_shared_header_only_digest_cells_change(self):
        canonical = valid_storyboard()
        raw = storyboard_to_xlsx(canonical, {'issues': []}, assume_compact=True)
        parts = package(raw)
        book = ET.fromstring(parts['xl/workbook.xml'])
        rels = {r.attrib['Id']: r.attrib['Target'].lstrip('/')
                for r in ET.fromstring(parts['xl/_rels/workbook.xml.rels'])}
        sheets = {s.attrib['name']: rels[s.attrib[f'{{{DOC_REL}}}id']]
                  for s in book.find(f'{{{NS}}}sheets')}
        side_path, audit_path = sheets[CANONICAL_SIDECAR_SHEET], sheets['00_范围与审计']
        side = ET.fromstring(parts[side_path])
        header = side.find(f'{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c[@r="K1"]')
        header.clear(); header.attrib.update({'r': 'K1', 't': 's'})
        ET.SubElement(header, f'{{{NS}}}v').text = '0'
        binding = side.find(f'{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c[@r="K2"]')
        binding.set('{urn:synthetic:metadata}keep', 'original binding metadata')
        parts[side_path] = ET.tostring(side)
        with_shared_strings(parts, ['audit_page_sha256'])
        raw = repack(parts)
        final = self.overlay(raw, canonical=canonical)
        self.assert_unchanged_parts(raw, final, {side_path, audit_path, 'xl/styles.xml'})
        self.assertEqual(extract_xlsx(final)['canonical_storyboard'], canonical)
        before = ET.fromstring(parts[side_path])
        after = ET.fromstring(package(final)[side_path])
        a_cells = {c.attrib['r']: c for c in before.iter(f'{{{NS}}}c')}
        b_cells = {c.attrib['r']: c for c in after.iter(f'{{{NS}}}c')}
        self.assertEqual(set(a_cells), set(b_cells))
        for address, cell in a_cells.items():
            if not address.startswith('K') or address == 'K1':
                self.assertEqual(ET.tostring(cell), ET.tostring(b_cells[address]), address)
        self.assertNotEqual(ET.tostring(a_cells['K2']), ET.tostring(b_cells['K2']))
        self.assertEqual(b_cells['K2'].get('{urn:synthetic:metadata}keep'), 'original binding metadata')
        self.assertEqual(package(raw)['xl/sharedStrings.xml'], package(final)['xl/sharedStrings.xml'])

    def test_overlay_still_requires_source_bytes_canonical_and_capability(self):
        raw = workbook_bytes()
        with self.assertRaisesRegex(ValueError, 'source_fingerprint_mismatch'):
            self.overlay(raw, payload_updates={'source_artifact_raw_sha256': '0' * 64})
        canonical = extract_xlsx(raw, audit_only=True)['canonical_storyboard']
        changed = dict(canonical, shots=[])
        with self.assertRaisesRegex(ValueError, 'source_canonical_mismatch'):
            self.overlay(raw, canonical=changed)
        with self.assertRaisesRegex(ValueError, 'unverified'):
            audit_saved_workbook_overlay(raw, canonical, {'issues': []}, audit_overlay={})


if __name__ == '__main__':
    unittest.main()
