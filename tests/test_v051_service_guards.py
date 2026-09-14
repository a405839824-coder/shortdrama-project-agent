from __future__ import annotations

import copy
import io
import tempfile
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any
from unittest.mock import patch

from openpyxl import load_workbook

import shortdrama_mcp.file_io as file_io
import shortdrama_mcp.service as service_module
import shortdrama_mcp.storage as storage_module
from shortdrama_mcp.app import create_mcp
from shortdrama_mcp.config import Settings
from shortdrama_mcp.file_io import storyboard_to_xlsx

from semantic_audit_helpers import semantic_audit_for_artifact
from test_efficiency import valid_storyboard


class V051ServiceGuardTests(unittest.TestCase):
    namespace = "local-single-user"

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="shortdrama-v051-service-")
        self.settings = Settings(
            data_dir=Path(self.temp.name),
            public_base_url="http://127.0.0.1:8765",
            host="127.0.0.1",
            port=8765,
            deployment_mode="local",
            auth_mode="none",
            oauth_issuer=None,
            oauth_audience=None,
            oauth_jwks_url=None,
            oauth_scopes=("storyboard:read", "storyboard:write"),
            allow_local_file_urls=False,
            max_file_bytes=50 * 1024 * 1024,
        )
        _server, self.service = create_mcp(self.settings)
        from historical_entry_fixture import freeze_legacy_entry
        freeze_legacy_entry(self.service)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _project(self, label: str) -> str:
        project_id = str(
            self.service.create_project(self.namespace, label).project_id
        )
        status = self.service.agent_status(
            self.namespace, project_id, full_self_test=True
        )
        self.assertTrue(status.data.get("runtime_test_certificate_valid"))
        return project_id

    def _begin_production(self, project_id: str, label: str) -> tuple[str, str]:
        session_id = f"production-{label}"
        response = self.service.begin_task(
            self.namespace,
            project_id=project_id,
            operation="storyboard_create",
            mode="execute",
            authorization="explicit_execute",
            request_text=f"v0.5.1 service guard: {label}",
            targets=["第一集"],
            authoritative_source="测试剧本",
            latest_version="v1",
            target_boundary="第一集",
            must_keep=["剧情、对白、连续性"],
            must_remove=[],
            protected_scope=[],
            expected_outputs=["正式分镜"],
            window_key=session_id,
            request_session_id=session_id,
        )
        return str(response.work_order_id), session_id

    def _begin_audit(
        self, project_id: str, artifact_id: str, label: str
    ) -> tuple[str, str]:
        session_id = f"audit-{label}"
        response = self.service.begin_task(
            self.namespace,
            project_id=project_id,
            operation="storyboard_audit",
            mode="audit_only",
            authorization="read_only",
            request_text=f"v0.5.1 independent audit: {label}",
            targets=[artifact_id],
            authoritative_source="已锁定分镜及其handoff",
            latest_version="current",
            target_boundary="第一集",
            must_keep=["只读独立审计"],
            must_remove=[],
            protected_scope=["全部源内容"],
            expected_outputs=["independent_audit_certificate"],
            window_key=session_id,
            request_session_id=session_id,
        )
        return str(response.work_order_id), session_id

    def test_unknown_certificate_cannot_authorize_an_audit_overlay(self) -> None:
        project_id = self._project("unknown-overlay-certificate")

        with self.assertRaisesRegex(
            ValueError,
            "independent_audit_overlay_certificate_unknown_or_tampered",
        ):
            self.service._load_verified_independent_audit_certificate(
                self.namespace,
                project_id,
                certificate_id="CERT-UNKNOWN",
                artifact_id="ART-UNKNOWN",
                canonical_sha256="a" * 64,
                source_artifact_raw_sha256="b" * 64,
            )

    def test_fake_hmac_on_persisted_certificate_cannot_authorize_overlay(self) -> None:
        project_id = self._project("tampered-overlay-certificate")
        certificate_id = "CERT-TAMPERED"
        certificate = self.service.store.save_project_record(
            self.namespace,
            project_id,
            kind="certificates",
            record_id=certificate_id,
            value={
                "certificate_id": certificate_id,
                "audit_class": "independent_audit",
            },
        )
        certificate["integrity_hmac"] = "0" * 64
        certificate_path = (
            self.service.store.project_dir(self.namespace, project_id)
            / "certificates"
            / f"{certificate_id}.json"
        )
        self.service.store.atomic_write_json(certificate_path, certificate)

        with self.assertRaisesRegex(
            ValueError,
            "independent_audit_overlay_certificate_unknown_or_tampered",
        ):
            self.service._load_verified_independent_audit_certificate(
                self.namespace,
                project_id,
                certificate_id=certificate_id,
                artifact_id="ART-UNKNOWN",
                canonical_sha256="a" * 64,
                source_artifact_raw_sha256="b" * 64,
            )

    def _commit(
        self,
        project_id: str,
        work_order_id: str,
        session_id: str,
        storyboard: dict,
        label: str,
    ):
        return self.service.commit_storyboard(
            self.namespace,
            project_id=project_id,
            work_order_id=work_order_id,
            storyboard=storyboard,
            file_name=f"{label}.xlsx",
            source_artifact_id=None,
            request_session_id=session_id,
        )

    def _assert_no_publication_side_effects(self, project_id: str) -> None:
        self.assertEqual(self.service.store.list_artifacts(self.namespace, project_id), [])
        self.assertEqual(
            self.service.store.list_project_records(
                self.namespace, project_id, kind="handoffs"
            ),
            [],
        )
        hub = self.service.store.load_hub(self.namespace, project_id)
        self.assertEqual(hub.get("deliverables", []), [])
        self.assertEqual(hub.get("handoffs", []), [])
        self.assertNotIn("storyboard", hub["project_state"]["latest_artifacts"])

    def _authority_snapshot(self, project_id: str) -> dict:
        return {
            "hub": self.service.store.load_hub(self.namespace, project_id),
            "ledger": self.service.store.load_ledger(
                self.namespace, project_id
            ),
            "handoffs": self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="handoffs"
            ),
            "certificates": self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="certificates"
            ),
            "repair_packages": self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="repair_packages"
            ),
            "telemetry": self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="telemetry"
            ),
            "artifacts": self.service.store.list_artifacts(
                self.namespace, project_id
            ),
        }

    def _project_raw_snapshot(
        self,
        project_id: str,
        *,
        exclude: tuple[Path, ...] = (),
    ) -> dict[str, bytes]:
        project_dir = self.service.store.project_dir(
            self.namespace, project_id
        )
        excluded = {path.resolve() for path in exclude}
        return {
            str(path.relative_to(project_dir)): path.read_bytes()
            for path in sorted(project_dir.rglob("*"))
            if path.is_file() and path.resolve() not in excluded
        }

    def test_storage_rollback_rejects_forged_migration_authority_without_mutation(
        self,
    ) -> None:
        project_id = self._project("storage-rollback-authority-guard")
        store = self.service.store
        hub_before = store.load_hub(self.namespace, project_id)
        ledger_before = store.load_ledger(self.namespace, project_id)
        hub_after = copy.deepcopy(hub_before)
        ledger_after = copy.deepcopy(ledger_before)
        hub_after["storage_rollback_guard_marker"] = "migrated"
        ledger_after["storage_rollback_guard_marker"] = "migrated"
        migration_id = "MIG-STORAGE-ROLLBACK-AUTHORITY-GUARD"
        runtime_certificate = self.service._require_runtime_test_certificate(
            self.namespace, project_id
        )
        journal = store.commit_project_state_migration(
            self.namespace,
            project_id,
            migration_id=migration_id,
            migration_key="storage-rollback-authority-guard",
            work_order_id="WO-STORAGE-ROLLBACK-AUTHORITY-GUARD",
            from_version="0.5.4",
            target_version="0.5.5",
            hub_before=hub_before,
            ledger_before=ledger_before,
            hub_after=hub_after,
            ledger_after=ledger_after,
            runtime_test_certificate_id=str(
                runtime_certificate["certificate_id"]
            ),
            runtime_test_certificate_key=self.service._runtime_certificate_key(),
            ruleset_version=self.service._ruleset_version,
            validator_version=self.service._validator_version(),
            changed_paths=["hub.storage_rollback_guard_marker"],
        )
        self.assertEqual(journal["status"], "committed")
        self.assertTrue(store.project_record_is_signed(journal))
        journal_path = (
            store.project_dir(self.namespace, project_id)
            / "migrations"
            / f"{migration_id}.json"
        )

        for defect in (
            "record_id",
            "migration_id",
            "project_id",
            "hub_before",
            "hub_after",
            "ledger_before",
            "ledger_after",
        ):
            with self.subTest(defect=defect):
                tampered = copy.deepcopy(journal)
                if defect == "record_id":
                    tampered[defect] = "MIG-FORGED-RECORD-ID"
                elif defect == "migration_id":
                    tampered[defect] = "MIG-FORGED-PRIMARY-ID"
                elif defect == "project_id":
                    tampered[defect] = "PRJ-FORGED-TARGET"
                else:
                    tampered[defect]["forged_body_marker"] = defect
                tampered["integrity_hmac"] = store._record_integrity_hmac(
                    tampered
                )
                self.assertTrue(store.project_record_is_signed(tampered))
                store.atomic_write_json(journal_path, tampered)
                before = self._project_raw_snapshot(project_id)
                try:
                    with self.assertRaisesRegex(
                        ValueError,
                        "project_record_identity_mismatch|"
                        "migration_record_authority_conflict",
                    ):
                        store.rollback_project_state_migration(
                            self.namespace, project_id, migration_id
                        )
                    self.assertEqual(
                        self._project_raw_snapshot(project_id), before
                    )
                finally:
                    store.atomic_write_json(journal_path, journal)

    @staticmethod
    def _failed_semantic(clean_semantic: dict) -> dict:
        failed = copy.deepcopy(clean_semantic)
        failed_layer = failed["layers"][0]
        failed_layer["status"] = "fail_current"
        failed_layer["summary"] = "独立复核发现原文语义错位。"
        failed["issues"] = [
            {
                "issue_id": "SEM-NEW-FAIL-001",
                "layer": failed_layer["layer"],
                "error_code": "E-SOURCE-SELECT-001",
                "severity": "blocking",
                "blocking": True,
                "shot_ids": [
                    str(valid_storyboard()["shots"][0]["镜头编号"])
                ],
                "batch_ids": [],
                "field": "原文支撑句",
                "evidence_ids": [failed_layer["evidence_ids"][0]],
                "evidence": "原文与镜头语义不一致。",
                "repair_action": "重建该镜语义绑定后重新独立审计。",
            }
        ]
        failed["conclusion"] = {
            "status": "fail_current",
            "summary": "独立复核发现阻塞问题。",
        }
        return failed

    def _passed_audit_fixture(self, label: str) -> dict:
        fixture = self._unaudited_fixture(label)
        first = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            semantic_audit=fixture["clean_semantic"],
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(first.status, "pass")
        return {**fixture, "first": first}

    def _unaudited_fixture(self, label: str) -> dict:
        project_id = self._project(label)
        production_id, production_session = self._begin_production(
            project_id, label
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            label,
        )
        artifact_id = str(committed.artifact_id)
        root_handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, label
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        return {
            "project_id": project_id,
            "production_id": production_id,
            "production_session": production_session,
            "artifact_id": artifact_id,
            "root_handoff_id": root_handoff_id,
            "audit_id": audit_id,
            "audit_session": audit_session,
            "clean_semantic": clean_semantic,
        }

    def _leave_prepared_reaudit_successor(
        self, fixture: dict, semantic_audit: dict
    ) -> dict:
        with patch.object(
            self.service.store,
            "claim_reaudit_successor",
            side_effect=RuntimeError("simulated loss before first claim"),
        ), patch.object(
            self.service.store,
            "rollback_reaudit_successor",
            return_value=None,
        ):
            with self.assertRaisesRegex(RuntimeError, "before first claim"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=semantic_audit,
                    request_session_id=fixture["audit_session"],
                )
        return next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, fixture["project_id"], kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id")
            == fixture["root_handoff_id"]
        )

    def _direct_reaudit_successor_candidate(
        self,
        fixture: dict,
        *,
        successor_id: str,
        route_trace_id: str,
        event_trace_id: str,
    ) -> dict:
        root = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=fixture["root_handoff_id"],
        )
        certificate = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="certificates",
            record_id=str(root["audit_certificate_id"]),
        )
        event = next(
            row
            for row in self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"]
            if row.get("event_id") == certificate.get("event_id")
        )
        audit_order = next(
            row
            for row in self.service.store.load_hub(
                self.namespace, fixture["project_id"]
            )["work_orders"]
            if row.get("work_order_id") == fixture["audit_id"]
        )
        successor = self.service._handoff_record_projection(root)
        for key in (
            "handoff_id",
            "created_at",
            "status",
            "claimed_by_work_order_id",
            "claim_token",
            "claim_generation",
            "claim_epoch_seconds",
            "claim_lease_seconds",
            "claim_lease_expires_at",
            "claim_lease_expires_epoch_seconds",
            "consumed_by_work_order_id",
            "consumed_at",
            "consumption_claim",
            "audit_certificate_id",
            "audit_certificate_ids",
            "reaudit_count",
            "reaudited_at",
            "latest_audit_certificate_id",
            "reaudit_preclaim_projection_sha256",
            "reaudit_prerelease_projection_sha256",
        ):
            successor.pop(key, None)
        successor.update(
            {
                "handoff_id": successor_id,
                "created_at": service_module.utc_stamp(),
                "status": "prepared",
                "independent_audit_status": "PENDING_AUDIT",
                "stage_trace_id": route_trace_id,
                "production_stage_trace_id": str(root["stage_trace_id"]),
                "reaudit_kind": "independent_audit_successor",
                "reaudit_root_handoff_id": str(root["handoff_id"]),
                "reaudit_parent_handoff_id": str(root["handoff_id"]),
                "reaudit_sequence": 1,
                "reaudit_route_trace_id": route_trace_id,
                "reaudit_event_trace_id": event_trace_id,
                "previous_audit_certificate_id": str(
                    certificate["certificate_id"]
                ),
                "previous_audit_certificate_hmac": str(
                    certificate["integrity_hmac"]
                ),
                "previous_audit_event_sha256": storage_module.sha256_bytes(
                    storage_module.canonical_json_bytes(event)
                ),
                "claim_authority_work_order_id": fixture["audit_id"],
                "reaudit_audit_window_binding_id": audit_order[
                    "server_window_binding_id"
                ],
                "reaudit_audit_actor_binding_id": audit_order[
                    "server_actor_binding_id"
                ],
                "reaudit_audit_session_fingerprint": audit_order[
                    "transport_session_fingerprint"
                ],
                "server_minted_lineage": True,
            }
        )
        successor["reaudit_lineage_sha256"] = service_module.sha256_bytes(
            service_module._stable_json_bytes(
                self.service._reaudit_lineage_payload(
                    root=root,
                    parent=root,
                    successor=successor,
                )
            )
        )
        return successor

    def _leave_prerelease_half_write(
        self, fixture: dict, semantic_audit: dict
    ) -> tuple[dict, dict]:
        """Model signed-prepared / Hub-exact-consuming process loss."""

        prepared = self._leave_prepared_reaudit_successor(
            fixture, semantic_audit
        )
        successor_id = str(prepared["handoff_id"])
        consuming = self.service.store.claim_reaudit_successor(
            self.namespace,
            fixture["project_id"],
            successor_id,
            consumer_work_order_id=fixture["audit_id"],
        )
        released = self.service.store.release_reaudit_successor_claim(
            self.namespace,
            fixture["project_id"],
            successor_id,
            consumer_work_order_id=fixture["audit_id"],
            claim_token=str(consuming["claim_token"]),
            reason="simulate handled failure before Hub release write",
        )
        self.assertEqual(released["status"], "prepared")
        self.assertEqual(
            released["reaudit_prerelease_projection_sha256"],
            self.service.store.project_state_sha256(
                self.service._handoff_record_projection(consuming)
            ),
        )
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (hub, _ledger):
            hub["handoffs"] = [
                self.service._handoff_record_projection(consuming)
                if row.get("handoff_id") == successor_id
                else row
                for row in hub["handoffs"]
            ]
        return released, consuming

    def test_preflight_guards_create_no_artifact_version_export_or_handoff(self) -> None:
        cases = (
            (
                "aspect",
                lambda shot: shot.__setitem__(
                    "生成限制", "16:9横屏；禁止字幕；禁止额外配乐"
                ),
                "aspect_ratio_contract_violation",
                None,
                0,
            ),
            (
                "bgm",
                lambda shot: shot.__setitem__(
                    "环境音效", str(shot["环境音效"]) + "；BGM渐入"
                ),
                None,
                "THG-AUDIO",
                0,
            ),
            (
                "dynamic-keyframe",
                lambda shot: shot.__setitem__(
                    "起始状态", str(shot["起始状态"]) + "；随后抬手"
                ),
                None,
                "THG-03",
                0,
            ),
            (
                "duplicate-transition",
                lambda shot: shot.__setitem__(
                    "切镜/转场提示词", "手停时切；；手停时切"
                ),
                "transition_contract_violation",
                None,
                0,
            ),
        )
        for label, mutate, expected_error, expected_gate, projector_calls in cases:
            with self.subTest(label=label):
                project_id = self._project(f"preflight-{label}")
                work_order_id, session_id = self._begin_production(
                    project_id, label
                )
                storyboard = valid_storyboard()
                mutate(storyboard["shots"][0])
                with patch.object(
                    service_module,
                    "_export_compact_storyboard",
                    wraps=service_module._export_compact_storyboard,
                ) as projector:
                    if expected_error:
                        with self.assertRaisesRegex(ValueError, expected_error):
                            self._commit(
                                project_id,
                                work_order_id,
                                session_id,
                                storyboard,
                                label,
                            )
                    else:
                        response = self._commit(
                            project_id,
                            work_order_id,
                            session_id,
                            storyboard,
                            label,
                        )
                        self.assertEqual(response.status, "fail_current")
                        self.assertFalse(response.data["artifact_created"])
                        self.assertFalse(response.data["handoff_created"])
                        gate_ids = {
                            issue.get("template_gate_id")
                            for issue in response.data["report"]["issues"]
                        }
                        self.assertIn(expected_gate, gate_ids)
                self.assertEqual(projector.call_count, projector_calls)
                self._assert_no_publication_side_effects(project_id)

                # A rejected attempt cannot consume a storyboard version.
                # A failed draft scope must be corrected through the same
                # work order; opening a new batch would bypass the learned
                # prevention profile and is intentionally blocked in v0.5.5.
                committed = self._commit(
                    project_id,
                    work_order_id,
                    session_id,
                    valid_storyboard(),
                    f"{label}-retry",
                )
                self.assertEqual(committed.status, "production_preflight_pass")
                self.assertEqual(committed.data["artifact"]["version"], "v1")

    def test_projector_exception_leaves_no_orphan_artifact(self) -> None:
        project_id = self._project("projector-transaction")
        work_order_id, session_id = self._begin_production(
            project_id, "projector-transaction"
        )
        with patch.object(
            service_module,
            "_export_compact_storyboard",
            side_effect=RuntimeError("projector exploded"),
        ):
            with self.assertRaisesRegex(RuntimeError, "projector exploded"):
                self._commit(
                    project_id,
                    work_order_id,
                    session_id,
                    valid_storyboard(),
                    "projector-transaction",
                )
        self._assert_no_publication_side_effects(project_id)

        committed = self._commit(
            project_id,
            work_order_id,
            session_id,
            valid_storyboard(),
            "projector-retry",
        )
        self.assertEqual(committed.status, "production_preflight_pass")
        self.assertEqual(committed.data["artifact"]["version"], "v1")

    def test_new_fail_after_pass_revokes_latest_and_writes_fail_overlay(self) -> None:
        project_id = self._project("pass-then-fail")
        production_id, production_session = self._begin_production(
            project_id, "pass-then-fail"
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            "pass-then-fail",
        )
        artifact_id = str(committed.artifact_id)
        handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, "pass-then-fail"
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        first = self.service.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=handoff_id,
            semantic_audit=clean_semantic,
            request_session_id=audit_session,
        )
        self.assertEqual(first.status, "pass")
        pass_overlay_id = first.data["audited_workbook"]["artifact_id"]
        hub = self.service.store.load_hub(self.namespace, project_id)
        self.assertEqual(
            hub["project_state"]["latest_artifacts"]["storyboard"][
                "deliverable_id"
            ],
            pass_overlay_id,
        )

        failed_semantic = self._failed_semantic(clean_semantic)

        root_before_reaudit = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=handoff_id,
        )
        root_hub_before_reaudit = copy.deepcopy(
            next(
                row
                for row in self.service.store.load_hub(
                    self.namespace, project_id
                )["handoffs"]
                if row["handoff_id"] == handoff_id
            )
        )

        # A semantic-audit contract change invalidates the old cache while the
        # producer bytes remain identical, exercising authority revocation.
        with patch.object(
            service_module,
            "SEMANTIC_AUDIT_CONTRACT_FINGERPRINT",
            "f" * 64,
        ):
            second = self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=handoff_id,
                force_machine_recheck=True,
                semantic_audit=failed_semantic,
                request_session_id=audit_session,
            )

        self.assertEqual(second.status, "fail_current")
        self.assertEqual(second.data["audited_workbook"]["status"], "FAIL_CURRENT")
        fail_overlay_id = second.data["audited_workbook"]["artifact_id"]
        self.assertNotEqual(fail_overlay_id, pass_overlay_id)
        fail_record, _path = self.service.store.get_artifact(
            self.namespace, project_id, fail_overlay_id
        )
        self.assertEqual(fail_record["artifact_type"], "storyboard_audit_overlay")
        self.assertEqual(
            fail_record["metadata"]["independent_audit_status"], "FAIL_CURRENT"
        )

        hub = self.service.store.load_hub(self.namespace, project_id)
        self.assertNotIn("storyboard", hub["project_state"]["latest_artifacts"])
        producer = next(
            row
            for row in hub["work_orders"]
            if row["work_order_id"] == production_id
        )
        self.assertEqual(producer["status"], "failed_current")
        self.assertTrue(producer["awaiting_independent_audit"])
        self.assertNotIn("independent_audit_certificate_id", producer)
        self.assertNotIn("verified_independent_audit_certificate_id", producer)
        audit_order = next(
            row
            for row in hub["work_orders"]
            if row["work_order_id"] == audit_id
        )
        self.assertEqual(audit_order["status"], "failed_current")
        self.assertEqual(
            audit_order["reopened_by_independent_audit"],
            second.data["certificate"]["certificate_id"],
        )
        registered_ids = set(audit_order.get("deliverable_ids", []))
        self.assertIn(
            second.data["certificate"]["certificate_id"], registered_ids
        )
        self.assertIn(fail_overlay_id, registered_ids)
        hub_deliverable_ids = {
            row.get("deliverable_id") for row in hub.get("deliverables", [])
        }
        self.assertIn(
            second.data["certificate"]["certificate_id"],
            hub_deliverable_ids,
        )
        self.assertIn(fail_overlay_id, hub_deliverable_ids)
        persisted_handoff = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=handoff_id,
        )
        successor_id = str(second.data["handoff"]["handoff_id"])
        self.assertNotEqual(successor_id, handoff_id)
        self.assertEqual(
            persisted_handoff["audit_certificate_id"],
            first.data["certificate"]["certificate_id"],
        )
        self.assertEqual(persisted_handoff, root_before_reaudit)
        self.assertNotIn("reaudit_count", persisted_handoff)
        self.assertNotIn("audit_certificate_ids", persisted_handoff)
        successor = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=successor_id,
        )
        self.service._validate_schema_record(
            self.service._project_schema, "handoff_manifest", successor
        )
        for field in (
            "reaudit_preclaim_projection_sha256",
            "reaudit_prerelease_projection_sha256",
            "latest_audit_certificate_id",
            "consumption_claim",
        ):
            with self.subTest(partial_successor_marker=field):
                partial_successor = copy.deepcopy(root_before_reaudit)
                partial_successor[field] = (
                    {}
                    if field == "consumption_claim"
                    else "0" * 64
                )
                with self.assertRaisesRegex(
                    ValueError, "lifecycle_schema_mismatch: handoff_manifest"
                ):
                    self.service._validate_schema_record(
                        self.service._project_schema,
                        "handoff_manifest",
                        partial_successor,
                    )
        successor_authority_fields = (
            "reaudit_kind",
            "reaudit_root_handoff_id",
            "reaudit_parent_handoff_id",
            "reaudit_sequence",
            "reaudit_route_trace_id",
            "reaudit_event_trace_id",
            "production_stage_trace_id",
            "previous_audit_certificate_id",
            "previous_audit_certificate_hmac",
            "previous_audit_event_sha256",
            "claim_authority_work_order_id",
            "reaudit_audit_window_binding_id",
            "reaudit_audit_actor_binding_id",
            "reaudit_audit_session_fingerprint",
            "server_minted_lineage",
            "reaudit_lineage_sha256",
        )
        for field in successor_authority_fields:
            with self.subTest(missing_successor_authority=field):
                invalid_successor = copy.deepcopy(successor)
                invalid_successor.pop(field)
                with self.assertRaisesRegex(
                    ValueError, "lifecycle_schema_mismatch: handoff_manifest"
                ):
                    self.service._validate_schema_record(
                        self.service._project_schema,
                        "handoff_manifest",
                        invalid_successor,
                    )
        for field, invalid_value in (
            ("reaudit_sequence", 0),
            ("server_minted_lineage", False),
            ("previous_audit_certificate_hmac", "not-a-sha256"),
            ("previous_audit_event_sha256", "not-a-sha256"),
            ("reaudit_audit_session_fingerprint", ""),
            ("status", "verified"),
        ):
            with self.subTest(invalid_successor_authority=field):
                invalid_successor = copy.deepcopy(successor)
                invalid_successor[field] = invalid_value
                with self.assertRaisesRegex(
                    ValueError, "lifecycle_schema_mismatch: handoff_manifest"
                ):
                    self.service._validate_schema_record(
                        self.service._project_schema,
                        "handoff_manifest",
                        invalid_successor,
                    )
        consumed_authority_fields = (
            "consumed_by_work_order_id",
            "audit_certificate_id",
            "latest_audit_certificate_id",
            "consumed_at",
            "consumption_claim",
            "claim_generation",
            "reaudit_preclaim_projection_sha256",
        )
        for field in consumed_authority_fields:
            with self.subTest(missing_consumed_successor_binding=field):
                invalid_successor = copy.deepcopy(successor)
                invalid_successor.pop(field)
                with self.assertRaisesRegex(
                    ValueError, "lifecycle_schema_mismatch: handoff_manifest"
                ):
                    self.service._validate_schema_record(
                        self.service._project_schema,
                        "handoff_manifest",
                        invalid_successor,
                    )
        for field in (
            "claim_token",
            "claim_generation",
            "claim_epoch_seconds",
            "claim_lease_seconds",
            "claim_lease_expires_at",
            "claim_lease_expires_epoch_seconds",
        ):
            with self.subTest(missing_consumption_claim_binding=field):
                invalid_successor = copy.deepcopy(successor)
                invalid_successor["consumption_claim"].pop(field)
                with self.assertRaisesRegex(
                    ValueError, "lifecycle_schema_mismatch: handoff_manifest"
                ):
                    self.service._validate_schema_record(
                        self.service._project_schema,
                        "handoff_manifest",
                        invalid_successor,
                    )
        consuming_successor = copy.deepcopy(successor)
        consumption_claim = consuming_successor.pop("consumption_claim")
        consuming_successor["status"] = "consuming"
        consuming_successor["claimed_by_work_order_id"] = (
            consuming_successor.pop("consumed_by_work_order_id")
        )
        for field in (
            "audit_certificate_id",
            "latest_audit_certificate_id",
            "consumed_at",
        ):
            consuming_successor.pop(field)
        consuming_successor.update(copy.deepcopy(consumption_claim))
        self.service._validate_schema_record(
            self.service._project_schema,
            "handoff_manifest",
            consuming_successor,
        )
        for field in (
            "claimed_by_work_order_id",
            "claim_token",
            "claim_generation",
            "claim_epoch_seconds",
            "claim_lease_seconds",
            "claim_lease_expires_at",
            "claim_lease_expires_epoch_seconds",
            "reaudit_preclaim_projection_sha256",
        ):
            with self.subTest(missing_consuming_successor_binding=field):
                invalid_successor = copy.deepcopy(consuming_successor)
                invalid_successor.pop(field)
                with self.assertRaisesRegex(
                    ValueError, "lifecycle_schema_mismatch: handoff_manifest"
                ):
                    self.service._validate_schema_record(
                        self.service._project_schema,
                        "handoff_manifest",
                        invalid_successor,
                    )
        self.assertEqual(successor["status"], "consumed")
        self.assertEqual(successor["reaudit_root_handoff_id"], handoff_id)
        self.assertEqual(successor["reaudit_parent_handoff_id"], handoff_id)
        self.assertEqual(successor["reaudit_sequence"], 1)
        self.assertEqual(
            successor["previous_audit_certificate_id"],
            first.data["certificate"]["certificate_id"],
        )
        self.assertEqual(
            successor["audit_certificate_id"],
            second.data["certificate"]["certificate_id"],
        )
        self.assertEqual(
            second.data["certificate"]["handoff_id"], successor_id
        )
        self.assertEqual(
            second.data["certificate"]["production_trace_id"],
            successor["reaudit_route_trace_id"],
        )
        self.assertEqual(
            second.data["certificate"]["stage_trace_id"],
            successor["reaudit_event_trace_id"],
        )
        repair_package = second.data["repair_package"]
        repair_handoff_summary = second.data["audit_to_repair_handoff"]
        self.assertIsInstance(repair_package, dict)
        self.assertIsInstance(repair_handoff_summary, dict)
        self.assertTrue(
            self.service.store.project_record_is_signed(repair_package)
        )
        self.assertEqual(
            repair_package["audit_certificate_id"],
            second.data["certificate"]["certificate_id"],
        )
        self.assertEqual(
            repair_package["audit_event_id"],
            second.data["certificate"]["event_id"],
        )
        self.assertEqual(
            repair_package["production_handoff_id"], successor_id
        )
        self.assertEqual(
            repair_package["source_audit_handoff_id"], successor_id
        )
        self.assertEqual(
            repair_package["source_audit_event_trace_id"],
            successor["reaudit_event_trace_id"],
        )
        self.assertEqual(
            repair_package["source_audit_route_trace_id"],
            successor["reaudit_route_trace_id"],
        )
        expected_shot_id = str(valid_storyboard()["shots"][0]["镜头编号"])
        self.assertEqual(
            repair_package["authorized_scope"]["shot_ids"],
            [expected_shot_id],
        )
        self.assertEqual(
            repair_package["issue_targets"][0]["issue_id"],
            "SEM-NEW-FAIL-001",
        )
        repair_handoff = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=str(repair_package["audit_to_repair_handoff_id"]),
        )
        self.assertEqual(repair_handoff["handoff_type"], "audit_to_repair")
        self.assertEqual(repair_handoff["source_audit_handoff_id"], successor_id)
        for forbidden in (
            "reaudit_kind",
            "reaudit_root_handoff_id",
            "reaudit_parent_handoff_id",
            "reaudit_sequence",
            "reaudit_route_trace_id",
            "reaudit_event_trace_id",
            "reaudit_lineage_sha256",
            "server_minted_lineage",
            "claim_token",
            "consumption_claim",
            "consumed_by_work_order_id",
        ):
            self.assertNotIn(forbidden, repair_handoff)
        self.assertNotEqual(
            successor["reaudit_route_trace_id"],
            successor["reaudit_event_trace_id"],
        )
        hub_successor = next(
            row
            for row in hub["handoffs"]
            if row["handoff_id"] == successor_id
        )
        self.assertEqual(
            hub_successor,
            self.service._handoff_record_projection(successor),
        )
        self.assertEqual(
            next(
                row
                for row in hub["handoffs"]
                if row["handoff_id"] == handoff_id
            ),
            root_hub_before_reaudit,
        )
        trace_ids = {
            str(persisted_handoff["stage_trace_id"]),
            str(first.data["certificate"]["stage_trace_id"]),
            str(successor["reaudit_route_trace_id"]),
            str(second.data["certificate"]["stage_trace_id"]),
        }
        self.assertEqual(len(trace_ids), 4)
        route_trace = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="telemetry",
            record_id=str(successor["reaudit_route_trace_id"]),
        )
        self.assertEqual(
            route_trace["operation"], "audit_storyboard_reaudit_route"
        )
        # Simulate the only recoverable hard-crash boundary: the signed child
        # reached ``consumed`` while Hub still has the exact prior consuming
        # projection.  Exact replay repairs Hub; arbitrary divergence remains
        # fail-closed.
        consumption_claim = successor["consumption_claim"]
        half_written = copy.deepcopy(successor)
        half_written["status"] = "consuming"
        half_written["claimed_by_work_order_id"] = audit_id
        half_written.update(copy.deepcopy(consumption_claim))
        for key in (
            "consumed_by_work_order_id",
            "audit_certificate_id",
            "latest_audit_certificate_id",
            "consumed_at",
            "consumption_claim",
        ):
            half_written.pop(key, None)
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (half_written_hub, _ledger):
            half_written_hub["handoffs"] = [
                self.service._handoff_record_projection(half_written)
                if row["handoff_id"] == successor_id
                else row
                for row in half_written_hub["handoffs"]
            ]
        recovered = self.service.store.finish_reaudit_successor(
            self.namespace,
            project_id,
            successor_id,
            consumer_work_order_id=audit_id,
            audit_certificate_id=str(
                second.data["certificate"]["certificate_id"]
            ),
            claim_token=str(consumption_claim["claim_token"]),
        )
        self.assertEqual(recovered, successor)
        self.assertEqual(
            next(
                row
                for row in self.service.store.load_hub(
                    self.namespace, project_id
                )["handoffs"]
                if row["handoff_id"] == successor_id
            ),
            self.service._handoff_record_projection(successor),
        )
        binding_before = self._authority_snapshot(project_id)
        with self.assertRaisesRegex(
            ValueError, "reaudit_certificate_binding_invalid"
        ):
            self.service.store.finish_reaudit_successor(
                self.namespace,
                project_id,
                successor_id,
                consumer_work_order_id=audit_id,
                audit_certificate_id=str(
                    first.data["certificate"]["certificate_id"]
                ),
                claim_token=str(consumption_claim["claim_token"]),
            )
        self.assertEqual(self._authority_snapshot(project_id), binding_before)
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (tampered_hub, _ledger):
            next(
                row
                for row in tampered_hub["handoffs"]
                if row["handoff_id"] == successor_id
            )["status"] = "prepared"
        tampered_before = self._authority_snapshot(project_id)
        with self.assertRaisesRegex(
            ValueError, "reaudit_successor_projection_mismatch"
        ):
            self.service.store.finish_reaudit_successor(
                self.namespace,
                project_id,
                successor_id,
                consumer_work_order_id=audit_id,
                audit_certificate_id=str(
                    second.data["certificate"]["certificate_id"]
                ),
                claim_token=str(consumption_claim["claim_token"]),
            )
        self.assertEqual(self._authority_snapshot(project_id), tampered_before)
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (restored_hub, _ledger):
            restored_hub["handoffs"] = [
                self.service._handoff_record_projection(successor)
                if row["handoff_id"] == successor_id
                else row
                for row in restored_hub["handoffs"]
            ]
        replay_before = self._authority_snapshot(project_id)
        # Replay under the same semantic-contract environment that issued the
        # successor certificate.  Returning to the default contract here
        # would correctly invalidate *both* certificates as environment-stale
        # and would test a different rule.
        with patch.object(
            service_module,
            "SEMANTIC_AUDIT_CONTRACT_FINGERPRINT",
            "f" * 64,
        ):
            replay = self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=handoff_id,
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=audit_session,
            )
        self.assertEqual(replay.status, "fail_current")
        self.assertTrue(replay.data["cache_hit"])
        self.assertEqual(
            replay.data["certificate"]["certificate_id"],
            second.data["certificate"]["certificate_id"],
        )
        self.assertEqual(replay.data["handoff"]["handoff_id"], successor_id)
        self.assertEqual(self._authority_snapshot(project_id), replay_before)
        with patch.object(
            service_module,
            "SEMANTIC_AUDIT_CONTRACT_FINGERPRINT",
            "f" * 64,
        ):
            second_replay = self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=handoff_id,
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=audit_session,
            )
        self.assertEqual(second_replay.status, "fail_current")
        self.assertEqual(
            second_replay.data["certificate"]["certificate_id"],
            second.data["certificate"]["certificate_id"],
        )
        self.assertEqual(self._authority_snapshot(project_id), replay_before)
        ledger = self.service.store.load_ledger(self.namespace, project_id)
        self.assertEqual(ledger["human_signoffs"], [])
        self.assertEqual(len(ledger["authority_reopens"]), 1)
        self.assertEqual(
            ledger["authority_reopens"][0]["authority"],
            "independent_audit",
        )
        source_record, _path = self.service.store.get_artifact(
            self.namespace, project_id, artifact_id
        )
        self.assertEqual(source_record["metadata"]["audit_status"], "FAIL_CURRENT")

    def test_reaudit_semantic_rejection_has_zero_authority_side_effects(self) -> None:
        project_id = self._project("reaudit-semantic-zero-side-effects")
        production_id, production_session = self._begin_production(
            project_id, "reaudit-semantic-zero-side-effects"
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            "reaudit-semantic-zero-side-effects",
        )
        artifact_id = str(committed.artifact_id)
        handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, "reaudit-semantic-zero-side-effects"
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        first = self.service.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=handoff_id,
            semantic_audit=clean_semantic,
            request_session_id=audit_session,
        )
        self.assertEqual(first.status, "pass")
        invalid_semantic = copy.deepcopy(clean_semantic)
        invalid_semantic["canonical_sha256"] = "0" * 64
        before = self._authority_snapshot(project_id)
        with self.assertRaisesRegex(ValueError, "semantic_audit.*canonical"):
            self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=handoff_id,
                force_machine_recheck=True,
                semantic_audit=invalid_semantic,
                request_session_id=audit_session,
            )
        self.assertEqual(self._authority_snapshot(project_id), before)

    def test_semantic_repair_targets_reject_empty_unknown_and_alias_duplicates(self) -> None:
        for fault in ("empty", "unknown", "alias-duplicate"):
            with self.subTest(fault=fault):
                fixture = self._unaudited_fixture(
                    f"semantic-repair-target-{fault}"
                )
                failed = self._failed_semantic(
                    fixture["clean_semantic"]
                )
                storyboard, _record = self.service._storyboard_from_artifact(
                    self.namespace,
                    fixture["project_id"],
                    fixture["artifact_id"],
                )
                shot = storyboard["shots"][0]
                if fault == "empty":
                    failed["issues"][0]["shot_ids"] = []
                    expected = "semantic_audit_repair_target_required"
                elif fault == "unknown":
                    failed["issues"][0]["shot_ids"] = ["SHOT-UNKNOWN"]
                    expected = "semantic_audit_reference_invalid"
                else:
                    primary = str(shot["镜头编号"])
                    alias = str(shot.get("shot_uid") or "")
                    self.assertTrue(alias and alias != primary)
                    failed["issues"][0]["shot_ids"] = [primary, alias]
                    expected = "semantic_audit_reference_invalid"
                before = self._authority_snapshot(fixture["project_id"])
                with self.assertRaisesRegex(ValueError, expected):
                    self.service.audit_storyboard(
                        self.namespace,
                        project_id=fixture["project_id"],
                        artifact_id=fixture["artifact_id"],
                        audit_cycle_id=None,
                        work_order_id=fixture["audit_id"],
                        handoff_artifact_id=fixture["root_handoff_id"],
                        semantic_audit=failed,
                        request_session_id=fixture["audit_session"],
                    )
                self.assertEqual(
                    self._authority_snapshot(fixture["project_id"]), before
                )

    def test_repair_target_location_matrix_fails_closed(self) -> None:
        storyboard = valid_storyboard()
        second = copy.deepcopy(storyboard["shots"][0])
        second["镜头编号"] = "EP01-002"
        storyboard["shots"].append(second)
        base_issue = {
            "error_code": "E-TARGET-001",
            "blocking": True,
            "shot_ids": ["EP01-001"],
            "field": "画面内容",
        }

        same = {**base_issue, "location": "shots[1].画面内容"}
        self.assertEqual(
            self.service._repair_issue_targets(storyboard, [same]),
            [
                {
                    "issue_id": "E-TARGET-001",
                    "issue_code": "E-TARGET-001",
                    "shot_id": "EP01-001",
                    "field": "画面内容",
                    "location": "shots[1].画面内容",
                }
            ],
        )
        for issue, expected in (
            (
                {**base_issue, "location": "shots[2].画面内容"},
                "repair_package_target_conflict",
            ),
            (
                {**base_issue, "location": "shots[999].画面内容"},
                "repair_package_target_invalid",
            ),
            (
                {
                    **base_issue,
                    "shot_ids": [],
                    "location": "shots[999].画面内容",
                },
                "repair_package_target_invalid",
            ),
            (
                {
                    **base_issue,
                    "shot_ids": [],
                    "location": "semantic_audit.no_target",
                },
                "repair_package_target_required",
            ),
        ):
            with self.subTest(location=issue["location"]):
                with self.assertRaisesRegex(ValueError, expected):
                    self.service._repair_issue_targets(
                        storyboard, [issue]
                    )

    def test_reaudit_successor_hub_write_failure_rolls_back_to_zero(self) -> None:
        project_id = self._project("reaudit-successor-write-rollback")
        production_id, production_session = self._begin_production(
            project_id, "reaudit-successor-write-rollback"
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            "reaudit-successor-write-rollback",
        )
        artifact_id = str(committed.artifact_id)
        handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, "reaudit-successor-write-rollback"
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        self.service.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=handoff_id,
            semantic_audit=clean_semantic,
            request_session_id=audit_session,
        )
        before = self._authority_snapshot(project_id)
        before_handoff_count = len(before["hub"]["handoffs"])
        original_atomic_write = self.service.store.atomic_write_json
        injected = {"raised": False}

        def fail_successor_hub_write(path, value):
            if (
                not injected["raised"]
                and path.name == "hub.json"
                and len(value.get("handoffs", [])) > before_handoff_count
            ):
                injected["raised"] = True
                raise OSError("simulated re-audit Hub write failure")
            return original_atomic_write(path, value)

        with patch.object(
            self.service.store,
            "atomic_write_json",
            side_effect=fail_successor_hub_write,
        ):
            with self.assertRaisesRegex(OSError, "re-audit Hub write failure"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=project_id,
                    artifact_id=artifact_id,
                    audit_cycle_id=None,
                    work_order_id=audit_id,
                    handoff_artifact_id=handoff_id,
                    force_machine_recheck=True,
                    semantic_audit=clean_semantic,
                    request_session_id=audit_session,
                )
        self.assertTrue(injected["raised"])
        self.assertEqual(self._authority_snapshot(project_id), before)

        # A handled failure after the route and its telemetry are committed,
        # but before the one-time claim succeeds, must compensate both.
        with patch.object(
            self.service.store,
            "claim_reaudit_successor",
            side_effect=RuntimeError("simulated successor claim failure"),
        ):
            with self.assertRaisesRegex(
                RuntimeError, "successor claim failure"
            ):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=project_id,
                    artifact_id=artifact_id,
                    audit_cycle_id=None,
                    work_order_id=audit_id,
                    handoff_artifact_id=handoff_id,
                    force_machine_recheck=True,
                    semantic_audit=clean_semantic,
                    request_session_id=audit_session,
                )
        self.assertEqual(self._authority_snapshot(project_id), before)

        # Crash after rollback has removed Hub but before it can unlink the
        # signed prepared route.  A later explicit rollback must recognize
        # the same signed-only state and finish cleanup idempotently.
        original_unlink = Path.unlink
        unlink_crash = {"raised": False}

        def fail_successor_unlink(path, *args, **kwargs):
            if (
                not unlink_crash["raised"]
                and path.parent.name == "handoffs"
                and path.name.startswith("HND-")
                and path.name.endswith(".json")
                and path.name != f"{handoff_id}.json"
            ):
                unlink_crash["raised"] = True
                raise OSError("simulated rollback unlink crash")
            return original_unlink(path, *args, **kwargs)

        with patch.object(
            self.service.store,
            "claim_reaudit_successor",
            side_effect=RuntimeError("simulated successor claim failure"),
        ), patch.object(Path, "unlink", new=fail_successor_unlink):
            with self.assertRaisesRegex(
                RuntimeError, "successor claim failure"
            ):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=project_id,
                    artifact_id=artifact_id,
                    audit_cycle_id=None,
                    work_order_id=audit_id,
                    handoff_artifact_id=handoff_id,
                    force_machine_recheck=True,
                    semantic_audit=clean_semantic,
                    request_session_id=audit_session,
                )
        self.assertTrue(unlink_crash["raised"])
        stranded = [
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id") == handoff_id
        ]
        self.assertEqual(len(stranded), 1)
        stranded_id = str(stranded[0]["handoff_id"])
        stranded_trace_id = str(stranded[0]["reaudit_route_trace_id"])
        self.assertFalse(
            any(
                row.get("handoff_id") == stranded_id
                for row in self.service.store.load_hub(
                    self.namespace, project_id
                )["handoffs"]
            )
        )
        self.service.store.rollback_reaudit_successor(
            self.namespace,
            project_id,
            stranded_id,
            route_trace_id=stranded_trace_id,
        )
        self.assertEqual(self._authority_snapshot(project_id), before)
        self.service.store.rollback_reaudit_successor(
            self.namespace,
            project_id,
            stranded_id,
            route_trace_id=stranded_trace_id,
        )
        self.assertEqual(self._authority_snapshot(project_id), before)

    def test_restart_resumes_signed_prepared_successor_missing_hub(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-restart-signed-only")
        project_id = fixture["project_id"]
        artifact_id = fixture["artifact_id"]
        root_handoff_id = fixture["root_handoff_id"]
        audit_id = fixture["audit_id"]
        audit_session = fixture["audit_session"]
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])
        pending = self._leave_prepared_reaudit_successor(
            fixture, failed_semantic
        )
        self.assertEqual(pending["status"], "prepared")
        pending_id = str(pending["handoff_id"])
        route_trace_id = str(pending["reaudit_route_trace_id"])
        event_trace_id = str(pending["reaudit_event_trace_id"])

        # Narrow the persisted state to the more severe signed-first crash:
        # signed route exists, Hub and route telemetry do not.
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (hub, _ledger):
            hub["handoffs"] = [
                row
                for row in hub["handoffs"]
                if row.get("handoff_id") != pending_id
            ]
        route_trace_path = (
            self.service.store.project_dir(self.namespace, project_id)
            / "telemetry"
            / f"{route_trace_id}.json"
        )
        route_trace_path.unlink()

        _server, restarted = create_mcp(self.settings)
        recovered = restarted.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=root_handoff_id,
            force_machine_recheck=True,
            semantic_audit=failed_semantic,
            request_session_id=audit_session,
        )
        self.assertEqual(recovered.status, "fail_current")
        self.assertEqual(recovered.data["handoff"]["handoff_id"], pending_id)
        completed = restarted.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=pending_id,
        )
        self.assertEqual(completed["status"], "consumed")
        self.assertEqual(completed["reaudit_route_trace_id"], route_trace_id)
        self.assertEqual(completed["reaudit_event_trace_id"], event_trace_id)
        all_successors = [
            row
            for row in restarted.store.list_project_records_strict(
                self.namespace, project_id, kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id") == root_handoff_id
        ]
        self.assertEqual([row["handoff_id"] for row in all_successors], [pending_id])
        hub_successors = [
            row
            for row in restarted.store.load_hub(
                self.namespace, project_id
            )["handoffs"]
            if row.get("handoff_id") == pending_id
        ]
        self.assertEqual(
            hub_successors,
            [restarted._handoff_record_projection(completed)],
        )
        child_events = [
            row
            for row in restarted.store.load_ledger(
                self.namespace, project_id
            )["events"]
            if row.get("handoff_id") == pending_id
        ]
        child_certificates = [
            row
            for row in restarted.store.list_project_records_strict(
                self.namespace, project_id, kind="certificates"
            )
            if row.get("handoff_id") == pending_id
        ]
        self.assertEqual(len(child_events), 1)
        self.assertEqual(len(child_certificates), 1)
        self.assertEqual(child_events[0]["stage_trace_id"], event_trace_id)
        self.assertEqual(
            child_certificates[0]["production_trace_id"], route_trace_id
        )
        self.assertEqual(
            restarted.store.get_project_record(
                self.namespace,
                project_id,
                kind="telemetry",
                record_id=route_trace_id,
            )["operation"],
            "audit_storyboard_reaudit_route",
        )
        root_replay = restarted.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=root_handoff_id,
            force_machine_recheck=False,
            semantic_audit=None,
            request_session_id=audit_session,
        )
        self.assertEqual(root_replay.status, "fail_current")
        self.assertEqual(root_replay.data["handoff"]["handoff_id"], pending_id)
        self.assertEqual(
            root_replay.data["certificate"]["certificate_id"],
            child_certificates[0]["certificate_id"],
        )
        self.assertEqual(
            next(
                row
                for row in restarted.store.load_hub(
                    self.namespace, project_id
                )["handoffs"]
                if row.get("handoff_id") == pending_id
            ),
            restarted._handoff_record_projection(completed),
        )

    def test_restart_recovers_exact_prerelease_half_write(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-prerelease-recovery")
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])
        released, consuming = self._leave_prerelease_half_write(
            fixture, failed_semantic
        )
        successor_id = str(released["handoff_id"])

        _server, restarted = create_mcp(self.settings)
        recovered = restarted.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=failed_semantic,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(recovered.status, "fail_current")
        self.assertEqual(
            recovered.data["handoff"]["handoff_id"], successor_id
        )
        completed = restarted.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=successor_id,
        )
        self.assertEqual(completed["status"], "consumed")
        self.assertEqual(
            completed["reaudit_route_trace_id"],
            consuming["reaudit_route_trace_id"],
        )
        self.assertEqual(
            completed["reaudit_event_trace_id"],
            consuming["reaudit_event_trace_id"],
        )
        self.assertEqual(
            [
                row
                for row in restarted.store.load_ledger(
                    self.namespace, fixture["project_id"]
                )["events"]
                if row.get("handoff_id") == successor_id
            ].__len__(),
            1,
        )
        self.assertEqual(
            [
                row
                for row in restarted.store.list_project_records_strict(
                    self.namespace,
                    fixture["project_id"],
                    kind="certificates",
                )
                if row.get("handoff_id") == successor_id
            ].__len__(),
            1,
        )
        self.assertEqual(
            next(
                row
                for row in restarted.store.load_hub(
                    self.namespace, fixture["project_id"]
                )["handoffs"]
                if row.get("handoff_id") == successor_id
            ),
            restarted._handoff_record_projection(completed),
        )

    def test_prerelease_half_write_tamper_rejects_without_mutation(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-prerelease-tamper")
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])
        released, _consuming = self._leave_prerelease_half_write(
            fixture, failed_semantic
        )
        successor_id = str(released["handoff_id"])
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (hub, _ledger):
            successor_projection = next(
                row
                for row in hub["handoffs"]
                if row.get("handoff_id") == successor_id
            )
            successor_projection["claim_generation"] = (
                int(successor_projection.get("claim_generation") or 0) + 99
            )
        before = self._authority_snapshot(fixture["project_id"])

        _server, restarted = create_mcp(self.settings)
        with self.assertRaisesRegex(
            ValueError, "reaudit_successor_projection_mismatch"
        ):
            restarted.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=True,
                semantic_audit=failed_semantic,
                request_session_id=fixture["audit_session"],
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), before
        )
        with self.assertRaisesRegex(
            ValueError, "reaudit_successor_projection_mismatch"
        ):
            restarted.store.rollback_reaudit_successor(
                self.namespace,
                fixture["project_id"],
                successor_id,
                route_trace_id=str(released["reaudit_route_trace_id"]),
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), before
        )

    def test_rollback_telemetry_unlink_failure_is_retryable(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-rollback-trace-retry")
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])
        prepared = self._leave_prepared_reaudit_successor(
            fixture, failed_semantic
        )
        successor_id = str(prepared["handoff_id"])
        route_trace_id = str(prepared["reaudit_route_trace_id"])
        trace_path = (
            self.service.store.project_dir(
                self.namespace, fixture["project_id"]
            )
            / "telemetry"
            / f"{route_trace_id}.json"
        )
        before = self._authority_snapshot(fixture["project_id"])
        original_unlink = Path.unlink

        def fail_trace_unlink(path, *args, **kwargs):
            if path == trace_path:
                raise OSError("simulated route telemetry unlink failure")
            return original_unlink(path, *args, **kwargs)

        with patch.object(Path, "unlink", new=fail_trace_unlink):
            with self.assertRaisesRegex(OSError, "telemetry unlink"):
                self.service.store.rollback_reaudit_successor(
                    self.namespace,
                    fixture["project_id"],
                    successor_id,
                    route_trace_id=route_trace_id,
                )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), before
        )

        self.service.store.rollback_reaudit_successor(
            self.namespace,
            fixture["project_id"],
            successor_id,
            route_trace_id=route_trace_id,
        )
        self.assertFalse(trace_path.exists())
        self.assertFalse(
            any(
                row.get("handoff_id") == successor_id
                for row in self.service.store.load_hub(
                    self.namespace, fixture["project_id"]
                )["handoffs"]
            )
        )
        self.assertFalse(
            (
                self.service.store.project_dir(
                    self.namespace, fixture["project_id"]
                )
                / "handoffs"
                / f"{successor_id}.json"
            ).exists()
        )
        self.service.store.rollback_reaudit_successor(
            self.namespace,
            fixture["project_id"],
            successor_id,
            route_trace_id=route_trace_id,
        )

    def test_consuming_preclaim_half_write_waits_for_lease_then_resumes(self) -> None:
        project_id = self._project("reaudit-consuming-lease-recovery")
        production_id, production_session = self._begin_production(
            project_id, "reaudit-consuming-lease-recovery"
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            "reaudit-consuming-lease-recovery",
        )
        artifact_id = str(committed.artifact_id)
        root_handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, "reaudit-consuming-lease-recovery"
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        self.service.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=root_handoff_id,
            semantic_audit=clean_semantic,
            request_session_id=audit_session,
        )
        failed_semantic = self._failed_semantic(clean_semantic)
        with patch.object(
            self.service.store,
            "claim_reaudit_successor",
            side_effect=RuntimeError("simulated loss before first claim"),
        ), patch.object(
            self.service.store,
            "rollback_reaudit_successor",
            return_value=None,
        ):
            with self.assertRaisesRegex(RuntimeError, "before first claim"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=project_id,
                    artifact_id=artifact_id,
                    audit_cycle_id=None,
                    work_order_id=audit_id,
                    handoff_artifact_id=root_handoff_id,
                    force_machine_recheck=True,
                    semantic_audit=failed_semantic,
                    request_session_id=audit_session,
                )
        prepared = next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id") == root_handoff_id
        )
        prepared_id = str(prepared["handoff_id"])
        prepared_projection = self.service._handoff_record_projection(
            prepared
        )
        consuming = self.service.store.claim_reaudit_successor(
            self.namespace,
            project_id,
            prepared_id,
            consumer_work_order_id=audit_id,
            lease_seconds=1,
        )
        self.assertEqual(consuming["status"], "consuming")
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (hub, _ledger):
            hub["handoffs"] = [
                copy.deepcopy(prepared_projection)
                if row.get("handoff_id") == prepared_id
                else row
                for row in hub["handoffs"]
            ]

        # A live lease cannot be stolen.  The request may repair the exact Hub
        # pre-claim half-write, but it creates no event/certificate and keeps
        # the original fenced claim.
        with self.assertRaisesRegex(ValueError, "being consumed"):
            self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=root_handoff_id,
                force_machine_recheck=True,
                semantic_audit=failed_semantic,
                request_session_id=audit_session,
            )
        repaired_consuming = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=prepared_id,
        )
        self.assertEqual(repaired_consuming, consuming)
        self.assertEqual(
            next(
                row
                for row in self.service.store.load_hub(
                    self.namespace, project_id
                )["handoffs"]
                if row.get("handoff_id") == prepared_id
            ),
            self.service._handoff_record_projection(consuming),
        )
        self.assertFalse(
            any(
                row.get("handoff_id") == prepared_id
                for row in self.service.store.load_ledger(
                    self.namespace, project_id
                )["events"]
            )
        )
        self.assertFalse(
            any(
                row.get("handoff_id") == prepared_id
                for row in self.service.store.list_project_records_strict(
                    self.namespace, project_id, kind="certificates"
                )
            )
        )

        after_expiry = float(
            consuming["claim_lease_expires_epoch_seconds"]
        ) + 1.0
        with patch.object(storage_module.time, "time", return_value=after_expiry):
            recovered = self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=root_handoff_id,
                force_machine_recheck=True,
                semantic_audit=failed_semantic,
                request_session_id=audit_session,
            )
        self.assertEqual(recovered.status, "fail_current")
        self.assertEqual(recovered.data["handoff"]["handoff_id"], prepared_id)
        completed = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=prepared_id,
        )
        self.assertEqual(completed["status"], "consumed")
        self.assertEqual(completed["claim_generation"], 2)
        self.assertEqual(completed["claim_history"][-1]["state"], "lease_expired")
        self.assertEqual(
            len(
                [
                    row
                    for row in self.service.store.load_ledger(
                        self.namespace, project_id
                    )["events"]
                    if row.get("handoff_id") == prepared_id
                ]
            ),
            1,
        )

    def test_restart_recovers_event_only_reaudit_with_exact_semantic_identity(self) -> None:
        project_id = self._project("reaudit-incomplete-event-certificate")
        production_id, production_session = self._begin_production(
            project_id, "reaudit-incomplete-event-certificate"
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            "reaudit-incomplete-event-certificate",
        )
        artifact_id = str(committed.artifact_id)
        root_handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, "reaudit-incomplete-event-certificate"
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        first = self.service.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=root_handoff_id,
            semantic_audit=clean_semantic,
            request_session_id=audit_session,
        )
        failed_semantic = self._failed_semantic(clean_semantic)
        with patch.object(
            self.service,
            "_persist_audit_certificate",
            side_effect=RuntimeError("simulated loss before certificate"),
        ), patch.object(
            self.service.store,
            "release_reaudit_successor_claim",
            return_value=None,
        ), patch.object(
            self.service.store,
            "rollback_reaudit_successor",
            return_value=None,
        ):
            with self.assertRaisesRegex(RuntimeError, "before certificate"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=project_id,
                    artifact_id=artifact_id,
                    audit_cycle_id=None,
                    work_order_id=audit_id,
                    handoff_artifact_id=root_handoff_id,
                    force_machine_recheck=True,
                    semantic_audit=failed_semantic,
                    request_session_id=audit_session,
                )
        pending = next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id") == root_handoff_id
        )
        pending_id = str(pending["handoff_id"])
        child_events = [
            row
            for row in self.service.store.load_ledger(
                self.namespace, project_id
            )["events"]
            if row.get("handoff_id") == pending_id
        ]
        self.assertEqual(len(child_events), 1)
        self.assertFalse(
            any(
                row.get("handoff_id") == pending_id
                for row in self.service.store.list_project_records_strict(
                    self.namespace, project_id, kind="certificates"
                )
            )
        )
        event_only_before = self._authority_snapshot(project_id)
        _server, restarted = create_mcp(self.settings)

        different_semantic = copy.deepcopy(failed_semantic)
        different_semantic["issues"][0]["evidence"] = (
            "不同的独立复审证据不得接管已落盘事件。"
        )
        with self.assertRaisesRegex(
            ValueError, "reaudit_event_recovery_mismatch"
        ):
            restarted.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=root_handoff_id,
                force_machine_recheck=True,
                semantic_audit=different_semantic,
                request_session_id=audit_session,
            )
        self.assertEqual(self._authority_snapshot(project_id), event_only_before)

        original_event = copy.deepcopy(child_events[0])
        variants: list[tuple[str, dict[str, Any], bool]] = []
        for label, field, value in (
            ("source-version", "source_version", "tampered-source-version"),
            ("file-name", "file_name", "storyboard-independent_audit-CHANGED"),
            ("episode", "episode_id", "EP-TAMPERED"),
            ("round", "audit_round", int(original_event["audit_round"]) + 1),
            ("handoff", "handoff_id", "HND-TAMPERED-EVENT"),
        ):
            variant = copy.deepcopy(original_event)
            variant[field] = value
            variants.append((label, variant, False))
        issue_variant = copy.deepcopy(original_event)
        issue_variant["issues"][0]["evidence"] += "（已篡改）"
        variants.append(("issues", issue_variant, False))
        summary_variant = copy.deepcopy(original_event)
        summary_variant["audit_summary"]["blocking_issue_count"] = 0
        variants.append(("audit-summary", summary_variant, False))
        quality_variant = copy.deepcopy(original_event)
        quality_variant["first_draft_quality"][
            "direct_usability_percent"
        ] = 99.0
        variants.append(("first-quality", quality_variant, False))
        blockers_variant = copy.deepcopy(original_event)
        blockers_variant["blockers"] = ["E-RECOVERY-TAMPER-001"]
        variants.append(("blockers", blockers_variant, False))
        counts_variant = copy.deepcopy(original_event)
        counts_variant["counts"]["blocking_error_count"] += 1
        variants.append(("counts", counts_variant, False))
        unknown_variant = copy.deepcopy(original_event)
        unknown_variant["unexpected_authority"] = "not-produced"
        variants.append(("unknown-field", unknown_variant, False))
        scope_variant = copy.deepcopy(original_event)
        scope_variant["scope_key"] = "SCOPE-TAMPERED-PAIR"
        variants.append(("scope-key-pair", scope_variant, True))

        for label, variant, alter_scope in variants:
            with self.subTest(event_field=label):
                with self.service.store.mutate_project(
                    self.namespace, project_id
                ) as (_hub, ledger):
                    ledger["events"] = [
                        copy.deepcopy(variant)
                        if row.get("event_id") == original_event["event_id"]
                        else row
                        for row in ledger["events"]
                    ]
                    if alter_scope:
                        next(
                            row
                            for row in ledger["scope_states"]
                            if row.get("last_event_id")
                            == original_event["event_id"]
                        )["scope_key"] = variant["scope_key"]
                tampered_before = self._authority_snapshot(project_id)
                with self.assertRaisesRegex(
                    ValueError,
                    "reaudit_recovery_invalid|reaudit_event_recovery_mismatch",
                ):
                    restarted.audit_storyboard(
                        self.namespace,
                        project_id=project_id,
                        artifact_id=artifact_id,
                        audit_cycle_id=None,
                        work_order_id=audit_id,
                        handoff_artifact_id=root_handoff_id,
                        force_machine_recheck=True,
                        semantic_audit=failed_semantic,
                        request_session_id=audit_session,
                    )
                self.assertEqual(
                    self._authority_snapshot(project_id), tampered_before
                )
                with self.service.store.mutate_project(
                    self.namespace, project_id
                ) as (_hub, ledger):
                    ledger["events"] = [
                        copy.deepcopy(original_event)
                        if row.get("event_id") == original_event["event_id"]
                        else row
                        for row in ledger["events"]
                    ]
                    for scope in ledger["scope_states"]:
                        if scope.get("last_event_id") == original_event["event_id"]:
                            scope["scope_key"] = original_event["scope_key"]

        live_before = self._authority_snapshot(project_id)
        with self.assertRaisesRegex(ValueError, "successor is being consumed"):
            restarted.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=root_handoff_id,
                force_machine_recheck=True,
                semantic_audit=failed_semantic,
                request_session_id=audit_session,
            )
        self.assertEqual(self._authority_snapshot(project_id), live_before)
        after_expiry = float(
            pending["claim_lease_expires_epoch_seconds"]
        ) + 1.0
        with patch.object(
            storage_module.time, "time", return_value=after_expiry
        ):
            recovered = restarted.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=root_handoff_id,
                force_machine_recheck=True,
                semantic_audit=failed_semantic,
                request_session_id=audit_session,
            )
        self.assertEqual(recovered.status, "fail_current")
        self.assertEqual(
            recovered.data["certificate"]["event_id"],
            original_event["event_id"],
        )
        self.assertEqual(
            len(
                [
                    row
                    for row in restarted.store.load_ledger(
                        self.namespace, project_id
                    )["events"]
                    if row.get("handoff_id") == pending_id
                ]
            ),
            1,
        )
        self.assertEqual(
            len(
                [
                    row
                    for row in restarted.store.list_project_records_strict(
                        self.namespace, project_id, kind="certificates"
                    )
                    if row.get("handoff_id") == pending_id
                ]
            ),
            1,
        )
        stable = self._authority_snapshot(project_id)
        replay = restarted.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=root_handoff_id,
            force_machine_recheck=True,
            semantic_audit=failed_semantic,
            request_session_id=audit_session,
        )
        self.assertEqual(replay.status, "fail_current")
        self.assertEqual(self._authority_snapshot(project_id), stable)

    def test_restart_rejects_certificate_only_reaudit_state(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-certificate-only")
        pending = self._leave_prepared_reaudit_successor(
            fixture, self._failed_semantic(fixture["clean_semantic"])
        )
        pending_id = str(pending["handoff_id"])
        root_certificate = fixture["first"].data["certificate"]
        certificate_only = copy.deepcopy(root_certificate)
        for key in ("record_id", "project_id", "integrity_hmac"):
            certificate_only.pop(key, None)
        certificate_only.update(
            {
                "certificate_id": "CERT-REAUDIT-CERTIFICATE-ONLY",
                "handoff_id": pending_id,
                "event_id": "EVT-REAUDIT-CERTIFICATE-ONLY",
                "stage_trace_id": pending["reaudit_event_trace_id"],
                "production_trace_id": pending["reaudit_route_trace_id"],
                "production_stage_trace_id": pending[
                    "production_stage_trace_id"
                ],
            }
        )
        self.service.store.save_project_record(
            self.namespace,
            fixture["project_id"],
            kind="certificates",
            record_id="CERT-REAUDIT-CERTIFICATE-ONLY",
            value=certificate_only,
        )
        certificate_only_before = self._authority_snapshot(
            fixture["project_id"]
        )
        with self.assertRaisesRegex(
            ValueError, "reaudit_recovery_incomplete"
        ):
            self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=True,
                semantic_audit=self._failed_semantic(
                    fixture["clean_semantic"]
                ),
                request_session_id=fixture["audit_session"],
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]),
            certificate_only_before,
        )

    def test_checkpointed_legacy_complete_pair_recovers_but_partial_hash_rejects(self) -> None:
        from test_v055_quality_migration import (
            V055QualityAndMigrationTests,
        )

        fixture = self._passed_audit_fixture("reaudit-legacy-pair")
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])
        with patch.object(
            service_module,
            "_export_compact_storyboard",
            side_effect=RuntimeError("simulated legacy pair boundary"),
        ):
            with self.assertRaisesRegex(RuntimeError, "legacy pair"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=failed_semantic,
                    request_session_id=fixture["audit_session"],
                )
        pending = next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, fixture["project_id"], kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id")
            == fixture["root_handoff_id"]
        )
        pending_id = str(pending["handoff_id"])
        event = next(
            row
            for row in self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"]
            if row.get("handoff_id") == pending_id
        )
        original_event = copy.deepcopy(event)
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (_hub, ledger):
            target = next(
                row
                for row in ledger["events"]
                if row.get("event_id") == event["event_id"]
            )
            target["semantic_report_input_sha256"] = "b" * 64
        mismatch_before = self._authority_snapshot(fixture["project_id"])
        with self.assertRaisesRegex(
            ValueError,
            "reaudit_event_recovery_mismatch|reaudit_recovery_invalid",
        ):
            self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=True,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]),
            mismatch_before,
        )
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (_hub, ledger):
            ledger["events"] = [
                copy.deepcopy(original_event)
                if row.get("event_id") == original_event["event_id"]
                else row
                for row in ledger["events"]
            ]

        migration_session = "v051-legacy-pair-migration"
        migration_order = self.service.begin_task(
            self.namespace,
            project_id=fixture["project_id"],
            operation="version_management",
            mode="execute",
            authorization="explicit_execute",
            request_text="封存旧独审完整证据对",
            targets=["0.5.4→0.5.5"],
            authoritative_source="签名迁移台账",
            latest_version="0.5.4",
            target_boundary="项目控制态",
            must_keep=["历史独审事件与证书"],
            must_remove=[],
            protected_scope=[],
            expected_outputs=["签名replay checkpoint"],
            window_key=migration_session,
            request_session_id=migration_session,
        )
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (_hub, ledger):
            target = next(
                row
                for row in ledger["events"]
                if row.get("event_id") == original_event["event_id"]
            )
            target.pop("semantic_audit_sha256", None)
            target.pop("semantic_report_input_sha256", None)
        legacy_event = next(
            row
            for row in self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"]
            if row.get("event_id") == original_event["event_id"]
        )
        pending_certificate = next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace,
                fixture["project_id"],
                kind="certificates",
            )
            if row.get("handoff_id") == pending_id
        )
        pending_certificate["event_sha256"] = (
            service_module.sha256_bytes(
                service_module._stable_json_bytes(legacy_event)
            )
        )
        self.service.store.save_project_record(
            self.namespace,
            fixture["project_id"],
            kind="certificates",
            record_id=str(pending_certificate["certificate_id"]),
            value=pending_certificate,
        )
        V055QualityAndMigrationTests._make_legacy_state(
            self.service, self.namespace, fixture["project_id"]
        )
        migrated = self.service.migrate_project_state(
            self.namespace,
            project_id=fixture["project_id"],
            work_order_id=str(migration_order.work_order_id),
            action="execute",
            target_version="0.5.5",
            expected_current_version="0.5.4",
            migration_id=None,
            request_session_id=migration_session,
        )
        self.assertEqual(migrated.status, "migrated")
        checkpoint = self.service.store.load_ledger(
            self.namespace, fixture["project_id"]
        )["replay_checkpoint"]
        self.assertGreaterEqual(
            int(checkpoint["event_count"]),
            int(original_event["sequence"]),
        )
        _server, restarted = create_mcp(self.settings)
        recovered = restarted.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=None,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(recovered.status, "fail_current")
        self.assertEqual(
            recovered.data["handoff"]["handoff_id"], pending_id
        )

    def test_restart_finishes_unique_existing_event_and_certificate(self) -> None:
        project_id = self._project("reaudit-restart-after-certificate")
        production_id, production_session = self._begin_production(
            project_id, "reaudit-restart-after-certificate"
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            "reaudit-restart-after-certificate",
        )
        artifact_id = str(committed.artifact_id)
        root_handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, "reaudit-restart-after-certificate"
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        self.service.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=root_handoff_id,
            semantic_audit=clean_semantic,
            request_session_id=audit_session,
        )
        failed_semantic = self._failed_semantic(clean_semantic)

        # The audit event, certificate, overlay, telemetry and Hub updates are
        # already durable when final handoff consumption begins.  Suppressing
        # handled compensation models a process death at that exact boundary.
        with patch.object(
            self.service.store,
            "finish_reaudit_successor",
            side_effect=RuntimeError("simulated process loss before finish"),
        ), patch.object(
            self.service.store,
            "release_reaudit_successor_claim",
            return_value=None,
        ), patch.object(
            self.service.store,
            "rollback_reaudit_successor",
            return_value=None,
        ):
            with self.assertRaisesRegex(RuntimeError, "before finish"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=project_id,
                    artifact_id=artifact_id,
                    audit_cycle_id=None,
                    work_order_id=audit_id,
                    handoff_artifact_id=root_handoff_id,
                    force_machine_recheck=True,
                    semantic_audit=failed_semantic,
                    request_session_id=audit_session,
                )
        pending = next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id") == root_handoff_id
        )
        pending_id = str(pending["handoff_id"])
        self.assertEqual(pending["status"], "consuming")
        events_before = [
            row
            for row in self.service.store.load_ledger(
                self.namespace, project_id
            )["events"]
            if row.get("handoff_id") == pending_id
        ]
        certificates_before = [
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, project_id, kind="certificates"
            )
            if row.get("handoff_id") == pending_id
        ]
        self.assertEqual(len(events_before), 1)
        self.assertEqual(len(certificates_before), 1)
        certificate_id = str(certificates_before[0]["certificate_id"])
        audit_trace_id = str(pending["reaudit_event_trace_id"])
        trace_before = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="telemetry",
            record_id=audit_trace_id,
        )

        _server, restarted = create_mcp(self.settings)
        live_before = self._authority_snapshot(project_id)
        with self.assertRaisesRegex(ValueError, "successor is being consumed"):
            restarted.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=root_handoff_id,
                force_machine_recheck=True,
                semantic_audit=None,
                request_session_id=audit_session,
            )
        self.assertEqual(self._authority_snapshot(project_id), live_before)
        after_expiry = float(
            pending["claim_lease_expires_epoch_seconds"]
        ) + 1.0
        with patch.object(
            storage_module.time, "time", return_value=after_expiry
        ):
            recovered = restarted.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=root_handoff_id,
                force_machine_recheck=True,
                semantic_audit=None,
                request_session_id=audit_session,
            )
        self.assertEqual(recovered.status, "fail_current")
        self.assertTrue(recovered.data["cache_hit"])
        self.assertEqual(
            recovered.data["certificate"]["certificate_id"], certificate_id
        )
        completed = restarted.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=pending_id,
        )
        self.assertEqual(completed["status"], "consumed")
        self.assertEqual(completed["audit_certificate_id"], certificate_id)
        events_after = [
            row
            for row in restarted.store.load_ledger(
                self.namespace, project_id
            )["events"]
            if row.get("handoff_id") == pending_id
        ]
        certificates_after = [
            row
            for row in restarted.store.list_project_records_strict(
                self.namespace, project_id, kind="certificates"
            )
            if row.get("handoff_id") == pending_id
        ]
        self.assertEqual(events_after, events_before)
        self.assertEqual(certificates_after, certificates_before)
        self.assertEqual(
            restarted.store.get_project_record(
                self.namespace,
                project_id,
                kind="telemetry",
                record_id=audit_trace_id,
            ),
            trace_before,
        )
        self.assertEqual(
            sum(
                row.get("telemetry_id") == audit_trace_id
                for row in restarted.store.load_ledger(
                    self.namespace, project_id
                )["performance_traces"]
            ),
            1,
        )

    def test_post_certificate_failures_resume_same_successor(self) -> None:
        for fault in (
            "certificate-postwrite",
            "overlay",
            "telemetry",
            "metadata",
            "download",
            "finish",
        ):
            with self.subTest(fault=fault):
                fixture = self._passed_audit_fixture(
                    f"reaudit-post-certificate-{fault}"
                )
                failed_semantic = self._failed_semantic(
                    fixture["clean_semantic"]
                )
                error = f"simulated post-certificate {fault} failure"
                if fault == "certificate-postwrite":
                    original_persist_certificate = (
                        self.service._persist_audit_certificate
                    )

                    def persist_then_fail(*args, **kwargs):
                        original_persist_certificate(*args, **kwargs)
                        raise RuntimeError(error)

                    fault_patcher = patch.object(
                        self.service,
                        "_persist_audit_certificate",
                        side_effect=persist_then_fail,
                    )
                elif fault == "overlay":
                    fault_patcher = patch.object(
                        service_module,
                        "_export_compact_storyboard",
                        side_effect=RuntimeError(error),
                    )
                elif fault == "telemetry":
                    original_save_telemetry = self.service._save_telemetry

                    def fail_audit_telemetry(*args, **kwargs):
                        if kwargs.get("operation") == "audit_storyboard":
                            raise RuntimeError(error)
                        return original_save_telemetry(*args, **kwargs)

                    fault_patcher = patch.object(
                        self.service,
                        "_save_telemetry",
                        side_effect=fail_audit_telemetry,
                    )
                elif fault == "metadata":
                    fault_patcher = patch.object(
                        self.service.store,
                        "update_artifact_metadata",
                        side_effect=RuntimeError(error),
                    )
                elif fault == "download":
                    fault_patcher = patch.object(
                        self.service.store,
                        "signed_download_url",
                        side_effect=RuntimeError(error),
                    )
                else:
                    fault_patcher = patch.object(
                        self.service.store,
                        "finish_reaudit_successor",
                        side_effect=RuntimeError(error),
                    )

                with fault_patcher:
                    with self.assertRaisesRegex(RuntimeError, fault):
                        self.service.audit_storyboard(
                            self.namespace,
                            project_id=fixture["project_id"],
                            artifact_id=fixture["artifact_id"],
                            audit_cycle_id=None,
                            work_order_id=fixture["audit_id"],
                            handoff_artifact_id=fixture["root_handoff_id"],
                            force_machine_recheck=True,
                            semantic_audit=failed_semantic,
                            request_session_id=fixture["audit_session"],
                        )
                pending = next(
                    row
                    for row in self.service.store.list_project_records_strict(
                        self.namespace,
                        fixture["project_id"],
                        kind="handoffs",
                    )
                    if row.get("reaudit_root_handoff_id")
                    == fixture["root_handoff_id"]
                )
                successor_id = str(pending["handoff_id"])
                self.assertEqual(pending["status"], "prepared")
                events_before = [
                    row
                    for row in self.service.store.load_ledger(
                        self.namespace, fixture["project_id"]
                    )["events"]
                    if row.get("handoff_id") == successor_id
                ]
                certificates_before = [
                    row
                    for row in self.service.store.list_project_records_strict(
                        self.namespace,
                        fixture["project_id"],
                        kind="certificates",
                    )
                    if row.get("handoff_id") == successor_id
                ]
                self.assertEqual(len(events_before), 1)
                self.assertEqual(len(certificates_before), 1)

                _server, restarted = create_mcp(self.settings)
                recovered = restarted.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=failed_semantic,
                    request_session_id=fixture["audit_session"],
                )
                self.assertEqual(recovered.status, "fail_current")
                self.assertEqual(
                    recovered.data["handoff"]["handoff_id"], successor_id
                )
                completed = restarted.store.get_project_record(
                    self.namespace,
                    fixture["project_id"],
                    kind="handoffs",
                    record_id=successor_id,
                )
                self.assertEqual(completed["status"], "consumed")
                self.assertEqual(
                    [
                        row
                        for row in restarted.store.load_ledger(
                            self.namespace, fixture["project_id"]
                        )["events"]
                        if row.get("handoff_id") == successor_id
                    ],
                    events_before,
                )
                self.assertEqual(
                    [
                        row
                        for row in restarted.store.list_project_records_strict(
                            self.namespace,
                            fixture["project_id"],
                            kind="certificates",
                        )
                        if row.get("handoff_id") == successor_id
                    ],
                    certificates_before,
                )
                self.assertEqual(
                    next(
                        row
                        for row in restarted.store.load_hub(
                            self.namespace, fixture["project_id"]
                        )["handoffs"]
                        if row.get("handoff_id") == successor_id
                    ),
                    restarted._handoff_record_projection(completed),
                )

    def test_completed_force_retry_and_ordinary_replay_are_authority_noops(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-completion-receipt")
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])

        # Model loss after successor consumption but before the response is
        # returned.  The signed certificate is the durable completion receipt.
        with patch.object(
            service_module,
            "ToolResponse",
            side_effect=RuntimeError("simulated post-finish response loss"),
        ):
            with self.assertRaisesRegex(RuntimeError, "post-finish"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=failed_semantic,
                    request_session_id=fixture["audit_session"],
                )
        completed = next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, fixture["project_id"], kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id")
            == fixture["root_handoff_id"]
        )
        successor_id = str(completed["handoff_id"])
        self.assertEqual(completed["status"], "consumed")
        completed_before = self._authority_snapshot(fixture["project_id"])

        _server, restarted = create_mcp(self.settings)
        force_replay = restarted.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=failed_semantic,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(force_replay.status, "fail_current")
        self.assertTrue(force_replay.data["cache_hit"])
        self.assertEqual(
            force_replay.data["handoff"]["handoff_id"], successor_id
        )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]),
            completed_before,
        )

        for _index in range(2):
            ordinary = restarted.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )
            self.assertEqual(ordinary.status, "fail_current")
            self.assertTrue(ordinary.data["cache_hit"])
            self.assertEqual(
                ordinary.data["handoff"]["handoff_id"], successor_id
            )
            self.assertEqual(
                self._authority_snapshot(fixture["project_id"]),
                completed_before,
            )
        overlay_id = str(
            force_replay.data["audited_workbook"]["artifact_id"]
        )
        self.service.store.update_artifact_metadata(
            self.namespace,
            fixture["project_id"],
            overlay_id,
            {"audit_status": "TAMPERED"},
        )
        tampered_before = self._authority_snapshot(fixture["project_id"])
        with self.assertRaisesRegex(
            ValueError, "authoritative_audit_overlay_metadata_binding_invalid"
        ):
            restarted.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=True,
                semantic_audit=failed_semantic,
                request_session_id=fixture["audit_session"],
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]),
            tampered_before,
        )

    def test_prepared_event_certificate_pair_has_one_concurrent_finisher(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-pair-concurrent")
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])
        with patch.object(
            service_module,
            "_export_compact_storyboard",
            side_effect=RuntimeError("simulated overlay interruption"),
        ):
            with self.assertRaisesRegex(RuntimeError, "overlay interruption"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=failed_semantic,
                    request_session_id=fixture["audit_session"],
                )
        successor = next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, fixture["project_id"], kind="handoffs"
            )
            if row.get("reaudit_root_handoff_id")
            == fixture["root_handoff_id"]
        )
        successor_id = str(successor["handoff_id"])
        self.assertEqual(successor["status"], "prepared")
        certificate = next(
            row
            for row in self.service.store.list_project_records_strict(
                self.namespace, fixture["project_id"], kind="certificates"
            )
            if row.get("handoff_id") == successor_id
        )
        original_claim = self.service.store.claim_reaudit_successor
        first_claimed = threading.Event()
        release_first = threading.Event()
        claim_guard = threading.Lock()
        first_seen = False

        def fenced_claim(*args: Any, **kwargs: Any) -> dict[str, Any]:
            nonlocal first_seen
            try:
                claimed = original_claim(*args, **kwargs)
            except ValueError:
                raise
            with claim_guard:
                is_first = not first_seen
                if is_first:
                    first_seen = True
            if is_first:
                first_claimed.set()
                if not release_first.wait(timeout=10):
                    raise RuntimeError("concurrent finisher release timeout")
            return claimed

        def retry() -> Any:
            return self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=True,
                semantic_audit=failed_semantic,
                request_session_id=fixture["audit_session"],
            )

        with patch.object(
            self.service.store,
            "claim_reaudit_successor",
            side_effect=fenced_claim,
        ), ThreadPoolExecutor(max_workers=1) as executor:
            first_future = executor.submit(retry)
            self.assertTrue(first_claimed.wait(timeout=10))
            claimed_snapshot = self._project_raw_snapshot(
                fixture["project_id"]
            )
            with self.assertRaisesRegex(
                ValueError, "authority_operation_in_progress"
            ):
                retry()
            self.assertEqual(
                self._project_raw_snapshot(fixture["project_id"]),
                claimed_snapshot,
            )
            release_first.set()
            first = first_future.result(timeout=15)
        self.assertEqual(first.status, "fail_current")
        completed_snapshot = self._project_raw_snapshot(
            fixture["project_id"]
        )
        replay = retry()
        self.assertEqual(replay.status, "fail_current")
        self.assertTrue(replay.data["cache_hit"])
        self.assertEqual(
            replay.data["handoff"]["handoff_id"], successor_id
        )
        self.assertEqual(
            self._project_raw_snapshot(fixture["project_id"]),
            completed_snapshot,
        )
        completed = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=successor_id,
        )
        self.assertEqual(completed["status"], "consumed")
        self.assertEqual(
            len(
                [
                    row
                    for row in self.service.store.load_ledger(
                        self.namespace, fixture["project_id"]
                    )["events"]
                    if row.get("handoff_id") == successor_id
                ]
            ),
            1,
        )
        self.assertEqual(
            len(
                [
                    row
                    for row in self.service.store.list_project_records_strict(
                        self.namespace,
                        fixture["project_id"],
                        kind="certificates",
                    )
                    if row.get("handoff_id") == successor_id
                ]
            ),
            1,
        )
        self.assertEqual(
            len(
                [
                    row
                    for row in self.service.store.list_artifacts(
                        self.namespace, fixture["project_id"]
                    )
                    if row.get("artifact_type")
                    == "storyboard_audit_overlay"
                    and row.get("metadata", {}).get(
                        "independent_audit_certificate_id"
                    )
                    == certificate["certificate_id"]
                ]
            ),
            1,
        )

    def test_stale_root_replay_cannot_overwrite_new_successor_fail(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-stale-root-replay")
        root_certificate_id = fixture["first"].data["certificate"][
            "certificate_id"
        ]
        entered = threading.Event()
        release = threading.Event()
        original_response = self.service._completed_audit_response

        def paused_response(*args: Any, **kwargs: Any) -> Any:
            if kwargs["certificate"]["certificate_id"] == root_certificate_id:
                entered.set()
                if not release.wait(timeout=10):
                    raise RuntimeError("stale replay release timeout")
            return original_response(*args, **kwargs)

        def stale_replay() -> Any:
            return self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )

        with patch.object(
            self.service,
            "_completed_audit_response",
            side_effect=paused_response,
        ), ThreadPoolExecutor(max_workers=1) as executor:
            stale_future = executor.submit(stale_replay)
            self.assertTrue(entered.wait(timeout=10))
            replay_snapshot = self._authority_snapshot(
                fixture["project_id"]
            )
            with self.assertRaisesRegex(
                ValueError, "authority_operation_in_progress"
            ):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=self._failed_semantic(
                        fixture["clean_semantic"]
                    ),
                    request_session_id=fixture["audit_session"],
                )
            self.assertEqual(
                self._authority_snapshot(fixture["project_id"]),
                replay_snapshot,
            )
            release.set()
            replayed = stale_future.result(timeout=15)
            self.assertEqual(replayed.status, "pass")
        latest = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=self._failed_semantic(
                fixture["clean_semantic"]
            ),
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(latest.status, "fail_current")
        latest_snapshot = self._authority_snapshot(fixture["project_id"])
        replayed_latest = stale_replay()
        self.assertEqual(replayed_latest.status, "fail_current")
        self.assertTrue(replayed_latest.data["cache_hit"])
        self.assertEqual(
            replayed_latest.data["certificate"]["certificate_id"],
            latest.data["certificate"]["certificate_id"],
        )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]),
            latest_snapshot,
        )
        hub = self.service.store.load_hub(
            self.namespace, fixture["project_id"]
        )
        self.assertEqual(
            hub["project_state"]["current_stage"], "failed_current"
        )
        source_record, _path = self.service.store.get_artifact(
            self.namespace,
            fixture["project_id"],
            fixture["artifact_id"],
        )
        self.assertEqual(
            source_record["metadata"]["independent_audit_certificate_id"],
            latest.data["certificate"]["certificate_id"],
        )

    def test_completed_replay_linearizes_after_response_construction(self) -> None:
        fixture = self._passed_audit_fixture(
            "reaudit-late-completed-replay"
        )
        root_certificate_id = str(
            fixture["first"].data["certificate"]["certificate_id"]
        )
        entered = threading.Event()
        release = threading.Event()
        original_assert_latest = (
            self.service.store.assert_latest_reaudit_receipt
        )

        def pause_final_receipt_check(
            *args: Any, **kwargs: Any
        ) -> None:
            if kwargs.get("certificate_id") == root_certificate_id:
                entered.set()
                if not release.wait(timeout=15):
                    raise RuntimeError(
                        "late completed replay release timeout"
                    )
            original_assert_latest(*args, **kwargs)

        def replay_root() -> Any:
            return self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )

        with patch.object(
            self.service.store,
            "assert_latest_reaudit_receipt",
            side_effect=pause_final_receipt_check,
        ), ThreadPoolExecutor(max_workers=1) as executor:
            stale_future = executor.submit(replay_root)
            self.assertTrue(entered.wait(timeout=15))
            paused_raw = self._project_raw_snapshot(
                fixture["project_id"]
            )
            try:
                with self.assertRaisesRegex(
                    ValueError, "authority_operation_in_progress"
                ):
                    self.service.audit_storyboard(
                        self.namespace,
                        project_id=fixture["project_id"],
                        artifact_id=fixture["artifact_id"],
                        audit_cycle_id=None,
                        work_order_id=fixture["audit_id"],
                        handoff_artifact_id=fixture["root_handoff_id"],
                        force_machine_recheck=True,
                        semantic_audit=self._failed_semantic(
                            fixture["clean_semantic"]
                        ),
                        request_session_id=fixture["audit_session"],
                    )
                self.assertEqual(
                    self._project_raw_snapshot(fixture["project_id"]),
                    paused_raw,
                )
            finally:
                release.set()
            replayed = stale_future.result(timeout=20)
            self.assertEqual(replayed.status, "pass")
            latest = self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=True,
                semantic_audit=self._failed_semantic(
                    fixture["clean_semantic"]
                ),
                request_session_id=fixture["audit_session"],
            )
            self.assertEqual(latest.status, "fail_current")
            latest_raw = self._project_raw_snapshot(
                fixture["project_id"]
            )
            replayed_latest = replay_root()
            self.assertEqual(replayed_latest.status, "fail_current")
            self.assertTrue(replayed_latest.data["cache_hit"])
            self.assertEqual(
                replayed_latest.data["certificate"]["certificate_id"],
                latest.data["certificate"]["certificate_id"],
            )
        self.assertEqual(
            self._project_raw_snapshot(fixture["project_id"]),
            latest_raw,
        )

    def test_completed_replay_rejects_late_resigned_root_replacement(self) -> None:
        fixture = self._passed_audit_fixture(
            "reaudit-late-resigned-root-replacement"
        )
        latest = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=self._failed_semantic(fixture["clean_semantic"]),
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(latest.status, "fail_current")
        latest_certificate_id = str(
            latest.data["certificate"]["certificate_id"]
        )
        entered = threading.Event()
        release = threading.Event()
        original_assert_latest = (
            self.service.store.assert_latest_reaudit_receipt
        )

        def pause_final_receipt_check(
            *args: Any, **kwargs: Any
        ) -> None:
            if kwargs.get("certificate_id") == latest_certificate_id:
                entered.set()
                if not release.wait(timeout=15):
                    raise RuntimeError(
                        "resigned root replacement release timeout"
                    )
            original_assert_latest(*args, **kwargs)

        def replay_latest() -> Any:
            return self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )

        with patch.object(
            self.service.store,
            "assert_latest_reaudit_receipt",
            side_effect=pause_final_receipt_check,
        ), ThreadPoolExecutor(max_workers=1) as executor:
            replay_future = executor.submit(replay_latest)
            self.assertTrue(entered.wait(timeout=15))
            root = self.service.store.get_project_record(
                self.namespace,
                fixture["project_id"],
                kind="handoffs",
                record_id=fixture["root_handoff_id"],
            )
            tampered_root = copy.deepcopy(root)
            tampered_root["scope"]["target_ids"] = [
                "EP-TAMPERED-AFTER-RESPONSE"
            ]
            resigned_root = self.service.store.save_project_record(
                self.namespace,
                fixture["project_id"],
                kind="handoffs",
                record_id=fixture["root_handoff_id"],
                value=tampered_root,
            )
            with self.service.store.mutate_project(
                self.namespace, fixture["project_id"]
            ) as (hub, _ledger):
                hub["handoffs"] = [
                    self.service._handoff_record_projection(resigned_root)
                    if row.get("handoff_id") == fixture["root_handoff_id"]
                    else row
                    for row in hub.get("handoffs", [])
                ]
            tampered_raw = self._project_raw_snapshot(
                fixture["project_id"]
            )
            release.set()
            with self.assertRaisesRegex(
                ValueError, "reaudit_replay_superseded"
            ):
                replay_future.result(timeout=20)
        self.assertEqual(
            self._project_raw_snapshot(fixture["project_id"]),
            tampered_raw,
        )

    def test_completed_replay_rejects_late_event_replacement(self) -> None:
        fixture = self._passed_audit_fixture(
            "reaudit-late-event-replacement"
        )
        certificate = fixture["first"].data["certificate"]
        certificate_id = str(certificate["certificate_id"])
        event_id = str(certificate["event_id"])
        entered = threading.Event()
        release = threading.Event()
        original_assert_latest = (
            self.service.store.assert_latest_reaudit_receipt
        )

        def pause_final_receipt_check(
            *args: Any, **kwargs: Any
        ) -> None:
            if kwargs.get("certificate_id") == certificate_id:
                entered.set()
                if not release.wait(timeout=15):
                    raise RuntimeError(
                        "late event replacement release timeout"
                    )
            original_assert_latest(*args, **kwargs)

        def replay_completed() -> Any:
            return self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )

        with patch.object(
            self.service.store,
            "assert_latest_reaudit_receipt",
            side_effect=pause_final_receipt_check,
        ), ThreadPoolExecutor(max_workers=1) as executor:
            replay_future = executor.submit(replay_completed)
            self.assertTrue(entered.wait(timeout=15))
            with self.service.store.mutate_project(
                self.namespace, fixture["project_id"]
            ) as (_hub, ledger):
                event = next(
                    row
                    for row in ledger["events"]
                    if row.get("event_id") == event_id
                )
                event["timestamp"] = "2026-08-31T23:59:59Z"
            tampered_raw = self._project_raw_snapshot(
                fixture["project_id"]
            )
            release.set()
            with self.assertRaisesRegex(
                ValueError, "reaudit_replay_superseded"
            ):
                replay_future.result(timeout=20)

        self.assertEqual(
            self._project_raw_snapshot(fixture["project_id"]),
            tampered_raw,
        )

    def test_completed_replay_rejects_late_resigned_intermediate_successor(self) -> None:
        fixture = self._passed_audit_fixture(
            "reaudit-late-resigned-intermediate-successor"
        )
        first_successor = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=self._failed_semantic(fixture["clean_semantic"]),
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(first_successor.status, "fail_current")
        second_failed_semantic = self._failed_semantic(
            fixture["clean_semantic"]
        )
        second_failed_semantic["issues"][0]["evidence"] = (
            "第二轮独立复核以新证据确认同一阻塞仍未修复。"
        )
        latest = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=second_failed_semantic,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(latest.status, "fail_current")
        latest_certificate_id = str(
            latest.data["certificate"]["certificate_id"]
        )
        intermediate_id = str(
            first_successor.data["handoff"]["handoff_id"]
        )
        entered = threading.Event()
        release = threading.Event()
        original_assert_latest = (
            self.service.store.assert_latest_reaudit_receipt
        )

        def pause_final_receipt_check(
            *args: Any, **kwargs: Any
        ) -> None:
            if kwargs.get("certificate_id") == latest_certificate_id:
                entered.set()
                if not release.wait(timeout=15):
                    raise RuntimeError(
                        "resigned intermediate release timeout"
                    )
            original_assert_latest(*args, **kwargs)

        def replay_latest() -> Any:
            return self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )

        with patch.object(
            self.service.store,
            "assert_latest_reaudit_receipt",
            side_effect=pause_final_receipt_check,
        ), ThreadPoolExecutor(max_workers=1) as executor:
            replay_future = executor.submit(replay_latest)
            self.assertTrue(entered.wait(timeout=15))
            intermediate = self.service.store.get_project_record(
                self.namespace,
                fixture["project_id"],
                kind="handoffs",
                record_id=intermediate_id,
            )
            tampered_intermediate = copy.deepcopy(intermediate)
            tampered_intermediate["scope"]["target_ids"] = [
                "EP-TAMPERED-INTERMEDIATE"
            ]
            resigned_intermediate = self.service.store.save_project_record(
                self.namespace,
                fixture["project_id"],
                kind="handoffs",
                record_id=intermediate_id,
                value=tampered_intermediate,
            )
            with self.service.store.mutate_project(
                self.namespace, fixture["project_id"]
            ) as (hub, _ledger):
                hub["handoffs"] = [
                    self.service._handoff_record_projection(
                        resigned_intermediate
                    )
                    if row.get("handoff_id") == intermediate_id
                    else row
                    for row in hub.get("handoffs", [])
                ]
            tampered_raw = self._project_raw_snapshot(
                fixture["project_id"]
            )
            release.set()
            with self.assertRaisesRegex(
                ValueError, "reaudit_replay_superseded"
            ):
                replay_future.result(timeout=20)
        self.assertEqual(
            self._project_raw_snapshot(fixture["project_id"]),
            tampered_raw,
        )

    def test_completed_force_receipt_normalizes_canonical_sha256_case(self) -> None:
        for verdict in ("pass", "fail_current"):
            with self.subTest(verdict=verdict):
                fixture = self._passed_audit_fixture(
                    f"reaudit-receipt-sha-case-{verdict}"
                )
                semantic = (
                    copy.deepcopy(fixture["clean_semantic"])
                    if verdict == "pass"
                    else self._failed_semantic(
                        fixture["clean_semantic"]
                    )
                )
                completed = self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=semantic,
                    request_session_id=fixture["audit_session"],
                )
                self.assertEqual(completed.status, verdict)
                completed_before = self._authority_snapshot(
                    fixture["project_id"]
                )
                uppercase_semantic = copy.deepcopy(semantic)
                uppercase_semantic["canonical_sha256"] = str(
                    uppercase_semantic["canonical_sha256"]
                ).upper()

                replay = self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=uppercase_semantic,
                    request_session_id=fixture["audit_session"],
                )
                self.assertEqual(replay.status, verdict)
                self.assertTrue(replay.data["cache_hit"])
                self.assertEqual(
                    replay.data["certificate"]["certificate_id"],
                    completed.data["certificate"]["certificate_id"],
                )
                self.assertEqual(
                    replay.data["handoff"]["handoff_id"],
                    completed.data["handoff"]["handoff_id"],
                )
                self.assertEqual(
                    replay.data["audit"]["event"]["event_id"],
                    completed.data["audit"]["event"]["event_id"],
                )
                self.assertEqual(
                    self._authority_snapshot(fixture["project_id"]),
                    completed_before,
                )

    def test_repair_route_commit_rejects_resigned_certificate_after_validation(self) -> None:
        for field in ("audit_actor_binding_id", "issue_count"):
            with self.subTest(certificate_field=field):
                fixture = self._unaudited_fixture(
                    f"repair-route-certificate-cas-{field}"
                )
                failed_semantic = self._failed_semantic(
                    fixture["clean_semantic"]
                )

                def publication_tail() -> dict[str, Any]:
                    return {
                        "repair_packages": (
                            self.service.store.list_project_records_strict(
                                self.namespace,
                                fixture["project_id"],
                                kind="repair_packages",
                            )
                        ),
                        "repair_handoffs": [
                            row
                            for row in self.service.store.list_project_records_strict(
                                self.namespace,
                                fixture["project_id"],
                                kind="handoffs",
                            )
                            if row.get("handoff_type") == "audit_to_repair"
                        ],
                        "overlays": [
                            row
                            for row in self.service.store.list_artifacts(
                                self.namespace, fixture["project_id"]
                            )
                            if row.get("artifact_type")
                            == "storyboard_audit_overlay"
                        ],
                        "telemetry": (
                            self.service.store.list_project_records_strict(
                                self.namespace,
                                fixture["project_id"],
                                kind="telemetry",
                            )
                        ),
                    }

                before_publication = publication_tail()
                original_commit = (
                    self.service.store.commit_repair_package_route
                )
                storage_rejected = False

                def resign_then_commit(
                    *args: Any, **kwargs: Any
                ) -> tuple[dict[str, Any], dict[str, Any]]:
                    nonlocal storage_rejected
                    certificate_id = str(
                        kwargs["expected_latest_certificate_id"]
                    )
                    certificate = self.service.store.get_project_record(
                        self.namespace,
                        fixture["project_id"],
                        kind="certificates",
                        record_id=certificate_id,
                    )
                    if field == "audit_actor_binding_id":
                        certificate[field] = (
                            str(certificate.get(field) or "")
                            + "-RESIGNED"
                        )
                    else:
                        certificate[field] = int(
                            certificate.get(field) or 0
                        ) + 1
                    certificate["integrity_hmac"] = (
                        self.service.store._record_integrity_hmac(
                            certificate
                        )
                    )
                    certificate_path = (
                        self.service.store.project_dir(
                            self.namespace, fixture["project_id"]
                        )
                        / "certificates"
                        / f"{certificate_id}.json"
                    )
                    self.service.store.atomic_write_json(
                        certificate_path, certificate
                    )
                    before_storage = self._project_raw_snapshot(
                        fixture["project_id"],
                        exclude=(certificate_path,),
                    )
                    try:
                        return original_commit(*args, **kwargs)
                    except ValueError as exc:
                        self.assertRegex(
                            str(exc), "repair_route_freshness_invalid"
                        )
                        self.assertEqual(
                            self._project_raw_snapshot(
                                fixture["project_id"],
                                exclude=(certificate_path,),
                            ),
                            before_storage,
                        )
                        storage_rejected = True
                        raise

                with patch.object(
                    self.service.store,
                    "commit_repair_package_route",
                    side_effect=resign_then_commit,
                ):
                    with self.assertRaisesRegex(
                        ValueError, "repair_route_freshness_invalid"
                    ):
                        self.service.audit_storyboard(
                            self.namespace,
                            project_id=fixture["project_id"],
                            artifact_id=fixture["artifact_id"],
                            audit_cycle_id=None,
                            work_order_id=fixture["audit_id"],
                            handoff_artifact_id=fixture[
                                "root_handoff_id"
                            ],
                            semantic_audit=failed_semantic,
                            request_session_id=fixture["audit_session"],
                        )
                self.assertTrue(storage_rejected)
                self.assertEqual(publication_tail(), before_publication)

    def test_completed_fail_replay_rejects_resigned_route_misbindings(self) -> None:
        for record_kind, field in (
            ("repair_packages", "audit_work_order_id"),
            ("handoffs", "from_actor_binding_id"),
        ):
            with self.subTest(record_kind=record_kind, field=field):
                fixture = self._unaudited_fixture(
                    f"completed-route-misbinding-{record_kind}-{field}"
                )
                completed = self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    semantic_audit=self._failed_semantic(
                        fixture["clean_semantic"]
                    ),
                    request_session_id=fixture["audit_session"],
                )
                self.assertEqual(completed.status, "fail_current")
                package_id = str(
                    completed.data["repair_package"]["repair_package_id"]
                )
                repair_handoff_id = str(
                    completed.data["repair_package"][
                        "audit_to_repair_handoff_id"
                    ]
                )
                record_id = (
                    package_id
                    if record_kind == "repair_packages"
                    else repair_handoff_id
                )
                record = self.service.store.get_project_record(
                    self.namespace,
                    fixture["project_id"],
                    kind=record_kind,
                    record_id=record_id,
                )
                tampered = copy.deepcopy(record)
                tampered[field] = "VALID-HMAC-BUT-WRONG-BINDING"
                resigned = self.service.store.save_project_record(
                    self.namespace,
                    fixture["project_id"],
                    kind=record_kind,
                    record_id=record_id,
                    value=tampered,
                )
                with self.service.store.mutate_project(
                    self.namespace, fixture["project_id"]
                ) as (hub, _ledger):
                    collection = (
                        "repair_packages"
                        if record_kind == "repair_packages"
                        else "handoffs"
                    )
                    identity_field = (
                        "repair_package_id"
                        if record_kind == "repair_packages"
                        else "handoff_id"
                    )
                    hub[collection] = [
                        copy.deepcopy(resigned)
                        if row.get(identity_field) == record_id
                        else row
                        for row in hub.get(collection, [])
                    ]
                tampered_before = self._authority_snapshot(
                    fixture["project_id"]
                )
                with self.assertRaisesRegex(
                    ValueError, "reaudit_recovery_invalid"
                ):
                    self.service.audit_storyboard(
                        self.namespace,
                        project_id=fixture["project_id"],
                        artifact_id=fixture["artifact_id"],
                        audit_cycle_id=None,
                        work_order_id=fixture["audit_id"],
                        handoff_artifact_id=fixture["root_handoff_id"],
                        force_machine_recheck=False,
                        semantic_audit=None,
                        request_session_id=fixture["audit_session"],
                    )
                self.assertEqual(
                    self._authority_snapshot(fixture["project_id"]),
                    tampered_before,
                )

    def test_completed_fail_replay_rejects_resigned_fake_consumed_package(
        self,
    ) -> None:
        fixture = self._unaudited_fixture(
            "completed-fake-consumed-repair-package"
        )
        completed = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            semantic_audit=self._failed_semantic(
                fixture["clean_semantic"]
            ),
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(completed.status, "fail_current")
        package_id = str(
            completed.data["repair_package"]["repair_package_id"]
        )
        package = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="repair_packages",
            record_id=package_id,
        )
        fake_consumed = {
            **copy.deepcopy(package),
            "status": "consumed",
            "consumed_by_work_order_id": "WO-FAKE-REPAIR",
            "successor_artifact_id": "ART-FAKE-SUCCESSOR",
            "successor_canonical_sha256": "f" * 64,
            "repair_to_audit_handoff_id": "HND-FAKE-RETURN",
        }
        resigned = self.service.store.save_project_record(
            self.namespace,
            fixture["project_id"],
            kind="repair_packages",
            record_id=package_id,
            value=fake_consumed,
        )
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (hub, _ledger):
            hub["repair_packages"] = [
                copy.deepcopy(resigned)
                if row.get("repair_package_id") == package_id
                else row
                for row in hub.get("repair_packages", [])
            ]
        tampered_before = self._authority_snapshot(
            fixture["project_id"]
        )

        with self.assertRaisesRegex(
            ValueError, "reaudit_recovery_invalid"
        ):
            self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]),
            tampered_before,
        )

    def test_root_fail_completed_ordinary_replay_is_read_only(self) -> None:
        fixture = self._unaudited_fixture(
            "root-fail-completed-ordinary-replay"
        )
        failed_semantic = self._failed_semantic(
            fixture["clean_semantic"]
        )
        completed = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            semantic_audit=failed_semantic,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(completed.status, "fail_current")
        completed_before = self._authority_snapshot(
            fixture["project_id"]
        )

        replay = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=False,
            semantic_audit=None,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(replay.status, "fail_current")
        self.assertTrue(replay.data["cache_hit"])
        self.assertEqual(
            replay.data["handoff"]["handoff_id"],
            completed.data["handoff"]["handoff_id"],
        )
        self.assertEqual(
            replay.data["audit"]["event"]["event_id"],
            completed.data["audit"]["event"]["event_id"],
        )
        self.assertEqual(
            replay.data["certificate"]["certificate_id"],
            completed.data["certificate"]["certificate_id"],
        )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]),
            completed_before,
        )

    def test_root_fail_missing_repair_route_recovers_same_event_certificate(self) -> None:
        fixture = self._unaudited_fixture(
            "root-fail-missing-repair-route"
        )
        completed = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            semantic_audit=self._failed_semantic(
                fixture["clean_semantic"]
            ),
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(completed.status, "fail_current")
        event_id = str(completed.data["audit"]["event"]["event_id"])
        certificate_id = str(
            completed.data["certificate"]["certificate_id"]
        )
        package_id = str(
            completed.data["repair_package"]["repair_package_id"]
        )
        repair_handoff_id = str(
            completed.data["repair_package"][
                "audit_to_repair_handoff_id"
            ]
        )
        events_before = copy.deepcopy(
            self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"]
        )
        certificates_before = (
            self.service.store.list_project_records_strict(
                self.namespace,
                fixture["project_id"],
                kind="certificates",
            )
        )
        project_dir = self.service.store.project_dir(
            self.namespace, fixture["project_id"]
        )
        package_path = (
            project_dir / "repair_packages" / f"{package_id}.json"
        )
        repair_handoff_path = (
            project_dir / "handoffs" / f"{repair_handoff_id}.json"
        )
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (hub, _ledger):
            hub["repair_packages"] = [
                row
                for row in hub.get("repair_packages", [])
                if row.get("repair_package_id") != package_id
            ]
            hub["handoffs"] = [
                row
                for row in hub.get("handoffs", [])
                if row.get("handoff_id") != repair_handoff_id
            ]
        package_path.unlink()
        repair_handoff_path.unlink()

        recovered = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=False,
            semantic_audit=None,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(recovered.status, "fail_current")
        self.assertTrue(recovered.data["cache_hit"])
        self.assertEqual(
            recovered.data["audit"]["event"]["event_id"], event_id
        )
        self.assertEqual(
            recovered.data["certificate"]["certificate_id"],
            certificate_id,
        )
        self.assertEqual(
            recovered.data["repair_package"]["repair_package_id"],
            package_id,
        )
        self.assertEqual(
            recovered.data["repair_package"][
                "audit_to_repair_handoff_id"
            ],
            repair_handoff_id,
        )
        self.assertTrue(package_path.is_file())
        self.assertTrue(repair_handoff_path.is_file())
        self.assertEqual(
            self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"],
            events_before,
        )
        self.assertEqual(
            self.service.store.list_project_records_strict(
                self.namespace,
                fixture["project_id"],
                kind="certificates",
            ),
            certificates_before,
        )
        hub = self.service.store.load_hub(
            self.namespace, fixture["project_id"]
        )
        self.assertEqual(
            sum(
                row.get("repair_package_id") == package_id
                for row in hub.get("repair_packages", [])
            ),
            1,
        )
        self.assertEqual(
            sum(
                row.get("handoff_id") == repair_handoff_id
                for row in hub.get("handoffs", [])
            ),
            1,
        )

    def test_incomplete_fail_projection_recovery_is_superseded_atomically(self) -> None:
        for incomplete_state in ("hub_missing", "package_only"):
            with self.subTest(incomplete_state=incomplete_state):
                fixture = self._unaudited_fixture(
                    f"repair-route-stale-recovery-{incomplete_state}"
                )
                old_semantic = self._failed_semantic(
                    fixture["clean_semantic"]
                )
                old = self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    semantic_audit=old_semantic,
                    request_session_id=fixture["audit_session"],
                )
                self.assertEqual(old.status, "fail_current")
                old_certificate_id = str(
                    old.data["certificate"]["certificate_id"]
                )
                old_package = old.data["repair_package"]
                old_package_id = str(
                    old_package["repair_package_id"]
                )
                old_repair_handoff_id = str(
                    old_package["audit_to_repair_handoff_id"]
                )
                project_dir = self.service.store.project_dir(
                    self.namespace, fixture["project_id"]
                )
                old_package_path = (
                    project_dir
                    / "repair_packages"
                    / f"{old_package_id}.json"
                )
                old_repair_handoff_path = (
                    project_dir
                    / "handoffs"
                    / f"{old_repair_handoff_id}.json"
                )
                old_package_bytes = old_package_path.read_bytes()
                with self.service.store.mutate_project(
                    self.namespace, fixture["project_id"]
                ) as (hub, _ledger):
                    hub["repair_packages"] = [
                        row
                        for row in hub.get("repair_packages", [])
                        if row.get("repair_package_id")
                        != old_package_id
                    ]
                    hub["handoffs"] = [
                        row
                        for row in hub.get("handoffs", [])
                        if row.get("handoff_id")
                        != old_repair_handoff_id
                    ]
                if incomplete_state == "package_only":
                    old_repair_handoff_path.unlink()
                self.assertEqual(
                    old_package_path.read_bytes(), old_package_bytes
                )

                commit_entered = threading.Event()
                release_old_commit = threading.Event()
                original_commit = (
                    self.service.store.commit_repair_package_route
                )

                def pause_old_repair_commit(
                    *args: Any, **kwargs: Any
                ) -> tuple[dict[str, Any], dict[str, Any]]:
                    if (
                        kwargs.get("expected_latest_certificate_id")
                        == old_certificate_id
                    ):
                        commit_entered.set()
                        if not release_old_commit.wait(timeout=20):
                            raise RuntimeError(
                                "stale repair projection release timeout"
                            )
                    return original_commit(*args, **kwargs)

                def recover_old_projection() -> Any:
                    return self.service.audit_storyboard(
                        self.namespace,
                        project_id=fixture["project_id"],
                        artifact_id=fixture["artifact_id"],
                        audit_cycle_id=None,
                        work_order_id=fixture["audit_id"],
                        handoff_artifact_id=fixture[
                            "root_handoff_id"
                        ],
                        force_machine_recheck=False,
                        semantic_audit=None,
                        request_session_id=fixture["audit_session"],
                    )

                newer_semantic = copy.deepcopy(old_semantic)
                newer_semantic["issues"][0].update(
                    {
                        "issue_id": (
                            "SEM-NEW-FAIL-SUPERSEDING-"
                            + incomplete_state.upper()
                        ),
                        "evidence": (
                            "后续独立复核形成新的阻塞证据："
                            + incomplete_state
                        ),
                    }
                )
                newer_semantic["layers"][0]["summary"] = (
                    "后续独立复核确认新的阻塞证据。"
                )
                newer_semantic["conclusion"]["summary"] = (
                    "后续独立复核未通过。"
                )

                with patch.object(
                    self.service.store,
                    "commit_repair_package_route",
                    side_effect=pause_old_repair_commit,
                ), ThreadPoolExecutor(max_workers=1) as executor:
                    stale_future = executor.submit(
                        recover_old_projection
                    )
                    self.assertTrue(commit_entered.wait(timeout=15))
                    paused_raw = self._project_raw_snapshot(
                        fixture["project_id"]
                    )
                    try:
                        with self.assertRaisesRegex(
                            ValueError,
                            "authority_operation_in_progress",
                        ):
                            self.service.audit_storyboard(
                                self.namespace,
                                project_id=fixture["project_id"],
                                artifact_id=fixture["artifact_id"],
                                audit_cycle_id=None,
                                work_order_id=fixture["audit_id"],
                                handoff_artifact_id=fixture[
                                    "root_handoff_id"
                                ],
                                force_machine_recheck=True,
                                semantic_audit=newer_semantic,
                                request_session_id=fixture[
                                    "audit_session"
                                ],
                            )
                        self.assertEqual(
                            self._project_raw_snapshot(
                                fixture["project_id"]
                            ),
                            paused_raw,
                        )
                    finally:
                        release_old_commit.set()
                    recovered = stale_future.result(timeout=20)
                    self.assertEqual(recovered.status, "fail_current")
                    self.assertEqual(
                        recovered.data["certificate"]["certificate_id"],
                        old_certificate_id,
                    )
                newer = self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=newer_semantic,
                    request_session_id=fixture["audit_session"],
                )
                self.assertEqual(newer.status, "fail_current")
                self.assertNotEqual(
                    newer.data["certificate"]["certificate_id"],
                    old_certificate_id,
                )
                newer_raw = self._project_raw_snapshot(
                    fixture["project_id"]
                )
                replayed = recover_old_projection()
                self.assertEqual(replayed.status, "fail_current")
                self.assertTrue(replayed.data["cache_hit"])
                self.assertEqual(
                    replayed.data["certificate"]["certificate_id"],
                    newer.data["certificate"]["certificate_id"],
                )
                self.assertEqual(
                    self._project_raw_snapshot(fixture["project_id"]),
                    newer_raw,
                )

    def test_same_second_pass_successor_is_latest_signoff_authority(self) -> None:
        fixed_timestamp = "2026-08-31T12:00:00Z"
        with patch.object(
            service_module, "utc_stamp", return_value=fixed_timestamp
        ):
            fixture = self._passed_audit_fixture(
                "reaudit-pass-successor-signoff"
            )
            second_semantic = copy.deepcopy(fixture["clean_semantic"])
            second_semantic["layers"][0]["summary"] += (
                "｜第二次独立复核使用新证据。"
            )
            second = self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=True,
                semantic_audit=second_semantic,
                request_session_id=fixture["audit_session"],
            )
        self.assertEqual(second.status, "pass")
        first_certificate_id = str(
            fixture["first"].data["certificate"]["certificate_id"]
        )
        latest_certificate_id = str(
            second.data["certificate"]["certificate_id"]
        )
        self.assertNotEqual(first_certificate_id, latest_certificate_id)
        first_certificate = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="certificates",
            record_id=first_certificate_id,
        )
        latest_certificate = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="certificates",
            record_id=latest_certificate_id,
        )
        self.assertEqual(
            first_certificate["created_at"], latest_certificate["created_at"]
        )
        self.assertEqual(first_certificate["created_at"], fixed_timestamp)

        # If the latest certificate is unavailable, sign-off must fail closed
        # instead of falling back to the older clean root certificate.
        original_get_record = self.service.store.get_project_record

        def reject_latest_certificate(*args, **kwargs):
            if (
                kwargs.get("kind") == "certificates"
                and kwargs.get("record_id") == latest_certificate_id
            ):
                raise ValueError("latest successor certificate unavailable")
            return original_get_record(*args, **kwargs)

        before_rejection = self._authority_snapshot(fixture["project_id"])
        with patch.object(
            self.service.store,
            "get_project_record",
            side_effect=reject_latest_certificate,
        ):
            with self.assertRaisesRegex(
                ValueError, "independent_audit_required"
            ):
                self.service.signoff_task(
                    self.namespace,
                    project_id=fixture["project_id"],
                    work_order_id=fixture["production_id"],
                    decision="approve",
                    signer="successor signoff negative probe",
                )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), before_rejection
        )

        signed = self.service.signoff_task(
            self.namespace,
            project_id=fixture["project_id"],
            work_order_id=fixture["production_id"],
            decision="approve",
            signer="successor signoff regression",
        )
        self.assertEqual(signed.status, "archived")
        self.assertEqual(
            signed.data["independent_audit_certificate_id"],
            latest_certificate_id,
        )
        ledger = self.service.store.load_ledger(
            self.namespace, fixture["project_id"]
        )
        latest_scope = next(
            row
            for row in ledger["scope_states"]
            if row.get("last_event_id") == latest_certificate["event_id"]
        )
        self.assertEqual(latest_scope["closure_state"], "human_signed_off")
        self.assertEqual(
            ledger["human_signoffs"][-1]["scope_key"],
            latest_scope["scope_key"],
        )

    def test_legacy_certificate_requires_modern_successor_bridge(
        self,
    ) -> None:
        fixture = self._passed_audit_fixture(
            "legacy-certificate-successor-bridge"
        )
        root_certificate_id = str(
            fixture["first"].data["certificate"]["certificate_id"]
        )
        root_certificate = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="certificates",
            record_id=root_certificate_id,
        )
        root_event = next(
            row
            for row in self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"]
            if row.get("event_id") == root_certificate["event_id"]
        )
        legacy_certificate = copy.deepcopy(root_certificate)
        legacy_certificate.pop("event_sha256")
        legacy_certificate = self.service.store.save_project_record(
            self.namespace,
            fixture["project_id"],
            kind="certificates",
            record_id=root_certificate_id,
            value=legacy_certificate,
        )
        legacy_before = self._authority_snapshot(fixture["project_id"])

        with self.assertRaisesRegex(
            ValueError, "fresh successor audit"
        ):
            self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )
        with self.assertRaisesRegex(
            ValueError, "independent_audit_required"
        ):
            self.service.signoff_task(
                self.namespace,
                project_id=fixture["project_id"],
                work_order_id=fixture["production_id"],
                decision="approve",
                signer="legacy certificate negative probe",
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]),
            legacy_before,
        )

        successor_semantic = copy.deepcopy(fixture["clean_semantic"])
        successor_semantic["layers"][0]["summary"] += (
            "｜升级后使用新证据完成独立复核。"
        )
        successor = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=successor_semantic,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(successor.status, "pass")
        successor_handoff = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=str(successor.data["handoff"]["handoff_id"]),
        )
        self.assertEqual(
            successor_handoff["previous_audit_certificate_id"],
            root_certificate_id,
        )
        self.assertEqual(
            successor_handoff["previous_audit_certificate_hmac"],
            legacy_certificate["integrity_hmac"],
        )
        self.assertEqual(
            successor_handoff["previous_audit_event_sha256"],
            service_module.sha256_bytes(
                service_module._stable_json_bytes(root_event)
            ),
        )

        stable = self._authority_snapshot(fixture["project_id"])
        replay = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=False,
            semantic_audit=None,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(replay.status, "pass")
        self.assertTrue(replay.data["cache_hit"])
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), stable
        )
        signed = self.service.signoff_task(
            self.namespace,
            project_id=fixture["project_id"],
            work_order_id=fixture["production_id"],
            decision="approve",
            signer="legacy bridge signoff regression",
        )
        self.assertEqual(signed.status, "archived")
        self.assertEqual(
            signed.data["independent_audit_certificate_id"],
            successor.data["certificate"]["certificate_id"],
        )

    def test_signoff_rejects_resigned_certificate_path_identity_mismatch(
        self,
    ) -> None:
        fixture = self._passed_audit_fixture(
            "reaudit-certificate-path-identity"
        )
        certificate_id = str(
            fixture["first"].data["certificate"]["certificate_id"]
        )
        certificate_path = (
            self.service.store.project_dir(
                self.namespace, fixture["project_id"]
            )
            / "certificates"
            / f"{certificate_id}.json"
        )
        certificate = self.service.store.read_json(certificate_path)
        forged_id = "CERT-VALID-HMAC-WRONG-ID"
        certificate["record_id"] = forged_id
        certificate["certificate_id"] = forged_id
        certificate["integrity_hmac"] = (
            self.service.store._record_integrity_hmac(certificate)
        )
        self.service.store.atomic_write_json(certificate_path, certificate)
        before = self._project_raw_snapshot(fixture["project_id"])

        with self.assertRaisesRegex(
            ValueError, "project_record_identity_mismatch"
        ):
            self.service.store.get_project_record(
                self.namespace,
                fixture["project_id"],
                kind="certificates",
                record_id=certificate_id,
            )
        with self.assertRaisesRegex(
            ValueError, "independent_audit_required"
        ):
            self.service.signoff_task(
                self.namespace,
                project_id=fixture["project_id"],
                work_order_id=fixture["production_id"],
                decision="approve",
                signer="certificate path identity negative probe",
            )

        self.assertEqual(
            self._project_raw_snapshot(fixture["project_id"]), before
        )

    def test_signoff_rejects_resigned_inconsistent_ancestor_certificate(
        self,
    ) -> None:
        for field, forged_value in (
            ("result", "fail_current"),
            ("issue_count", 1),
        ):
            with self.subTest(field=field):
                fixture = self._passed_audit_fixture(
                    f"reaudit-ancestor-certificate-{field}"
                )
                second_semantic = copy.deepcopy(
                    fixture["clean_semantic"]
                )
                second_semantic["layers"][0]["summary"] += (
                    "｜第二轮独立证据。"
                )
                second = self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=second_semantic,
                    request_session_id=fixture["audit_session"],
                )
                third_semantic = copy.deepcopy(
                    fixture["clean_semantic"]
                )
                third_semantic["layers"][0]["summary"] += (
                    "｜第三轮独立证据。"
                )
                third = self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    force_machine_recheck=True,
                    semantic_audit=third_semantic,
                    request_session_id=fixture["audit_session"],
                )
                self.assertEqual(second.status, "pass")
                self.assertEqual(third.status, "pass")
                ancestor_certificate_id = str(
                    second.data["certificate"]["certificate_id"]
                )
                ancestor_path = (
                    self.service.store.project_dir(
                        self.namespace, fixture["project_id"]
                    )
                    / "certificates"
                    / f"{ancestor_certificate_id}.json"
                )
                ancestor = self.service.store.read_json(ancestor_path)
                ancestor[field] = forged_value
                ancestor["integrity_hmac"] = (
                    self.service.store._record_integrity_hmac(ancestor)
                )
                self.service.store.atomic_write_json(
                    ancestor_path, ancestor
                )
                before = self._project_raw_snapshot(
                    fixture["project_id"]
                )

                with self.assertRaisesRegex(
                    ValueError, "independent_audit_required"
                ):
                    self.service.signoff_task(
                        self.namespace,
                        project_id=fixture["project_id"],
                        work_order_id=fixture["production_id"],
                        decision="approve",
                        signer="ancestor certificate negative probe",
                    )

                self.assertEqual(
                    self._project_raw_snapshot(fixture["project_id"]),
                    before,
                )

    def test_fail_repair_route_recovers_package_only_and_hub_missing_states(self) -> None:
        fixture = self._passed_audit_fixture("reaudit-repair-route-recovery")
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])
        failed = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=failed_semantic,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(failed.status, "fail_current")
        package = failed.data["repair_package"]
        self.assertIsInstance(package, dict)
        package_id = str(package["repair_package_id"])
        repair_handoff_id = str(package["audit_to_repair_handoff_id"])
        project_dir = self.service.store.project_dir(
            self.namespace, fixture["project_id"]
        )
        package_path = project_dir / "repair_packages" / f"{package_id}.json"
        repair_handoff_path = (
            project_dir / "handoffs" / f"{repair_handoff_id}.json"
        )
        package_bytes = package_path.read_bytes()
        events_before = copy.deepcopy(
            self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"]
        )
        certificates_before = self.service.store.list_project_records_strict(
            self.namespace, fixture["project_id"], kind="certificates"
        )

        # Hard crash after the signed package write but before the signed
        # audit-to-repair handoff or either Hub projection.
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (hub, _ledger):
            hub["repair_packages"] = [
                row
                for row in hub.get("repair_packages", [])
                if row.get("repair_package_id") != package_id
            ]
            hub["handoffs"] = [
                row
                for row in hub.get("handoffs", [])
                if row.get("handoff_id") != repair_handoff_id
            ]
        repair_handoff_path.unlink()
        recovered_package_only = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=False,
            semantic_audit=None,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(recovered_package_only.status, "fail_current")
        self.assertEqual(
            recovered_package_only.data["repair_package"]["repair_package_id"],
            package_id,
        )
        self.assertEqual(package_path.read_bytes(), package_bytes)
        self.assertTrue(repair_handoff_path.is_file())

        # A later crash may leave both exact signed records while Hub lacks
        # their projections.  Replay reconciles projections without a new
        # package, audit event, certificate, or successor.
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (hub, _ledger):
            hub["repair_packages"] = [
                row
                for row in hub.get("repair_packages", [])
                if row.get("repair_package_id") != package_id
            ]
            hub["handoffs"] = [
                row
                for row in hub.get("handoffs", [])
                if row.get("handoff_id") != repair_handoff_id
            ]
        recovered_hub = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=False,
            semantic_audit=None,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(recovered_hub.status, "fail_current")
        hub = self.service.store.load_hub(
            self.namespace, fixture["project_id"]
        )
        self.assertEqual(
            sum(
                row.get("repair_package_id") == package_id
                for row in hub["repair_packages"]
            ),
            1,
        )
        self.assertEqual(
            sum(
                row.get("handoff_id") == repair_handoff_id
                for row in hub["handoffs"]
            ),
            1,
        )
        self.assertEqual(
            self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"],
            events_before,
        )
        self.assertEqual(
            self.service.store.list_project_records_strict(
                self.namespace,
                fixture["project_id"],
                kind="certificates",
            ),
            certificates_before,
        )
        stable = self._authority_snapshot(fixture["project_id"])
        self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=False,
            semantic_audit=None,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), stable
        )

    def test_latest_successor_fail_supersedes_root_repair_package(self) -> None:
        fixture = self._unaudited_fixture("latest-repair-authority")
        first_semantic = self._failed_semantic(fixture["clean_semantic"])
        first = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            semantic_audit=first_semantic,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(first.status, "fail_current")
        old_package_id = str(
            first.data["repair_package"]["repair_package_id"]
        )

        second_semantic = copy.deepcopy(first_semantic)
        second_semantic["issues"][0].update(
            {
                "issue_id": "SEM-NEW-FAIL-002",
                "error_code": "E-SOURCE-SELECT-002",
                "evidence": "复审发现同一镜头存在不同的原文语义错位。",
                "repair_action": "按最新复审证据重建该镜后再次审计。",
            }
        )
        second_semantic["layers"][0]["summary"] = (
            "复审确认新的独立语义缺陷。"
        )
        second_semantic["conclusion"]["summary"] = (
            "复审确认新的阻塞问题。"
        )
        second = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=second_semantic,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(second.status, "fail_current")
        new_package_id = str(
            second.data["repair_package"]["repair_package_id"]
        )
        self.assertNotEqual(new_package_id, old_package_id)

        repair_session = "repair-latest-authority"
        repair = self.service.begin_task(
            self.namespace,
            project_id=fixture["project_id"],
            operation="storyboard_repair",
            mode="audit_and_repair",
            authorization="explicit_execute",
            request_text="使用最新独立审计证据执行结构修复",
            targets=[fixture["artifact_id"]],
            authoritative_source="最新独立审计证书与修复包",
            latest_version="current",
            target_boundary="第一集",
            must_keep=["未授权镜头不变"],
            must_remove=[],
            protected_scope=[],
            expected_outputs=["完整修复版"],
            window_key=repair_session,
            request_session_id=repair_session,
        )
        hub = self.service.store.load_hub(
            self.namespace, fixture["project_id"]
        )
        repair_order = next(
            row
            for row in hub["work_orders"]
            if row["work_order_id"] == repair.work_order_id
        )
        before = self._authority_snapshot(fixture["project_id"])
        with self.assertRaisesRegex(
            ValueError, "repair_package_superseded"
        ):
            self.service._full_repair_authorization(
                self.namespace,
                fixture["project_id"],
                work_order=repair_order,
                source_artifact_id=fixture["artifact_id"],
                repair_package_id=old_package_id,
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), before
        )
        latest = self.service._full_repair_authorization(
            self.namespace,
            fixture["project_id"],
            work_order=repair_order,
            source_artifact_id=fixture["artifact_id"],
            repair_package_id=new_package_id,
        )
        self.assertEqual(
            latest["package"]["repair_package_id"], new_package_id
        )
        self.assertEqual(
            latest["certificate"]["certificate_id"],
            second.data["certificate"]["certificate_id"],
        )

    def test_root_fail_repair_route_write_failures_resume_exact_completion(self) -> None:
        for fault in ("package", "handoff", "hub"):
            with self.subTest(fault=fault):
                fixture = self._unaudited_fixture(
                    f"root-repair-route-fault-{fault}"
                )
                failed_semantic = self._failed_semantic(
                    fixture["clean_semantic"]
                )
                original_write = self.service.store.atomic_write_json
                fired = False

                def flaky_write(path, value):
                    nonlocal fired
                    target = Path(path)
                    is_target = (
                        fault == "package"
                        and target.parent.name == "repair_packages"
                    ) or (
                        fault == "handoff"
                        and isinstance(value, dict)
                        and value.get("handoff_type") == "audit_to_repair"
                    ) or (
                        fault == "hub"
                        and target.name == "hub.json"
                        and any(
                            isinstance(row, dict)
                            and row.get("handoff_type")
                            == "audit_to_repair"
                            for row in value.get("handoffs", [])
                        )
                    )
                    if is_target and not fired:
                        fired = True
                        raise OSError(
                            f"simulated repair route {fault} write failure"
                        )
                    return original_write(path, value)

                with patch.object(
                    self.service.store,
                    "atomic_write_json",
                    side_effect=flaky_write,
                ):
                    with self.assertRaisesRegex(OSError, fault):
                        self.service.audit_storyboard(
                            self.namespace,
                            project_id=fixture["project_id"],
                            artifact_id=fixture["artifact_id"],
                            audit_cycle_id=None,
                            work_order_id=fixture["audit_id"],
                            handoff_artifact_id=fixture[
                                "root_handoff_id"
                            ],
                            semantic_audit=failed_semantic,
                            request_session_id=fixture["audit_session"],
                        )
                self.assertTrue(fired)
                root = self.service.store.get_project_record(
                    self.namespace,
                    fixture["project_id"],
                    kind="handoffs",
                    record_id=fixture["root_handoff_id"],
                )
                self.assertEqual(root["status"], "prepared")
                events = [
                    row
                    for row in self.service.store.load_ledger(
                        self.namespace, fixture["project_id"]
                    )["events"]
                    if row.get("handoff_id")
                    == fixture["root_handoff_id"]
                ]
                certificates = [
                    row
                    for row in self.service.store.list_project_records_strict(
                        self.namespace,
                        fixture["project_id"],
                        kind="certificates",
                    )
                    if row.get("handoff_id")
                    == fixture["root_handoff_id"]
                ]
                self.assertEqual(len(events), 1)
                self.assertEqual(len(certificates), 1)
                self.assertEqual(
                    self.service.store.list_project_records_strict(
                        self.namespace,
                        fixture["project_id"],
                        kind="repair_packages",
                    ),
                    [],
                )
                self.assertFalse(
                    any(
                        row.get("handoff_type") == "audit_to_repair"
                        for row in self.service.store.list_project_records_strict(
                            self.namespace,
                            fixture["project_id"],
                            kind="handoffs",
                        )
                    )
                )
                with self.assertRaisesRegex(ValueError, "unknown telemetry"):
                    self.service.store.get_project_record(
                        self.namespace,
                        fixture["project_id"],
                        kind="telemetry",
                        record_id=str(certificates[0]["stage_trace_id"]),
                    )

                _server, restarted = create_mcp(self.settings)
                recovered = restarted.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    semantic_audit=None,
                    request_session_id=fixture["audit_session"],
                )
                self.assertEqual(recovered.status, "fail_current")
                self.assertTrue(recovered.data["cache_hit"])
                self.assertEqual(
                    recovered.data["certificate"]["certificate_id"],
                    certificates[0]["certificate_id"],
                )
                self.assertEqual(
                    [
                        row
                        for row in restarted.store.load_ledger(
                            self.namespace, fixture["project_id"]
                        )["events"]
                        if row.get("handoff_id")
                        == fixture["root_handoff_id"]
                    ],
                    events,
                )
                self.assertEqual(
                    [
                        row
                        for row in restarted.store.list_project_records_strict(
                            self.namespace,
                            fixture["project_id"],
                            kind="certificates",
                        )
                        if row.get("handoff_id")
                        == fixture["root_handoff_id"]
                    ],
                    certificates,
                )
                package = recovered.data["repair_package"]
                self.assertIsInstance(package, dict)
                self.assertEqual(
                    package["audit_certificate_id"],
                    certificates[0]["certificate_id"],
                )
                restarted.store.get_project_record(
                    self.namespace,
                    fixture["project_id"],
                    kind="telemetry",
                    record_id=str(certificates[0]["stage_trace_id"]),
                )

    def test_root_fail_hard_crash_waits_for_lease_then_completes_first_retry(self) -> None:
        fixture = self._unaudited_fixture("root-fail-consuming-recovery")
        failed_semantic = self._failed_semantic(fixture["clean_semantic"])
        original_write = self.service.store.atomic_write_json
        fired = False

        def fail_package_once(path, value):
            nonlocal fired
            if Path(path).parent.name == "repair_packages" and not fired:
                fired = True
                raise OSError("simulated hard crash during repair package write")
            return original_write(path, value)

        with patch.object(
            self.service.store,
            "atomic_write_json",
            side_effect=fail_package_once,
        ), patch.object(
            self.service.store,
            "release_handoff_claim",
            return_value=None,
        ):
            with self.assertRaisesRegex(OSError, "hard crash"):
                self.service.audit_storyboard(
                    self.namespace,
                    project_id=fixture["project_id"],
                    artifact_id=fixture["artifact_id"],
                    audit_cycle_id=None,
                    work_order_id=fixture["audit_id"],
                    handoff_artifact_id=fixture["root_handoff_id"],
                    semantic_audit=failed_semantic,
                    request_session_id=fixture["audit_session"],
                )
        self.assertTrue(fired)
        consuming = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=fixture["root_handoff_id"],
        )
        self.assertEqual(consuming["status"], "consuming")
        events_before = copy.deepcopy(
            self.service.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"]
        )
        certificates_before = (
            self.service.store.list_project_records_strict(
                self.namespace,
                fixture["project_id"],
                kind="certificates",
            )
        )

        _server, restarted = create_mcp(self.settings)
        busy_before = self._authority_snapshot(fixture["project_id"])
        with self.assertRaisesRegex(
            ValueError, "handoff_state_conflict: handoff is being consumed"
        ):
            restarted.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=fixture["root_handoff_id"],
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), busy_before
        )

        consuming["claim_lease_expires_epoch_seconds"] = 1.0
        consuming["claim_lease_expires_at"] = "1970-01-01T00:00:01Z"
        consuming["integrity_hmac"] = (
            restarted.store._record_integrity_hmac(consuming)
        )
        handoff_path = (
            restarted.store.project_dir(
                self.namespace, fixture["project_id"]
            )
            / "handoffs"
            / f"{fixture['root_handoff_id']}.json"
        )
        restarted.store.atomic_write_json(handoff_path, consuming)
        recovered = restarted.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            semantic_audit=None,
            request_session_id=fixture["audit_session"],
        )
        self.assertEqual(recovered.status, "fail_current")
        completed = restarted.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=fixture["root_handoff_id"],
        )
        self.assertEqual(completed["status"], "consumed")
        self.assertEqual(
            completed["claim_history"][-1]["state"], "lease_expired"
        )
        self.assertEqual(
            restarted.store.load_ledger(
                self.namespace, fixture["project_id"]
            )["events"],
            events_before,
        )
        self.assertEqual(
            restarted.store.list_project_records_strict(
                self.namespace,
                fixture["project_id"],
                kind="certificates",
            ),
            certificates_before,
        )

    def test_explicit_hub_only_cleanup_preserves_untrusted_trace_target(self) -> None:
        project_id = self._project("reaudit-hub-only-trace-forensics")
        production_id, production_session = self._begin_production(
            project_id, "reaudit-hub-only-trace-forensics"
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            "reaudit-hub-only-trace-forensics",
        )
        artifact_id = str(committed.artifact_id)
        root_handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, "reaudit-hub-only-trace-forensics"
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        first = self.service.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=root_handoff_id,
            semantic_audit=clean_semantic,
            request_session_id=audit_session,
        )
        root = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=root_handoff_id,
        )
        protected_trace_id = str(first.data["certificate"]["stage_trace_id"])
        protected_trace_path = (
            self.service.store.project_dir(self.namespace, project_id)
            / "telemetry"
            / f"{protected_trace_id}.json"
        )
        protected_trace_bytes = protected_trace_path.read_bytes()
        orphan_id = "HND-HUB-ONLY-UNTRUSTED-TRACE"
        hub_only = self.service._handoff_record_projection(root)
        for key in (
            "audit_certificate_id",
            "consumed_by_work_order_id",
            "consumed_at",
        ):
            hub_only.pop(key, None)
        hub_only.update(
            {
                "handoff_id": orphan_id,
                "status": "prepared",
                "reaudit_kind": "independent_audit_successor",
                "reaudit_root_handoff_id": root_handoff_id,
                "reaudit_parent_handoff_id": root_handoff_id,
                "reaudit_sequence": 1,
                # The unsigned row maliciously targets an unrelated,
                # authoritative audit trace.
                "reaudit_route_trace_id": protected_trace_id,
                "reaudit_event_trace_id": "TRC-HUB-ONLY-EVENT",
                "server_minted_lineage": True,
            }
        )
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (hub, _ledger):
            hub["handoffs"].append(hub_only)
        rejected_before = self._authority_snapshot(project_id)
        with self.assertRaisesRegex(ValueError, "Hub-only route"):
            self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=root_handoff_id,
                force_machine_recheck=True,
                semantic_audit=clean_semantic,
                request_session_id=audit_session,
            )
        self.assertEqual(self._authority_snapshot(project_id), rejected_before)
        self.assertEqual(protected_trace_path.read_bytes(), protected_trace_bytes)
        removed = self.service.store.cleanup_reaudit_hub_orphans(
            self.namespace,
            project_id,
            root_handoff_id=root_handoff_id,
        )
        self.assertEqual(removed, [orphan_id])
        self.assertFalse(
            any(
                row.get("handoff_id") == orphan_id
                for row in self.service.store.load_hub(
                    self.namespace, project_id
                )["handoffs"]
            )
        )
        self.assertEqual(protected_trace_path.read_bytes(), protected_trace_bytes)

    def test_forged_reaudit_lineage_is_rejected_without_mutation(self) -> None:
        project_id = self._project("reaudit-forged-lineage")
        production_id, production_session = self._begin_production(
            project_id, "reaudit-forged-lineage"
        )
        committed = self._commit(
            project_id,
            production_id,
            production_session,
            valid_storyboard(),
            "reaudit-forged-lineage",
        )
        artifact_id = str(committed.artifact_id)
        handoff_id = str(committed.data["handoff"]["handoff_id"])
        audit_id, audit_session = self._begin_audit(
            project_id, artifact_id, "reaudit-forged-lineage"
        )
        clean_semantic = semantic_audit_for_artifact(
            self.service, self.settings, project_id, artifact_id
        )
        self.service.audit_storyboard(
            self.namespace,
            project_id=project_id,
            artifact_id=artifact_id,
            audit_cycle_id=None,
            work_order_id=audit_id,
            handoff_artifact_id=handoff_id,
            semantic_audit=clean_semantic,
            request_session_id=audit_session,
        )
        root = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=handoff_id,
        )
        parent_certificate = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="certificates",
            record_id=str(root["audit_certificate_id"]),
        )
        audit_order = next(
            row
            for row in self.service.store.load_hub(
                self.namespace, project_id
            )["work_orders"]
            if row["work_order_id"] == audit_id
        )
        forged_id = "HND-FORGED-REAUDIT-LINEAGE"
        forged = self.service._handoff_record_projection(root)
        forged.update(
            {
                "handoff_id": forged_id,
                "created_at": service_module.utc_stamp(),
                "status": "consumed",
                "stage_trace_id": "TRC-FORGED-ROUTE",
                "production_stage_trace_id": root["stage_trace_id"],
                "reaudit_kind": "independent_audit_successor",
                "reaudit_root_handoff_id": handoff_id,
                "reaudit_parent_handoff_id": handoff_id,
                "reaudit_sequence": 1,
                "reaudit_route_trace_id": "TRC-FORGED-ROUTE",
                "reaudit_event_trace_id": "TRC-FORGED-EVENT",
                "previous_audit_certificate_id": root[
                    "audit_certificate_id"
                ],
                "previous_audit_certificate_hmac": parent_certificate[
                    "integrity_hmac"
                ],
                "previous_audit_event_sha256": parent_certificate[
                    "event_sha256"
                ],
                "audit_certificate_id": root["audit_certificate_id"],
                "latest_audit_certificate_id": root[
                    "audit_certificate_id"
                ],
                "consumed_by_work_order_id": audit_id,
                "consumed_at": service_module.utc_stamp(),
                "claim_generation": 1,
                "reaudit_preclaim_projection_sha256": "1" * 64,
                "consumption_claim": {
                    "claim_token": "HCL-FORGED-REAUDIT-LINEAGE",
                    "claim_generation": 1,
                    "claim_epoch_seconds": 1.0,
                    "claim_lease_seconds": 300,
                    "claim_lease_expires_at": "2026-08-31T00:05:00Z",
                    "claim_lease_expires_epoch_seconds": 301.0,
                },
                "claim_authority_work_order_id": audit_id,
                "reaudit_audit_window_binding_id": audit_order[
                    "server_window_binding_id"
                ],
                "reaudit_audit_actor_binding_id": audit_order[
                    "server_actor_binding_id"
                ],
                "reaudit_audit_session_fingerprint": audit_order[
                    "transport_session_fingerprint"
                ],
                "server_minted_lineage": True,
                "reaudit_lineage_sha256": "0" * 64,
            }
        )
        persisted_forged = self.service.store.save_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=forged_id,
            value=forged,
        )
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (hub, _ledger):
            hub["handoffs"].append(
                self.service._handoff_record_projection(persisted_forged)
            )
            hub_only = self.service._handoff_record_projection(
                persisted_forged
            )
            hub_only["handoff_id"] = "HND-HUB-ONLY-BESIDE-FORGERY"
            hub_only["reaudit_route_trace_id"] = str(root["stage_trace_id"])
            hub["handoffs"].append(hub_only)
        before = self._authority_snapshot(project_id)
        with self.assertRaisesRegex(ValueError, "reaudit_lineage_invalid"):
            self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=artifact_id,
                audit_cycle_id=None,
                work_order_id=audit_id,
                handoff_artifact_id=forged_id,
                force_machine_recheck=True,
                semantic_audit=clean_semantic,
                request_session_id=audit_session,
            )
        self.assertEqual(self._authority_snapshot(project_id), before)

    def test_storage_commit_rejects_equal_or_reused_reaudit_traces_without_mutation(
        self,
    ) -> None:
        fixture = self._passed_audit_fixture(
            "storage-reaudit-trace-freshness"
        )
        project_id = fixture["project_id"]
        telemetry_trace_id = "TRC-USED-SIGNED-TELEMETRY"
        self.service.store.save_project_record(
            self.namespace,
            project_id,
            kind="telemetry",
            record_id=telemetry_trace_id,
            value={
                "telemetry_id": telemetry_trace_id,
                "trace_id": telemetry_trace_id,
                "operation": "trace-freshness-fixture",
            },
        )
        hub_trace_id = "TRC-USED-HUB-PERFORMANCE"
        ledger_trace_id = "TRC-USED-LEDGER-PERFORMANCE"
        event_trace_id = "TRC-USED-LEDGER-EVENT"
        with self.service.store.mutate_project(
            self.namespace, project_id
        ) as (hub, ledger):
            hub.setdefault("performance_traces", []).append(
                {
                    "telemetry_id": hub_trace_id,
                    "trace_id": hub_trace_id,
                }
            )
            ledger.setdefault("performance_traces", []).append(
                {
                    "telemetry_id": ledger_trace_id,
                    "trace_id": ledger_trace_id,
                }
            )
            ledger.setdefault("events", []).append(
                {
                    "event_id": "EVT-USED-TRACE-FIXTURE",
                    "stage_trace_id": event_trace_id,
                }
            )
        root = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=fixture["root_handoff_id"],
        )
        cases = (
            ("equal", "TRC-FRESH-EQUAL", "TRC-FRESH-EQUAL"),
            ("handoff", str(root["stage_trace_id"]), "TRC-FRESH-EVENT-1"),
            ("telemetry", telemetry_trace_id, "TRC-FRESH-EVENT-2"),
            ("hub-performance", hub_trace_id, "TRC-FRESH-EVENT-3"),
            (
                "ledger-performance",
                ledger_trace_id,
                "TRC-FRESH-EVENT-4",
            ),
            ("ledger-event", "TRC-FRESH-ROUTE-5", event_trace_id),
        )
        for index, (label, route_trace_id, audit_trace_id) in enumerate(
            cases, start=1
        ):
            with self.subTest(trace_conflict=label):
                successor = self._direct_reaudit_successor_candidate(
                    fixture,
                    successor_id=f"HND-DIRECT-TRACE-CONFLICT-{index}",
                    route_trace_id=route_trace_id,
                    event_trace_id=audit_trace_id,
                )
                before = self._authority_snapshot(project_id)
                with self.assertRaisesRegex(
                    ValueError, "reaudit_trace_conflict"
                ):
                    self.service.store.commit_reaudit_successor(
                        self.namespace,
                        project_id,
                        root_handoff_id=fixture["root_handoff_id"],
                        parent_handoff_id=fixture["root_handoff_id"],
                        successor_handoff_id=str(successor["handoff_id"]),
                        consumer_work_order_id=fixture["audit_id"],
                        successor=successor,
                    )
                self.assertEqual(
                    self._authority_snapshot(project_id), before
                )

    def test_storage_commit_rejects_forged_reaudit_context_without_mutation(
        self,
    ) -> None:
        fixture = self._passed_audit_fixture(
            "storage-reaudit-context-authority"
        )
        project_id = fixture["project_id"]
        root = self.service.store.get_project_record(
            self.namespace,
            project_id,
            kind="handoffs",
            record_id=fixture["root_handoff_id"],
        )

        def reseal_lineage(candidate: dict) -> None:
            candidate["reaudit_lineage_sha256"] = (
                service_module.sha256_bytes(
                    service_module._stable_json_bytes(
                        self.service._reaudit_lineage_payload(
                            root=root,
                            parent=root,
                            successor=candidate,
                        )
                    )
                )
            )

        def mutate_route(candidate: dict) -> None:
            candidate["handoff_type"] = "audit_intake"
            candidate["from_window_role"] = "external_input"
            reseal_lineage(candidate)

        def mutate_lineage(candidate: dict) -> None:
            candidate["reaudit_lineage_sha256"] = "0" * 64

        def mutate_authority(candidate: dict) -> None:
            candidate["claim_authority_work_order_id"] = "WO-FORGED-AUDIT"
            reseal_lineage(candidate)

        def mutate_root(candidate: dict) -> None:
            candidate["reaudit_root_handoff_id"] = "HND-FORGED-ROOT"
            reseal_lineage(candidate)

        def mutate_parent(candidate: dict) -> None:
            candidate["reaudit_parent_handoff_id"] = "HND-FORGED-PARENT"
            reseal_lineage(candidate)

        def mutate_sequence(candidate: dict) -> None:
            candidate["reaudit_sequence"] = 2
            reseal_lineage(candidate)

        def mutate_window(candidate: dict) -> None:
            candidate["reaudit_audit_window_binding_id"] = "WIN-FORGED"
            reseal_lineage(candidate)

        def mutate_identity(candidate: dict) -> None:
            candidate["canonical_sha256"] = "f" * 64
            reseal_lineage(candidate)

        def mutate_production_trace(candidate: dict) -> None:
            candidate["production_stage_trace_id"] = "TRC-FORGED-PRODUCTION"
            reseal_lineage(candidate)

        def mutate_route_stage(candidate: dict) -> None:
            candidate["stage_trace_id"] = "TRC-FORGED-ROUTE-STAGE"
            reseal_lineage(candidate)

        mutations = (
            ("root_route_role", mutate_route),
            ("lineage_sha256", mutate_lineage),
            ("audit_authority", mutate_authority),
            ("root_id", mutate_root),
            ("parent_id", mutate_parent),
            ("sequence", mutate_sequence),
            ("audit_window", mutate_window),
            ("immutable_identity", mutate_identity),
            ("production_trace", mutate_production_trace),
            ("route_stage", mutate_route_stage),
        )
        for index, (label, mutate) in enumerate(mutations, start=1):
            with self.subTest(context_defect=label):
                candidate = self._direct_reaudit_successor_candidate(
                    fixture,
                    successor_id=f"HND-DIRECT-CONTEXT-{index}",
                    route_trace_id=f"TRC-DIRECT-CONTEXT-ROUTE-{index}",
                    event_trace_id=f"TRC-DIRECT-CONTEXT-EVENT-{index}",
                )
                mutate(candidate)
                before = self._authority_snapshot(project_id)
                with self.assertRaisesRegex(
                    ValueError, "reaudit_lineage_invalid"
                ):
                    self.service.store.commit_reaudit_successor(
                        self.namespace,
                        project_id,
                        root_handoff_id=fixture["root_handoff_id"],
                        parent_handoff_id=fixture["root_handoff_id"],
                        successor_handoff_id=str(candidate["handoff_id"]),
                        consumer_work_order_id=fixture["audit_id"],
                        successor=candidate,
                    )
                self.assertEqual(
                    self._authority_snapshot(project_id), before
                )

    def test_reaudit_lifecycle_shapes_reject_signed_tamper_without_mutation(
        self,
    ) -> None:
        prepared_fixture = self._passed_audit_fixture(
            "reaudit-prepared-shape-tamper"
        )
        prepared_candidate = self._direct_reaudit_successor_candidate(
            prepared_fixture,
            successor_id="HND-DIRECT-PREPARED-SHAPE-TAMPER",
            route_trace_id="TRC-DIRECT-PREPARED-SHAPE-ROUTE",
            event_trace_id="TRC-DIRECT-PREPARED-SHAPE-EVENT",
        )
        prepared_candidate["claimed_by_work_order_id"] = (
            prepared_fixture["audit_id"]
        )
        prepared_before = self._authority_snapshot(
            prepared_fixture["project_id"]
        )
        with self.assertRaisesRegex(
            ValueError, "reaudit_claim_binding_invalid"
        ):
            self.service.store.commit_reaudit_successor(
                self.namespace,
                prepared_fixture["project_id"],
                root_handoff_id=prepared_fixture["root_handoff_id"],
                parent_handoff_id=prepared_fixture["root_handoff_id"],
                successor_handoff_id=str(
                    prepared_candidate["handoff_id"]
                ),
                consumer_work_order_id=prepared_fixture["audit_id"],
                successor=prepared_candidate,
            )
        self.assertEqual(
            self._authority_snapshot(prepared_fixture["project_id"]),
            prepared_before,
        )

        consuming_fixture = self._passed_audit_fixture(
            "reaudit-consuming-shape-tamper"
        )
        consuming_semantic = copy.deepcopy(
            consuming_fixture["clean_semantic"]
        )
        consuming_semantic["layers"][0]["summary"] += (
            "｜生成待领取复审路由。"
        )
        prepared = self._leave_prepared_reaudit_successor(
            consuming_fixture, consuming_semantic
        )
        consuming = self.service.store.claim_reaudit_successor(
            self.namespace,
            consuming_fixture["project_id"],
            str(prepared["handoff_id"]),
            consumer_work_order_id=consuming_fixture["audit_id"],
        )
        forged_consuming = copy.deepcopy(consuming)
        forged_consuming["audit_certificate_id"] = "CERT-STRAY-CONSUMING"
        forged_consuming = self.service.store.save_project_record(
            self.namespace,
            consuming_fixture["project_id"],
            kind="handoffs",
            record_id=str(consuming["handoff_id"]),
            value=forged_consuming,
        )
        with self.service.store.mutate_project(
            self.namespace, consuming_fixture["project_id"]
        ) as (hub, _ledger):
            index = next(
                index
                for index, row in enumerate(hub["handoffs"])
                if row.get("handoff_id") == consuming["handoff_id"]
            )
            hub["handoffs"][index] = (
                self.service._handoff_record_projection(forged_consuming)
            )
        consuming_before = self._authority_snapshot(
            consuming_fixture["project_id"]
        )
        with self.assertRaisesRegex(
            ValueError, "reaudit_claim_binding_invalid"
        ):
            self.service.store.release_reaudit_successor_claim(
                self.namespace,
                consuming_fixture["project_id"],
                str(consuming["handoff_id"]),
                consumer_work_order_id=consuming_fixture["audit_id"],
                claim_token=str(consuming["claim_token"]),
                reason="signed shape tamper",
            )
        self.assertEqual(
            self._authority_snapshot(consuming_fixture["project_id"]),
            consuming_before,
        )

        consumed_fixture = self._passed_audit_fixture(
            "reaudit-consumed-shape-tamper"
        )
        consumed_semantic = copy.deepcopy(
            consumed_fixture["clean_semantic"]
        )
        consumed_semantic["layers"][0]["summary"] += (
            "｜完成第二轮独立复审。"
        )
        completed = self.service.audit_storyboard(
            self.namespace,
            project_id=consumed_fixture["project_id"],
            artifact_id=consumed_fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=consumed_fixture["audit_id"],
            handoff_artifact_id=consumed_fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=consumed_semantic,
            request_session_id=consumed_fixture["audit_session"],
        )
        successor_id = str(completed.data["handoff"]["handoff_id"])
        store = self.service.store
        baseline = store.get_project_record(
            self.namespace,
            consumed_fixture["project_id"],
            kind="handoffs",
            record_id=successor_id,
        )

        def replace_successor(value: dict) -> dict:
            signed = store.save_project_record(
                self.namespace,
                consumed_fixture["project_id"],
                kind="handoffs",
                record_id=successor_id,
                value=value,
            )
            with store.mutate_project(
                self.namespace, consumed_fixture["project_id"]
            ) as (hub, _ledger):
                index = next(
                    index
                    for index, row in enumerate(hub["handoffs"])
                    if row.get("handoff_id") == successor_id
                )
                hub["handoffs"][index] = (
                    self.service._handoff_record_projection(signed)
                )
            return signed

        def latest_pointer(value: dict) -> None:
            value["latest_audit_certificate_id"] = "CERT-FORGED-LATEST"

        def claim_generation(value: dict) -> None:
            value["consumption_claim"]["claim_generation"] += 1

        def lease_arithmetic(value: dict) -> None:
            value["consumption_claim"][
                "claim_lease_expires_epoch_seconds"
            ] += 1

        def non_finite_claim_epoch(value: dict) -> None:
            value["consumption_claim"]["claim_epoch_seconds"] = float("nan")

        def non_finite_lease_expiry(value: dict) -> None:
            value["consumption_claim"][
                "claim_lease_expires_epoch_seconds"
            ] = float("inf")

        def preclaim_binding(value: dict) -> None:
            value["reaudit_preclaim_projection_sha256"] = "0" * 64

        def live_claim(value: dict) -> None:
            value["claimed_by_work_order_id"] = consumed_fixture[
                "audit_id"
            ]

        def consumed_authority(value: dict) -> None:
            value["consumed_by_work_order_id"] = "WO-FORGED-CONSUMER"

        mutations = (
            ("latest_certificate", latest_pointer),
            ("claim_generation", claim_generation),
            ("lease_arithmetic", lease_arithmetic),
            ("non_finite_claim_epoch", non_finite_claim_epoch),
            ("non_finite_lease_expiry", non_finite_lease_expiry),
            ("preclaim_binding", preclaim_binding),
            ("live_claim", live_claim),
            ("consumed_authority", consumed_authority),
        )
        for label, mutate in mutations:
            with self.subTest(consumed_defect=label):
                forged = copy.deepcopy(baseline)
                mutate(forged)
                replace_successor(forged)
                before = self._authority_snapshot(
                    consumed_fixture["project_id"]
                )
                try:
                    with self.assertRaisesRegex(
                        ValueError, "reaudit_lineage_invalid"
                    ):
                        self.service.audit_storyboard(
                            self.namespace,
                            project_id=consumed_fixture["project_id"],
                            artifact_id=consumed_fixture["artifact_id"],
                            audit_cycle_id=None,
                            work_order_id=consumed_fixture["audit_id"],
                            handoff_artifact_id=consumed_fixture[
                                "root_handoff_id"
                            ],
                            force_machine_recheck=True,
                            semantic_audit=consumed_semantic,
                            request_session_id=consumed_fixture[
                                "audit_session"
                            ],
                        )
                    self.assertEqual(
                        self._authority_snapshot(
                            consumed_fixture["project_id"]
                        ),
                        before,
                    )
                finally:
                    replace_successor(copy.deepcopy(baseline))

    def test_equal_reaudit_route_and_event_trace_is_rejected_without_mutation(
        self,
    ) -> None:
        fixture = self._passed_audit_fixture(
            "reaudit-equal-route-event-trace"
        )
        second_semantic = copy.deepcopy(fixture["clean_semantic"])
        second_semantic["layers"][0]["summary"] += (
            "｜第二轮使用独立的新证据。"
        )
        completed = self.service.audit_storyboard(
            self.namespace,
            project_id=fixture["project_id"],
            artifact_id=fixture["artifact_id"],
            audit_cycle_id=None,
            work_order_id=fixture["audit_id"],
            handoff_artifact_id=fixture["root_handoff_id"],
            force_machine_recheck=True,
            semantic_audit=second_semantic,
            request_session_id=fixture["audit_session"],
        )
        successor_id = str(completed.data["handoff"]["handoff_id"])
        root = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=fixture["root_handoff_id"],
        )
        successor = self.service.store.get_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=successor_id,
        )
        tampered = copy.deepcopy(successor)
        tampered["reaudit_event_trace_id"] = tampered[
            "reaudit_route_trace_id"
        ]
        tampered["reaudit_lineage_sha256"] = service_module.sha256_bytes(
            service_module._stable_json_bytes(
                self.service._reaudit_lineage_payload(
                    root=root,
                    parent=root,
                    successor=tampered,
                )
            )
        )
        persisted = self.service.store.save_project_record(
            self.namespace,
            fixture["project_id"],
            kind="handoffs",
            record_id=successor_id,
            value=tampered,
        )
        self.assertTrue(self.service.store.project_record_is_signed(persisted))
        # Standard JSON Schema enforces the complete structural field set.
        # Cross-field inequality is enforced by runtime/storage below.
        self.service._validate_schema_record(
            self.service._project_schema,
            "handoff_manifest",
            persisted,
        )
        with self.service.store.mutate_project(
            self.namespace, fixture["project_id"]
        ) as (hub, _ledger):
            hub["handoffs"] = [
                self.service._handoff_record_projection(persisted)
                if row.get("handoff_id") == successor_id
                else row
                for row in hub["handoffs"]
            ]
        before = self._authority_snapshot(fixture["project_id"])
        with self.assertRaisesRegex(
            ValueError,
            "reaudit_lineage_invalid: successor .* trace binding",
        ):
            self.service.audit_storyboard(
                self.namespace,
                project_id=fixture["project_id"],
                artifact_id=fixture["artifact_id"],
                audit_cycle_id=None,
                work_order_id=fixture["audit_id"],
                handoff_artifact_id=successor_id,
                force_machine_recheck=False,
                semantic_audit=None,
                request_session_id=fixture["audit_session"],
            )
        self.assertEqual(
            self._authority_snapshot(fixture["project_id"]), before
        )

    def test_xlsx_without_canonical_sidecar_cannot_be_independently_audited(self) -> None:
        project_id = self._project("no-sidecar")
        payload = storyboard_to_xlsx(
            valid_storyboard(), {"status": "pass", "issues": []}, assume_compact=True
        )
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        self.assertIn(file_io.CANONICAL_SIDECAR_SHEET, workbook.sheetnames)
        workbook.remove(workbook[file_io.CANONICAL_SIDECAR_SHEET])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        record = self.service.store.save_artifact_bytes(
            self.namespace,
            project_id,
            output.getvalue(),
            file_name="legacy-no-sidecar.xlsx",
            artifact_type="storyboard",
            version="v1",
            metadata={
                "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            },
        )
        audit_id, audit_session = self._begin_audit(
            project_id, str(record["artifact_id"]), "no-sidecar"
        )
        artifact_count = len(
            self.service.store.list_artifacts(self.namespace, project_id)
        )
        certificates_before = self.service.store.list_project_records(
            self.namespace, project_id, kind="certificates"
        )

        with self.assertRaisesRegex(
            ValueError, "canonical_projection_migration_required"
        ):
            self.service.audit_storyboard(
                self.namespace,
                project_id=project_id,
                artifact_id=str(record["artifact_id"]),
                audit_cycle_id=None,
                work_order_id=audit_id,
                semantic_audit=None,
                request_session_id=audit_session,
            )

        self.assertEqual(
            len(self.service.store.list_artifacts(self.namespace, project_id)),
            artifact_count,
        )
        self.assertEqual(
            self.service.store.list_project_records(
                self.namespace, project_id, kind="handoffs"
            ),
            [],
        )
        self.assertEqual(
            self.service.store.list_project_records(
                self.namespace, project_id, kind="certificates"
            ),
            certificates_before,
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
