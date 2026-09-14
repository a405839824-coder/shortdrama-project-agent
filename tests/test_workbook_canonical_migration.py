from __future__ import annotations

import copy
import hashlib
import io
import re
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from shortdrama_mcp.app import create_mcp
from shortdrama_mcp.config import Settings
from shortdrama_mcp.file_io import CANONICAL_SIDECAR_SHEET, storyboard_to_xlsx
from shortdrama_mcp.models import OpenAIFile
import shortdrama_mcp.service as runtime
from shortdrama_mcp.workbook_migration import (
    WORKBOOK_MIGRATION_PHASE, WORKBOOK_MIGRATION_TARGET,
    recover_legacy_workbook, shot_atoms,
)
from semantic_audit_helpers import semantic_audit_for_artifact
from test_efficiency import valid_storyboard


def legacy_bytes(candidate=None, edit=None):
    candidate = candidate or valid_storyboard()
    book = load_workbook(io.BytesIO(storyboard_to_xlsx(candidate, {"issues": []}, assume_compact=True)))
    del book[CANONICAL_SIDECAR_SHEET]
    sheet = book["正式分镜母版"]
    for row, shot in enumerate(candidate["shots"], 2):
        # This fixture represents the earlier workbook with visible clocks.
        # New exports deliberately omit those clocks; pin the legacy source
        # field rather than silently changing our migration input with them.
        formal = sheet.cell(row, 5).value
        match = re.search(r"(?ms)(^画面内容：)(.*?)(?=^音色：)", formal)
        assert match is not None
        sheet.cell(row, 5).value = formal[:match.start(2)] + shot["画面内容"] + "\n\n" + formal[match.end(2):]
        sheet.cell(row, 4).value = shot["功能类型"] + " / " + shot["场景连续性ID"]
        sheet.cell(row, 3).value = f"=ROUND(B{row},1)" if row == 2 else f"=ROUND(C{row-1}+B{row},1)"
    if edit:
        edit(book, sheet)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class WorkbookRecoveryTests(unittest.TestCase):
    def test_lossless_multiline_atoms_and_formula_proof(self):
        candidate = valid_storyboard()
        candidate["shots"][0]["画面内容"] += "\n甲说道：“甲。”\n\n乙说道：“乙？”"
        candidate["shots"][0]["音色"] = "甲声线：\n\n完整多行。\n乙声线：\n\n保留标点！"
        recovered, proof = recover_legacy_workbook(legacy_bytes(candidate))
        self.assertEqual(shot_atoms(candidate), shot_atoms(recovered))
        self.assertEqual(proof["formula_evidence"][0]["verified_value"], 4.0)
        self.assertFalse(proof["sidecar_present"])
        self.assertFalse(proof["formal_certificate_issued"])

    def test_ambiguous_or_changed_workbook_rejected(self):
        edits = [
            lambda b, s: setattr(s["C2"], "value", "=B3"),
            lambda b, s: setattr(s["C2"], "value", 9),
            lambda b, s: setattr(s["B2"], "value", "=2+2"),
            lambda b, s: setattr(s["D2"], "value", "问答 / 不确定场景"),
            lambda b, s: setattr(s["A2"], "value", "EP02-001"),
            lambda b, s: setattr(s["E2"], "value", "不应丢失的前缀\n" + s["E2"].value),
            lambda b, s: setattr(s["E2"], "value", s["E2"].value + "\n机位：重复字段"),
            lambda b, s: b.copy_worksheet(s),
            lambda b, s: b.create_sheet(CANONICAL_SIDECAR_SHEET),
        ]
        for edit in edits:
            with self.subTest(edit=edits.index(edit)), self.assertRaisesRegex(ValueError, "workbook_migration_input_invalid"):
                recover_legacy_workbook(legacy_bytes(edit=edit))


class WorkbookMigrationLifecycleTests(unittest.TestCase):
    namespace = "local-single-user"
    session = "workbook-migration-test"

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="workbook-migration-")
        self.settings = Settings(data_dir=Path(self.temp.name), public_base_url="http://127.0.0.1:8765",
            host="127.0.0.1", port=8765, deployment_mode="local", auth_mode="none",
            oauth_issuer=None, oauth_audience=None, oauth_jwks_url=None,
            oauth_scopes=("storyboard:read", "storyboard:write"), allow_local_file_urls=False,
            max_file_bytes=50*1024*1024)
        self.server, self.service = create_mcp(self.settings)
        from historical_entry_fixture import freeze_legacy_entry
        freeze_legacy_entry(self.service)
        self.project = self.service.create_project(self.namespace, "migration fixture").project_id
        status = self.service.agent_status(self.namespace, self.project, full_self_test=True)
        self.assertTrue(status.data["runtime_test_certificate_valid"])
        self.payload = legacy_bytes()
        intake = self.begin("project_intake", ["EP01"])
        with patch.object(runtime, "download_openai_file", return_value=(self.payload, "legacy.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")):
            imported = self.service.import_source_file(self.namespace, project_id=self.project,
                work_order_id=intake, file=OpenAIFile(download_url="https://example.invalid/fixture", file_id="fixture"),
                source_role="storyboard", request_session_id=self.session)
        self.source = imported.artifact_id
        gap_order = self.begin("capability_gap", [self.source], mode="plan_only", authorization="plan_only")
        gap = self.service.record_capability_gap(self.namespace, project_id=self.project, work_order_id=gap_order,
            request="Bind exact legacy workbook", missing_dependency="canonical sidecar", risk="source drift",
            safe_fallback="read-only audit", proposed_implementation="lossless controlled migration",
            required_tests=["source immutability", "full preflight", "independent sign-off"], request_session_id=self.session)
        self.targets = [self.source, gap.data["gap_id"], "EP01", "controlled_phase:"+WORKBOOK_MIGRATION_PHASE,
                        "source_sha256:"+hashlib.sha256(self.payload).hexdigest()]
        self.rule_order = self.begin("rule_improvement", self.targets)

    def tearDown(self):
        self.temp.cleanup()

    def begin(self, operation, targets, mode="execute", authorization="explicit_execute"):
        return self.service.begin_task(self.namespace, project_id=self.project, operation=operation,
            mode=mode, authorization=authorization, request_text="User explicitly approved exact controlled migration",
            targets=targets, authoritative_source="original imported workbook", latest_version="current",
            target_boundary="EP01", must_keep=["all source bytes and dialogue"], must_remove=[], protected_scope=[],
            expected_outputs=["verified result"], window_key=self.session, request_session_id=self.session,
            action="create_window_neutral").work_order_id

    def migrate(self, action="plan", plan=None, order=None):
        return self.service.migrate_project_state(self.namespace, project_id=self.project,
            work_order_id=order or self.rule_order, action=action, target_version=WORKBOOK_MIGRATION_TARGET,
            expected_current_version=plan.data["plan_sha256"] if plan else None,
            migration_id=plan.data["migration_id"] if plan else None, request_session_id=self.session)

    def records(self, kind):
        return self.service.store.list_project_records(self.namespace, self.project, kind=kind)

    def tree(self):
        root = self.service.store.project_dir(self.namespace, self.project)
        return {str(p.relative_to(root)):p.read_bytes() for p in root.rglob("*") if p.is_file()}

    def candidate(self):
        parsed, _ = recover_legacy_workbook(self.payload)
        candidate = valid_storyboard()
        candidate["shots"] = parsed["shots"]
        return candidate

    def preflight(self, order, candidate, scope_kind="full", scope_id="EP01"):
        return self.service.preflight_storyboard_draft(self.namespace, project_id=self.project, work_order_id=order,
            storyboard=candidate, scope_kind=scope_kind, scope_id=scope_id, request_session_id=self.session)

    def clean_preflight(self, order, candidate):
        scene = self.preflight(order, candidate, "scene", "EP01-SC01")
        self.assertEqual(scene.status, runtime.DRAFT_PREFLIGHT_PASS, scene.data.get("report"))
        full = self.preflight(order, candidate)
        self.assertTrue(full.data["commit_ready"], full.data.get("report"))
        return full.data["draft_preflight"]["draft_preflight_id"]

    def commit(self, order, candidate, preflight_id=None, source=None):
        return self.service.commit_storyboard(self.namespace, project_id=self.project, work_order_id=order,
            storyboard=candidate, file_name="migrated.xlsx", source_artifact_id=source,
            draft_preflight_id=preflight_id, request_session_id=self.session)

    def test_plan_is_read_only_and_exact_execution_is_idempotent(self):
        before = self.tree()
        plan = self.migrate()
        self.assertEqual(before, self.tree())
        execution = self.migrate("execute", plan)
        after = self.tree()
        replay = self.migrate("execute", plan)
        self.assertEqual(execution.work_order_id, replay.work_order_id)
        self.assertEqual(after, self.tree())
        self.assertEqual(len(self.records("workbook_migrations")), 1)
        self.assertFalse(self.records("handoffs"))
        self.assertFalse(execution.data["certificate_created"])

    def test_full_pipeline_preserves_source_and_unrelated_pending_handoff(self):
        # An unrelated normal producer remains pending throughout this migration.
        producer = self.begin("storyboard_create", ["EP01"])
        unrelated = valid_storyboard()
        prior = self.commit(producer, unrelated, self.clean_preflight(producer, unrelated))
        prior_handoffs = copy.deepcopy(self.records("handoffs"))
        with self.assertRaisesRegex(ValueError, "independent_audit_pending"):
            self.begin("storyboard_create", ["EP01"])
        source_record, source_path = self.service.store.get_artifact(self.namespace, self.project, self.source)
        source_before = copy.deepcopy(source_record)
        execution = self.migrate("execute", self.migrate())
        candidate = self.candidate()
        draft_id = self.clean_preflight(execution.work_order_id, candidate)
        committed = self.commit(execution.work_order_id, candidate, draft_id, self.source)
        self.assertNotEqual(committed.artifact_id, prior.artifact_id)
        actual_record, _ = self.service.store.get_artifact(self.namespace, self.project, self.source)
        self.assertEqual(source_before, actual_record)
        self.assertEqual(source_path.read_bytes(), self.payload)
        bound, bound_record = self.service._storyboard_from_artifact(self.namespace, self.project, committed.artifact_id,
            persist_recovered_canonical=False, require_bound_projection=True)
        self.assertEqual(bound_record["metadata"]["workbook_migration_id"], execution.data["migration_id"])
        audit_order = self.begin("storyboard_audit", [committed.artifact_id], mode="audit_only", authorization="read_only")
        result = self.service.audit_storyboard(self.namespace, project_id=self.project, artifact_id=committed.artifact_id,
            work_order_id=audit_order, audit_cycle_id=None, force_machine_recheck=True,
            semantic_audit=semantic_audit_for_artifact(self.service, self.settings, self.project, committed.artifact_id),
            request_session_id=self.session)
        self.assertEqual(result.status, "pass", result.data)
        self.assertEqual(result.data["certificate"]["certificate_type"], "independent_audit")
        self.assertEqual(result.data["certificate"]["issue_count"], 0)
        self.assertTrue(result.data["audited_workbook"]["artifact_id"])
        all_handoffs = {row["handoff_id"]:row for row in self.records("handoffs")}
        for old in prior_handoffs:
            self.assertEqual(old, all_handoffs[old["handoff_id"]])
        self.assertEqual(source_before, self.service.store.get_artifact(self.namespace, self.project, self.source)[0])

    def test_wrong_scope_and_stale_plan_have_no_side_effects(self):
        wrong = self.begin("rule_improvement", [item for item in self.targets if item != "EP01"] + ["EP02"])
        before = self.tree()
        with self.assertRaisesRegex(ValueError, "workbook_migration_scope_mismatch"):
            self.migrate(order=wrong)
        self.assertEqual(before, self.tree())
        plan = self.migrate()
        plan.data["plan_sha256"] = "0"*64
        with self.assertRaisesRegex(ValueError, "workbook_migration_plan_cas_mismatch"):
            self.migrate("execute", plan)
        self.assertEqual(before, self.tree())

    def test_semantic_changes_and_commit_without_full_preflight_are_rejected(self):
        execution = self.migrate("execute", self.migrate())
        candidate = self.candidate()
        with self.assertRaisesRegex(ValueError, "workbook_migration_full_preflight_required"):
            self.commit(execution.work_order_id, candidate, source=self.source)
        candidate["shots"][0]["画面内容"] += "新增剧情。"
        result = self.preflight(execution.work_order_id, candidate, "scene", "EP01-SC01")
        self.assertEqual(result.status, runtime.DRAFT_PREFLIGHT_FAIL)
        self.assertIn("workbook_migration_semantic_change_forbidden", str(result.data))
        self.assertFalse(self.records("handoffs"))

    def test_crash_after_hub_write_replays_same_authority(self):
        plan = self.migrate()
        original = self.service.store.save_project_record
        def crash(namespace, project_id, **kwargs):
            if kwargs.get("kind") == "workbook_migrations" and kwargs["value"].get("status") == "authorized":
                raise RuntimeError("simulated migration response crash")
            return original(namespace, project_id, **kwargs)
        with patch.object(self.service.store, "save_project_record", side_effect=crash):
            with self.assertRaisesRegex(RuntimeError, "simulated migration response crash"):
                self.migrate("execute", plan)
        prepared = self.records("workbook_migrations")[0]
        self.assertEqual(prepared["status"], "prepared")
        result = self.migrate("execute", plan)
        self.assertEqual(result.work_order_id, prepared["production_request"]["work_order_id"])
        hub = self.service.store.load_hub(self.namespace, self.project)
        self.assertEqual(sum(row["work_order_id"] == result.work_order_id for row in hub["work_orders"]), 1)
        self.assertEqual(self.records("workbook_migrations")[0]["status"], "authorized")

    def test_source_authority_change_invalidates_bound_work_order(self):
        execution = self.migrate("execute", self.migrate())
        with self.service.store.mutate_project(self.namespace, self.project) as (hub, ledger):
            hub["project_state"]["authoritative_sources"][0]["fingerprint"] = "0"*64
        before = self.tree()
        with self.assertRaisesRegex(ValueError, "workbook_migration_source_authority_mismatch"):
            self.preflight(execution.work_order_id, self.candidate())
        self.assertEqual(before, self.tree())

    def test_signed_wal_tampering_is_rejected(self):
        execution = self.migrate("execute", self.migrate())
        path = self.service.store.project_dir(self.namespace, self.project) / "workbook_migrations" / (execution.data["migration_id"]+".json")
        data = self.service.store.read_json(path)
        data["production_request"]["targets"] = ["EP02"]
        self.service.store.atomic_write_json(path, data)
        before = self.tree()
        with self.assertRaises(ValueError):
            self.preflight(execution.work_order_id, self.candidate())
        self.assertEqual(before, self.tree())

    def test_missing_scene_evidence_prevents_full_commit(self):
        execution = self.migrate("execute", self.migrate())
        result = self.preflight(execution.work_order_id, self.candidate())
        self.assertEqual(result.status, runtime.DRAFT_PREFLIGHT_FAIL)
        self.assertFalse(result.data["commit_ready"])
        self.assertIn("E-DRAFT-SCENE-PREFLIGHT-001", str(result.data["report"]))
        self.assertFalse(self.records("handoffs"))


if __name__ == "__main__":
    unittest.main()
