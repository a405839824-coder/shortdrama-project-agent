"""R5 reading cards derived from the same 27 fields and exact E/F cells.

No generated prose or audit verdict is authored here. The caller signs and
checks the existing audit page separately and verifies every serialized card.
"""
from collections import OrderedDict
import math
import re
import unicodedata

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .field_content import voice_parts
from .picture_presentation import display_picture

MASTER_NAME = "正式分镜母版"
GUIDE_NAME = "使用与来源"
WIDTHS = (15, 9, 11, 30, 150, 105, 52, 27, 30, 66)


def _label_value(label, value):
    return [str(label), None, str(value) if value is not None else "", *([None] * 7)]


def _span(value):
    return [str(value), *([None] * 9)]


def _episode_name(shot):
    raw = str(shot.get("episode_id") or "")
    match = re.fullmatch(r"(?:EP|第)?0*(\d+)(?:集)?", raw, flags=re.I)
    if match is None:
        match = re.match(r"EP0*(\d+)(?:-|_)", str(shot.get("镜头编号") or ""), flags=re.I)
    if match:
        return f"{int(match.group(1))}集_逐镜阅读"
    # Do not infer an episode from a non-episode shot number.
    safe = re.sub(r"[\\/*?:\[\]]", "_", raw)[:23]
    return f"{safe or '未标集'}_逐镜阅读"


def reading_matrices(storyboard, formal_rows):
    groups = OrderedDict()
    for shot, projected in zip(storyboard.get("shots", []), formal_rows, strict=True):
        groups.setdefault(_episode_name(shot), []).append((shot, projected))
    sheets = OrderedDict()
    for name, items in groups.items():
        rows = [_span(name.replace("_逐镜阅读", "") + "｜完整分镜脚本"),
                _span(f"{len(items)}镜｜{sum(float(s.get('时长') or 0) for s, _ in items):g}秒"),
                _span("审计结论以“00_范围与审计”页绑定的当前版本记录为准。"),
                _span(f"可复制的十列母版见“{MASTER_NAME}”页；累计时长按分集计算。"),
                _span("")]
        for shot, projected in items:
            rows.append(_span(f"{shot['镜头编号']}　{shot['时长']}s / 累计 {shot['累计时长']}s　｜　{shot.get('功能类型', '')}"))
            rows.append(_label_value("场景连续性ID", shot.get("场景连续性ID")))
            rows.append(_label_value("景别／可见主体", f"{shot.get('镜头类型', '')}｜{shot.get('可见主体', '')}"))
            for field in ("空间站位描述", "起始状态", "机位", "运镜方式"):
                rows.append(_label_value(field, shot.get(field)))
            picture = display_picture(shot)
            # Existing authored paragraphs remain paragraphs; never restore clocks
            # or invent three stages merely to fill a template.
            paragraphs = picture.split("\n")
            for index, part in enumerate(paragraphs, 1):
                label = "画面内容" if len(paragraphs) == 1 else f"画面内容 {index}/{len(paragraphs)}"
                rows.append(_label_value(label, part))
            voice = voice_parts(shot.get("音色"))
            if voice:
                # Same grouping as the reference: six delivery subfields share
                # one reading row, their exact labels and bodies are preserved.
                delivery, after_group = [], False
                for label, value in voice:
                    if label == "发声方式":
                        after_group = True
                    elif label == "禁止语气":
                        rows.append(_label_value("音色 · 发声方式", "\n".join(delivery)))
                        rows.append(_label_value("音色 · " + label, value)); after_group = False
                    elif after_group:
                        delivery.append(label + "：" + value)
                    else:
                        rows.append(_label_value("音色 · " + label, value))
            else:
                rows.append(_label_value("音色", shot.get("音色")))
            for field in ("心理音轨", "环境音效", "环境动态", "光影", "光源来源", "光源方向", "色温倾向",
                          "生成限制", "转场类型", "切镜/转场提示词", "参考图锚定", "覆盖原文", "原文依据类型", "原文支撑句"):
                rows.append(_label_value(field, shot.get(field)))
            prompt_lines = str(projected["GPT关键帧生成指令"]).splitlines()
            if prompt_lines and prompt_lines[0] == "GPT关键帧生成指令":
                prompt_lines = prompt_lines[1:]
            for index, value in enumerate(prompt_lines, 1):
                rows.append(_label_value(f"GPT关键帧 {index}/{len(prompt_lines)}", value))
            rows.append(_span(""))
        sheets[name] = rows
    sheets[GUIDE_NAME] = [
        _span("使用与来源"),
        _label_value("正式母版", "十列固定；逐镜阅读页与正式母版、GPT关键帧使用同一份分镜数据。"),
        _label_value("画面内容", "自然动作与准确对白；制作时间标记保存在内部绑定中，时长和累计时长列继续保留。"),
        _label_value("审计状态", "查看“00_范围与审计”页；生成或导出完成不表示独立审计通过。"),
        _label_value("原文来源", "逐镜核对“覆盖原文”“原文依据类型”“原文支撑句”；本页不补造缺失来源。"),
        _label_value("模板", "用户提供的第36—37集R5模板结构；原件及原文均不改写。"),
    ]
    return sheets


def _text_width(value):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in str(value or ""))


def write_reading_sheets(workbook, matrices):
    for name, matrix in matrices.items():
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        for col in range(1, 11):
            sheet.column_dimensions[get_column_letter(col)].width = 13 if col < 3 else 12
        for row_index, values in enumerate(matrix, 1):
            sheet.append(values)
            body = values[2]
            span = body is None
            if span:
                sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=10)
            else:
                sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=2)
                sheet.merge_cells(start_row=row_index, start_column=3, end_row=row_index, end_column=10)
            for col in (1,) if span else (1, 3):
                cell = sheet.cell(row_index, col)
                # Formula-like source/dialogue is literal text, never executed.
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(name="微软雅黑", size=11, color="FFFFFF" if span and values[0] else "25363B", bold=span or col == 1)
                if span and values[0]: cell.fill = PatternFill("solid", fgColor="243D44" if row_index == 1 else "315F65")
                elif col == 1: cell.fill = PatternFill("solid", fgColor="E7EFEB")
            text = values[0] if span else body
            lines = sum(max(1, math.ceil(_text_width(line) / (117 if span else 85))) for line in str(text or "").split("\n"))
            sheet.row_dimensions[row_index].height = min(409, max(18, 16 * lines + 10))
    return matrices
