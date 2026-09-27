from __future__ import annotations

import base64
import binascii
import csv
import datetime as dt
import hashlib
import hmac
import io
import ipaddress
import json
import mimetypes
import re
import socket
import ssl
import urllib.error
import urllib.parse
import urllib.request
import unicodedata
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from .config import Settings
from .materializer import (
    COMPACT_SCHEMA_VERSION,
    COMPACT_SHOT_FIELDS,
    FORMAL_MERGED_FIELD,
    LEGACY_STANDALONE_FIELDS,
    MOTHER_TEMPLATE_COLUMNS,
    LEGACY_MOTHER_TEMPLATE_COLUMNS,
    formal_mother_template_rows,
    legacy_v050_mother_template_rows,
    legacy_v053_mother_template_rows,
    migrate_legacy_shot,
    migrate_storyboard_to_compact,
)
from .models import OpenAIFile
from .storage import slug_filename


SHOT_FIELDS = list(COMPACT_SHOT_FIELDS)

BATCH_FIELDS = [
    "批次编号", "覆盖镜头", "时长", "主运镜", "生成限制", "起始画面", "起始表演状态", "动作过程",
    "听话反应", "结束状态", "摄影机", "声音", "环境动态", "光影", "独立执行提示词",
]

MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAX_BULK_RESPONSE_CHARS = 16_000
MAX_BULK_RANGES = 16
DEFAULT_BULK_RESPONSE_CHARS = 16_000
DEFAULT_BULK_CHARACTER_LIMIT = 16_000
DEFAULT_BULK_ITEM_LIMITS = {"paragraphs": 80, "pages": 8, "rows": 50}

CANONICAL_SIDECAR_SHEET = "__shortdrama_canonical_v1"
CANONICAL_SIDECAR_MARKER = "SHORTDRAMA_CANONICAL_V1"
CANONICAL_SIDECAR_FORMAT = "canonical-storyboard-json-v1"
CANONICAL_SIDECAR_ENCODING = "base64url-utf8-json"
CANONICAL_SIDECAR_PAYLOAD_PREFIX = "b64u:"
CANONICAL_SIDECAR_CHUNK_CHARS = 24_000
CANONICAL_SIDECAR_HEADERS = [
    "record_type",
    "marker",
    "format",
    "encoding",
    "sha256",
    "chunk_index",
    "chunk_count",
    "payload",
    "projection_contract",
    "projection_sha256",
    "audit_page_sha256",
]
LEGACY_CANONICAL_SIDECAR_HEADERS = CANONICAL_SIDECAR_HEADERS[:-3]
CURRENT_PROJECTION_CONTRACT = "shortdrama-projector/0.5.56"
FROZEN_0555_PROJECTION_CONTRACT = "shortdrama-projector/0.5.46"
R5_LEGACY_PROJECTION_CONTRACT = "shortdrama-projector/0.5.32"
PREVIOUS_PROJECTION_CONTRACT = "shortdrama-projector/0.5.5"
from . import frozen_projector_0531, frozen_projector_0555
from .r5_workbook import (MASTER_NAME, WIDTHS, is_reading_sheet, reading_matrices,
                          workbook_audit_scope, write_reading_sheets)
LEGACY_PROJECTION_CONTRACT = "shortdrama-projector/0.5.0"
CONTROLLED_HISTORICAL_PROJECTION_CONTRACTS = frozenset(
    {CURRENT_PROJECTION_CONTRACT, FROZEN_0555_PROJECTION_CONTRACT, R5_LEGACY_PROJECTION_CONTRACT,
     PREVIOUS_PROJECTION_CONTRACT, "shortdrama-projector/0.5.4"}
)


_VERIFIED_AUDIT_OVERLAY_SEAL = object()


class _VerifiedAuditOverlay:
    """Opaque capability for a service-verified independent-audit overlay.

    Workbook projection is also used by local library callers, so accepting a
    plain mapping here would let arbitrary code turn a forged HMAC-looking
    value into a visible PASS page.  Only the service path that has reloaded
    and verified a persisted certificate may mint this internal capability.
    """

    __slots__ = ("_payload",)

    def __init__(self, payload: dict[str, Any], *, seal: object) -> None:
        if seal is not _VERIFIED_AUDIT_OVERLAY_SEAL:
            raise TypeError(
                "independent_audit_overlay_unverified: "
                "a service-verified overlay capability is required"
            )
        self._payload = dict(payload)

    def _as_dict(self) -> dict[str, Any]:
        return dict(self._payload)


def _seal_service_verified_audit_overlay(
    payload: dict[str, Any],
) -> _VerifiedAuditOverlay:
    """Mint an internal overlay capability after service certificate checks."""

    if not isinstance(payload, dict):
        raise TypeError(
            "independent_audit_overlay_unverified: payload must be a mapping"
        )
    return _VerifiedAuditOverlay(payload, seal=_VERIFIED_AUDIT_OVERLAY_SEAL)


# The v0.5.1 workbook contract renames the merged formal-master column to the
# shorter ``正式分镜母版`` label.  Keep the previously shipped label readable so
# an intact v0.5.0 workbook does not become an unparseable legacy artifact.
FORMAL_MASTER_FIELD_ALIASES = tuple(
    dict.fromkeys(
        (
            FORMAL_MERGED_FIELD,
            "正式分镜母版",
            "正式分镜母版·字段合并",
        )
    )
)


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return [json_safe(item) for item in value]
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    return str(value)


def _is_local_host(host: str) -> bool:
    return host.casefold() in {"localhost", "127.0.0.1", "::1"}


def validate_download_url(url: str, settings: Settings) -> None:
    parsed = urllib.parse.urlsplit(url)
    if parsed.username or parsed.password:
        raise ValueError("file download URL must not contain credentials")
    if parsed.fragment:
        raise ValueError("file download URL must not contain a fragment")
    if not parsed.hostname:
        raise ValueError("file download URL has no host")
    local_allowed = settings.allow_local_file_urls and _is_local_host(parsed.hostname)
    if parsed.scheme != "https" and not (local_allowed and parsed.scheme == "http"):
        raise ValueError("file download URL must use HTTPS")
    if parsed.port not in {None, 80, 443} and not local_allowed:
        raise ValueError("file download URL uses a disallowed port")
    try:
        addresses = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror as exc:
        raise ValueError("file download host cannot be resolved") from exc
    for address in addresses:
        ip_value = ipaddress.ip_address(address[4][0])
        if not ip_value.is_global and not local_allowed:
            raise ValueError("file download URL resolves to a private or reserved address")


class _SafeRedirectHandler(urllib.request.HTTPRedirectHandler):
    def __init__(self, settings: Settings):
        self.settings = settings

    def redirect_request(self, req: Any, fp: Any, code: int, msg: str, headers: Any, newurl: str) -> Any:
        validate_download_url(newurl, self.settings)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def download_openai_file(file: OpenAIFile, settings: Settings) -> tuple[bytes, str, str]:
    validate_download_url(file.download_url, settings)
    request = urllib.request.Request(
        file.download_url,
        headers={"User-Agent": "shortdrama-project-agent/0.3"},
        method="GET",
    )
    opener = urllib.request.build_opener(_SafeRedirectHandler(settings), urllib.request.HTTPSHandler(context=ssl.create_default_context()))
    try:
        with opener.open(request, timeout=30) as response:
            content_length = response.headers.get("Content-Length")
            if content_length and int(content_length) > settings.max_file_bytes:
                raise ValueError("uploaded file exceeds MAX_FILE_BYTES")
            buffer = bytearray()
            while True:
                chunk = response.read(1024 * 1024)
                if not chunk:
                    break
                buffer.extend(chunk)
                if len(buffer) > settings.max_file_bytes:
                    raise ValueError("uploaded file exceeds MAX_FILE_BYTES")
            mime_type = file.mime_type or response.headers.get_content_type() or "application/octet-stream"
    except urllib.error.HTTPError as exc:
        raise ValueError(f"file download failed with HTTP {exc.code}") from exc
    except urllib.error.URLError as exc:
        raise ValueError(f"file download failed: {exc.reason}") from exc
    name = slug_filename(file.file_name or Path(urllib.parse.urlsplit(file.download_url).path).name, "uploaded-file")
    return bytes(buffer), name, mime_type


def _validate_zip_archive(
    data: bytes,
    *,
    max_entries: int = MAX_ARCHIVE_ENTRIES,
    max_uncompressed_bytes: int = MAX_ARCHIVE_UNCOMPRESSED_BYTES,
) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > max_entries:
                raise ValueError("archive contains too many entries")
            total = sum(max(0, int(item.file_size)) for item in entries)
            if total > max_uncompressed_bytes:
                raise ValueError("archive expands beyond the safe limit")
            if any(item.filename.startswith(("/", "\\")) or ".." in Path(item.filename).parts for item in entries):
                raise ValueError("archive contains an unsafe path")
    except zipfile.BadZipFile as exc:
        raise ValueError("uploaded Office file is not a valid ZIP archive") from exc


def _unique_headers(headers: list[str]) -> list[str]:
    """Keep every column, including collisions with generated suffixes."""
    used: set[str] = set()
    result: list[str] = []
    for index, header in enumerate(headers, 1):
        base = header or f"列{index}"
        name, suffix = base, 2
        while name in used:
            name, suffix = f"{base}_{suffix}", suffix + 1
        used.add(name)
        result.append(name)
    return result


def _rows_from_sheet(sheet: Any) -> tuple[list[str], list[dict[str, Any]]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return [], []
    header_index = 0
    for index, row in enumerate(values[:20]):
        normalized = [str(value).strip() if value is not None else "" for value in row]
        if any(normalized):
            header_index = index
            break
    raw_headers = [str(value).strip() if value is not None else "" for value in values[header_index]]
    last = max((index for index, value in enumerate(raw_headers) if value), default=-1)
    unique_headers = _unique_headers(raw_headers[: last + 1])
    rows = []
    for raw in values[header_index + 1 :]:
        row = {header: json_safe(raw[index] if index < len(raw) else None) for index, header in enumerate(unique_headers)}
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return unique_headers, rows


def _sheet_value_matrix(sheet: Any) -> list[list[Any]]:
    """Return the complete used value range, including non-table columns."""

    raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    return _used_value_matrix(raw_rows)


def _used_value_matrix(raw_rows: list[list[Any]]) -> list[list[Any]]:
    last_row = max(
        (
            index
            for index, row in enumerate(raw_rows)
            if any(value not in (None, "") for value in row)
        ),
        default=-1,
    )
    if last_row < 0:
        return []
    last_column = max(
        (
            index
            for row in raw_rows[: last_row + 1]
            for index, value in enumerate(row)
            if value not in (None, "")
        ),
        default=-1,
    )
    return [
        [json_safe(value) for value in row[: last_column + 1]]
        for row in raw_rows[: last_row + 1]
    ]


def _json_cell(value: Any, fallback: Any) -> Any:
    if value in (None, ""):
        return fallback
    if isinstance(value, (dict, list, int, float, bool)):
        return value
    try:
        return json.loads(str(value))
    except json.JSONDecodeError:
        return value


def _canonical_sidecar_error(reason: str) -> ValueError:
    return ValueError(f"canonical sidecar invalid: {reason}")


def _canonical_positive_int(value: Any, field: str) -> int:
    if isinstance(value, bool):
        raise _canonical_sidecar_error(f"{field} must be a positive integer")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float) and value.is_integer():
        parsed = int(value)
    elif isinstance(value, str) and value.strip().isdigit():
        parsed = int(value.strip())
    else:
        raise _canonical_sidecar_error(f"{field} must be a positive integer")
    if parsed <= 0:
        raise _canonical_sidecar_error(f"{field} must be a positive integer")
    return parsed


def _canonical_storyboard_from_sidecar(
    sheets: list[dict[str, Any]],
    *,
    _allowed_projection_contracts: frozenset[str] | None = None,
) -> tuple[dict[str, Any], str, str | None, str | None] | None:
    sidecars = [sheet for sheet in sheets if sheet.get("name") == CANONICAL_SIDECAR_SHEET]
    if not sidecars:
        return None
    if len(sidecars) != 1:
        raise _canonical_sidecar_error("duplicate reserved sidecar sheets")

    sidecar = sidecars[0]
    sidecar_headers = sidecar.get("headers")
    if sidecar_headers not in (
        CANONICAL_SIDECAR_HEADERS,
        LEGACY_CANONICAL_SIDECAR_HEADERS,
    ):
        raise _canonical_sidecar_error("sidecar header or marker is missing")
    rows = sidecar.get("rows")
    if not isinstance(rows, list):
        raise _canonical_sidecar_error("sidecar rows are missing")

    manifests = [row for row in rows if row.get("record_type") == "manifest"]
    chunks = [row for row in rows if row.get("record_type") == "chunk"]
    if len(manifests) != 1:
        raise _canonical_sidecar_error("exactly one manifest is required")
    if len(manifests) + len(chunks) != len(rows):
        raise _canonical_sidecar_error("unknown sidecar record type")

    manifest = manifests[0]
    if sidecar_headers == CANONICAL_SIDECAR_HEADERS:
        projection_contract = manifest.get("projection_contract")
        allowed_projection_contracts = (
            _allowed_projection_contracts
            if _allowed_projection_contracts is not None
            else frozenset({CURRENT_PROJECTION_CONTRACT, FROZEN_0555_PROJECTION_CONTRACT, R5_LEGACY_PROJECTION_CONTRACT, PREVIOUS_PROJECTION_CONTRACT})
        )
        if projection_contract not in allowed_projection_contracts:
            raise _canonical_sidecar_error("projection contract marker mismatch")
        projection_sha256 = manifest.get("projection_sha256")
        if not isinstance(projection_sha256, str) or re.fullmatch(
            r"[0-9a-f]{64}", projection_sha256
        ) is None:
            raise _canonical_sidecar_error("projection SHA-256 is malformed")
        audit_page_sha256 = manifest.get("audit_page_sha256")
        if not isinstance(audit_page_sha256, str) or re.fullmatch(
            r"[0-9a-f]{64}", audit_page_sha256
        ) is None:
            raise _canonical_sidecar_error("audit page SHA-256 is malformed")
    else:
        projection_contract = LEGACY_PROJECTION_CONTRACT
        projection_sha256 = None
        audit_page_sha256 = None
    if manifest.get("marker") != CANONICAL_SIDECAR_MARKER:
        raise _canonical_sidecar_error("manifest marker mismatch")
    if manifest.get("format") != CANONICAL_SIDECAR_FORMAT:
        raise _canonical_sidecar_error("manifest format mismatch")
    if manifest.get("encoding") != CANONICAL_SIDECAR_ENCODING:
        raise _canonical_sidecar_error("manifest encoding mismatch")
    digest = manifest.get("sha256")
    if not isinstance(digest, str) or re.fullmatch(r"[0-9a-f]{64}", digest) is None:
        raise _canonical_sidecar_error("manifest SHA-256 is malformed")
    chunk_count = _canonical_positive_int(manifest.get("chunk_count"), "chunk_count")
    if len(chunks) != chunk_count:
        raise _canonical_sidecar_error("canonical payload chunk count mismatch")

    payload_by_index: dict[int, str] = {}
    for chunk in chunks:
        if (
            sidecar_headers == CANONICAL_SIDECAR_HEADERS
            and (
                chunk.get("projection_contract") != projection_contract
                or chunk.get("projection_sha256") != projection_sha256
                or chunk.get("audit_page_sha256") != audit_page_sha256
            )
        ):
            raise _canonical_sidecar_error(
                "chunk projection contract manifest mismatch"
            )
        if chunk.get("marker") != CANONICAL_SIDECAR_MARKER:
            raise _canonical_sidecar_error("chunk marker mismatch")
        if chunk.get("format") != CANONICAL_SIDECAR_FORMAT:
            raise _canonical_sidecar_error("chunk format mismatch")
        if chunk.get("encoding") != CANONICAL_SIDECAR_ENCODING:
            raise _canonical_sidecar_error("chunk encoding mismatch")
        if chunk.get("sha256") != digest:
            raise _canonical_sidecar_error("chunk SHA-256 manifest mismatch")
        if _canonical_positive_int(chunk.get("chunk_count"), "chunk_count") != chunk_count:
            raise _canonical_sidecar_error("chunk count manifest mismatch")
        index = _canonical_positive_int(chunk.get("chunk_index"), "chunk_index")
        if index > chunk_count:
            raise _canonical_sidecar_error("chunk index is out of range")
        if index in payload_by_index:
            raise _canonical_sidecar_error("duplicate canonical payload chunk")
        payload = chunk.get("payload")
        if not isinstance(payload, str) or not payload.startswith(CANONICAL_SIDECAR_PAYLOAD_PREFIX):
            raise _canonical_sidecar_error("chunk payload encoding marker is missing")
        encoded = payload[len(CANONICAL_SIDECAR_PAYLOAD_PREFIX) :]
        if not encoded:
            raise _canonical_sidecar_error("canonical payload chunk is empty")
        payload_by_index[index] = encoded

    expected_indexes = set(range(1, chunk_count + 1))
    if set(payload_by_index) != expected_indexes:
        raise _canonical_sidecar_error("canonical payload chunk sequence is incomplete")
    encoded_payload = "".join(payload_by_index[index] for index in range(1, chunk_count + 1))
    try:
        serialized = base64.b64decode(encoded_payload, altchars=b"-_", validate=True)
    except (binascii.Error, ValueError) as exc:
        raise _canonical_sidecar_error("canonical payload is not valid base64url") from exc
    actual_digest = hashlib.sha256(serialized).hexdigest()
    if not hmac.compare_digest(actual_digest, digest):
        raise _canonical_sidecar_error("canonical payload SHA-256 mismatch")
    try:
        canonical = json.loads(serialized.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise _canonical_sidecar_error("canonical payload is not valid UTF-8 JSON") from exc
    if not isinstance(canonical, dict):
        raise _canonical_sidecar_error("canonical storyboard root must be an object")
    return (
        canonical,
        str(projection_contract),
        projection_sha256,
        audit_page_sha256,
    )


def _metadata_from_sheets(sheets: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]]:
    metadata_sheet = next(
        (sheet for sheet in sheets if sheet.get("name") in {"元数据", "metadata"}),
        None,
    )
    evidence_sheet = next(
        (
            sheet
            for sheet in sheets
            if sheet.get("name") in {"流程证据", "prevention_evidence"}
        ),
        None,
    )
    metadata: dict[str, Any] = {}
    if metadata_sheet:
        for row in metadata_sheet.get("rows", []):
            values = list(row.values())
            if values and values[0] not in (None, ""):
                metadata[str(values[0])] = _json_cell(
                    values[1] if len(values) > 1 else None,
                    None,
                )
    evidence: dict[str, Any] = {}
    if evidence_sheet:
        for row in evidence_sheet.get("rows", []):
            values = list(row.values())
            if values and values[0] not in (None, ""):
                evidence[str(values[0])] = _json_cell(
                    values[1] if len(values) > 1 else None,
                    None,
                )
    return metadata, evidence


def _parse_labeled_cell(value: Any) -> dict[str, str]:
    if value in (None, ""):
        return {}
    parsed: dict[str, str] = {}
    accepted = set(COMPACT_SHOT_FIELDS) | LEGACY_STANDALONE_FIELDS | {
        "切镜转场提示词",
        "光源",
        "光线方向",
        "色温",
    }
    current_label: str | None = None
    voice_nested_labels = {
        "基础声线", "言语行为", "当前表达", "发声方式", "音量", "语速与节奏",
        "呼吸", "停顿", "重音", "尾音", "禁止语气",
    }
    # Legacy formal cells used both newlines and full-/half-width pipes as
    # structural separators. Unlabelled fragments remain continuations of the
    # preceding field (for example movement modifiers and picture time stages).
    for raw_line in re.split(r"[\n|｜]", str(value)):
        line = raw_line.strip()
        if not line:
            continue
        candidate = ""
        text = ""
        labeled = re.match(r"^([^:：]+)\s*[:：]\s*(.*)$", line)
        if labeled:
            candidate = labeled.group(1).strip()
            text = labeled.group(2)
        if (
            current_label == "音色"
            and candidate in voice_nested_labels
            and (candidate != "发声方式" or not text.strip())
        ):
            prior = parsed.get("音色", "")
            parsed["音色"] = f"{prior}\n{line}".strip()
            continue
        if candidate in accepted:
            current_label = candidate
            if text.strip():
                # A legacy formal cell may repeat a structural label for every
                # dialogue or voice line.  Assignment here used to overwrite
                # all earlier lines, which made XLSX -> canonical recovery
                # silently lose dialogue. Preserve repeated atoms in order.
                prior = parsed.get(current_label, "")
                parsed[current_label] = f"{prior}\n{text.strip()}".strip()
            else:
                parsed.setdefault(current_label, "")
            continue
        if current_label:
            prior = parsed.get(current_label, "")
            parsed[current_label] = f"{prior}\n{line}".strip()
    for old, current in (
        ("光源", "光源来源"),
        ("光线方向", "光源方向"),
        ("色温", "色温倾向"),
    ):
        if not parsed.get(current) and parsed.get(old):
            parsed[current] = parsed[old]
        parsed.pop(old, None)
    return parsed


def _parse_function_scene_cell(value: Any) -> tuple[str, str]:
    text = str(value or "").strip()
    function = ""
    scene = ""
    function_match = re.search(
        r"功能类型\s*[:：]\s*(.*?)(?=\s*(?:[|｜]\s*)?场景连续性ID\s*[:：]|$)",
        text,
    )
    scene_match = re.search(r"场景连续性ID\s*[:：]\s*(.+)$", text)
    if function_match:
        function = function_match.group(1).strip()
    elif text and "场景连续性ID" not in text:
        function = text
    if scene_match:
        scene = scene_match.group(1).strip()
    elif not function_match and "场景连续性ID" not in text:
        # Legacy ten-column workbooks also store this cell as
        # "function / SCENE-ID". Recover only an unambiguous, explicitly
        # written identifier. A full-width slash may be part of the function;
        # prose, paths and multiple ASCII separators must not invent a scene.
        legacy_parts = re.split(r"\s+/\s+", text)
        if (
            len(legacy_parts) == 2
            and legacy_parts[0].strip()
            and re.fullmatch(
                r"[A-Za-z][A-Za-z0-9]*(?:[-_][A-Za-z0-9]+)+",
                legacy_parts[1],
            )
        ):
            function = legacy_parts[0].strip()
            scene = legacy_parts[1]
    return function, scene


def _normalized_mother_headers(headers: Any) -> list[str]:
    normalized: list[str] = []
    for value in list(headers or []):
        header = str(value or "").strip()
        normalized.append(
            FORMAL_MERGED_FIELD
            if header in FORMAL_MASTER_FIELD_ALIASES
            else header
        )
    return normalized


def _mother_sheet_candidates(sheets: list[dict[str, Any]]) -> list[dict[str, Any]]:
    expected = [MOTHER_TEMPLATE_COLUMNS, LEGACY_MOTHER_TEMPLATE_COLUMNS]
    return [
        sheet
        for sheet in sheets
        if not is_reading_sheet(sheet.get("name"))
        and _normalized_mother_headers(sheet.get("headers", [])) in expected
    ]


def _formal_master_cell(source: dict[str, Any]) -> Any:
    for header in FORMAL_MASTER_FIELD_ALIASES:
        if header in source:
            return source.get(header)
    return None


def _projection_cell_identity(value: Any) -> str:
    """Return the workbook-visible identity of one deterministic cell."""

    visible = _cell_value(value)
    if isinstance(visible, str):
        visible = unicodedata.normalize(
            "NFC",
            visible.replace("\r\n", "\n").replace("\r", "\n"),
        )
    if visible in (None, ""):
        visible = None
    elif isinstance(visible, float) and visible.is_integer():
        # Excel may return ``3`` for an authored ``3.0``.  The displayed and
        # mathematical cell are identical, while strings and booleans remain
        # type-distinct through JSON serialization.
        visible = int(visible)
    return json.dumps(
        json_safe(visible),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    )


def projection_rows_sha256(rows: list[dict[str, Any]]) -> str:
    """Hash trusted projected rows under their exact versioned column layout."""

    identities = [
        [_projection_cell_identity(row.get(header)) for header in _projection_columns(rows)]
        for row in rows
    ]
    payload = json.dumps(
        identities,
        ensure_ascii=False,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _visible_sheet_sha256(name: str, matrix: list[list[Any]]) -> str:
    identities = {
        "name": name,
        "matrix": [
            [_projection_cell_identity(value) for value in row]
            for row in matrix
        ],
    }
    payload = json.dumps(
        identities,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _assert_visible_projection_matches(
    sheets: list[dict[str, Any]],
    canonical: dict[str, Any],
) -> None:
    """Verify the visible formal master against its canonical sidecar.

    The hidden canonical is an integrity source, not permission to ignore the
    formal workbook fields. A stale or incorrectly generated master projection
    must therefore block import/audit even when the sidecar itself is intact.
    """

    expected_rows = formal_mother_template_rows(canonical, assume_compact=True)
    _assert_visible_projection_rows_match(sheets, expected_rows)


def _rows_for_projection(storyboard, contract):
    if contract in {PREVIOUS_PROJECTION_CONTRACT, "shortdrama-projector/0.5.4"}:
        return frozen_projector_0531.formal_mother_template_rows(storyboard)
    if contract == R5_LEGACY_PROJECTION_CONTRACT:
        return frozen_projector_0555.formal_mother_template_rows(storyboard, _legacy_layout=True)
    if contract == FROZEN_0555_PROJECTION_CONTRACT:
        return frozen_projector_0555.formal_mother_template_rows(storyboard)
    return formal_mother_template_rows(storyboard)


def _projection_columns(rows):
    return MOTHER_TEMPLATE_COLUMNS if rows and "光影" in rows[0] else LEGACY_MOTHER_TEMPLATE_COLUMNS


def _assert_auxiliary_sheets(sheets, storyboard, expected_rows, contract, *, include_reading=False):
    current = contract in {CURRENT_PROJECTION_CONTRACT, FROZEN_0555_PROJECTION_CONTRACT, R5_LEGACY_PROJECTION_CONTRACT}
    reader = (frozen_projector_0555.reading_matrices
              if contract in {FROZEN_0555_PROJECTION_CONTRACT, R5_LEGACY_PROJECTION_CONTRACT} else reading_matrices)
    expected_aux = reader(storyboard, expected_rows,
        _legacy_layout=contract == R5_LEGACY_PROJECTION_CONTRACT) if current else {}
    # Reading copies are outside audit. Only the exporter verifies its own
    # freshly generated copies; import, audit and recovery use the master.
    if not include_reading:
        sheets = [sheet for sheet in sheets if not is_reading_sheet(sheet.get("name"))]
        expected_aux = {name: rows for name, rows in expected_aux.items() if not is_reading_sheet(name)}
    names = [str(sheet.get("name") or "") for sheet in sheets
             if sheet.get("name") != CANONICAL_SIDECAR_SHEET]
    expected_names = ["00_范围与审计", MASTER_NAME if current else "正式生产母版", *expected_aux]
    if len(names) != len(set(names)) or sorted(names) != sorted(expected_names):
        raise _canonical_sidecar_error("current workbook has unexpected visible sheets")
    for name, expected in expected_aux.items():
        actual = next(item for item in sheets if item.get("name") == name)
        if _visible_sheet_sha256(name, actual.get("matrix") or []) != _visible_sheet_sha256(name, _used_value_matrix(expected)):
            raise _canonical_sidecar_error("visible reading/source projection mismatch: " + name)


def _visible_projection_contract(sheets: list[dict[str, Any]]) -> str:
    """Return ``current`` or ``legacy`` for the one compact mother sheet."""

    candidates = _mother_sheet_candidates(sheets)
    if len(candidates) != 1:
        raise _canonical_sidecar_error(
            "visible formal-master projection must contain exactly one versioned master sheet"
        )
    headers = [str(value or "").strip() for value in candidates[0].get("headers", [])]
    return "current" if FORMAL_MERGED_FIELD in headers else "legacy"


def _assert_visible_projection_rows_match(
    sheets: list[dict[str, Any]],
    expected_rows: list[dict[str, Any]],
) -> None:
    """Compare already-projected rows without compiling them a second time."""

    candidates = _mother_sheet_candidates(sheets)
    if len(candidates) != 1:
        raise _canonical_sidecar_error(
            "visible formal-master projection must contain exactly one versioned master sheet"
        )
    sheet = candidates[0]
    actual_headers = list(sheet.get("headers", []))
    expected_headers = list(_projection_columns(expected_rows))
    if _normalized_mother_headers(actual_headers) != expected_headers:
        raise _canonical_sidecar_error("visible columns differ from the signed projection contract")
    actual_rows = list(sheet.get("rows", []))
    if len(actual_rows) != len(expected_rows):
        raise _canonical_sidecar_error(
            "visible formal-master projection row count mismatch"
        )

    for row_index, (actual, expected) in enumerate(
        zip(actual_rows, expected_rows),
        start=2,
    ):
        for column_index, (actual_header, expected_header) in enumerate(
            zip(actual_headers, expected_headers),
            start=1,
        ):
            actual_value = actual.get(actual_header)
            expected_value = expected.get(expected_header)
            if _projection_cell_identity(actual_value) != _projection_cell_identity(
                expected_value
            ):
                raise _canonical_sidecar_error(
                    "visible formal-master projection mismatch at "
                    f"row {row_index}, column {column_index}"
                )


def _compact_shots_from_mother_sheet(sheet: dict[str, Any]) -> list[dict[str, Any]]:
    shots: list[dict[str, Any]] = []
    for source in sheet.get("rows", []):
        shot = _parse_labeled_cell(_formal_master_cell(source))
        shot["镜头编号"] = source.get("分镜序号")
        shot["时长"] = source.get("时长")
        shot["累计时长"] = source.get("累计时长")
        function, scene = _parse_function_scene_cell(
            source.get("功能类型/场景连续性ID")
        )
        shot["功能类型"] = function
        shot["场景连续性ID"] = scene
        for field in ("参考图锚定", "覆盖原文", "原文依据类型", "原文支撑句"):
            shot[field] = source.get(field)
        migrated = migrate_legacy_shot(shot, legacy=True)
        if not str(migrated.get("光影") or "").strip():
            raise ValueError(
                "materialization_input_missing: legacy migration cannot derive "
                f"光影 for {migrated.get('镜头编号') or '<unknown>'}"
            )
        shots.append(migrated)
    return shots


def workbook_to_storyboard(
    sheets: list[dict[str, Any]],
    *,
    migrate: bool = False,
    _allow_controlled_transitional_short_header: bool = False,
    _allowed_projection_contracts: frozenset[str] | None = None,
) -> dict[str, Any] | None:
    # Excluded reading pages cannot supply shots, metadata, or duplicate masters.
    sheets = [sheet for sheet in sheets if not is_reading_sheet(sheet.get("name"))]
    sidecar_result = _canonical_storyboard_from_sidecar(
        sheets,
        _allowed_projection_contracts=_allowed_projection_contracts,
    )
    if sidecar_result is not None:
        (
            sidecar_storyboard,
            projection_contract,
            projection_sha256,
            audit_page_sha256,
        ) = sidecar_result
        # v0.5.1 artifacts use the short formal-master header and current
        # projector.  A v0.5.0 workbook uses the legacy header and is checked
        # against a frozen read-only v0.5.0 projector.  The editable header is
        # therefore never enough to bypass visible/canonical verification.
        visible_contract = _visible_projection_contract(sheets)
        expected_visible_contract = (
            "current" if projection_sha256 is not None else "legacy"
        )
        controlled_transitional_short_header = bool(
            _allow_controlled_transitional_short_header
            and projection_contract == LEGACY_PROJECTION_CONTRACT
            and projection_sha256 is None
            and audit_page_sha256 is None
            and visible_contract == "current"
        )
        if (
            visible_contract != expected_visible_contract
            and not controlled_transitional_short_header
        ):
            raise _canonical_sidecar_error(
                "visible header does not match signed projection contract"
            )
        # The transitional branch uses the frozen v0.5.3 projector.  It never
        # treats those rows as v0.5.0 and never invokes today's projector.
        if projection_sha256 is not None:
            expected_rows = _rows_for_projection(sidecar_storyboard, projection_contract)
            _assert_auxiliary_sheets(sheets, sidecar_storyboard, expected_rows, projection_contract)
            scope_sheet = next(
                sheet for sheet in sheets if sheet.get("name") == "00_范围与审计"
            )
            actual_audit_page_sha256 = _visible_sheet_sha256(
                "00_范围与审计",
                list(scope_sheet.get("matrix") or []),
            )
            if actual_audit_page_sha256 != audit_page_sha256:
                raise _canonical_sidecar_error(
                    "visible audit page SHA-256 mismatch"
                )
            _assert_visible_projection_rows_match(sheets, expected_rows)
        elif controlled_transitional_short_header:
            visible_names = [
                str(sheet.get("name") or "")
                for sheet in sheets
                if sheet.get("name") != CANONICAL_SIDECAR_SHEET
            ]
            if sorted(visible_names) != sorted(
                ["00_范围与审计", "正式生产母版"]
            ):
                raise _canonical_sidecar_error(
                    "transitional workbook has unexpected visible sheets"
                )
            expected_rows = legacy_v053_mother_template_rows(
                sidecar_storyboard
            )
            _assert_visible_projection_rows_match(sheets, expected_rows)
        else:
            expected_rows = legacy_v050_mother_template_rows(sidecar_storyboard)
            _assert_visible_projection_rows_match(
                sheets,
                expected_rows,
            )
        if (
            projection_sha256 is not None
            and projection_rows_sha256(expected_rows) != projection_sha256
        ):
            raise _canonical_sidecar_error(
                "canonical projection SHA-256 mismatch"
            )
        # Audit callers must verify the artifact's original canonical bytes and
        # SHA before any migration. New-production callers opt into the one-way
        # compact migration explicitly.
        if migrate:
            return migrate_storyboard_to_compact(sidecar_storyboard)
        return sidecar_storyboard

    metadata, evidence = _metadata_from_sheets(sheets)
    mother_sheets = _mother_sheet_candidates(sheets)
    if len(mother_sheets) > 1:
        raise ValueError("workbook contains duplicate formal-master sheets")
    mother_sheet = mother_sheets[0] if mother_sheets else None
    if mother_sheet:
        if _normalized_mother_headers(mother_sheet.get("headers", [])) == MOTHER_TEMPLATE_COLUMNS:
            raise _canonical_sidecar_error("the current eleven-column workbook requires its canonical sidecar")
        shots = _compact_shots_from_mother_sheet(mother_sheet)
        total_duration = metadata.get("total_duration")
        if total_duration in (None, "") and shots:
            total_duration = round(
                sum(float(shot.get("时长") or 0) for shot in shots),
                1,
            )
        return {
            "project_version": metadata.get(
                "project_version",
                "V16.3-148-R10-StableRuntime",
            ),
            "schema_version": COMPACT_SCHEMA_VERSION,
            "workflow_profile": metadata.get("workflow_profile", "R10-PG1"),
            "prevention_evidence": evidence,
            "field_order": list(COMPACT_SHOT_FIELDS),
            "source_lock": metadata.get("source_lock", {}),
            "shots": shots,
            "total_duration": total_duration or 0,
        }

    # Legacy R10.0 workbook import is read-only. Its atomic and batch sheets are
    # accepted only long enough to migrate each shot; batches are discarded.
    shots_sheet = next(
        (sheet for sheet in sheets if "镜头编号" in sheet.get("headers", [])),
        None,
    )
    if not shots_sheet:
        return None
    batches_sheet = next(
        (sheet for sheet in sheets if "批次编号" in sheet.get("headers", [])),
        None,
    )
    identity_sheet = next(
        (sheet for sheet in sheets if sheet.get("name") == "内部标识"),
        None,
    )
    shots = list(shots_sheet.get("rows", []))
    batches = list(batches_sheet.get("rows", [])) if batches_sheet else []
    for row in batches:
        if isinstance(row.get("覆盖镜头"), str):
            row["覆盖镜头"] = [
                item.strip()
                for item in row["覆盖镜头"].replace("，", ",").split(",")
                if item.strip()
            ]
    if identity_sheet:
        shot_identity = {
            str(row.get("显示编号") or ""): row
            for row in identity_sheet.get("rows", [])
            if row.get("类型") == "shot"
        }
        batch_identity = {
            str(row.get("显示编号") or ""): row
            for row in identity_sheet.get("rows", [])
            if row.get("类型") == "batch"
        }
        for row in shots:
            identity = shot_identity.get(str(row.get("镜头编号") or ""))
            if identity:
                for field in ("shot_uid", "episode_id", "scene_uid"):
                    if identity.get(field) not in (None, ""):
                        row[field] = identity[field]
        for row in batches:
            identity = batch_identity.get(str(row.get("批次编号") or ""))
            if identity:
                for field in ("batch_uid", "episode_id"):
                    if identity.get(field) not in (None, ""):
                        row[field] = identity[field]
    total_duration = metadata.get("total_duration")
    if total_duration in (None, "") and shots:
        total_duration = round(
            sum(float(shot.get("时长") or 0) for shot in shots),
            1,
        )
    legacy = {
        "project_version": metadata.get(
            "project_version",
            "V16.3-148-R10-StableRuntime",
        ),
        "schema_version": "R10.0",
        "workflow_profile": metadata.get("workflow_profile", "R10-PG1"),
        "prevention_evidence": evidence,
        "field_order": shots_sheet.get("headers", []),
        "source_lock": metadata.get("source_lock", {}),
        "shots": shots,
        "batches": batches,
        "total_duration": total_duration or 0,
    }
    # Without a canonical sidecar there is no historical canonical SHA to
    # preserve. The legacy atom sheet is therefore migrated exactly once on
    # import and never re-exported in its old shape.
    return migrate_storyboard_to_compact(legacy)


def _external_workbook_audit(sheets: list[dict[str, Any]], raw_sha256: str) -> dict[str, Any] | None:
    """Read a saved master without materializing or certifying production history.

    This adapter is deliberately separate from production/sidecar recovery. All
    original cells remain evidence, including malformed or missing fields.
    """
    masters = _mother_sheet_candidates(sheets)
    if not masters:
        return None
    if len(masters) != 1:
        raise ValueError("workbook contains duplicate formal-master sheets")
    sheet = masters[0]
    headers = sheet["headers"]
    labels = set(COMPACT_SHOT_FIELDS) | {"镜头类型"}
    pattern = re.compile(r"(?m)^(" + "|".join(map(re.escape, labels)) + r")\s*[:：]")

    def fields(text):
        text = str(text or "")
        matches = list(pattern.finditer(text))
        result = {}
        for index, match in enumerate(matches):
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            value = text[match.end():end].strip()
            # Preserve repeated sections rather than silently discarding one.
            result[match[1]] = "\n".join(filter(None, (result.get(match[1]), value)))
        return result

    shots, rows = [], []
    matrix = sheet.get("matrix", [])
    header_index = next(index for index, values in enumerate(matrix[:20]) if any(value not in (None, "") for value in values))
    extra_cells = [f"{get_column_letter(column)}{row_number}"
                   for row_number, values in enumerate(matrix, 1)
                   for column, value in enumerate(values, 1)
                   if column > len(headers) and value not in (None, "")]
    if extra_cells:
        raise ValueError("external_workbook_unmapped_master_cells: " + ", ".join(extra_cells))
    for row_number, values in enumerate(matrix[header_index + 1:], header_index + 2):
        if not any(value is not None and value != "" for value in values):
            continue
        source = dict(zip(headers, values))
        shot = fields(_formal_master_cell(source))
        kind, separator, visible = str(shot.get("镜头类型") or "").partition("｜")
        if separator:
            shot["镜头类型"] = kind.strip()
            shot["可见主体"] = visible.removeprefix("可见主体：").strip()
        shot.update(fields(source.get("光影")))
        shot.update({"镜头编号": source.get("分镜序号"), "时长": source.get("时长"),
                     "累计时长": source.get("累计时长")})
        shot["功能类型"], shot["场景连续性ID"] = _parse_function_scene_cell(source.get("功能类型/场景连续性ID"))
        for field in ("参考图锚定", "覆盖原文", "原文依据类型", "原文支撑句"):
            shot[field] = source.get(field)
        missing = [field for field in COMPACT_SHOT_FIELDS if field not in shot]
        shots.append({field: shot.get(field, "") for field in COMPACT_SHOT_FIELDS})
        cells = {}
        for column, header in enumerate(headers, 1):
            value = source.get(header)
            cells[header] = {"address": f"{get_column_letter(column)}{row_number}", "value": value,
                             "sha256": hashlib.sha256(json.dumps(value, ensure_ascii=False,
                                 sort_keys=True, default=str).encode("utf-8")).hexdigest()}
        rows.append({"shot_id": shot["镜头编号"], "row": row_number, "cells": cells,
                     "missing_fields": missing})
    metadata, evidence = _metadata_from_sheets(
        [s for s in sheets if not is_reading_sheet(s.get("name"))])
    return {
        "schema_version": COMPACT_SCHEMA_VERSION, "field_order": list(COMPACT_SHOT_FIELDS),
        "shots": shots, "prevention_evidence": evidence,
        "source_lock": metadata.get("source_lock", {}),
        "audit_workbook_evidence": {
            "contract": "external-workbook-audit-v1", "raw_sha256": raw_sha256,
            "master_sheet": sheet["name"], "column_contract": len(headers), "rows": rows,
            "workbook_audit_scope": workbook_audit_scope(),
            "historical_canonical_available": False,
            "production_provenance": "unknown", "production_eligible": False,
            "identity_basis": "current_read_only_mapping_of_original_cells",
        },
    }


def extract_xlsx(data: bytes, *, audit_only: bool = False) -> dict[str, Any]:
    _validate_zip_archive(data)
    workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        headers, rows = _rows_from_sheet(sheet)
        sheets.append({
            "name": sheet.title,
            "headers": headers,
            "row_count": len(rows),
            "rows": rows,
            "matrix": _sheet_value_matrix(sheet),
        })
    canonical_sidecar_present = any(
        sheet.get("name") == CANONICAL_SIDECAR_SHEET for sheet in sheets
    )
    # A present sidecar must always validate. An invalid sidecar is never
    # downgraded into an unsigned external mapping by the audit option.
    canonical = (_external_workbook_audit(sheets, hashlib.sha256(data).hexdigest())
                 if audit_only and not canonical_sidecar_present else workbook_to_storyboard(sheets))
    saved_evidence = None
    if audit_only and canonical_sidecar_present:
        mapped = _external_workbook_audit(sheets, hashlib.sha256(data).hexdigest())
        if mapped is not None:
            saved_evidence = mapped["audit_workbook_evidence"]
            saved_evidence.update(historical_canonical_available=True,
                identity_basis="verified_original_sidecar_with_separate_current_audit_cells",
                projection_integrity_verified=True,
                projection_contract=next((row.get("projection_contract") or LEGACY_PROJECTION_CONTRACT
                    for sheet in sheets if sheet.get("name") == CANONICAL_SIDECAR_SHEET
                    for row in sheet.get("rows", []) if row.get("record_type") == "manifest"), None))
    workbook.close()
    public_sheets = [sheet for sheet in sheets if sheet.get("name") != CANONICAL_SIDECAR_SHEET]
    return {
        "kind": "workbook",
        "sheet_count": len(public_sheets),
        "sheets": public_sheets,
        "workbook_audit_scope": workbook_audit_scope(),
        "canonical_storyboard": canonical,
        "canonical_sidecar_present": canonical_sidecar_present,
        **({"audit_workbook_evidence": saved_evidence} if saved_evidence is not None else {}),
    }


def extract_xlsx_for_controlled_handoff_revalidation(
    data: bytes,
) -> dict[str, Any]:
    """Verify one immutable workbook for controlled handoff revalidation.

    This is deliberately narrower than :func:`extract_xlsx`.  General reads
    keep rejecting the transitional eight-column sidecar paired with the short
    formal-master header.  The recovery service may use this helper only after
    its production/session/work-order gates.  The sole unsigned-projector
    compatibility case must match the frozen v0.5.3 projector exactly; current
    eleven-column sidecars retain their projection and audit-page SHA guards.
    """

    _validate_zip_archive(data)
    workbook = load_workbook(
        io.BytesIO(data), data_only=False, read_only=True
    )
    sheets = []
    for sheet in workbook.worksheets:
        headers, rows = _rows_from_sheet(sheet)
        sheets.append(
            {
                "name": sheet.title,
                "headers": headers,
                "row_count": len(rows),
                "rows": rows,
                "matrix": _sheet_value_matrix(sheet),
                "sheet_state": sheet.sheet_state,
            }
        )

    sidecar_result = _canonical_storyboard_from_sidecar(
        sheets,
        _allowed_projection_contracts=(
            CONTROLLED_HISTORICAL_PROJECTION_CONTRACTS
        ),
    )
    if sidecar_result is None:
        raise _canonical_sidecar_error(
            "controlled handoff revalidation requires a canonical sidecar"
        )
    sidecar_sheet = next(
        sheet
        for sheet in sheets
        if sheet.get("name") == CANONICAL_SIDECAR_SHEET
    )
    if sidecar_sheet.get("sheet_state") != "veryHidden":
        raise _canonical_sidecar_error(
            "controlled handoff sidecar must remain veryHidden"
        )
    (
        sidecar_storyboard,
        projection_contract,
        projection_sha256,
        audit_page_sha256,
    ) = sidecar_result
    visible_contract = _visible_projection_contract(sheets)
    expected_visible_contract = (
        "current" if projection_sha256 is not None else "legacy"
    )
    compatibility_applied = bool(
        projection_contract == LEGACY_PROJECTION_CONTRACT
        and projection_sha256 is None
        and audit_page_sha256 is None
        and visible_contract == "current"
    )
    if projection_sha256 is None and not compatibility_applied:
        raise _canonical_sidecar_error(
            "controlled handoff revalidation supports only the exact v0.5.3 unsigned projection"
        )
    if (
        visible_contract != expected_visible_contract
        and not compatibility_applied
    ):
        raise _canonical_sidecar_error(
            "visible header does not match signed projection contract"
        )

    canonical = workbook_to_storyboard(
        sheets,
        _allow_controlled_transitional_short_header=(
            compatibility_applied
        ),
        _allowed_projection_contracts=(
            CONTROLLED_HISTORICAL_PROJECTION_CONTRACTS
        ),
    )
    if canonical != sidecar_storyboard:
        raise _canonical_sidecar_error(
            "controlled handoff canonical round-trip mismatch"
        )

    mother_sheet = _mother_sheet_candidates(sheets)[0]
    exact_headers = [
        str(value or "").strip()
        for value in mother_sheet.get("headers", [])
    ]
    exact_header_row = list(
        (mother_sheet.get("matrix") or [[]])[0]
    )
    if compatibility_applied and exact_header_row != list(
        LEGACY_MOTHER_TEMPLATE_COLUMNS
    ):
        raise _canonical_sidecar_error(
            "transitional workbook header is not the exact short ten-column contract"
        )
    header_sha256 = hashlib.sha256(
        json.dumps(
            exact_headers,
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    canonical_payload_sha256 = hashlib.sha256(
        json.dumps(
            json_safe(sidecar_storyboard),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")
    ).hexdigest()

    expected_rows = (
        legacy_v053_mother_template_rows(sidecar_storyboard)
        if compatibility_applied
        else (
            _rows_for_projection(sidecar_storyboard, projection_contract)
            if projection_sha256 is not None
            else legacy_v050_mother_template_rows(sidecar_storyboard)
        )
    )
    visible_projection_sha256 = projection_rows_sha256(expected_rows)

    _assert_auxiliary_sheets(sheets, sidecar_storyboard, expected_rows,
                             PREVIOUS_PROJECTION_CONTRACT if compatibility_applied else projection_contract)
    scope_sheet = next(
        sheet for sheet in sheets if sheet.get("name") == "00_范围与审计"
    )
    visible_audit_page_sha256 = _visible_sheet_sha256(
        "00_范围与审计", list(scope_sheet.get("matrix") or [])
    )

    public_sheets = [
        sheet
        for sheet in sheets
        if sheet.get("name") != CANONICAL_SIDECAR_SHEET
    ]
    return {
        "kind": "workbook",
        "sheet_count": len(public_sheets),
        "sheets": public_sheets,
        "canonical_storyboard": canonical,
        "canonical_sidecar_present": True,
        "controlled_projection_revalidation": {
            "spec": "controlled-workbook-projection-revalidation-v1",
            "mode": (
                "legacy_envelope_short_header_v053_rows"
                if compatibility_applied
                else "strict_projection_contract"
            ),
            "compatibility_applied": compatibility_applied,
            "sidecar_projection_contract": projection_contract,
            "visible_header_contract": visible_contract,
            "visible_header_sha256": header_sha256,
            "sidecar_projection_sha256": projection_sha256,
            "visible_projection_sha256": (
                visible_projection_sha256
            ),
            "sidecar_audit_page_sha256": audit_page_sha256,
            "visible_audit_page_sha256": (
                visible_audit_page_sha256
            ),
            "canonical_payload_sha256": canonical_payload_sha256,
            "sidecar_sheet_state": "veryHidden",
            "selected_projector": (
                "shortdrama-projector/0.5.3-frozen"
                if compatibility_applied
                else projection_contract
            ),
            "selected_projector_source_commit": (
                "9deb528db3f9eb27e0f406634036e98b8a3d1f80"
                if compatibility_applied
                else None
            ),
            "selected_projector_source_git_blob": (
                "d8d4cce11a50c40d5594867a9de8ae89d717fde1"
                if compatibility_applied
                else None
            ),
            "visible_rows_match_selected_projector": True,
            "visible_projection_matches_sidecar": bool(
                projection_sha256 is not None
                and visible_projection_sha256 == projection_sha256
            ),
            "visible_audit_page_matches_sidecar": bool(
                audit_page_sha256 is not None
                and visible_audit_page_sha256 == audit_page_sha256
            ),
        },
    }


def _extract_docx(data: bytes) -> dict[str, Any]:
    _validate_zip_archive(data)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs = []
    for paragraph in root.iter(f"{namespace}p"):
        text = "".join(node.text or "" for node in paragraph.iter(f"{namespace}t")).strip()
        if text:
            paragraphs.append(text)
    return {
        "kind": "text",
        "format": "docx",
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "text": "\n".join(paragraphs),
    }


def _extract_pdf(data: bytes) -> dict[str, Any]:
    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return {
        "kind": "text",
        "format": "pdf",
        "page_count": len(pages),
        "pages": pages,
        "text": "\n\n".join(pages),
    }


def _extract_csv(data: bytes) -> dict[str, Any]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    reader.fieldnames = _unique_headers(reader.fieldnames or [])
    rows = [{str(key): json_safe(value) for key, value in row.items()} for row in reader]
    return {"kind": "table", "format": "csv", "headers": reader.fieldnames or [], "row_count": len(rows), "rows": rows}


def _extract_mm(data: bytes) -> dict[str, Any]:
    root = ElementTree.fromstring(data)
    lines: list[str] = []

    def visit(node: Any, depth: int) -> None:
        text = node.attrib.get("TEXT") or node.attrib.get("text")
        if text:
            lines.append(f"{'  ' * depth}- {text}")
        for child in list(node):
            visit(child, depth + (1 if text else 0))

    visit(root, 0)
    return {
        "kind": "text",
        "format": "freemind",
        "node_count": len(lines),
        "paragraphs": lines,
        "text": "\n".join(lines),
    }


def extract_file(data: bytes, file_name: str, mime_type: str | None = None,
                 *, audit_only: bool = False) -> dict[str, Any]:
    suffix = Path(file_name).suffix.casefold()
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx(data, audit_only=audit_only)
    if suffix == ".docx":
        return _extract_docx(data)
    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".csv":
        return _extract_csv(data)
    if suffix == ".json":
        return {"kind": "json", "format": "json", "value": json.loads(data.decode("utf-8-sig"))}
    if suffix == ".mm":
        return _extract_mm(data)
    if suffix in {".md", ".txt", ".srt", ".vtt"} or (mime_type or "").startswith("text/"):
        return {"kind": "text", "format": suffix.lstrip(".") or "text", "text": data.decode("utf-8-sig")}
    guessed = mimetypes.guess_type(file_name)[0]
    raise ValueError(f"unsupported input format: {suffix or mime_type or guessed or 'unknown'}")


def _cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, str):
        # Normalize before XML serialization so Windows text-mode writers do
        # not expand an authored CRLF into CRCRLF and add blank lines.
        # Only display cells are normalized; canonical sidecars remain exact.
        value = value.replace("\r\n", "\n").replace("\r", "\n")
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        # Storyboard fields are prose, never formulas. Escape spreadsheet
        # formula prefixes to prevent execution when a generated XLSX is opened.
        return f"'{value}"
    return value


def _write_table(workbook: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> Any:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header)) for header in headers])
    sheet.freeze_panes = "A2"
    if headers:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6F5F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, header in enumerate(headers, 1):
        width = 14
        if header in {
            "画面内容",
            "动作与表演",
            "台词",
            "发声方式",
            "独立执行提示词",
            "正式分镜母版·字段合并",
            "GPT关键帧生成指令",
            "Seedance 2.5独立正式执行分镜",
            "证据",
            "修复动作",
        }:
            width = 34
        sheet.column_dimensions[get_column_letter(index)].width = width
    return sheet


def _write_canonical_sidecar(
    workbook: Workbook,
    storyboard: dict[str, Any],
    projection_rows: list[dict[str, Any]],
    audit_page_sha256: str,
) -> None:
    canonical = json_safe(storyboard)
    if not isinstance(canonical, dict):
        raise ValueError("canonical storyboard root must be an object")
    serialized = json.dumps(
        canonical,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")
    digest = hashlib.sha256(serialized).hexdigest()
    projection_digest = projection_rows_sha256(projection_rows)
    encoded = base64.urlsafe_b64encode(serialized).decode("ascii")
    chunks = [
        encoded[offset : offset + CANONICAL_SIDECAR_CHUNK_CHARS]
        for offset in range(0, len(encoded), CANONICAL_SIDECAR_CHUNK_CHARS)
    ]
    if not chunks:
        raise ValueError("canonical storyboard serialization is empty")

    sheet = workbook.create_sheet(CANONICAL_SIDECAR_SHEET)
    sheet.append(CANONICAL_SIDECAR_HEADERS)
    sheet.append(
        [
            "manifest",
            CANONICAL_SIDECAR_MARKER,
            CANONICAL_SIDECAR_FORMAT,
            CANONICAL_SIDECAR_ENCODING,
            digest,
            None,
            len(chunks),
            None,
            CURRENT_PROJECTION_CONTRACT,
            projection_digest,
            audit_page_sha256,
        ]
    )
    for index, chunk in enumerate(chunks, 1):
        sheet.append(
            [
                "chunk",
                CANONICAL_SIDECAR_MARKER,
                CANONICAL_SIDECAR_FORMAT,
                CANONICAL_SIDECAR_ENCODING,
                digest,
                index,
                len(chunks),
                f"{CANONICAL_SIDECAR_PAYLOAD_PREFIX}{chunk}",
                CURRENT_PROJECTION_CONTRACT,
                projection_digest,
                audit_page_sha256,
            ]
        )
    sheet.sheet_state = "veryHidden"


def _write_scope_audit_sheet(
    workbook: Workbook,
    storyboard: dict[str, Any],
    audit_report: dict[str, Any],
    audit_overlay: _VerifiedAuditOverlay | None = None,
) -> Any:
    """Write the sample-compatible single visible scope/audit page."""

    if audit_overlay is not None and not isinstance(
        audit_overlay, _VerifiedAuditOverlay
    ):
        raise ValueError(
            "independent_audit_overlay_unverified: plain mappings and caller-authored "
            "certificate claims are not accepted"
        )
    overlay = audit_overlay._as_dict() if audit_overlay is not None else {}
    signed = audit_overlay is not None
    if signed:
        required = {
            "status",
            "conclusion",
            "initial_issue_count",
            "residual_issue_count",
            "introduced_issue_count",
            "blocking_issue_count",
            "audit_trace_id",
            "certificate_id",
            "certificate_integrity_hmac",
            "canonical_sha256",
            "source_artifact_raw_sha256",
        }
        missing = sorted(
            key for key in required if overlay.get(key) in (None, "")
        )
        if missing:
            raise ValueError(
                "independent_audit_overlay_incomplete: " + ",".join(missing)
            )
        status = str(overlay["status"]).upper()
        if status not in {"PASS", "FAIL_CURRENT", "BLOCKED"}:
            raise ValueError("independent_audit_overlay_status_invalid")
        for key in (
            "initial_issue_count",
            "residual_issue_count",
            "introduced_issue_count",
            "blocking_issue_count",
        ):
            if not isinstance(overlay.get(key), int) or int(overlay[key]) < 0:
                raise ValueError(
                    f"independent_audit_overlay_count_invalid: {key}"
                )
        for key in ("canonical_sha256", "source_artifact_raw_sha256"):
            if not re.fullmatch(r"[0-9a-f]{64}", str(overlay[key])):
                raise ValueError(
                    f"independent_audit_overlay_fingerprint_invalid: {key}"
                )
        if status == "PASS" and any(
            int(overlay[key])
            for key in (
                "residual_issue_count",
                "introduced_issue_count",
                "blocking_issue_count",
            )
        ):
            raise ValueError(
                "independent_audit_overlay_pass_requires_zero_open_issues"
            )
        report_issues = list(audit_report.get("issues") or [])
        if status == "PASS" and report_issues:
            raise ValueError(
                "independent_audit_overlay_pass_report_has_issues"
            )
        if status in {"FAIL_CURRENT", "BLOCKED"} and not report_issues:
            raise ValueError(
                "independent_audit_overlay_failure_report_has_no_issues"
            )
        if status in {"FAIL_CURRENT", "BLOCKED"} and not any(
            int(overlay[key])
            for key in (
                "initial_issue_count",
                "residual_issue_count",
                "introduced_issue_count",
                "blocking_issue_count",
            )
        ):
            raise ValueError(
                "independent_audit_overlay_failure_requires_open_issues"
            )
        report_summary = audit_report.get("audit_summary")
        if isinstance(report_summary, dict):
            for key in (
                "initial_issue_count",
                "residual_issue_count",
                "introduced_issue_count",
                "blocking_issue_count",
            ):
                if key in report_summary and int(overlay[key]) != int(
                    report_summary[key]
                ):
                    raise ValueError(
                        f"independent_audit_overlay_report_count_mismatch: {key}"
                    )
            report_final = str(report_summary.get("final_status") or "").upper()
            if report_final and report_final != status:
                raise ValueError(
                    "independent_audit_overlay_report_status_mismatch"
                )
    else:
        status = "PENDING_AUDIT"

    sheet = workbook.create_sheet("00_范围与审计")
    sheet.append(
        [
            (
                f"00_范围与审计｜{status}｜独立审计已签发"
                if signed
                else "00_范围与审计｜PENDING_AUDIT｜待独立审计"
            ),
            None,
            None,
            None,
            None,
        ]
    )
    sheet.append([])
    sheet.append(["范围字段", "值", None, None, None])
    scope_rows = [
        ("independent_audit_status", status),
        (
            "independent_audit_conclusion",
            overlay.get("conclusion")
            if signed
            else "待独立审计／未签发结论",
        ),
        ("initial_issue_count", overlay.get("initial_issue_count")),
        ("residual_issue_count", overlay.get("residual_issue_count")),
        ("introduced_issue_count", overlay.get("introduced_issue_count")),
        ("blocking_issue_count", overlay.get("blocking_issue_count")),
        ("audit_trace_id", overlay.get("audit_trace_id")),
        ("certificate_id", overlay.get("certificate_id")),
        (
            "certificate_integrity_hmac",
            overlay.get("certificate_integrity_hmac"),
        ),
        ("canonical_sha256", overlay.get("canonical_sha256")),
        (
            "source_artifact_raw_sha256",
            overlay.get("source_artifact_raw_sha256"),
        ),
        ("project_version", storyboard.get("project_version")),
        ("schema_version", storyboard.get("schema_version")),
        ("workflow_profile", storyboard.get("workflow_profile")),
        ("total_duration", storyboard.get("total_duration")),
        (
            "production_evidence_location",
            "canonical sidecar（生产证据不作为独立审计结论）",
        ),
    ]
    for label, value in scope_rows:
        sheet.append([label, _cell_value(value), None, None, None])

    # ``source_lock`` and ``prevention_evidence`` are producer-authored
    # canonical data.  They stay recoverable in the integrity-bound hidden
    # sidecar, but are intentionally not echoed onto the visible audit page:
    # otherwise a producer could inject PASS/zero-issue/certificate text before
    # independent audit has issued any conclusion.

    sheet.append([])
    issue_header_row = sheet.max_row + 1
    sheet.append([
        "独立审计错误码" if signed else "生产预检错误码",
        "严重度",
        "位置",
        "证据",
        "修复动作",
    ])
    issue_rows = [
        [
            issue.get("error_code"),
            issue.get("severity"),
            issue.get("location"),
            issue.get("evidence"),
            issue.get("repair_action"),
        ]
        for issue in audit_report.get("issues", [])
    ]
    if not issue_rows:
        if signed:
            sentinel = (
                "INDEPENDENT_AUDIT_PASS"
                if status == "PASS"
                else f"INDEPENDENT_AUDIT_{status}"
            )
            issue_rows = [[
                sentinel,
                "info" if status == "PASS" else "fail",
                "全表",
                (
                    "独立审计证书已签发；残留0、新增0、阻塞0"
                    if status == "PASS"
                    else f"独立审计证书状态为{status}"
                ),
                "无需修复" if status == "PASS" else "按审计问题修复后重审",
            ]]
        else:
            issue_rows = [[
                "PENDING_AUDIT",
                "info",
                "全表",
                "生产预检未发现阻塞项；不等于独立审计结论",
                "等待独立审计签发结论",
            ]]
    for row in issue_rows:
        sheet.append([_cell_value(value) for value in row])

    green = PatternFill("solid", fgColor="1F6F5F")
    for row_index in (1, 3, issue_header_row):
        for cell in sheet[row_index]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = green
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            if cell.row not in {1, 3}:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate((24, 28, 24, 48, 36), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.merge_cells("A1:E1")
    sheet.row_dimensions[1].height = 28
    sheet.column_dimensions["A"].width = 39
    sheet.column_dimensions["B"].width = 58
    sheet.freeze_panes = "A3"
    return sheet


def storyboard_to_xlsx(
    storyboard: dict[str, Any],
    audit_report: dict[str, Any],
    *,
    assume_compact: bool = False,
    audit_overlay: _VerifiedAuditOverlay | None = None,
) -> bytes:
    # Export is a one-way compact projection. Legacy atom sheets, batch sheets,
    # model-authored derived prose and Seedance prompt copies are never emitted.
    compact = storyboard if assume_compact else migrate_storyboard_to_compact(storyboard)
    workbook = Workbook()
    workbook.remove(workbook.active)
    scope_sheet = _write_scope_audit_sheet(
        workbook,
        compact,
        audit_report,
        audit_overlay=audit_overlay,
    )
    formal_rows = formal_mother_template_rows(compact, assume_compact=True)
    formal_sheet = _write_table(
        workbook,
        MASTER_NAME,
        list(MOTHER_TEMPLATE_COLUMNS),
        formal_rows,
    )
    for index, width in enumerate(WIDTHS, 1):
        formal_sheet.column_dimensions[get_column_letter(index)].width = width
    formal_sheet.row_dimensions[1].height = 30
    for row_index in range(2, formal_sheet.max_row + 1):
        formal_sheet.row_dimensions[row_index].height = 409

    matrices = reading_matrices(compact, formal_rows)
    write_reading_sheets(workbook, matrices)
    # R5 reading pages first, then copyable master and source guide; retain the
    # existing signed audit page and very-hidden canonical evidence unchanged.
    reading_names = [name for name in matrices if name != "使用与来源"]
    names = [*reading_names, MASTER_NAME, "使用与来源", scope_sheet.title]
    workbook._sheets = [workbook[name] for name in names]

    audit_page_sha256 = _visible_sheet_sha256(
        scope_sheet.title,
        _sheet_value_matrix(scope_sheet),
    )
    _write_canonical_sidecar(
        workbook,
        compact,
        formal_rows,
        audit_page_sha256,
    )

    stream = io.BytesIO()
    workbook.save(stream)
    payload = stream.getvalue()

    # Re-open the serialized workbook before handing it to persistence.  This
    # catches projection drift introduced by cell conversion, writer behavior,
    # or a future exporter refactor while the canonical sidecar is still valid.
    reopened = load_workbook(io.BytesIO(payload), data_only=False, read_only=True)
    sheets = []
    for sheet in reopened.worksheets:
        headers, rows = _rows_from_sheet(sheet)
        sheets.append(
            {
                "name": sheet.title,
                "headers": headers,
                "row_count": len(rows),
                "rows": rows,
                "matrix": _sheet_value_matrix(sheet),
            }
        )
    reopened.close()
    roundtrip_result = _canonical_storyboard_from_sidecar(sheets)
    roundtrip = roundtrip_result[0] if roundtrip_result is not None else None
    _assert_visible_projection_rows_match(sheets, formal_rows)
    _assert_auxiliary_sheets(sheets, compact, formal_rows, CURRENT_PROJECTION_CONTRACT,
                             include_reading=True)
    if roundtrip != json_safe(compact):
        raise ValueError(
            "canonical sidecar invalid: saved workbook canonical round-trip mismatch"
        )
    return payload


def preview_extracted(value: dict[str, Any], section: str | None, offset: int, limit: int) -> dict[str, Any]:
    offset = max(0, int(offset))
    limit = max(1, min(int(limit), 200))
    kind = value.get("kind")
    if kind == "text":
        text = str(value.get("text") or "")
        return {"kind": kind, "offset": offset, "limit": limit, "total": len(text), "content": text[offset : offset + limit]}
    if kind == "workbook":
        if section == "canonical_storyboard":
            return {"kind": "json", "content": value.get("canonical_storyboard")}
        sheets = value.get("sheets", [])
        sheet = _workbook_sheet(value, section)
        if not sheet:
            return {"kind": kind, "available_sections": [item.get("name") for item in sheets], "rows": []}
        rows = sheet.get("rows", [])
        return {
            "kind": kind,
            "section": sheet.get("name"),
            "audit_included": not is_reading_sheet(sheet.get("name")),
            "workbook_audit_scope": workbook_audit_scope(),
            "headers": sheet.get("headers", []),
            "offset": offset,
            "limit": limit,
            "total": len(rows),
            "rows": rows[offset : offset + limit],
        }
    if kind == "table":
        rows = value.get("rows", [])
        return {"kind": kind, "headers": value.get("headers", []), "offset": offset, "limit": limit, "total": len(rows), "rows": rows[offset : offset + limit]}
    if kind == "json":
        content = value.get("value")
        if isinstance(content, dict):
            if not section:
                return {
                    "kind": kind,
                    "available_sections": sorted(content.keys()),
                    "summary": {
                        key: (len(item) if isinstance(item, (list, dict, str)) else item)
                        for key, item in content.items()
                    },
                }
            if section not in content:
                return {"kind": kind, "available_sections": sorted(content.keys()), "content": None}
            content = content[section]
        if isinstance(content, list):
            return {"kind": kind, "offset": offset, "limit": limit, "total": len(content), "content": content[offset : offset + limit]}
        if isinstance(content, dict):
            return {"kind": kind, "content": content}
        text = str(content or "")
        return {"kind": kind, "offset": offset, "limit": limit, "total": len(text), "content": text[offset : offset + limit]}
    return {"kind": kind, "content": value}


def _paragraphs_from_text(text: str) -> list[str]:
    blocks = [item.strip() for item in re.split(r"\n\s*\n", text) if item.strip()]
    if len(blocks) <= 1:
        lines = [item.strip() for item in text.splitlines() if item.strip()]
        if lines:
            return lines
    return blocks


def _range_dict(value: Any) -> dict[str, Any]:
    if hasattr(value, "model_dump") and callable(value.model_dump):
        return value.model_dump(exclude_none=True)
    if isinstance(value, dict):
        return dict(value)
    raise ValueError("bulk_ranges entries must be objects")


def _bulk_char_result(
    text: str,
    *,
    section: str | None,
    offset: int,
    limit: int | None,
    budget: int,
) -> tuple[dict[str, Any], int]:
    total = len(text)
    requested = limit if limit is not None else DEFAULT_BULK_CHARACTER_LIMIT
    take = min(max(0, requested), max(0, budget))
    content = text[offset : offset + take]
    next_offset = min(total, offset + len(content))
    result = {
        "section": section,
        "unit": "characters",
        "offset": offset,
        "limit": requested,
        "total": total,
        "content": content,
        "next_offset": next_offset,
        "complete": next_offset >= total,
    }
    return result, len(content)


def _bulk_text_items_result(
    items: list[Any],
    *,
    section: str | None,
    unit: str,
    offset: int,
    inner_offset: int,
    limit: int | None,
    budget: int,
) -> tuple[dict[str, Any], int]:
    total = len(items)
    requested = limit if limit is not None else DEFAULT_BULK_ITEM_LIMITS[unit]
    index = min(max(0, offset), total)
    remaining = max(0, budget)
    parts: list[str] = []
    completed_items = 0
    next_inner_offset = max(0, inner_offset) if index < total else 0

    while index < total and completed_items < requested and remaining > 0:
        item = str(items[index] or "")
        item_offset = min(next_inner_offset, len(item))
        segment = item[item_offset:]
        separator = "\n\n" if parts else ""
        required = len(separator) + len(segment)
        if required <= remaining:
            parts.append(separator + segment)
            remaining -= required
            index += 1
            completed_items += 1
            next_inner_offset = 0
            continue
        separator_size = len(separator)
        available = max(0, remaining - separator_size)
        if available:
            parts.append(separator + segment[:available])
            next_inner_offset = item_offset + available
        break

    content = "".join(parts)
    result = {
        "section": section,
        "unit": unit,
        "offset": offset,
        "inner_offset": inner_offset,
        "limit": requested,
        "total": total,
        "content": content,
        "next_offset": index,
        "next_inner_offset": next_inner_offset,
        "complete": index >= total,
        "partial_item": bool(next_inner_offset),
    }
    return result, len(content)


def _bulk_rows_result(
    rows: list[Any],
    *,
    section: str | None,
    offset: int,
    limit: int | None,
    budget: int,
) -> tuple[dict[str, Any], int]:
    total = len(rows)
    requested = limit if limit is not None else DEFAULT_BULK_ITEM_LIMITS["rows"]
    selected: list[Any] = []
    used = 0
    index = min(max(0, offset), total)
    oversized_item = False
    while index < total and len(selected) < requested:
        row = rows[index]
        row_size = len(json.dumps(row, ensure_ascii=False, separators=(",", ":")))
        separator_size = 1 if selected else 0
        if used + separator_size + row_size > max(0, budget):
            oversized_item = not selected and row_size > max(0, budget)
            break
        selected.append(row)
        used += separator_size + row_size
        index += 1
    result = {
        "section": section,
        "unit": "rows",
        "offset": offset,
        "limit": requested,
        "total": total,
        "rows": selected,
        "next_offset": index,
        "complete": index >= total,
        "oversized_item": oversized_item,
    }
    if oversized_item:
        result["hint"] = "single row exceeds max_chars; retry this section with unit=characters"
    return result, used


def _workbook_sheet(value: dict[str, Any], section: str | None) -> dict[str, Any] | None:
    sheets = value.get("sheets", [])
    if section:
        return next((item for item in sheets if item.get("name") == section), None)
    candidates = _mother_sheet_candidates(sheets)
    return candidates[0] if candidates else next(
        (sheet for sheet in sheets if not is_reading_sheet(sheet.get("name"))), None)


def _bulk_one_range(
    value: dict[str, Any],
    request: dict[str, Any],
    *,
    default_unit: str,
    budget: int,
) -> tuple[dict[str, Any], int]:
    kind = str(value.get("kind") or "unknown")
    section = request.get("section")
    offset = max(0, int(request.get("offset", 0)))
    inner_offset = max(0, int(request.get("inner_offset", 0)))
    limit_value = request.get("limit")
    limit = None if limit_value is None else max(1, int(limit_value))
    requested_unit = str(request.get("unit") or default_unit or "auto")

    if kind == "text":
        text = str(value.get("text") or "")
        if requested_unit == "auto":
            if value.get("format") == "pdf" and isinstance(value.get("pages"), list):
                requested_unit = "pages"
            elif isinstance(value.get("paragraphs"), list):
                requested_unit = "paragraphs"
            else:
                requested_unit = "characters"
        if requested_unit == "characters":
            return _bulk_char_result(text, section=section, offset=offset, limit=limit, budget=budget)
        if requested_unit == "paragraphs":
            paragraphs = value.get("paragraphs")
            if not isinstance(paragraphs, list):
                paragraphs = _paragraphs_from_text(text)
            return _bulk_text_items_result(
                paragraphs,
                section=section,
                unit="paragraphs",
                offset=offset,
                inner_offset=inner_offset,
                limit=limit,
                budget=budget,
            )
        if requested_unit == "pages":
            pages = value.get("pages")
            if not isinstance(pages, list):
                raise ValueError("pages unit is only available for page-indexed text")
            return _bulk_text_items_result(
                pages,
                section=section,
                unit="pages",
                offset=offset,
                inner_offset=inner_offset,
                limit=limit,
                budget=budget,
            )
        raise ValueError(f"unit {requested_unit!r} is not valid for text")

    if kind == "workbook":
        if section == "canonical_storyboard":
            content = value.get("canonical_storyboard")
            text = json.dumps(content, ensure_ascii=False, separators=(",", ":"))
            if requested_unit not in {"auto", "characters"}:
                raise ValueError("canonical_storyboard supports unit=characters")
            return _bulk_char_result(text, section=section, offset=offset, limit=limit, budget=budget)
        sheet = _workbook_sheet(value, section)
        if not sheet:
            return (
                {
                    "section": section,
                    "unit": "rows",
                    "offset": offset,
                    "limit": limit or DEFAULT_BULK_ITEM_LIMITS["rows"],
                    "total": 0,
                    "rows": [],
                    "next_offset": 0,
                    "complete": True,
                    "available_sections": [item.get("name") for item in value.get("sheets", [])],
                },
                0,
            )
        rows = list(sheet.get("rows") or [])
        if requested_unit == "characters":
            text = json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
            return _bulk_char_result(text, section=sheet.get("name"), offset=offset, limit=limit, budget=budget)
        if requested_unit not in {"auto", "rows"}:
            raise ValueError(f"unit {requested_unit!r} is not valid for a workbook sheet")
        result, used = _bulk_rows_result(
            rows,
            section=sheet.get("name"),
            offset=offset,
            limit=limit,
            budget=budget,
        )
        result["headers"] = sheet.get("headers", [])
        return result, used

    if kind == "table":
        rows = list(value.get("rows") or [])
        if requested_unit == "characters":
            text = json.dumps(rows, ensure_ascii=False, separators=(",", ":"))
            return _bulk_char_result(text, section=section, offset=offset, limit=limit, budget=budget)
        if requested_unit not in {"auto", "rows"}:
            raise ValueError(f"unit {requested_unit!r} is not valid for a table")
        result, used = _bulk_rows_result(rows, section=section, offset=offset, limit=limit, budget=budget)
        result["headers"] = value.get("headers", [])
        return result, used

    if kind == "json":
        content = value.get("value")
        if isinstance(content, dict) and section:
            if section not in content:
                return (
                    {
                        "section": section,
                        "unit": requested_unit,
                        "offset": offset,
                        "limit": limit,
                        "total": 0,
                        "content": "",
                        "next_offset": 0,
                        "complete": True,
                        "available_sections": sorted(content.keys()),
                    },
                    0,
                )
            content = content[section]
        if isinstance(content, list) and requested_unit in {"auto", "rows"}:
            return _bulk_rows_result(content, section=section, offset=offset, limit=limit, budget=budget)
        if requested_unit not in {"auto", "characters"}:
            raise ValueError(f"unit {requested_unit!r} is not valid for this JSON value")
        text = json.dumps(content, ensure_ascii=False, separators=(",", ":"))
        return _bulk_char_result(text, section=section, offset=offset, limit=limit, budget=budget)

    text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if requested_unit not in {"auto", "characters"}:
        raise ValueError(f"unit {requested_unit!r} is not valid for {kind}")
    return _bulk_char_result(text, section=section, offset=offset, limit=limit, budget=budget)


def preview_extracted_bulk(
    value: dict[str, Any],
    ranges: list[Any] | None,
    *,
    default_unit: str = "auto",
    max_chars: int = DEFAULT_BULK_RESPONSE_CHARS,
) -> dict[str, Any]:
    """Read several bounded ranges in one MCP round trip.

    `max_chars` limits returned source/table content across every requested
    range. Metadata is intentionally small and is not counted as source text.
    """

    budget = max(1, min(int(max_chars), MAX_BULK_RESPONSE_CHARS))
    raw_ranges = ranges or [{}]
    if len(raw_ranges) > MAX_BULK_RANGES:
        raise ValueError(f"bulk_ranges supports at most {MAX_BULK_RANGES} ranges per call")
    requests = [_range_dict(item) for item in raw_ranges]
    results: list[dict[str, Any]] = []
    used = 0
    for request in requests:
        remaining = max(0, budget - used)
        result, consumed = _bulk_one_range(value, request, default_unit=default_unit, budget=remaining)
        if value.get("kind") == "workbook":
            result["audit_included"] = not is_reading_sheet(result.get("section"))
        used += min(max(0, consumed), remaining)
        results.append(result)
    truncated = any(not bool(item.get("complete")) for item in results)
    return {
        "kind": value.get("kind"),
        **({"workbook_audit_scope": workbook_audit_scope()} if value.get("kind") == "workbook" else {}),
        "ranges": results,
        "max_chars": budget,
        "returned_chars": used,
        "truncated": truncated,
    }
