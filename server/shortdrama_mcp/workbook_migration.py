"""Lossless, read-only recovery of the current ten-column legacy workbook.

This module never writes a sidecar, artifact, certificate or project record.
Only the service's explicitly authorized migration route may consume its proof.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
import io
import re
from typing import Any

from openpyxl import load_workbook

from .canonical import normalize_json, stable_json_hash
from .file_io import (
    CANONICAL_SIDECAR_SHEET,
    _normalized_mother_headers, _parse_function_scene_cell, _validate_zip_archive,
)
from .materializer import (COMPACT_SHOT_FIELDS,
    LEGACY_MOTHER_TEMPLATE_COLUMNS as MOTHER_TEMPLATE_COLUMNS,
    LEGACY_FORMAL_MERGED_ATOMIC_FIELDS as FORMAL_MERGED_ATOMIC_FIELDS)
from .r5_workbook import is_reading_sheet

WORKBOOK_MIGRATION_TARGET = "workbook-canonical-v1"
WORKBOOK_MIGRATION_PHASE = "legacy_workbook_canonical_migration"
WORKBOOK_PARSER_CONTRACT = "lossless-current-formal-workbook-v1"
_LABELS = ("镜头序号", "镜头类型", *FORMAL_MERGED_ATOMIC_FIELDS)
_LABEL_RE = re.compile(r"(?m)^(" + "|".join(map(re.escape, _LABELS)) + r")：")
_HEADER_RE = re.compile(
    r"(?P<id>EP\d{2}-\d{3})（(?P<duration>\d+(?:\.\d+)?)s；累计(?P<cumulative>\d+(?:\.\d+)?)s）"
)


def _fail(message: str) -> ValueError:
    return ValueError("workbook_migration_input_invalid: " + message)


def _decimal(value: Any, location: str) -> Decimal:
    if isinstance(value, bool) or value in (None, ""):
        raise _fail(f"missing numeric value at {location}")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise _fail(f"nonnumeric value at {location}") from exc
    if not result.is_finite():
        raise _fail(f"nonfinite value at {location}")
    return result


def shot_atoms(storyboard: dict[str, Any]) -> list[dict[str, Any]]:
    rows = storyboard.get("shots")
    if not isinstance(rows, list) or not rows:
        raise _fail("shots must be a nonempty array")
    if any(not isinstance(row, dict) for row in rows):
        raise _fail("every shot must be an object")
    return normalize_json([{field: row.get(field) for field in COMPACT_SHOT_FIELDS} for row in rows])


def recover_legacy_workbook(payload: bytes) -> tuple[dict[str, Any], dict[str, Any]]:
    _validate_zip_archive(payload)
    book = load_workbook(io.BytesIO(payload), read_only=False, data_only=False)
    cached = load_workbook(io.BytesIO(payload), read_only=False, data_only=True)
    try:
        if CANONICAL_SIDECAR_SHEET in book.sheetnames:
            raise _fail("a workbook already containing a sidecar is outside this route")
        candidates = [sheet for sheet in book.worksheets if not is_reading_sheet(sheet.title) and
                      _normalized_mother_headers([cell.value for cell in sheet[1]]) ==
                      _normalized_mother_headers(MOTHER_TEMPLATE_COLUMNS)]
        if len(candidates) != 1:
            raise _fail("exactly one formal ten-column sheet is required")
        sheet = candidates[0]
        value_sheet = cached[sheet.title]
        rows: list[dict[str, Any]] = []
        formulas: list[dict[str, Any]] = []
        source_cells: list[list[Any]] = []
        episode_totals: dict[str, Decimal] = {}
        episode_start: dict[str, int] = {}
        seen_ids: set[str] = set()
        last_episode: str | None = None
        last_data_row: int | None = None
        for row_index in range(2, sheet.max_row + 1):
            cells = [sheet.cell(row_index, column) for column in range(1, 11)]
            values = [cell.value for cell in cells]
            if all(value in (None, "") for value in values):
                continue
            if any(cell.data_type == "f" for i, cell in enumerate(cells) if i != 2):
                raise _fail(f"only cumulative-time formulas are supported: row {row_index}")
            if any(cell.data_type == "e" for cell in cells):
                raise _fail(f"Excel error at row {row_index}")
            formal = normalize_json(str(values[4] or "")).strip()
            matches = list(_LABEL_RE.finditer(formal))
            if ([match[1] for match in matches] != list(_LABELS)
                    or matches[0].start() != 0):
                raise _fail(f"missing, duplicate, reordered or unsupported formal labels at row {row_index}")
            fields = {match[1]: formal[match.end():matches[i + 1].start() if i + 1 < len(matches) else len(formal)].strip()
                      for i, match in enumerate(matches)}
            heading = _HEADER_RE.fullmatch(fields.pop("镜头序号"))
            shot_id = str(values[0] or "")
            if not heading or heading["id"] != shot_id or shot_id in seen_ids:
                raise _fail(f"shot identity mismatch or duplicate at row {row_index}")
            seen_ids.add(shot_id)
            episode = shot_id[:4]
            if episode != last_episode and episode in episode_totals:
                raise _fail("episodes must be contiguous")
            if episode == last_episode and last_data_row != row_index - 1:
                raise _fail("blank rows inside an episode prevent formula-chain verification")
            if episode not in episode_totals:
                episode_start[episode] = row_index
                episode_totals[episode] = Decimal(0)
            last_episode = episode
            last_data_row = row_index
            duration = _decimal(values[1], f"B{row_index}")
            if duration <= 0 or duration * 10 != (duration * 10).to_integral_value():
                raise _fail(f"duration must be positive at 0.1-second precision: {shot_id}")
            cumulative = episode_totals[episode] + duration
            if Decimal(heading["duration"]) != duration or Decimal(heading["cumulative"]) != cumulative:
                raise _fail(f"heading and episode timing disagree: {shot_id}")
            if cells[2].data_type == "f":
                formula = str(values[2])
                normalized = re.sub(r"\s+", "", formula).replace("$", "").upper()
                start = episode_start[episode]
                allowed = {f"=SUM(B{start}:B{row_index})"}
                allowed.add(f"=B{row_index}" if start == row_index else f"=C{row_index - 1}+B{row_index}")
                if start != row_index:
                    allowed.add(f"=B{row_index}+C{row_index - 1}")
                allowed |= {f"=ROUND({item[1:]},1)" for item in tuple(allowed)}
                if normalized not in allowed:
                    raise _fail(f"unsupported or cross-episode cumulative formula: C{row_index}")
                cache = value_sheet.cell(row_index, 3).value
                if cache is not None and abs(_decimal(cache, f"cached C{row_index}") - cumulative) > Decimal("0.0000001"):
                    raise _fail(f"stale cumulative formula cache: C{row_index}")
                formulas.append({"cell":f"C{row_index}", "formula":formula, "cached_value":cache,
                                 "verified_value":float(cumulative)})
            elif abs(_decimal(values[2], f"C{row_index}") - cumulative) > Decimal("0.0000001"):
                raise _fail(f"cumulative value disagrees: {shot_id}")
            episode_totals[episode] = cumulative
            kind_subject = fields.pop("镜头类型")
            if kind_subject.count("｜可见主体：") != 1:
                raise _fail(f"ambiguous visible-subject separator: {shot_id}")
            kind, subject = kind_subject.split("｜可见主体：")
            function, scene = _parse_function_scene_cell(values[3])
            if not function or not scene:
                raise _fail(f"missing or ambiguous function/scene ID: {shot_id}")
            row = {"镜头编号":shot_id, "时长":float(duration), "累计时长":float(cumulative),
                   "功能类型":function, "场景连续性ID":scene, "镜头类型":kind, "可见主体":subject,
                   **fields, **dict(zip(COMPACT_SHOT_FIELDS[-4:], values[6:10]))}
            if set(row) != set(COMPACT_SHOT_FIELDS):
                raise _fail(f"exact 27-field recovery failed: {shot_id}")
            rows.append({field:row[field] for field in COMPACT_SHOT_FIELDS})
            source_cells.append(values)
        if not rows:
            raise _fail("formal sheet has no shots")
        canonical = {"project_version":"V16.3-148-R10-StableRuntime", "schema_version":"R10.1",
                     "workflow_profile":"R10-PG1", "field_order":list(COMPACT_SHOT_FIELDS),
                     "source_lock":{}, "prevention_evidence":{}, "shots":rows,
                     "total_duration":float(sum(episode_totals.values()))}
        proof = {"parser_contract":WORKBOOK_PARSER_CONTRACT, "formal_sheet":sheet.title,
                 "shot_ids":[row["镜头编号"] for row in rows], "episode_ids":list(episode_totals),
                 "episode_durations":{key:float(value) for key,value in episode_totals.items()},
                 "scene_ids":list(dict.fromkeys(row["场景连续性ID"] for row in rows)),
                 "atom_sha256":stable_json_hash(shot_atoms(canonical)),
                 "source_cells_sha256":stable_json_hash(source_cells),
                 "formula_evidence":formulas, "formula_evidence_sha256":stable_json_hash(formulas),
                 "sidecar_present":False, "formal_certificate_issued":False}
        return canonical, proof
    finally:
        book.close()
        cached.close()
